"""Unpublished, independently captured Yaohai and Feidong sources and Feidong fill.

Only the fixed 2026-10-29 public evidence is accepted. No network, database,
neighbouring locality, arbitrary template path, Word or model calls occur here.
"""
import hashlib
import io
import json
import re
import unicodedata
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT=Path(__file__).resolve().parents[1]
BASE=Path('reports/nationwide-verification-20261029/anhui')
LOCK_SHA='04b4f92ec378eae38e36d360777c10c78de2b54d9e064b7297ede75a954614b4'
FEIDONG_ADAPTER_KEY='anhui-feidong-official-application-20261029'
SPECS={
    'yaohai':('瑶海区','340102','b62d6ac30d27495b817e0b73bc4e3d2e','113401020029922420400201430200201','合肥市瑶海区人力资源和社会保障局'),
    'feidong':('肥东县','340122','94e006563cb441c1a2827939cafe265f','11340122003005878M400201430200201','肥东县人力资源和社会保障局'),
}
SECTIONS={'basic-expanded':'基本信息','conditions':'受理条件','materials':'申请材料','process':'办理流程','intermediary':'中介服务','faq':'常见问题','evaluation':'办事评价'}

def _json(root, relative):
    return json.loads((root/relative).read_text(encoding='utf8'))

def _locked(root, locality):
    lock_path=root/BASE/'source-lock.json'
    if hashlib.sha256(lock_path.read_bytes()).hexdigest()!=LOCK_SHA:
        raise ValueError('Official evidence lock was changed; a fresh source review is required')
    pins=_json(root,BASE/'source-lock.json')
    own={path:sha for path,sha in pins.items() if path.startswith((BASE/locality).as_posix()+'/')}
    if not own:
        raise ValueError('Missing independently captured locality')
    for relative,expected in own.items():
        path=(root/relative).resolve()
        if not path.is_relative_to(root.resolve()) or hashlib.sha256(path.read_bytes()).hexdigest()!=expected:
            raise ValueError('Own official evidence or original bytes differ: '+relative)
    return own

def load_guide(locality, *, root=ROOT):
    """Build an unpublished candidate from this locality's own pinned public sources."""
    if locality not in SPECS:
        raise ValueError('No independently verified source for requested locality')
    root=Path(root)
    district,region,code,service,authority=SPECS[locality]
    own=_locked(root,locality)
    url='https://hf.ahzwfw.gov.cn/bog-bsdt/static/workProcess/components/applicationMaterial.html?ssqdCode='+code
    public=_json(root,BASE/locality/'public-guide-viewmodel.json');m=public['data']
    if (public['source_url']!=url or [m.get(k) for k in ['ssqdZoneName','ssqdZoneCode','ssqdCode','ssqdBaseCode','ssqdDeptName']]
            !=[district,region+'000000',code,service,authority] or m['ssqdName']!='劳动人事争议仲裁申请'):
        raise ValueError('Own region, authority or service identity does not match')
    fields={};sections={};proof={}
    for folder,label in SECTIONS.items():
        relative=BASE/locality/folder/'public-source.json'
        page=_json(root,relative)
        if page['url']!=url or authority not in page['text'] or page['title']!=m['ssqdName'] or len(page['sections'])!=1:
            raise ValueError('Not the independently captured official panel: '+folder)
        sections[label]=page['sections'][0]['text']
        proof[label]={'capture_file':relative.as_posix(),'capture_sha256':own[relative.as_posix()], 'collected_at':page['captured_at'],'source_url':url}
        if folder=='basic-expanded':
            for table in page['tables']:
                for row in table:
                    if len(row) in (2,4):
                        for index in range(0,len(row),2):
                            key,value=row[index:index+2]
                            value=re.sub(r'\s*(说明|查看地图)$','',value).strip()
                            fields[key.strip()]={**proof[label],'raw_value':value,'normalized_value':value or None,'verification_status':'verified' if value else 'missing'}
    if fields['所属区划']['raw_value']!=district or fields['实施清单编码']['raw_value']!=service:
        raise ValueError('Rendered official identity differs from own public component')
    records=_json(root,BASE/locality/'downloads/manifest.json')
    materials=[];expected=[]
    for index,row in enumerate(m['clJson'],1):
        detail_folder=f'detail-{index}' if locality=='feidong' else ['detail-application','detail-id','detail-evidence-success'][index-1]
        detail=_json(root,BASE/locality/detail_folder/'public-source.json')
        normalized=lambda value:re.sub(r'\s+','',value)
        if not all(normalized(row[k]) in normalized(detail['text']) for k in ['materialEvidence','acceptStandard']):
            raise ValueError('Material legal basis or acceptance standard not actually captured')
        materials.append({'id':row['id'],'name':row['materialName'],'required':row['materialNeedTxt'],
                          'original_copy':row['json']['gbcllxTxt'],'copies':row['json']['ggfs'],'raw_copies':row['json']['ggfs'],
                          'format':row['json']['gbclxsTxt'],'verification_status':'conflict','applicable_conditions':None,
                          'source_url':url,'details':{k:row[k] for k in ['sourceChannelTxt','sourceChannelDescription','editNotice','materialEvidence','acceptStandard']},
                          'detail_evidence':(BASE/locality/detail_folder/'public-source.json').as_posix()})
        for field,kind in [('materialEmpty','blank'),('materialSample','sample')]:
            expected.extend((a,kind,row['id']) for a in row['json'][field])
    expected.extend((a,'instruction',None) for a in m['processImgJson'])
    if {a['id'] for a,_,_ in expected}!={r['id'] for r in records} or len(records)!=len(expected):
        raise ValueError('Own manifest does not exactly match official attachments')
    templates=[]
    for item,kind,material_id in expected:
        record=next(r for r in records if r['id']==item['id'])
        expected_path=(BASE/locality/'downloads'/(item['id']+'-'+item['fileName'])).as_posix()
        if (record['source_url']!=url or record['requested_url']!=item['filePath'] or record['final_url']!=item['filePath']
                or not item['filePath'].startswith('https://hf.ahzwfw.gov.cn/intranet/group')
                or record['classification']!=kind or record['path']!=expected_path or record['status']!='downloaded'
                or record['http_status']!=200 or own[expected_path]!=record['sha256']):
            raise ValueError('Own original scope, classification or download provenance changed')
        templates.append({'id':item['id'],'material_id':material_id,'official_kind':kind,'original_name':item['fileName'],
                          'file_type':item['fileExt'],'file':expected_path,'download_url':item['filePath'],
                          'source_url':url,'sha256':record['sha256'],'bytes':record['bytes'],'downloaded_at':record['downloaded_at'],
                          'download_status':'downloaded','document_types':[],
                          'content_classification':'blank_application' if kind=='blank' else 'official_sample' if kind=='sample' else 'workflow_instruction',
                          'content_verification_status':'original_visually_reviewed'})
    boundary=_json(root,BASE/locality/'online-boundary/public-summary.json')
    if boundary['source_url']!=url or boundary['origin']!='https://sso.ahzwfw.gov.cn' or boundary['path']!='/uccp-server/login' or not boundary['login_visible']:
        raise ValueError('Own native online authentication boundary missing')
    sections['时限说明']='\n'.join(x for x in [m['promiseTimeDesc'],m['legalTimeDesc']] if x)
    issues=[{'type':'OFFICIAL_SOURCE_CONFLICT','reason':'材料表复印件/原件或固定份数，与填报须知“窗口办理提供原件”及材料依据按被申请人人数提交副本并列保留，未统一推断。'},
            {'type':'OFFICIAL_SOURCE_CONFLICT','reason':'法定45个工作日、说明45日及最多延期15日、承诺1工作日分别保留，不表示案件1日裁决。'}]
    if locality=='yaohai':
        issues.append({'type':'OFFICIAL_SOURCE_CONFLICT','reason':'页面多点可办地址列民生中心2楼，公开字段hanleAddr及邮寄地址又列10楼；不静默选择。'})
    else:
        issues.append({'type':'OFFICIAL_SOURCE_CONFLICT','reason':'基本信息列网上/窗口且0次到场；流程仅写窗口收件。网办实际到统一认证，登录后能力未核验。'})
    return {'province':'anhui','province_name':'安徽省','city':'合肥市','district':district,'region_code':region,
            'official_service_region_code':region+'000000','service_name':m['ssqdName'],'service_code':service,'authority':authority,
            'source_url':url,'source_title':m['ssqdName'],'source_type':'official_service_guide','collected_at':public['captured_at'],
            'verified_at':public['captured_at'],'verification_status':'conflict','fields':fields,'materials':materials,'templates':templates,
            'sections':sections,'evidence':proof,'official_identity':{k:m[k] for k in ['ssqdCode','ssqdZoneCode','ssqdZoneName','ssqdDeptCode','ssqdDeptName','ssqdBaseCode','ssqdVersion','updateTime']},
            'online_evidence':boundary,'channel':{'form':'window_and_online','online_status':'auth_required','verified':False},
            'source_issues':issues,'template_search_status':'found_in_own_checked_source','publication_status':'unpublished',
            'can_generate':False,'candidate_web_qa':'NOT_RUN','warnings':['仅适用本区自身官方来源；不代表全市或全省。','实际网办止于统一认证，未登录或提交。','候选网页未验收，不发布、不增加181/166覆盖计数。']}

def _width(value):
    return sum(2 if unicodedata.east_asian_width(c) in ('W','F') else 1 for c in value)

def _clean(value):
    if not isinstance(value,str) or any(ord(c)<32 and c not in '\n' for c in value):
        raise ValueError('Plain text without control characters required')
    return value.strip()

def _wrap(value, first=56, later=56):
    value=_clean(value)
    result=[];line='';capacity=first
    for char in value:
        if char=='\n':
            result.append(line);line='';capacity=later
        elif _width(line+char)>capacity:
            result.append(line);line=char;capacity=later
        else:
            line+=char
    if line:result.append(line)
    return result

def _replace(paragraph,value):
    if paragraph.runs:
        paragraph.runs[0].text=value
        for run in paragraph.runs[1:]:run.text=''
    else:paragraph.add_run(value)
    # The official blank mixes 16/20/22pt in empty line runs. Populated runs use
    # its principal 16pt size, so the measured line capacity stays meaningful.
    for run in paragraph.runs:run.font.size=Pt(16)

def fill_feidong_application(guide,payload,*,root=ROOT):
    """Fill the own Feidong blank only, returning offline DOCX bytes.

    Explicit two-line party descriptions preserve every supplied character. The
    bounded field contract rejects overcapacity and unsupported fields, so no
    submitted fields can silently disappear. Signature/date/counts stay blank.
    """
    canonical=load_guide('feidong',root=root)
    if not isinstance(guide,dict) or guide!=canonical:
        raise ValueError('An exact, own-source Feidong candidate is required')
    allowed={'applicant_lines','agent_lines','respondent_lines','representative_lines','requests','facts'}
    if not isinstance(payload,dict) or set(payload)-allowed:
        raise ValueError('Unsupported fields for this own official blank')
    party={}
    for key in ['applicant_lines','agent_lines','respondent_lines','representative_lines']:
        lines=payload.get(key,[])
        if not isinstance(lines,list) or len(lines)>2 or any(not isinstance(x,str) or '\n' in x or _width(_clean(x))>(40 if i==0 else 50) for i,x in enumerate(lines)):
            raise ValueError('Each party has two lines: at most 20 then 25 full-width characters')
        party[key]=[_clean(x) for x in lines]
    requests=payload.get('requests',[])
    if not isinstance(requests,list) or any(not isinstance(x,str) for x in requests):
        raise ValueError('Requests must be a list of plain-text strings')
    request_lines=_wrap('\n'.join(requests),44,56)
    fact_lines=_wrap(payload.get('facts',''))
    if len(request_lines)>9 or len(fact_lines)>10:
        raise ValueError('Official blank capacity exceeded; no text was truncated')
    blank=next(t for t in canonical['templates'] if t['official_kind']=='blank')
    document=Document(Path(root)/blank['file']);p=document.paragraphs
    if len(p)!=47 or p[0].text!='劳动人事争议仲裁申请书' or p[19].text!='事实和理由（包括证据和证据来源，证人姓名和地址等情况）：':
        raise ValueError('Own official blank paragraph structure changed')
    for key,index in [('applicant_lines',2),('agent_lines',4),('respondent_lines',6),('representative_lines',8)]:
        lines=party[key]
        if lines:
            _replace(p[index],p[index].text+lines[0])
        if len(lines)==2:_replace(p[index+1],lines[1])
    for index,line in enumerate(request_lines):
        if index==0:
            p[10].runs[-1].text=line
            p[10].runs[-1].font.size=Pt(16)
        else:_replace(p[10+index],line)
    for index,line in enumerate(fact_lines):_replace(p[20+index],line)
    # Two terminal empty paragraphs create an entirely blank third source page.
    # Remove only those terminal empty paragraphs in the filled derivative.
    for paragraph in p[45:]:
        if paragraph.text.strip():raise ValueError('Unexpected terminal source content')
        paragraph._p.getparent().remove(paragraph._p)
    stream=io.BytesIO();document.save(stream)
    return stream.getvalue()
