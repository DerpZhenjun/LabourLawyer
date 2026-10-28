from __future__ import annotations

import hashlib
import io
import json
from typing import Any

from docx import Document


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _cell_text(cell) -> str:
    return "\n".join(paragraph.text for paragraph in cell.paragraphs).strip()


def analyze_docx(data: bytes) -> tuple[str, dict[str, Any]]:
    document = Document(io.BytesIO(data))
    sections = []
    for section in document.sections:
        sections.append(
            {
                "page_width": section.page_width,
                "page_height": section.page_height,
                "top_margin": section.top_margin,
                "bottom_margin": section.bottom_margin,
                "left_margin": section.left_margin,
                "right_margin": section.right_margin,
            }
        )
    shape = {
        "paragraphs": [paragraph.text.strip() for paragraph in document.paragraphs],
        "tables": [
            [[_cell_text(cell) for cell in row.cells] for row in table.rows]
            for table in document.tables
        ],
        "sections": sections,
    }
    encoded = json.dumps(shape, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest(), {
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "section_count": len(document.sections),
    }
