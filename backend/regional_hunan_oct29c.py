"""Rainflower/Yuhua and Wangcheng own captured sources; offline deterministic filling.

The frozen public source bundle is the sole authority. Equal attachment bytes
permit reuse of a form-editing algorithm, never another district's rules.
"""
import hashlib,json,math,re
from copy import deepcopy
from pathlib import Path
from urllib.parse import unquote,urlsplit,parse_qs
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from regional_documents import PreservedPackageDocument
from WenShu.template_fill import set_cell_text,replace_text,date_parts

ROOT=Path(__file__).resolve().parents[1]
REPORT='reports/nationwide-verification-20261029-c/hunan'
ADAPTATION='reports/nationwide-verification-20261029-c/hunan-adaptation'
ADAPTER='hunan_own_legacy_worker_oct29c_v1'
MANIFEST_SHA256='f0c2af59b09ea082d5e3d4678e0ed962415dc45ba03e8cedad98f218973bfdbc'
WHITELIST_SHA256='28f01defd91b9bfe9bd0796d4d3a4b44b85307e6e24c62106a063a19cd92a5e7'
DERIVED_MANIFEST_SHA256='05529145274cb0c8d1f630e9f508502e36d1b99d2189cddb65234d7316643a68'
WORKER_SHA='c6bbebf6eed3a11f1cbe5ee14cd6de9937f741c9fcda37165e06d0706e18ea97'
SPECS={
 'yuhua':{'code':'430111','district':'雨花区','approve':'E499C22DC24A06E7E053671515AC91D4','implementation':'12430111444932183M4002014302002','authority':'雨花区人社局','place':'朝晖路499号','phone':'0731-85880234'},
 'wangcheng':{'code':'430112','district':'望城区','approve':'f3060b49b9ed4a50b1b300254417b324','implementation':'1143012200615852XL4002014302002','authority':'望城区人力资源和社会保障局','place':'望府路200号信访接待中心111办公室','phone':'0731-88062246'},
}

def sha(body):return hashlib.sha256(body).hexdigest()
class Evidence:
 def __init__(self,root):
  self.root=Path(root).resolve()
  raw=(self.root/'SOURCE-MANIFEST.json').read_bytes();public=(self.root/'PUBLIC-WHITELIST.json').read_bytes()
  if sha(raw)!=MANIFEST_SHA256 or sha(public)!=WHITELIST_SHA256:raise ValueError('湖南C固定公开封存清单不匹配')
  self.files=json.loads(raw)['files'];self.public=json.loads(public)['files']
 def read(self,relative):
  path=(self.root/relative).resolve()
  if not path.is_relative_to(self.root):raise ValueError('来源路径越界')
  body=path.read_bytes();expected={'sha256':sha(body),'bytes':len(body)}
  if self.public.get(REPORT+'/'+relative)!=expected:raise ValueError('公开来源字节不匹配：'+relative)
  if relative in self.files and self.files[relative]!=expected:raise ValueError('原始来源清单不匹配')
  return body
 def load(self,relative):return json.loads(self.read(relative).decode('utf8'))

def parse_capture(name,capture_root=None):
 if name not in SPECS:raise ValueError('仅湖南C两区自身事项')
 e=Evidence(capture_root or ROOT/REPORT);s=SPECS[name];g=e.load(name+'/guide.json');source=g['source_url']
 if (g['region_code']!=s['code'] or g['district']!=s['district'] or g['authority']!=s['authority'] or g['service_code']!=s['implementation']
     or parse_qs(urlsplit(source).query).get('approve_id')!=[s['approve']] or g['service_name']!='劳动人事争议仲裁申请'):
  raise ValueError('自身地区、机关、事项或approve_id不匹配')
 cat=e.load(name+'/public-catalogue.json');c=cat['public_response'];request=cat['observed_request']
 if (c['approveId']!=s['approve'] or c['approveCode']!=s['implementation'] or c['areaCode']!=s['code']+'999000'
     or c['orgName']!=s['authority'] or c['approveName']!=g['service_name'] or request['status']!=200 or s['approve'] not in unquote(request['url'])):
  raise ValueError('自身匿名目录身份不匹配')
 for section in ['guide','materials-expanded','rights-duties','result-sample','faq','template-control','sample-control']:
  p=e.load(name+'/'+section+'/page.json');e.read(name+'/'+section+'/page.png')
  if p['url']!=source or s['authority'] not in p['text'] or g['service_name'] not in p['text']:raise ValueError('自身正文机关或事项缺失')
 page=e.load(name+'/guide/page.json')
 if s['place'] not in page['text'] or g['fields']['咨询服务']['raw_value']!=s['phone'] or s['phone'] not in page['text']:raise ValueError('自身窗口或电话缺失')
 if len(g['materials'])!=4 or [m['name'] for m in g['materials']]!=['劳动仲裁申请书','居民身份证','营业执照','证据材料']:raise ValueError('自身四材料不匹配')
 if [m['raw'] for m in g['materials']]!=c['materialList']:raise ValueError('材料与自身匿名目录不匹配')
 records=e.load(name+'/originals.json');observation=e.load(name+'/attachment-observation.json')
 if len(records)!=4 or observation['source_url']!=source:raise ValueError('自身四原件证据缺失')
 for t in g['templates']:
  r=next(x for x in records if x['official_annex_id']==t['official_annex_id']);e.read(r['file']+'.source.json');body=e.read(r['file'])
  annex=next((a for m in c['materialList'] for a in m.get('attachList',[]) if a['attachId']==r['official_annex_id']),None)
  if (not annex or r!=e.load(r['file']+'.source.json') or t['material_id']!=r['material_id'] or annex['businessId']!=r['material_id']
      or r['source_url']!=source or r['approve_id']!=s['approve'] or r['area_code']!=s['code']+'999000'
      or t['file']!=REPORT+'/'+r['file'] or t['sha256']!=r['sha256'] or sha(body)!=r['sha256'] or len(body)!=t['bytes']
      or t['original_name']!=annex['attachName'] or t['official_kind']!=r['official_kind']
      or r['attach_code']!=annex['attachCode'] or r['attach_code']!=('000002' if r['official_kind']=='blank' else '000003')
      or not any(unquote(x['url'])==unquote(r['download_url']) and x['status']==200 for x in observation['public_fetches'])):
   raise ValueError('自身附件、材料、原件来源或SHA不匹配')
 boundary=e.load(name+'/online-boundary.json')
 if boundary!=g['online_evidence'] or boundary['source_url']!=source or boundary['origin']!='https://auth.zwfw.hunan.gov.cn' or boundary['path']!='/oauth2/govLogin.html' or not boundary['login_visible'] or boundary['submitted'] or boundary['authenticated']:
  raise ValueError('自身网办认证边界不匹配')
 g['official_approve_id']=s['approve'];g['evidence_template_status']=g['separate_evidence_list_template_status']
 g['evidence']['manifest_sha256']=MANIFEST_SHA256
 g['channel']['entry_url']=source
 g['channel']['entry_check']={'url':boundary['origin']+boundary['path'],'official_link':True,'status':'auth_required','verified_at':boundary['captured_at'],'submitted':False,'post_login_flow':'unknown'}
 return g

def _derived(name,root=None):
 base=Path(root or ROOT);raw=(base/ADAPTATION/'DERIVED-MANIFEST.json').read_bytes()
 if sha(raw)!=DERIVED_MANIFEST_SHA256:raise ValueError('C转换清单已变化')
 item=next(x for x in json.loads(raw)['files'] if x['district']==name);source=base/item['source'];path=base/item['file']
 if item['source_sha256']!=WORKER_SHA or sha(source.read_bytes())!=WORKER_SHA or sha(path.read_bytes())!=item['sha256'] or item['kind']!='platform_conversion_of_own_official_doc_not_official_docx':raise ValueError('C自己原件至转换副本链条不匹配')
 return item,path

FACT_KEYS=['province','city','district','region_code','source_type','source_url','source_id','service_name','authority','service_code','service_ids','official_approve_id','fields','materials','sections','channel','online_evidence','evidence','evidence_template_status']
def blank_binding(guide,template,root=None):
 try:
  base=Path(root or ROOT);name=next(n for n,s in SPECS.items() if s['code']==guide.get('region_code'));retained=parse_capture(name,base/REPORT)
  if any(guide.get(k)!=retained[k] for k in FACT_KEYS):return None
  original=next(t for t in retained['templates'] if t['id']==template.get('id'))
  if original['official_kind']!='blank' or original['sha256']!=WORKER_SHA:return None
  if any(template.get(k)!=v for k,v in original.items() if k not in {'can_generate','document_types'}):return None
  _derived(name,base);return 'application'
 except (ValueError,OSError,KeyError,TypeError,StopIteration):return None

def with_capabilities(guide,root=None):
 result=deepcopy(guide)
 for t in result['templates']:
  kind=blank_binding(result,t,root);t.update(can_generate=bool(kind),document_types=[kind] if kind else [])
  if kind:t.update(adapter=ADAPTER,content_kind='filled_platform_conversion_of_own_official_doc',source_download_format='doc',generation_label='依据本区官方DOC空白表的平台转换副本填写（劳动者申请）')
 result['warnings']=[w for w in result.get('warnings',[]) if '尚未转换' not in w]
 result['warnings'].append('仅支持本区劳动者申请书的平台转换副本；官方原件为DOC，转换件不是官方DOCX原件。附随证据属于申请书，单位版与独立证据清单未适配。候选网页验收和发布由部署流程另行记录。')
 return result

ORDINARY=set('applicantName gender birthDate occupation nationality idAddress postalCode phone workUnit workUnitPostalCode workUnitPhone respondentName registeredAddress respondentPostalCode legalRepName legalRepPosition legalRepPhone hrContactName hrContactPosition hrContactPhone thirdParty committeeName applicationDate copyCount'.split())
def _validate(data):
 if not isinstance(data,dict) or set(data)-ORDINARY-{'requests','factsReason','evidenceList'}:raise ValueError('本表不支持的字段')
 def text(v,limit=None):
  if not isinstance(v,str) or re.search('[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]',v) or (limit is not None and len(v)>limit):raise ValueError('填写格式或容量不符合本表合同')
 for k,v in data.items():
  if k in ORDINARY:
   if k=='copyCount' and type(v) in {int,float} and math.isfinite(v):v=str(v)
   text(v,160)
 text(data.get('factsReason',''),1800)
 requests=data.get('requests',[]);items=data.get('evidenceList',[])
 if not isinstance(requests,list) or len(requests)>8 or not isinstance(items,list) or len(items)>10:raise ValueError('最多8项请求与10项附随证据')
 for value in requests:text(value)
 for item in items:
  if not isinstance(item,dict) or set(item)-{'name','source','purpose'}:raise ValueError('附随证据字段不匹配')
  for value in item.values():text(value,300)


def fill_document(guide,template,kind,path,data):
 if kind!='application' or blank_binding(guide,template)!=kind or template.get('adapter')!=ADAPTER or template.get('can_generate') is not True:
  raise ValueError('仅支持本区严格来源绑定的劳动者申请书转换副本')
 if Path(path).resolve()!=(ROOT/template['file']).resolve():raise ValueError('不得传入其他地区或任意原件路径')
 name=next(n for n,s in SPECS.items() if s['code']==guide['region_code']);item,derived_path=_derived(name)
 _validate(data)
 values=deepcopy(data);requests=values.get('requests') or [];items=values.get('evidenceList') or []
 if not isinstance(requests,list) or not isinstance(items,list) or len(requests)>8 or len(items)>10:raise ValueError('本表最多8项请求和10项附随证据')
 if len(str(values.get('factsReason') or ''))>1800:raise ValueError('事实理由超出已验收容量')
 for k,v in values.items():
  if k not in {'requests','evidenceList','factsReason'} and len(str(v or ''))>160:raise ValueError('本表字段过长：'+k)
 original=derived_path.read_bytes();document=Document(derived_path);table=document.tables[0];rows=list(table.rows)
 if (len(document.tables)!=1 or len(rows)!=57 or len(table.columns)!=1 or '申请人姓名' not in rows[0].cells[0].text
     or '请求事项' not in rows[10].cells[0].text or '事实与理由' not in rows[19].cells[0].text or '证据和证据来源' not in rows[34].cells[0].text):
  raise ValueError('该转换副本的已检查字段位置发生变化')
 def v(k):return str(values.get(k) or '')
 def put(i,text):set_cell_text(rows[i].cells[0],text)
 year,month,day=date_parts(values.get('birthDate'));birth=(f'{year}年{month}月{day}日' if year and month and day else v('birthDate'))
 slots={0:'申请人姓名：'+v('applicantName'),1:'性别：'+v('gender')+'    出生年月日：'+birth+'    职业：'+v('occupation')+'    民族：'+v('nationality'),
    2:'住所：'+v('idAddress')+'    邮编：'+v('postalCode')+'    电话：'+v('phone'),
    3:'工作单位：'+v('workUnit')+'    邮编：'+v('workUnitPostalCode')+'    电话：'+v('workUnitPhone'),
    4:'被申请人名称：'+v('respondentName'),5:'住所：'+v('registeredAddress')+'    邮编：'+v('respondentPostalCode'),
    6:'法定代表人（或主要负责人）：'+v('legalRepName')+'    职务：'+v('legalRepPosition')+'    电话：'+v('legalRepPhone'),
    7:'人力资源管理部门负责人：'+v('hrContactName')+'    职务：'+v('hrContactPosition')+'    电话：'+v('hrContactPhone'),8:'第三人：'+v('thirdParty')}
 for i,text in slots.items():put(i,text)
 for i,request in enumerate(requests):put(11+i,f'{i+1}. {request}')
 put(20,v('factsReason'))
 # Remove only unused fill rows from this exact source, preserving labels,
 # signatures, notes and the original page/table geometry and font sizes.
 for i in list(range(11+max(1,len(requests)),19))+list(range(21,33)):
  table._tbl.remove(rows[i]._tr)
 footer=rows[33].cells[0]
 if v('committeeName'):replace_text(footer.paragraphs[0],'×××劳动争议仲裁委员会',v('committeeName'))
 p=footer.paragraphs[1];text=p.text;year,month,day=date_parts(values.get('applicationDate'))
 if year and month and day:text=re.sub(r'年\s+月\s+日',f'{year}年 {month}月 {day}日',text,count=1)
 elif v('applicationDate'):text=re.sub(r'年\s+月\s+日',lambda m:v('applicationDate'),text,count=1)
 if values.get('copyCount') is not None:text=re.sub(r'申请书副本\s+份',lambda m:'申请书副本 '+str(values['copyCount'])+' 份',text,count=1)
 if text!=p.text:replace_text(p,p.text,text)
 for i,evidence in enumerate(items):
  if not isinstance(evidence,dict) or any(len(str(evidence.get(k) or ''))>300 for k in ['name','source','purpose']):raise ValueError('附随证据格式或文字长度不符合已验收容量')
  put(35+i,f"{i+1}. {evidence.get('name') or ''}；来源：{evidence.get('source') or ''}；证明目的：{evidence.get('purpose') or ''}")
 for i in range(35+max(1,len(items)),57):table._tbl.remove(rows[i]._tr)
 for i in [10,19,34]:
  for paragraph in rows[i].cells[0].paragraphs:paragraph.paragraph_format.keep_with_next=True
 for i in range(35,35+len(items)):
  props=rows[i]._tr.get_or_add_trPr()
  if props.find(qn('w:cantSplit')) is None:props.append(OxmlElement('w:cantSplit'))
 return PreservedPackageDocument(document,original)
