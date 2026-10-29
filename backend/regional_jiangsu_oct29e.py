"""Own-source Suzhou E worker forms; explicit roots never use private fallback."""
from copy import deepcopy
from datetime import date
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
from docx import Document
from docx.oxml.ns import qn
from docx.table import _Row
from regional_guides import digest
from regional_documents import PreservedPackageDocument
from WenShu.exact_template_fill import keep_body, fill_sections, fill_footer, remove_empty_edges
from WenShu.template_fill import replace_span, fill_label_slots, set_cell_text

ROOT = Path(__file__).resolve().parents[1]
REPORT = Path('reports/nationwide-verification-20261029-e/jiangsu')
ADAPTER = 'jiangsu_own_blank_oct29e_v1'
SUPPORTED = frozenset({'320507', '320509'})
MANIFEST_SHA = '7f92920807dc4ec74f26ace3be76adf67303cbd4877994777c1a8ca0b1e7bb5e'
WHITELIST_SHA = 'd3227b2755591c2be1375269a033bac6e62653d367b81150f89a0667091f9a8f'
SPECS = {
    '320507': dict(name='xiangcheng', district='相城区', authority='苏州市相城区人力资源和社会保障局',
        task='11320507555810103L400201430200201', webId='28',
        application='02a0e257-ba98-41d8-9d93-1f5c70d038f7', evidence='873e8c9a-fb43-466a-9a57-025a4c2629a9',
        guide_sha='1ffff398538093a61ba9e2c87d8c8a803e9b63a121a619d30f5724136c92fa26'),
    '320509': dict(name='wujiang', district='吴江区', authority='苏州市吴江区人力资源和社会保障局',
        task='1132058401420395XF400201430200201', webId='29',
        application='23baf5d4-a613-4e85-9893-95e556164af1', evidence='d51313d9-c124-4d31-96de-5e7fde669684',
        guide_sha='9c76ef7b6441e250eccc62167c49dc70438f39466034ad1c792557fac9a087db'),
}
APPLICATION_FIELDS = frozenset({'committeeName', 'applicantName', 'birthDate', 'gender', 'idNumber', 'idAddress',
    'otherDocumentType', 'otherDocumentNumber', 'applicantMailingAddress', 'phone', 'acceptMediation', 'legalAid',
    'respondentName', 'respondentType', 'registeredAddress', 'socialCreditCode', 'legalRepName', 'legalRepPosition',
    'departmentHead', 'departmentHeadPosition', 'mailingAddress', 'companyContact', 'companyContactPhone',
    'requests', 'factsReason', 'applicationDate', 'copyCount', 'evidenceCopyCount'})
EVIDENCE_FIELDS = frozenset({'committeeName', 'caseNumber', 'submitDate', 'evidenceList'})
EVIDENCE_ITEM_FIELDS = frozenset({'name', 'source', 'purpose', 'originalType', 'pageCount'})
DATA_FIELDS = {'application': APPLICATION_FIELDS, 'evidence': EVIDENCE_FIELDS, 'evidence_item': EVIDENCE_ITEM_FIELDS}
LIMITS = {'ordinary': 300, 'factsReason': 12000, 'requests': 4, 'request': 1200, 'evidenceList': 40,
          'evidence_name': 300, 'evidence_source': 300, 'evidence_originalType': 300,
          'evidence_pageCount': 300, 'evidence_purpose': 1200, 'copyCount_min': 1, 'copyCount_max': 30}
# Source fields are immutable; lifecycle/QA/capability annotations can evolve.
SOURCE_KEYS = ('province', 'province_name', 'city', 'district', 'region_code', 'service_name',
               'service_code', 'authority', 'source_url', 'source_title', 'source_type',
               'fields', 'materials', 'sections', 'scope_identity', 'evidence', 'channel',
               'online_boundary', 'source_issues', 'attachment_counts')
TEMPLATE_KEYS = ('id', 'annex_id', 'region_code', 'material_id', 'official_kind', 'original_name',
                 'source_url', 'download_url', 'file', 'file_type', 'sha256', 'bytes',
                 'download_status', 'public_release_status', 'content_role')


def _read_checked(path, checksum, boundary=None):
    path = Path(path).resolve()
    if boundary is not None and not path.is_relative_to(Path(boundary).resolve()):
        raise ValueError('自身公开证据路径超出给定根目录')
    try:
        raw = path.read_bytes()
    except OSError as exc:
        raise ValueError('自身公开来源或原件不可用') from exc
    if hashlib.sha256(raw).hexdigest() != checksum:
        raise ValueError('自身公开来源或原件SHA不匹配')
    return raw


def _baseline(code, report_root):
    spec = SPECS.get(str(code))
    if spec is None:
        raise ValueError('不是本批相城或吴江自身指南')
    report_root = Path(report_root).resolve()
    manifest = json.loads(_read_checked(report_root / 'SOURCE-MANIFEST.json', MANIFEST_SHA, report_root))
    white = json.loads(_read_checked(report_root / 'PUBLIC-WHITELIST.json', WHITELIST_SHA, report_root))['files']
    for relative, pin in manifest['files'].items():
        if relative.startswith(spec['name'] + '/'):
            if white.get((REPORT / relative).as_posix()) != pin:
                raise ValueError('自身来源文件不在冻结公开清单')
            if len(_read_checked(report_root / relative, pin['sha256'], report_root)) != pin['bytes']:
                raise ValueError('自身来源文件大小不匹配')
    guide_path = f"{spec['name']}/guide.json"
    if manifest['files'][guide_path]['sha256'] != spec['guide_sha']:
        raise ValueError('固定来源清单与指南不一致')
    if white[REPORT.as_posix() + '/' + guide_path] != manifest['files'][guide_path]:
        raise ValueError('自身指南不在冻结公开清单')
    g = json.loads(_read_checked(report_root / guide_path, spec['guide_sha'], report_root))
    expected = dict(region_code=code, district=spec['district'], city='苏州市', province='jiangsu',
                    authority=spec['authority'], service_code=spec['task'],
                    service_name='劳动人事争议仲裁申请', source_title='劳动人事争议仲裁申请')
    if any(g.get(k) != v for k, v in expected.items()):
        raise ValueError('固定指南自身身份不匹配')
    # Independently hash only public captures referenced by this exact guide.
    captures = [(g['evidence']['capture_file'], g['evidence']['capture_sha256']),
                (g['evidence']['search_capture_file'], g['evidence']['search_capture_sha256'])]
    captures.extend((m['detail_capture'], m['detail_capture_sha256']) for m in g['materials'])
    for source_file, checksum in captures:
        relative = Path(source_file).relative_to(REPORT)
        if manifest['files'][relative.as_posix()]['sha256'] != checksum:
            raise ValueError('自身控件或材料详情捕获绑定不匹配')
        if white[REPORT.as_posix() + '/' + relative.as_posix()] != manifest['files'][relative.as_posix()]:
            raise ValueError('自身控件或材料证据不在冻结公开清单')
        _read_checked(report_root / relative, checksum, report_root)
    g['source_id'] = digest(g['source_url'])[:24]
    return g


def parse_capture(code_or_name, capture_root=None):
    """Read frozen own public capture. capture_root is the sealed jiangsu directory."""
    key = str(code_or_name)
    code = next((c for c, s in SPECS.items() if key in {c, s['district'], s['name']}), None)
    return _baseline(code, Path(capture_root) if capture_root is not None else ROOT / REPORT)


def _validate_guide(guide, root=None, allow_missing_evidence=False):
    if not isinstance(guide, dict) or guide.get('generation_adapter') not in {None, ADAPTER}:
        raise ValueError('非本区指南或适配器标记')
    root = Path(root if root is not None else ROOT).resolve()
    if not (root / REPORT).resolve().is_relative_to(root):
        raise ValueError('自身公开来源目录超出给定根目录')
    baseline = _baseline(guide.get('region_code'), root / REPORT)
    lifecycle = {'verification_status', 'warnings', 'can_generate', 'generation_adapter', 'published',
                 'publication_status', 'candidate_web_qa', 'version', 'content_hash', 'region_id', 'templates'}
    if set(guide) - set(baseline) - lifecycle:
        raise ValueError('自身指南包含未知来源字段')
    keys = [k for k in baseline if k not in lifecycle and
            not (allow_missing_evidence and k == 'evidence' and k not in guide)]
    if guide.get('source_id') != baseline['source_id'] or any(guide.get(k) != baseline.get(k) for k in keys):
        raise ValueError('自身地区、指南字段、材料或来源控件绑定不匹配')
    if guide.get('verification_status') not in {'partial', 'verified', 'conflict'}:
        raise ValueError('自身指南尚未核验')
    actual = guide.get('templates', [])
    if len(actual) != len(baseline['templates']) or len({t.get('id') for t in actual}) != len(actual):
        raise ValueError('自身附件缺失或重复')
    for expected in baseline['templates']:
        matches = [t for t in actual if t.get('id') == expected['id']]
        if len(matches) != 1 or any(matches[0].get(k) != expected.get(k) for k in TEMPLATE_KEYS):
            raise ValueError('自身附件原始类别、材料、SHA或下载路径不匹配')
        if not _template_annotations_match(matches[0], expected, guide['region_code']):
            raise ValueError('自身附件事实或精确平台能力标记不匹配')
    return baseline, root


def _source(guide, kind, root=None):
    if kind not in {'application', 'evidence'}:
        raise ValueError('仅申请书与证据清单完成填写适配')
    baseline, root = _validate_guide(guide, root)
    annex = SPECS[guide['region_code']][kind]
    template = next(t for t in baseline['templates'] if t['id'] == annex)
    if template['official_kind'] != 'blank' or template['download_status'] != 'downloaded':
        raise ValueError('只允许自身已下载空白原件')
    path = (root / template['file']).resolve()
    if not path.is_relative_to((root / REPORT).resolve()):
        raise ValueError('自身原件超出公开来源目录')
    original = _read_checked(path, template['sha256'], root)
    return template, path, original


def blank_binding(guide, template, root=None):
    kind = next((k for k in ('application', 'evidence')
                 if SPECS.get(guide.get('region_code'), {}).get(k) == template.get('id')), None)
    if kind is None:
        raise ValueError('样例或占位材料不支持填写')
    expected, _, _ = _source(guide, kind, root)
    if any(template.get(k) != expected.get(k) for k in TEMPLATE_KEYS) or not _template_annotations_match(template, expected, guide['region_code']):
        raise ValueError('传入附件与自身空白绑定不一致')
    return _binding(expected, kind)


def _binding(template, kind):
    return {'adapter': ADAPTER, 'kind': kind, 'annex_id': template['id'],
            'source_sha256': template['sha256'], 'source_file': template['file'],
            'source_url': template['source_url']}


def _template_annotations_match(actual, canonical, code):
    if not isinstance(actual, dict):
        return False
    kind = next((k for k in ('application', 'evidence') if SPECS[code][k] == canonical['id']), None)
    markers = {'can_generate': bool(kind), 'document_types': [kind] if kind else []}
    if kind:
        markers.update(generation_adapter=ADAPTER, strict_binding=_binding(canonical, kind), generation_mode='blank',
                       content_verification_status='own_blank_and_synthetic_output_all_pages_viewed', render_status='PASS',
                       adapter=ADAPTER, content_kind='filled_official_blank', generation_label='依据本区自身官方空白表格填写')
    for key, value in canonical.items():
        if key not in actual or (actual[key] != value and (key not in markers or actual[key] != markers[key])):
            return False
    return all(key in markers and actual[key] == markers[key] for key in set(actual) - set(canonical))


def candidate_contract(guide, root=None):
    documents = {}
    for kind in ('application', 'evidence'):
        template, _, _ = _source(guide, kind, root)
        documents[kind] = blank_binding(guide, template, root)
    return {'adapter': ADAPTER, 'region_code': guide['region_code'], 'province': guide['province'],
            'city': guide['city'], 'district': guide['district'], 'source_id': guide['source_id'],
            'authority': guide['authority'], 'service_code': guide['service_code'],
            'documents': documents, 'scope': 'own_official_blank_only'}


def with_capabilities(guide, root=None):
    contract = candidate_contract(guide, root)
    result = deepcopy(guide)
    bindings = {v['annex_id']: (k, v) for k, v in contract['documents'].items()}
    for template in result['templates']:
        kind, binding = bindings.get(template['id'], (None, None))
        template['can_generate'] = bool(kind)
        template['document_types'] = [kind] if kind else []
        if kind:
            template['generation_adapter'] = ADAPTER
            template['strict_binding'] = binding
            template['generation_mode'] = 'blank'
            template['content_verification_status'] = 'own_blank_and_synthetic_output_all_pages_viewed'
            template['render_status'] = 'PASS'
            template.update(adapter=ADAPTER, content_kind='filled_official_blank', generation_label='依据本区自身官方空白表格填写')
    result['generation_adapter'] = ADAPTER
    result['can_generate'] = True
    original_notices = set(parse_capture(guide['region_code'], Path(root if root is not None else ROOT) / REPORT)['warnings'])
    result['warnings'] = [w for w in guide.get('warnings', []) if w not in original_notices] + [
        '仅支持本区自身空白申请书（劳动者、单被申请人）与六列证据清单。',
        '身份证明及证据复印件附件仅为标题占位，不支持填写生成；官方样例不分发、不作案件事实。',
        '指南申请书4份与原件副本按被申请人人数的语境分别保留；申请书副本与证据副本分别填写。',
        '在线办理已止步官方登录，未验证登录后提交；文书适配不代表已代交仲裁申请。',
        '独立证据清单没有份数槽，材料表证据清单2份仍保留在官方材料要求中。',
        '调解、法律援助、其他证件类型与原件或复印件栏均为自由填写槽，不自动代选或推定。',
    ]
    if result.get('published') and result.get('candidate_web_qa') == 'PASS':
        result['warnings'].append('本区填写版式及网页流程已验收；是否受理仍以本区仲裁机构审查为准。')
    return result


def _public_proof(guide, root):
    if root is None or not isinstance(guide, dict):
        raise ValueError('必须明确给定公开根目录和本区指南')
    canonical, root = _validate_guide(guide, root, allow_missing_evidence=True)
    metadata = {'version', 'content_hash', 'region_id', 'can_generate', 'published', 'publication_status',
                'candidate_web_qa', 'warnings', 'generation_adapter'}
    if set(guide) - set(canonical) - metadata:
        raise ValueError('指南包含非本区来源事实字段')
    if guide.get('generation_adapter') not in {None, ADAPTER}:
        raise ValueError('非本区适配器标记')
    report_root = root / REPORT
    manifest = json.loads(_read_checked(report_root / 'SOURCE-MANIFEST.json', MANIFEST_SHA, root))
    white = json.loads(_read_checked(report_root / 'PUBLIC-WHITELIST.json', WHITELIST_SHA, root))['files']
    prefix = SPECS[guide['region_code']]['name'] + '/'
    for relative, pin in manifest['files'].items():
        if not relative.startswith(prefix):
            continue
        path = report_root / relative
        if white.get(REPORT.as_posix() + '/' + relative) != pin:
            raise ValueError('自身证据不在冻结公开白名单')
        raw = _read_checked(path, pin['sha256'], root)
        if len(raw) != pin['bytes']:
            raise ValueError('自身公开证据大小不匹配')
    return canonical, root, white


def public_evidence(guide, root):
    """Rebuild only canonical public evidence, requiring an explicit root."""
    try:
        canonical, _, _ = _public_proof(guide, root)
        return deepcopy(canonical['evidence'])
    except (ValueError, OSError, KeyError, TypeError, StopIteration, AttributeError):
        return None


def approved_report_attachment(guide, template, root):
    try:
        canonical, root, white = _public_proof(guide, root)
        own = next(t for t in canonical['templates'] if t['id'] == template.get('id'))
        if (own['official_kind'] not in {'blank', 'instruction'} or own['download_status'] != 'downloaded'
                or not _template_annotations_match(template, own, canonical['region_code'])):
            return False
        row = white.get(own['file'])
        if row != {'sha256': own['sha256'], 'bytes': own['bytes']}:
            return False
        return len(_read_checked(root / own['file'], own['sha256'], root)) == own['bytes']
    except (ValueError, OSError, KeyError, TypeError, StopIteration, AttributeError):
        return False


def _text(value, limit=300):
    if value is None:
        return ''
    if not isinstance(value, str) or len(value) > limit:
        raise ValueError('填写内容类型或长度超出本区表格契约')
    if any(not (ch in '\t\n\r' or 0x20 <= ord(ch) <= 0xD7FF or
                0xE000 <= ord(ch) <= 0xFFFD or 0x10000 <= ord(ch) <= 0x10FFFF) for ch in value):
        raise ValueError('填写内容含无效XML字符')
    return value


def _validate_data(kind, data):
    if kind not in {'application', 'evidence'} or not isinstance(data, dict) or set(data) - DATA_FIELDS[kind]:
        raise ValueError('填写数据含本区未支持的字段')
    result = {}
    for key, value in data.items():
        if key in {'copyCount', 'evidenceCopyCount'}:
            if value is None or value == '':
                result[key] = None
            elif isinstance(value, bool) or not isinstance(value, (int, str)) or not re.fullmatch(r'[1-9]\d*', str(value)) or not 1 <= int(value) <= 30:
                raise ValueError('副本份数须留空或为1至30的整数')
            else:
                result[key] = int(value)
        elif key == 'requests':
            if not isinstance(value, list) or len(value) > LIMITS['requests']:
                raise ValueError('本区填写最多支持4项请求，不静默截断')
            result[key] = [_text(item, LIMITS['request']) for item in value]
        elif key == 'evidenceList':
            if not isinstance(value, list) or len(value) > LIMITS['evidenceList']:
                raise ValueError('本区填写最多支持40项证据，不静默截断')
            result[key] = []
            for item in value:
                if not isinstance(item, dict) or set(item) - EVIDENCE_ITEM_FIELDS:
                    raise ValueError('证据项包含本区未支持的字段')
                result[key].append({k: _text(v, LIMITS['evidence_' + k]) for k, v in item.items()})
        else:
            result[key] = _text(value, LIMITS['factsReason'] if key == 'factsReason' else LIMITS['ordinary'])
            if key in {'birthDate', 'applicationDate', 'submitDate'} and result[key]:
                try:
                    if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', result[key]):
                        raise ValueError
                    date.fromisoformat(result[key])
                except ValueError as exc:
                    raise ValueError('日期须留空或为有效的YYYY-MM-DD日期') from exc
    return result


# Audited against each E district's own 99df... worker original, paragraphs 1-12.
WORKER_SLOTS = {
    1: [('申请人', 'applicantName'), ('性别', 'gender'), ('出生日期', 'birthDate')],
    2: [('公民身份号码', 'idNumber'), ('居住地址', 'idAddress')],
    3: [('其他证件类型', 'otherDocumentType'), ('证件号码', 'otherDocumentNumber')],
    4: [('通讯地址', 'applicantMailingAddress'), ('移动电话', 'phone')],
    5: [('是否接受调解', 'acceptMediation'), ('是否法律援助', 'legalAid')],
    6: [('被申请人', 'respondentName'), ('性质', 'respondentType')],
    7: [('注册地址（住所）', 'registeredAddress')], 8: [('统一社会信用代码', 'socialCreditCode')],
    9: [('法定代表人或者主要负责人姓名', 'legalRepName'), ('职务', 'legalRepPosition')],
    10: [('劳动者所在部门负责人', 'departmentHead'), ('职务', 'departmentHeadPosition')],
    11: [('通讯地址', 'mailingAddress')], 12: [('联系人', 'companyContact'), ('移动电话', 'companyContactPhone')],
}


def _application(document, data):
    children = list(document._element.body)
    split = next(i for i, c in enumerate(children) if '仲裁申请书（用人单位申请）' in ''.join(c.itertext()))
    if split >= 2 and ''.join(children[split - 2].itertext()).strip().startswith('事项名称'):
        split -= 2
    keep_body(document, 0, split)
    paragraphs = list(document.paragraphs)
    if len(paragraphs) < 50 or paragraphs[0].text != '劳动人事争议仲裁申请书（劳动者申请）':
        raise ValueError('本区劳动者申请原件结构不符')
    for index, slots in WORKER_SLOTS.items():
        for label, _ in slots:
            if label not in paragraphs[index].text:
                raise ValueError('本区劳动者申请填写标签不符')
        fill_label_slots(paragraphs[index], [(label, data.get(key, '')) for label, key in slots])
    # This single-respondent contract removes only the optional second party.
    for paragraph in paragraphs[13:20]:
        paragraph._p.getparent().remove(paragraph._p)
    fill_sections(document, data)
    fill_footer(document, {**data, 'copyCount': data.get('copyCount') or '    '})
    closing = False
    for paragraph in document.paragraphs:
        match = re.search(r'相关证据(?P<slot>[\d \t]+)(?=份)', paragraph.text)
        if match:
            replace_span(paragraph, match.start('slot'), match.end('slot'), str(data.get('evidenceCopyCount') or '    '))
        compact = ''.join(paragraph.text.split())
        if compact == '此致':
            closing = True
        if closing:
            if compact.endswith('日') and '年' in compact and '月' in compact:
                closing = False
            else:
                paragraph.paragraph_format.keep_with_next = True
    remove_empty_edges(document)


def _evidence(document, data):
    if len(document.tables) != 1 or len(document.tables[0].columns) != 6 or len(document.paragraphs) != 9:
        raise ValueError('本区六列证据清单原件结构不符')
    paragraphs = document.paragraphs
    if data.get('committeeName'):
        replace_span(paragraphs[0], 0, len(paragraphs[0].text), data['committeeName'])
    if data.get('caseNumber'):
        replace_span(paragraphs[3], 0, len(paragraphs[3].text), data['caseNumber'])
    if data.get('submitDate'):
        value = date.fromisoformat(data['submitDate'])
        fill_label_slots(paragraphs[5], [('提交时间', f'{value.year}年{value.month}月{value.day}日')], pad=False)
    table = document.tables[0]
    row_template = deepcopy(table.rows[1]._tr)
    for row in list(table.rows)[1:]:
        table._tbl.remove(row._tr)
    from docx.oxml import OxmlElement
    properties = table.rows[0]._tr.get_or_add_trPr()
    if properties.find(qn('w:tblHeader')) is None:
        properties.append(OxmlElement('w:tblHeader'))
    items = data.get('evidenceList', []) or [{}]
    for index, item in enumerate(items, 1):
        node = deepcopy(row_template); table._tbl.append(node); row = _Row(node, table)
        for height in row._tr.iter(qn('w:trHeight')):
            height.set(qn('w:hRule'), 'atLeast')
        values = [str(index) if item else '', item.get('name', ''), item.get('source', ''), item.get('purpose', ''),
                  item.get('originalType', ''), item.get('pageCount', '')]
        for cell, value in zip(row.cells, values):
            set_cell_text(cell, value)
    remove_empty_edges(document)


def fill_candidate_document(guide, kind, data, root=None):
    """Patch the own blank package; no cross-region rule, identity or sample use."""
    _, _, original = _source(guide, kind, root)
    data = _validate_data(kind, data)
    document = Document(BytesIO(original))
    (_application if kind == 'application' else _evidence)(document, data)
    return PreservedPackageDocument(document, original)
