"""Concise Word deliverables; machine diagnostics stay in raw JSON files."""
import json
import re
from pathlib import Path
from urllib.parse import urlsplit
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION

MAIN='劳动仲裁政务服务门户观察报告.docx'
AI='劳动仲裁政务服务门户AI验证报告.docx'

def font(style,name,size):
    style.font.name=name;style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
    rpr=style._element.get_or_add_rPr();fonts=rpr.get_or_add_rFonts()
    for attr in list(fonts.attrib):
        if 'Theme' in attr:del fonts.attrib[attr]
    for attr in ('ascii','hAnsi','eastAsia','cs'):fonts.set(qn('w:'+attr),name)

def document(title,folder):
    d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=s.bottom_margin=Cm(1.6);s.left_margin=s.right_margin=Cm(2)
    for name in ('Normal','Caption','Hyperlink','Header','Footer'):
        if name in d.styles:font(d.styles[name],'宋体',10)
    for name,size in [('Title',21),('Heading 1',15),('Heading 2',12)]:
        font(d.styles[name],'黑体',size);d.styles[name].font.bold=False
        d.styles[name].paragraph_format.keep_with_next=True
    for border in list(d.styles.element.xpath('.//w:pBdr')):border.getparent().remove(border)
    for border in list(d.element.xpath('.//w:pBdr')):border.getparent().remove(border)
    normal=d.styles['Normal'].paragraph_format;normal.space_after=Pt(4);normal.line_spacing=1.1
    for name in ('Heading 1','Heading 2'):
        d.styles[name].paragraph_format.space_before=Pt(10);d.styles[name].paragraph_format.space_after=Pt(6)
    d.styles['Caption'].paragraph_format.space_after=Pt(8)
    d.add_heading(title,0);d.add_paragraph('报告编号：'+folder.name)
    f=s.footer.paragraphs[0];f.text=folder.name+'  |  '
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
    return d

def link(p,label,url):
    if not url:return
    h=OxmlElement('w:hyperlink')
    if url.startswith('#'):h.set(qn('w:anchor'),url[1:])
    else:
        rel=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
        h.set(qn('r:id'),rel)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    fonts=OxmlElement('w:rFonts')
    for attr in ('ascii','hAnsi','eastAsia'):fonts.set(qn('w:'+attr),'宋体')
    pr.append(fonts)
    color=OxmlElement('w:color');color.set(qn('w:val'),'2A5E97');pr.append(color)
    underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');pr.append(underline)
    r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def keep_evidence_block(d,start):
    # Keep the heading, conclusion, metadata and picture together. Captions end
    # the block so the next search cannot leave an orphan heading below it.
    block=d.paragraphs[start:]
    for p in block:p.paragraph_format.keep_with_next=True
    if block:block[-1].paragraph_format.keep_with_next=False

def url_line(d,label,url):
    shown=url or '未取得页面网址'
    if url:
        parsed=urlsplit(url);shown=parsed.netloc+'（打开完整网址）'
    p=d.add_paragraph(label+'：');link(p,shown,url)
    p.paragraph_format.keep_with_next=True

def material_label(item):
    # A material row may contain several filenames. Show the clicked leaf
    # control's name, while the full row and URL remain in the evidence JSON.
    control=re.sub(r'\s+',' ',item.get('control_text','')).strip()
    generic={'下载','下载查看','下载至本地','样例下载','示例下载','空白下载','查看附件','查看','示例','样例','[格式文本]','[示范文本]'}
    label=control if control and control not in generic else item.get('text','')
    if label.strip() in generic:
        label=item.get('material_name') or '文书附件'
        suffix={'blank':'（空白表格）','sample':'（填写样例）'}.get(item.get('kind'),'')
        label+=suffix
    label=re.sub(r'\s+',' ',label).strip() or '文书附件'
    return label if len(label)<=32 else label[:31]+'…'

def table(d,headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
    borders=OxmlElement('w:tblBorders')
    for name in ('top','left','bottom','right','insideH','insideV'):
        border=OxmlElement('w:'+name);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'D9D9D9');borders.append(border)
    t._tbl.tblPr.append(borders)
    for c,w,text in zip(t.rows[0].cells,widths,headers):c.width=Cm(w);c.text=text
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for c,w,value in zip(cells,widths,row):
            c.width=Cm(w)
            if isinstance(value,tuple):link(c.paragraphs[0],value[0],value[1])
            else:c.text=str(value)
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    for index,row in enumerate(t.rows):
        keep=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(keep)
        for c in row.cells:
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins=OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                m=OxmlElement('w:'+edge);m.set(qn('w:w'),'80');m.set(qn('w:type'),'dxa');margins.append(m)
            c._tc.get_or_add_tcPr().append(margins)
            if index==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EAF0F4');c._tc.get_or_add_tcPr().append(shade)
            for p in c.paragraphs:p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
    return t

def evidence(d,folder,x,obs,caption):
    url_line(d,'页面网址',obs.get('final_url') or obs.get('requested_url'))
    p=d.add_paragraph('采集时间：'+obs.get('finished_at','未记录')[:19].replace('T',' '));p.paragraph_format.keep_with_next=True
    if obs.get('evidence_error'):
        d.add_paragraph('本次截图校验失败，暂不能作为页面证据；待补拍。');return
    image=folder/x['slug']/obs.get('screenshot','absent.png')
    if image.is_file() and not obs.get('screenshot_blank'):
        from PIL import Image
        with Image.open(image) as im:width,height=im.size
        scale=min(17/width,10.6/height)
        d.add_picture(str(image),width=Cm(width*scale),height=Cm(height*scale));d.paragraphs[-1].paragraph_format.keep_with_next=True
        d.add_paragraph(caption,style='Caption')
    elif image.is_file():d.add_paragraph('已保存截图，但画面为空白；不能据此判断入口有无。原图和错误记录已保留。')
    else:d.add_paragraph('未取得截图文件；不能据此判断入口有无。错误记录已保留。')

def current_searches(item):
    observations={o['stage']:o for o in item.get('observations',[])}
    current=[]
    for search in item.get('searches',[]):
        search=observations.get(search['stage'],search)
        seen=set()
        while search.get('superseded_by') in observations and search['stage'] not in seen:
            seen.add(search['stage']);search=observations[search['superseded_by']]
        current.append(search)
    return current

def build_report(folder,items):
    # Older saved records may have mistaken a security interception page for
    # a result page. Reclassify from the preserved evidence, without recrawling.
    from portal_audit_capture import search_outcome
    manual_file=folder/'manual-reviews.json'
    manual=json.loads(manual_file.read_text(encoding='utf8')) if manual_file.is_file() else {}
    for x in items:
        x['searches']=current_searches(x)
        for obs in x['searches']:
            checked=manual.get(x['slug'],{}).get(obs['stage'],{})
            if 'entrance_visible' in checked and checked.get('screenshot_sha256')==obs.get('screenshot_sha256'):
                obs['visual_entrance_visible']=checked['entrance_visible']
                obs['visual_review_note']=checked['conclusion']
                if checked['entrance_visible']:obs['outcome']='截图中可见相关办事入口，链接未自动提取'
        changed=False
        for obs in x.get('searches',[]):
            path=folder/x['slug']/obs.get('text_file','absent.txt')
            text=path.read_text(encoding='utf8') if path.is_file() else ''
            if re.search('403 Forbidden|405.{0,8}Method Not Allowed|429 Too Many|安全防护.{0,8}拦截',text[:2000],re.I):
                obs.update(page_problem='blocked',result_page_confirmed=False,candidate_links=[])
                obs['outcome']=search_outcome(obs);changed=True
                matching=next((o for o in x['observations'] if o['stage']==obs['stage']),None)
                if matching is not None:matching.update(obs)
        # Recompute the summary after continuation; a historical conclusion
        # cannot override newly reached guide and material evidence.
        portal_materials=[o for o in x.get('observations',[]) if o['stage'].endswith('-templates') and not o.get('superseded_by') and not o['stage'].startswith('official-source') and not o.get('page_problem') and not o.get('non_target_service') and o.get('template_links')]
        if portal_materials:x.update(result_code='template_candidate_visible',result='已进入指南，检测到文书下载控件')
        elif any((s.get('candidate_links') or s.get('visual_entrance_visible')) and not s.get('page_problem') for s in x['searches']):x.update(result_code='service_candidate_visible',result='检索画面中可见相关入口候选')
        elif all(s.get('result_page_confirmed') and not s.get('page_problem') for s in x['searches']):x.update(result_code='not_located_in_limited_search',result='三个关键词的本次结果均未检测到入口')
        else:x.update(result_code='access_failed',result='部分检索未完成，不能判断入口有无')
        (folder/x['slug']/'observations.json').write_text(json.dumps(x,ensure_ascii=False,indent=2),encoding='utf8')
    (folder/'results.json').write_text(json.dumps(items,ensure_ascii=False,indent=2),encoding='utf8')
    # Do not embed a solid white image as evidence of an absent service.
    from PIL import Image
    quality=[]
    for x in items:
        for obs in x.get('observations',[]):
            p=folder/x['slug']/obs.get('screenshot','absent.png')
            if not p.is_file():continue
            with Image.open(p) as im:
                gray=im.convert('L');gray.thumbnail((160,100));hist=gray.histogram()
                obs['screenshot_blank']=sum(hist[250:])/sum(hist)>0.999
                quality.append(dict(region=x['slug'],stage=obs['stage'],blank=obs['screenshot_blank']))
        for s in x.get('searches',[]):
            corresponding=next((o for o in x['observations'] if o['stage']==s['stage']),{})
            s['screenshot_blank']=corresponding.get('screenshot_blank',False)
    (folder/'image-quality.json').write_text(json.dumps(quality,ensure_ascii=False,indent=2),encoding='utf8')
    d=document('劳动仲裁政务服务门户观察报告',folder)
    d.add_paragraph('检索词：劳动仲裁、劳动人事争议仲裁、劳动人事争议调解。')
    d.add_paragraph('以下结果仅对应本次公开页面和截图；检索未完成与搜索结果没有入口分别记录。')
    d.add_paragraph('“已有可填充文书”来自项目已下载模板，独立于本次门户检索结果。')
    d.add_heading('全国结果概况',1)
    rows=[]
    for x in items:
        available='、'.join('申请书' if k=='application' else '证据清单' for k in x.get('documents',{})) or '暂无已验证Word模板'
        rows.append([(x['name'],'#province_'+x['slug']),(urlsplit(x['portal_url']).netloc,x['portal_url']),x.get('result','待复核'),available])
    table(d,['地区','官方门户网址','本次检索结果','已有可填充文书'],rows,[2.3,4.8,5.6,4.3])
    for index,x in enumerate(items,1):
        section=d.add_section(WD_SECTION.NEW_PAGE)
        section.header.is_linked_to_previous=False
        section.header.paragraphs[0].text='劳动仲裁门户观察  /  '+x['name']
        heading=d.add_heading(f'{index:02d}  {x["name"]}',1)
        bookmark=OxmlElement('w:bookmarkStart');bookmark.set(qn('w:id'),str(index));bookmark.set(qn('w:name'),'province_'+x['slug']);heading._p.insert(0,bookmark)
        end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(index));heading._p.append(end)
        url_line(d,'省级门户',x['portal_url'])
        if x.get('service_platform_url'):url_line(d,'该门户指向的平台',x['service_platform_url'])
        d.add_paragraph('结论：'+x.get('result','待复核'))
        if x.get('access_note'):d.add_paragraph(x['access_note'])
        coverage=x.get('workflow_coverage',{})
        if coverage:
            d.add_paragraph(f"指南候选网址：{coverage.get('discovered_candidates',coverage.get('discovered_guides',0))} 个；已保存指南 {coverage.get('captured_guides',0)} 份，其中 {coverage.get('workflow_guides',0)} 份取得流程文字或步骤。")
            p=d.add_paragraph('地区流程数据：');link(p,'打开本省指南与覆盖记录',(folder/x['slug']/'workflows.json').resolve().as_uri())
            if coverage.get('pending_guides'):d.add_paragraph('仍有地区或分页未完成；详情见流程数据中的待续查记录。')
        for number,obs in enumerate(x.get('searches',[]),1):
            block_start=len(d.paragraphs)
            d.add_heading(f'搜索“{obs["keyword"]}”',2)
            d.add_paragraph(obs.get('outcome','待复核')+'。'+('已填写并按回车。' if obs.get('enter_submitted') else '未成功提交搜索。'))
            if obs.get('visual_review_note'):d.add_paragraph('截图复核：'+obs['visual_review_note'])
            problem=obs.get('page_problem')
            if problem:
                text={'certificate':'证书验证失败，不能认定门户不存在服务。','blank':'页面未完成有效渲染，不能作为无入口证据。','network':'本次连接异常，未完成搜索。','blocked':'页面显示访问限制或验证要求。','render_error':'搜索页返回非网页内容或源代码，未形成可用结果页。'}.get(problem,problem)
                d.add_paragraph(text)
            if obs.get('search_scope') and obs['search_scope']!='页面默认检索范围':d.add_paragraph('页面检索范围：'+obs['search_scope'])
            if (obs.get('http_status') or 200)>=400:d.add_paragraph('服务器返回 HTTP '+str(obs['http_status'])+'。')
            if obs.get('access_detail'):d.add_paragraph(obs['access_detail']+'。')
            evidence(d,folder,x,obs,'搜索后的实际页面' if obs.get('enter_submitted') else '无法完成检索时的实际页面')
            keep_evidence_block(d,block_start)
        # Show each reached material page once; raw guide screenshots remain in the run.
        seen=set()
        candidates=[o for o in x.get('observations',[]) if o['stage'].endswith('-templates') and not o.get('superseded_by') and (not o['stage'].startswith('district-') or o.get('report_include'))]
        candidates.sort(key=lambda o:(not bool(o.get('template_links')),o.get('material_capture_version',0)<5))
        for obs in candidates:
            if not obs['stage'].endswith('-templates') or obs.get('superseded_by'):continue
            key=obs.get('final_url')
            if key in seen:continue
            seen.add(key)
            block_start=len(d.paragraphs)
            source=obs['stage'].startswith('official-source')
            pending=not bool(obs.get('template_links'))
            if obs.get('non_target_service'):
                d.add_paragraph('该入口是其他事项：'+obs['non_target_service']+'。本项不计为劳动仲裁模板入口。');continue
            scope=' · '.join(s for s in obs.get('selected_jurisdictions',[]) if s)
            if not scope:scope=obs.get('service_name') or ('来源说明' if source else '办理指南 '+str(len(seen)))
            d.add_heading(('补充官方来源：' if source else '办理指南：')+scope[:80],2)
            if obs.get('page_problem'):d.add_paragraph('该入口本次访问异常，尚未完成模板检查。')
            elif pending:d.add_paragraph('本次没有检测到文书模板。')
            else:
                kinds={a.get('kind') for a in obs.get('template_links',[])}
                d.add_paragraph(('记录有文书下载控件，截图尚未显示文书区域。' if obs.get('template_visual_verified') is False else '检测到文书下载控件。')+('包含样表或示例；是否为空白模板以实际附件为准。' if 'sample' in kinds else ''))
            if obs.get('screenshot_scope_note'):d.add_paragraph(obs['screenshot_scope_note'])
            if source:d.add_paragraph('此来源独立于本次门户搜索结果；适用范围：'+x['scope'])
            if obs.get('selected_jurisdictions'):d.add_paragraph('本次选择区域：'+' → '.join(obs['selected_jurisdictions']))
            if not pending and not obs.get('template_links'):d.add_paragraph('可见材料说明，尚未确认模板下载链接。')
            names=list(dict.fromkeys(material_label(l) for l in obs.get('template_links',[])))
            if not names:names=list(dict.fromkeys(obs.get('material_names',[])))
            d.add_paragraph('页面材料名称：'+('、'.join(names)[:160] or '本次画面未确认申请书或证据清单'))
            evidence(d,folder,x,obs,'指南实际页面' if pending else '申请材料及文书下载区域')
            keep_evidence_block(d,block_start)
            for attachment in obs.get('template_links',[])[:4]:
                attachment_url=attachment.get('download_url') or attachment.get('url','')
                if attachment.get('file'):
                    p=d.add_paragraph('已取得附件：');link(p,material_label(attachment),(folder/x['slug']/attachment['file']).resolve().as_uri())
                elif attachment_url.startswith(('http://','https://')):
                    p=d.add_paragraph('附件：');link(p,material_label(attachment),attachment_url)
    output=folder/MAIN;d.save(output)
    if any(x.get('ai_review') for x in items):build_ai_report(folder,items)
    return output

def build_ai_report(folder,items):
    # Resolve provenance from the saved provider response, rather than from a
    # legacy region summary that may have guessed localhost for older records.
    provenance_changed=False
    for x in items:
        sources={}
        for path in (folder/x['slug']).glob('openclaw-*.json'):
            try:raw=json.loads(path.read_text(encoding='utf8'))
            except (OSError,ValueError):continue
            actual=raw.get('result',{})
            if raw.get('timed_out') or actual.get('ok') is not True or actual.get('provider')!='ollama' or actual.get('model')!='qwen3.8:27b':continue
            for entry in raw.get('verified_json',[]):
                digest=entry.get('screenshot_sha256')
                if entry.get('state')=='complete' and digest in raw.get('input_sha256',[]):
                    sources.setdefault(digest,[]).append((path.name,raw.get('base_url') or '历史记录未注明服务器'))
        for v in x.get('ai_review',{}).get('images',[]):
            matches=sources.get(v.get('screenshot_sha256'),[])
            exact=[server for name,server in matches if name==v.get('source_record')]
            unique={server for _,server in matches}
            server=exact[0] if exact else next(iter(unique)) if len(unique)==1 else '历史记录未注明服务器' if len(unique)>1 else None
            if server and v.get('model_server')!=server:
                v['model_server']=server;provenance_changed=True
    if provenance_changed:
        for x in items:
            region=folder/x['slug']/'observations.json'
            if region.is_file():
                temp=region.with_suffix('.tmp');temp.write_text(json.dumps(x,ensure_ascii=False,indent=2),encoding='utf8');temp.replace(region)
        manifest=folder/'results.json';temp=manifest.with_suffix('.tmp')
        temp.write_text(json.dumps(items,ensure_ascii=False,indent=2),encoding='utf8');temp.replace(manifest)
    d=document('劳动仲裁政务服务门户AI验证报告',folder)
    manual_path=folder/'manual-reviews.json'
    manual=json.loads(manual_path.read_text(encoding='utf8')) if manual_path.is_file() else {}
    d.add_paragraph('验证工具：OpenClaw；模型：Ollama / qwen3.8:27b。本次新增复核使用 Ubuntu 服务器 http://10.66.66.1:11434；保留此前已成功的历史复核记录。')
    prior=sum(v.get('state')=='complete' and v.get('model_server','历史记录未注明服务器')=='历史记录未注明服务器' for x in items for v in x.get('ai_review',{}).get('images',[]))
    if prior:d.add_paragraph(f'保留中断前已完成的 {prior} 条历史复核记录；原始响应未记载服务器地址，不作来源推断，也未重复调用。')
    info=folder/'model-info.json'
    if info.exists():
        p=d.add_paragraph('模型版本：');link(p,'Ubuntu 模型元数据与版本校验记录',info.resolve().as_uri())
    done=sum(x.get('ai_review',{}).get('state')=='complete' for x in items)
    d.add_paragraph(f'截图复核完整地区：{done}/{len(items)}。AI仅核对截图可见内容；不能替代实际点击或下载验证。')
    d.add_paragraph('截图项目可点击查看原图；对应截图也收录在同批次门户观察报告中。')
    for index,x in enumerate(items):
        section=d.add_section(WD_SECTION.NEW_PAGE);section.header.is_linked_to_previous=False
        section.header.paragraphs[0].text='OpenClaw 图像验证  /  '+x['name']
        d.add_heading(f'{index+1:02d}  '+x['name'],1)
        review=x.get('ai_review',{});d.add_paragraph('状态：'+('全部目标截图已复核' if review.get('state')=='complete' else '复核未完成'))
        rows=[]
        for v in review.get('images',[]):
            label=('检索：' if v['stage'].startswith('search-') else '入口材料：')+(v.get('keyword') or '') if v.get('keyword') else ('已有来源模板区域' if v['stage'].startswith('official-source') else '入口模板区域')
            observed=next((o for o in x['observations'] if o['stage']==v['stage']),{})
            picture=folder/x['slug']/observed.get('screenshot','absent.png')
            if picture.is_file():label=(label,picture.resolve().as_uri())
            if v.get('state')!='complete':rows.append([label,'调用未完成','不可判断',v.get('error','')[:100]]);continue
            names='、'.join(v.get('template_names',[]))
            result=('可见相关入口' if v.get('entrance_visible') else '截图中未检测到相关入口')
            if names:result='可见文书控件：'+names
            elif v['stage'].endswith('-templates'):result='本次截图没有检测到文书模板'
            warning='；'.join(v.get('consistency_warnings',[])) or v.get('warning','')
            checked=manual.get(x['slug'],{}).get(v['stage'])
            if checked:
                warning='人工对照：'+checked['conclusion']
                if not checked.get('accept_ai_result',True):
                    result='原AI判读未采纳；'+checked['conclusion']
            if not names and observed.get('template_links') and 'AI未从截图确认' not in warning:warning+='；记录有下载控件，但AI未从截图确认；待对照原图复核'
            if v.get('reused_from_stage'):warning+='；相同截图复用已完成的AI判读'
            server=v.get('model_server','历史记录未注明服务器')
            origin='Ubuntu 复核\n'+server.removeprefix('http://').removesuffix(':11434') if server in ('http://10.66.66.1:11434','http://192.168.110.89:11434') else ('前序本机复核' if server=='http://127.0.0.1:11434' else '历史复核\n服务器未记载')
            rows.append([label,v.get('page_kind','未知')+'\n'+origin,result,warning[:160]])
        table(d,['截图项目','AI页面识别','可见结果','复核提示'],rows,[3.7,2.5,5.2,5.6])
        records=list((folder/x['slug']).glob('openclaw-*.json'))
        if records:
            p=d.add_paragraph('AI原始记录：');link(p,'本地区证据目录（含模型原始输出）',(folder/x['slug']).as_uri())
        if manual.get(x['slug']):
            p=d.add_paragraph('人工复核记录：');link(p,'结论与第二模型复核依据',manual_path.resolve().as_uri())
    output=folder/AI;d.save(output);return output
