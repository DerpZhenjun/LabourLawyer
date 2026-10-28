"""Fill the official Zhejiang worker arbitration application template."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from WenShu.template_fill import date_parts, replace_text, set_cell_text


def _format_date(value: str) -> str:
    try:
        parsed = date.fromisoformat(value)
        return f"{parsed.year}年{parsed.month}月{parsed.day}日"
    except (TypeError, ValueError):
        return str(value or "")


def _keep_worker_form_only(doc):
    """The official file contains worker and employer forms; keep the worker form only."""
    body = doc._element.body
    first_table_seen = False
    for child in list(body):
        if child.tag == qn("w:tbl") and not first_table_seen:
            first_table_seen = True
            continue
        if first_table_seen and child.tag != qn("w:sectPr"):
            body.remove(child)


def generate_LAdocx(data: dict, template_docx: str | Path):
    doc = Document(str(template_docx))
    _keep_worker_form_only(doc)
    table = doc.tables[0]

    requests = [str(item).strip() for item in data.get("requests", []) if str(item).strip()]
    content_fields = (
        "committeeName", "applicantName", "birthDate", "gender", "nationality", "idNumber",
        "phone", "householdType", "idAddress", "respondentName", "socialCreditCode",
        "registeredAddress", "mailingAddress", "workAddress", "legalRepName", "legalRepPhone",
        "legalRepPosition", "companyContact", "companyContactPhone", "factsReason", "applicationDate",
    )
    if not requests and not any(str(data.get(key, "")).strip() for key in content_fields):
        return doc

    set_cell_text(table.rows[1].cells[1], data.get("applicantName", ""))
    set_cell_text(table.rows[1].cells[4], _format_date(data.get("birthDate", "")))
    set_cell_text(table.rows[2].cells[1], data.get("gender", ""))
    set_cell_text(table.rows[2].cells[4], data.get("nationality", ""))
    set_cell_text(table.rows[3].cells[1], data.get("idNumber", ""))
    set_cell_text(table.rows[3].cells[4], data.get("phone", ""))
    if data.get("householdType"):
        options = table.rows[4].cells[1].paragraphs
        selected = data["householdType"]
        for paragraph in options:
            if selected in paragraph.text:
                replace_text(paragraph, selected, "√" + selected)
    set_cell_text(table.rows[5].cells[2], data.get("idAddress", ""))

    set_cell_text(table.rows[7].cells[2], data.get("respondentName", ""))
    set_cell_text(table.rows[8].cells[2], data.get("socialCreditCode", ""))
    set_cell_text(table.rows[9].cells[2], data.get("registeredAddress", ""))
    set_cell_text(table.rows[10].cells[2], data.get("mailingAddress", ""))
    set_cell_text(table.rows[11].cells[2], data.get("workAddress", ""))
    set_cell_text(table.rows[12].cells[2], data.get("legalRepName", ""))
    set_cell_text(table.rows[12].cells[4], data.get("legalRepPhone", ""))
    set_cell_text(table.rows[13].cells[2], data.get("legalRepName", ""))
    set_cell_text(table.rows[13].cells[4], data.get("legalRepPosition", ""))
    set_cell_text(table.rows[14].cells[2], data.get("companyContact", ""))
    set_cell_text(table.rows[14].cells[4], data.get("companyContactPhone", ""))

    for offset, row_index in enumerate((16, 17, 18, 19)):
        if offset < len(requests):
            paragraph = table.rows[row_index].cells[0].paragraphs[0]
            if row_index == 19:
                replace_text(paragraph, "……", f"4、{requests[offset]}")
            else:
                paragraph.runs[0].text += requests[offset]

    set_cell_text(table.rows[21].cells[0], data.get("factsReason", ""))
    # The paper form reserves several blank handwriting lines. Typed content
    # expands naturally from the original single-line minimum instead.
    table.rows[21].height=table.rows[16].height

    footer = table.rows[22].cells[0].paragraphs
    committee = str(data.get("committeeName", ""))
    suffix = "劳动人事争议仲裁委员会"
    if committee:
        if committee.endswith(suffix):
            footer[2].runs[1].text = committee[:-len(suffix)]
        else:
            footer[2].runs[1].text = committee
            footer[2].runs[2].text = ""
    # Leave the official signature slot for handwriting.
    year, month, day = date_parts(data.get("applicationDate", ""))
    if year and month and day:
        replace_text(footer[4], "年   月    日", f"{year}年   {month}月    {day}日")
    footer[3].paragraph_format.keep_with_next=True
    tr_pr=table.rows[22]._tr.get_or_add_trPr()
    if tr_pr.find(qn('w:cantSplit')) is None:tr_pr.append(OxmlElement('w:cantSplit'))
    replace_text(footer[5], "      份", f"  {data.get('copyCount', 1)}   份")
    if len(requests)>4:
        raise ValueError("浙江官方模板最多容纳4项请求，请合并相关请求或附续页")
    for index in range(19,15+len(requests),-1):
        table._tbl.remove(table.rows[index]._tr)
    return doc
