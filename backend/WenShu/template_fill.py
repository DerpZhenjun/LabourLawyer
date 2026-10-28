"""Small edits to official Word templates without rebuilding their formatting."""

from __future__ import annotations

from datetime import date
from copy import deepcopy
import re
import unicodedata

from docx.text.run import Run
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def replace_span(paragraph, start: int, end: int, value: str, *, pad: bool = False, style_run=None) -> None:
    """Replace only a template slot; retain label runs and the slot's own rPr.

    A paragraph is not one style: labels may be bold while blank runs carry
    underline/font information. Never copy a label's style to all input text.
    """
    runs = paragraph.runs
    full = ''.join(r.text for r in runs)
    if not 0 <= start <= end <= len(full):
        raise ValueError('模板字段范围无效')
    if start == end:
        # Empty slots at the end of a labelled paragraph still inherit its font.
        if runs:
            node = deepcopy(runs[-1]._r)
            for child in list(node):
                if child.tag != qn('w:rPr'):
                    node.remove(child)
            paragraph._p.append(node)
            Run(node, paragraph).text = value
        else:
            paragraph.add_run(value)
        return
    offsets, position = [], 0
    for run in runs:
        offsets.append((position, position + len(run.text), run))
        position += len(run.text)
    affected = [(a, b, r) for a, b, r in offsets if a < end and b > start]
    if not affected:
        raise ValueError('模板字段不包含可编辑文本')
    first_a, _, first = affected[0]
    last_a, _, last = affected[-1]
    slot = full[start:end]
    styled = next((r for a, b, r in affected if r.underline), affected[0][2])
    # Prefer a dedicated blank run over the preceding label's bold run.
    if not any(r.underline for _, _, r in affected):
        styled = next((r for a, b, r in affected if not r.text.strip(' _＿\t\n\u3000')), styled)
    if style_run is not None:
        styled=style_run
    def make_run(source, text):
        node = deepcopy(source._r)
        for child in list(node):
            if child.tag != qn('w:rPr'):
                node.remove(child)
        Run(node, paragraph).text = text
        return node
    value = str(value or '')
    if pad and value:
        width = lambda s: sum(2 if unicodedata.east_asian_width(c) in 'WF' else 1 for c in s)
        value += ' ' * max(0, min(40, width(slot)) - width(value))
    inserted = make_run(styled, value)
    if ('_' in slot or '＿' in slot) and not Run(inserted, paragraph).underline:
        Run(inserted, paragraph).underline = True
    prefix = first.text[:start - first_a]
    suffix = last.text[end - last_a:]
    if prefix:
        first._r.addprevious(make_run(first, prefix))
    first._r.addprevious(inserted)
    if suffix:
        first._r.addprevious(make_run(last, suffix))
    for _, _, run in affected:
        run._r.getparent().remove(run._r)


def fill_label_slots(paragraph, fields: list[tuple[str, str]], *, pad: bool = True) -> None:
    """Fill exact labels in one paragraph, applying changes from right to left."""
    text = paragraph.text
    matches = []
    for label, value in fields:
        # Official conversions frequently insert spaces inside Chinese labels.
        pattern = r'\s*'.join(re.escape(char) for char in label.replace(' ', '')) + r'(?:[ \t]*[:：])?'
        match = re.search(pattern, text)
        if match:
            matches.append((match.start(), match.end(), value))
    matches.sort()
    for index in range(len(matches) - 1, -1, -1):
        _, start, value = matches[index]
        end = matches[index + 1][0] if index + 1 < len(matches) else len(text)
        # Retain a small separator before the following label, not a second
        # copy of the template's province/city/county address placeholders.
        spacer = '  ' if index + 1 < len(matches) else ''
        if str(value or '').strip():
            replace_span(paragraph, start, end, str(value) + spacer, pad=pad)


def all_paragraphs(doc):
    yield from doc.paragraphs
    seen = set()
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                yield from cell.paragraphs


def set_paragraph_text(paragraph, value) -> None:
    """Fill a blank field using its existing first run and paragraph properties."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(str(value or ""))
        return
    runs[0].text = str(value or "")
    for run in runs[1:]:
        run.text = ""


def set_cell_text(cell, value) -> None:
    """Fill a blank table cell, keeping its paragraph and cell layout."""
    set_paragraph_text(cell.paragraphs[0], value)
    for paragraph in cell.paragraphs[1:]:
        set_paragraph_text(paragraph, '')


def replace_text(paragraph, old: str, new) -> None:
    """Replace text across runs while retaining the surrounding run properties."""
    full_text = "".join(run.text for run in paragraph.runs)
    start = full_text.find(old)
    if start < 0:
        raise ValueError(f"模板占位符不存在：{old!r}")
    end = start + len(old)
    positions = []
    for run_index, run in enumerate(paragraph.runs):
        positions.extend((run_index, offset) for offset in range(len(run.text)))
    first_run, first_offset = positions[start]
    last_run, last_offset = positions[end - 1]
    runs = paragraph.runs
    before = runs[first_run].text[:first_offset]
    after = runs[last_run].text[last_offset + 1:]
    runs[first_run].text = before + str(new or "") + (after if first_run == last_run else "")
    for index in range(first_run + 1, last_run + 1):
        runs[index].text = after if index == last_run else ""


def date_parts(value) -> tuple[str, str, str]:
    if not value:
        return "", "", ""
    try:
        parsed = date.fromisoformat(str(value))
        return str(parsed.year), str(parsed.month), str(parsed.day)
    except ValueError:
        match = re.fullmatch(r"\s*(\d{4})[年/.-](\d{1,2})[月/.-](\d{1,2})日?\s*", str(value))
        if match:
            return match.groups()
        return str(value), "", ""


def fill_date_runs(paragraph, value, *, occurrence: int = 0) -> None:
    """Put numbers before existing 年/月/日 runs, leaving the units in place."""
    year, month, day = date_parts(value)
    if not year:
        return
    if not month or not day:
        set_paragraph_text(paragraph, value)
        return
    markers = [index for index, run in enumerate(paragraph.runs) if "年" in run.text]
    if occurrence >= len(markers):
        raise ValueError("模板日期占位符不足")
    year_index = markers[occurrence]
    if year_index == 0:
        paragraph.runs[year_index].text = year + paragraph.runs[year_index].text
    else:
        paragraph.runs[year_index - 1].text = year
    for marker, value_part in (("月", month), ("日", day)):
        index = next(index for index in range(year_index + 1, len(paragraph.runs)) if marker in paragraph.runs[index].text)
        paragraph.runs[index - 1].text = value_part
