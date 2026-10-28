"""Audited field locations in acquired official templates (no new blank forms)."""
from copy import deepcopy
from datetime import date
from pathlib import Path
import re
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from WenShu.template_fill import (replace_span, fill_label_slots, set_cell_text,
                                  set_paragraph_text, all_paragraphs, date_parts)


def text(value):
    return str(value or '').strip()


def formatted_date(value):
    y, m, d = date_parts(value)
    return f'{y}年{m}月{d}日' if y and m and d else text(value)


def value(data, key):
    if key.endswith('Date'):
        return formatted_date(data.get(key))
    if key == 'applicantMailingAddress':
        return text(data.get(key) or data.get('idAddress'))
    if key == 'mailingAddress':
        return text(data.get(key) or data.get('registeredAddress'))
    return text(data.get(key))


def body_text(element):
    return re.sub(r'\s+', '', ''.join(n.text or '' for n in element.iter(qn('w:t'))))


def normalize_smart_tags(doc):
    # Word 2003 smart tags can contain visible example dates that python-docx
    # excludes from paragraph.runs. Unwrap the metadata while retaining runs
    # and their original font properties, before any writable slot is cleared.
    for tag in list(doc._element.iter(qn('w:smartTag'))):
        parent=tag.getparent()
        for child in list(tag):
            if child.tag!=qn('w:smartTagPr'):tag.addprevious(child)
        parent.remove(tag)


def keep_body(doc, start, end):
    for index, child in enumerate(list(doc._element.body)):
        if child.tag != qn('w:sectPr') and (index < start or index >= end):
            doc._element.body.remove(child)
    first = next((c for c in doc._element.body if c.tag != qn('w:sectPr')), None)
    if first is not None:
        for node in list(first.iter(qn('w:br'))):
            if node.get(qn('w:type')) == 'page':
                node.getparent().remove(node)
        for node in list(first.iter(qn('w:pageBreakBefore'))):
            node.getparent().remove(node)


def select_form(doc, path, kind, *, profile=None):
    name = Path(path).name
    children = list(doc._element.body)
    if name == '湖北省劳动人事争议仲裁文书标准.docx':
        if kind == 'application':
            start = next(i for i, c in enumerate(children) if body_text(c) == '劳动人事争议仲裁申请书')
            end = next(i for i in range(start + 1, len(children)) if '供用人单位申请仲裁使用' in body_text(children[i]))
        else:
            start = next(i for i, c in enumerate(children) if body_text(c) in {'证据清单', '证据目录'})
            table = next(i for i in range(start, len(children)) if children[i].tag == qn('w:tbl'))
            end = table + 1
            while end < len(children) and (not body_text(children[end]) or body_text(children[end]).startswith(('提交人', '注：', '备注'))):
                end += 1
        keep_body(doc, start, end)
    elif name == '附录A（仲裁申请书和证据清单）.docx':
        split = next(i for i, c in enumerate(children) if body_text(c) == '证据清单')
        keep_body(doc, 0 if kind == 'application' else split, split if kind == 'application' else len(children))
    elif (profile or Path(path).parent.name) == 'JiangSuSheng' and kind == 'application':
        split = next(i for i, c in enumerate(children) if '仲裁申请书（用人单位申请）' in body_text(c))
        if split >= 2 and body_text(children[split - 2]).startswith('事项名称'):
            split -= 2
        keep_body(doc, 0, split)


def fill_paragraph_fields(doc, data, bindings):
    for index, fields in bindings.items():
        fill_label_slots(doc.paragraphs[index], [(label, value(data, key)) for label, key in fields])


# These positions are verified against each retained official file, not inferred
# from an arbitrary first occurrence of '申请人' or '电话'. Optional parties stay blank.
INLINE = {
 'JiangSuSheng': {
  1:[('申请人','applicantName'),('性别','gender'),('出生日期','birthDate')],
  2:[('公民身份号码','idNumber'),('居住地址','idAddress')],
  3:[('其他证件类型','otherDocumentType'),('证件号码','otherDocumentNumber')],
  4:[('通讯地址','applicantMailingAddress'),('移动电话','phone')],
  5:[('是否接受调解','acceptMediation'),('是否法律援助','legalAid')],
  6:[('被申请人','respondentName'),('性质','respondentType')],
  7:[('注册地址（住所）','registeredAddress')],8:[('统一社会信用代码','socialCreditCode')],
  9:[('法定代表人或者主要负责人姓名','legalRepName'),('职务','legalRepPosition')],
  10:[('劳动者所在部门负责人','departmentHead'),('职务','departmentHeadPosition')],
  11:[('通讯地址','mailingAddress')],12:[('联系人','companyContact'),('移动电话','companyContactPhone')]},
 'NeiMengGuZiZhiQu': {
  2:[('申请人','applicantName'),('性别','gender'),('民族','nationality'),('身份证号','idNumber')],
  3:[('户籍所在地','idAddress')],4:[('住址','currentAddress')],5:[('联系电话','phone')],
  7:[('被申请人（公司名称）','respondentName')],
  8:[('法定代表人（以营业执照为准）','legalRepName'),('职务（选填项）','legalRepPosition'),('统一社会信用代码','socialCreditCode')],
  9:[('住址','registeredAddress')],10:[('联系电话','companyContactPhone')]},
 'HuNanSheng': {
  1:[('申请人','applicantName')],2:[('性别','gender'),('出生年月','birthDate'),('民族','nationality')],
  3:[('住所','idAddress'),('电话','phone')],4:[('被申请人','respondentName')],5:[('单位地址','registeredAddress')],
  6:[('法定代表人(或主要负责人)','legalRepName'),('职务','legalRepPosition'),('电话','legalRepPhone')]},
 'SiChuanSheng': {
  3:[('申请人','applicantName'),('性别','gender'),('民族','nationality'),('身份证号码','idNumber')],
  4:[('住址','idAddress'),('联系电话','phone')],5:[('确认有效的通讯地址','applicantMailingAddress'),('邮政编码','applicantPostalCode')],
  6:[('被申请人','respondentName')],7:[('法定代表人(或主要负责人)','legalRepName'),('职务','legalRepPosition')],
  8:[('住所','registeredAddress')],9:[('联系电话','companyContactPhone')]},
 'GuiZhouSheng': {
  2:[('申请人','applicantName'),('性别','gender'),('民族','nationality')],
  3:[('出生日期','birthDate'),('原籍','nativePlace')],4:[('现住址','currentAddress'),('联系电话','phone')],
  5:[('确认有效的通讯地址','applicantMailingAddress'),('邮政编码','applicantPostalCode')],
  7:[('被申请人','respondentName')],8:[('法定代表人（主要负责人）','legalRepName'),('职务','legalRepPosition')],
  9:[('住所','registeredAddress')],10:[('联系电话','companyContactPhone')]},
 'GanSuSheng': {
  2:[('姓名','applicantName'),('性别','gender'),('出生年月日','birthDate'),('通讯地址','applicantMailingAddress'),('联系电话','phone')],
  4:[('单位名称','respondentName'),('法定代表人或负责人','legalRepName')],
  5:[('职务','legalRepPosition'),('地址','registeredAddress'),('联系电话','companyContactPhone')]},
}


TABLE_FIELDS = {
 'BeiJingShi': {(1,1):'applicantName',(2,1):'gender',(3,1):'birthDate',(4,1):'householdType',
  (5,1):'idAddress',(8,1):'currentAddress',(10,1):'identityDocumentType',(11,1):'idNumber',(13,1):'phone',
  (1,4):'respondentName',(3,4):'respondentType',(5,4):'legalRepName',(6,4):'legalRepPosition',
  (7,4):'legalRepPhone',(8,4):'registeredAddress',(9,4):'workAddress',(11,4):'officeAtRegisteredAddress',
  (12,4):'actualOfficeAddress',(13,4):'companyContact',(14,4):'companyContactPhone'},
 'LiaoNingSheng': {(0,2):'applicantName',(0,4):'idNumber',(1,2):'gender',(1,4):'age',
  (2,2):'occupation',(2,4):'respondentName',(3,2):'applicantMailingAddress',(4,2):'phone',(4,4):'applicantPostalCode',
  (5,2):'respondentName',(5,4):'respondentType',(6,2):'legalRepName',(6,4):'legalRepPosition',
  (7,2):'mailingAddress',(8,2):'companyContactPhone',(8,4):'respondentPostalCode'},
 'ShanDongSheng': {(0,2):'applicantName',(0,5):'idNumber',(1,2):'gender',(1,4):'nationality',(1,6):'phone',
  (2,2):'idAddress',(3,2):'respondentName',(3,6):'companyContactPhone',(4,2):'registeredAddress',
  (5,4):'legalRepName',(5,6):'legalRepPosition'},
 'ShaanXiSheng': {(1,1):'applicantName',(1,5):'respondentName',(2,1):'idNumber',(2,5):'socialCreditCode',
  (3,5):'respondentType',(4,5):'legalRepName',(5,1):'gender',(5,3):'age',(6,1):'nationality',
  (6,3):'employmentType',(7,1):'respondentName',(8,1):'applicantMailingAddress',(8,5):'mailingAddress',
  (9,1):'phone',(9,5):'companyContactPhone',(10,1):'applicantPostalCode',(10,5):'respondentPostalCode'},
 'YunNanSheng': {(0,2):'applicantName',(0,4):'gender',(0,6):'nationality',(0,8):'birthDate',
  (1,2):'idNumber',(1,6):'idAddress',(2,2):'applicantAddressWithPostalCode',(2,6):'phone',
  (3,2):'respondentName',(3,6):'socialCreditCode',(4,2):'respondentAddressWithPostalCode',(4,6):'companyContactPhone',
  (5,2):'legalRepName',(5,6):'legalRepPosition'},
 'XinJiangWeiWuErZiZhiQu': {(0,1):'applicantName',(0,3):'gender',(1,1):'birthDate',(1,3):'nationality',
  (2,1):'currentAddress',(2,3):'applicantPostalCode',(3,1):'applicantMailingAddress',(3,3):'applicantPostalCode',
  (4,1):'contractPeriod',(5,1):'phone',(6,1):'respondentName',(7,1):'legalRepName',(7,3):'legalRepPosition',
  (8,1):'mailingAddress',(8,3):'respondentPostalCode',(9,1):'companyContactPhone'},
}


def derived(data):
    result = dict(data)
    result.setdefault('currentAddress', value(data, 'idAddress'))
    result['applicantAddressWithPostalCode'] = ' '.join(filter(None,[value(data,'applicantMailingAddress'),value(data,'applicantPostalCode')]))
    result['respondentAddressWithPostalCode'] = ' '.join(filter(None,[value(data,'mailingAddress'),value(data,'respondentPostalCode')]))
    start, end = value(data, 'contractStartDate'), value(data, 'contractEndDate')
    result['contractPeriod'] = f'{start}至{end}' if start and end else ('未签订劳动合同' if data.get('signedContract') == '无' else '')
    if not result.get('age') and data.get('birthDate') and data.get('applicationDate'):
        try:
            birth, filing = date.fromisoformat(data['birthDate']), date.fromisoformat(data['applicationDate'])
            result['age'] = str(filing.year-birth.year-((filing.month,filing.day)<(birth.month,birth.day)))
        except ValueError:
            pass
    return result


def fill_sections(doc, data):
    """Use official heading/blank run styles; remove only unused writing slots."""
    ps = list(doc.paragraphs)
    request = next((p for p in ps if re.match(r'^(?:仲裁请求|请求事项|申请事项|请求事项/反映问题)',re.sub(r'\s+','',p.text))), None)
    facts = next((p for p in ps if re.match(r'^(?:[\d、.\s]*)(?:事实与理由|事实和理由|事实理由)',re.sub(r'\s+','',p.text))), None)
    if request is not None and facts is not None:
        a, b = ps.index(request), ps.index(facts)
        claims = [text(x) for x in data.get('requests',[]) if text(x)]
        if claims:
            template = ps[a+1] if a+1 < b else request
            body_run=request.runs[-1] if request.runs else None
            heading = re.match(r'^.*?[:：]',request.text)
            if heading:
                replace_span(request,heading.end(),len(request.text),f'1、{claims[0]}',style_run=body_run)
            else:
                request.add_run(f'：1、{claims[0]}')
            request.runs[-1].bold=False
            body_properties=deepcopy(request.runs[-1]._r.rPr)
            anchor = request._p
            for index, claim in enumerate(claims[1:],2):
                node = deepcopy(template._p)
                from docx.text.paragraph import Paragraph
                p = Paragraph(node,request._parent)
                set_paragraph_text(p,f'{index}、{claim}')
                first=p.runs[0]._r
                if first.rPr is not None:first.remove(first.rPr)
                if body_properties is not None:first.insert(0,deepcopy(body_properties))
                for numbering in list(node.iter(qn('w:numPr'))):
                    numbering.getparent().remove(numbering)
                for br in list(node.iter(qn('w:br'))):
                    if br.get(qn('w:type')) == 'page':
                        br.getparent().remove(br)
                anchor.addnext(node)
                anchor=node
            for p in ps[a+1:b]:
                p._p.getparent().remove(p._p)
    if facts is not None and text(data.get('factsReason')):
        index=ps.index(facts)
        colon=re.search(r'[:：]',facts.text)
        if colon:
            replace_span(facts,colon.end(),len(facts.text),text(data['factsReason']),style_run=facts.runs[-1] if facts.runs else None)
            facts.runs[-1].bold=False
        else:
            facts.add_run('：'+text(data['factsReason']))
        for p in ps[index+1:]:
            compact=body_text(p._p)
            if not compact or re.fullmatch(r'[\s_＿.。]+',p.text):
                if p._p.getparent() is not None:
                    p._p.getparent().remove(p._p)
            else:
                break
    for heading in (request,facts):
        if heading is not None:
            for br in list(heading._p.iter(qn('w:br'))):
                if br.get(qn('w:type'))=='page':br.getparent().remove(br)


def fill_footer(doc,data, *, handwritten_date=False):
    body_paragraphs={p._p for p in doc.paragraphs}
    for p in all_paragraphs(doc):
        compact=body_text(p._p)
        if re.fullmatch(r'[×xX]*.*(?:劳动人事争议|劳动争议)仲裁委员会(?:高新区派出庭)?',compact) and len(compact)<55:
            if value(data,'committeeName'):
                replace_span(p,0,len(p.text),value(data,'committeeName'))
        is_footer=p._p in body_paragraphs or any(t in getattr(p._parent,'text','') for t in ('此致','签名','签字'))
        if is_footer and not handwritten_date and re.fullmatch(r'(?:二[○〇]|[×xX\d]*)年[×xX\d]*月[×xX\d]*日',compact):
            if value(data,'applicationDate'):
                replace_span(p,0,len(p.text),value(data,'applicationDate'))
                from docx.enum.text import WD_ALIGN_PARAGRAPH
                p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        for match in reversed(list(re.finditer(r'(?:副本|证据材料|有关证据材料|相关证据|材料清单)(?P<slot>[_＿ \t]+)(?=份)',p.text))):
            replace_span(p,match.start('slot'),match.end('slot'),str(data.get('copyCount',1)))
    # Signing and receiving are human actions; no name/receiver/date is invented.


def fill_application(data, path, *, profile=None):
    doc=Document(path)
    if not any(text(data.get(k)) for k in ('applicantName','respondentName','factsReason','applicationDate')) and not data.get('requests'):
        return doc
    select_form(doc,path,'application',profile=profile)
    normalize_smart_tags(doc)
    data=derived(data)
    folder=profile or Path(path).parent.name
    # Official example documents contain real-looking example values. Clear
    # every writable example cell even when the caller leaves a field blank.
    if folder == 'XinJiangWeiWuErZiZhiQu':
        for (r,c) in TABLE_FIELDS[folder]:
            set_cell_text(doc.tables[0].rows[r].cells[c], '')
    if folder in INLINE:
        fill_paragraph_fields(doc,data,INLINE[folder])
    if folder=='JiangSuSheng':
        # This source's birth-date slot shares a narrow line with name/sex.
        # The expanded Chinese date wrapped its month/day onto a new line.
        # Preserve the run's font/underline and use the same date compactly.
        year,month,day=date_parts(data.get('birthDate'))
        if year and month and day:
            fill_label_slots(doc.paragraphs[1],[('出生日期',f'{year}-{int(month):02d}-{int(day):02d}')],pad=False)
        # The retained worker form offers a second respondent. This frontend
        # explicitly collects one respondent; remove that unused party block.
        for p in list(doc.paragraphs)[13:20]:
            p._p.getparent().remove(p._p)
    if folder=='GanSuSheng':
        for p in [doc.paragraphs[3],doc.paragraphs[6]]:
            p._p.getparent().remove(p._p)
    if folder in TABLE_FIELDS and TABLE_FIELDS[folder]:
        table=doc.tables[0]
        for (r,c),key in TABLE_FIELDS[folder].items():
            if value(data,key):
                set_cell_text(table.rows[r].cells[c],value(data,key))
    if folder == 'FuJianSheng':
        table=doc.tables[0]
        for row, specs in enumerate(([
          [('申请人','applicantName'),('出生','birthDate'),('性别','gender'),('民族','nationality')],
          [('户籍地址','idAddress'),('身份证号','idNumber')],[('联系电话','phone')]],[
          [('被申请人','respondentName')],[('法定代表人/负责人','legalRepName'),('职务','legalRepPosition')],
          [('住所','registeredAddress')],[('联系电话','companyContactPhone')]])):
            for p,fields in zip(table.rows[row].cells[0].paragraphs,specs):
                fill_label_slots(p,[(label,value(data,key)) for label,key in fields])
    if folder == 'HuBeiSheng':
        fill_paragraph_fields(doc,data,{2:[('申请人','applicantName'),('性别','gender')],
          3:[('出生日期','birthDate'),('民族','nationality')],4:[('身份证件号码','idNumber'),('户口性质','householdType')],
          5:[('户籍所在地','idAddress')],6:[('现住址','currentAddress'),('联系电话','phone')],
          7:[('确认有效的通讯地址','applicantMailingAddress')],8:[('邮政编码','applicantPostalCode')],
          10:[('被申请人','respondentName')],11:[('法定代表人（主要负责人）','legalRepName'),('职务','legalRepPosition')],
          12:[('住所','registeredAddress')],13:[('联系电话','companyContactPhone'),('邮政编码','respondentPostalCode')]})
        for p in list(doc.paragraphs)[15:22]:
            p._p.getparent().remove(p._p)
    if folder == 'YunNanSheng':
        table=doc.tables[0]
        claims=[text(x) for x in data.get('requests',[]) if text(x)]
        if len(claims)>4:
            raise ValueError('本模板最多容纳4项请求，请合并相关请求或附续页')
        for index,claim in enumerate(claims[:4]):
            set_cell_text(table.rows[9+index].cells[2],claim)
        set_cell_text(table.rows[13].cells[1],text(data.get('factsReason')))
        # Preserve the footer's individual paragraphs and alignments.
        for p in table.rows[14].cells[0].paragraphs:
            if '劳动人事争议仲裁委员会' in p.text:
                replace_span(p,0,len(p.text),value(data,'committeeName'))
        for row in reversed(list(table.rows[9+len(claims):13])):
            if len(claims)<4:
                table._tbl.remove(row._tr)
    elif folder == 'XinJiangWeiWuErZiZhiQu':
        for row in doc.tables[0].rows[10:14]:
            for c in range(1,len(row.cells),2):
                set_cell_text(row.cells[c],'')
        for p in list(doc.paragraphs):
            if p.text.strip().startswith('以及是否签订劳动合同'):
                p._p.getparent().remove(p._p)
        fill_sections(doc,data)
    elif folder == 'TianJinShi':
        p=doc.paragraphs
        replace_span(p[1],p[1].text.index('：')+1,len(p[1].text),
            f"{value(data,'applicantName')}，{value(data,'gender')}，{value(data,'birthDate')}出生，{value(data,'nationality')}，公民身份号码：{value(data,'idNumber')}，住址：{value(data,'idAddress')}，户籍类型：{value(data,'householdType')}")
        replace_span(p[5],p[5].text.index('：')+1,len(p[5].text),
            f"{value(data,'respondentName')}，法定代表人（主要负责人）：{value(data,'legalRepName')}，职务：{value(data,'legalRepPosition')}")
        fill_paragraph_fields(doc,data,{2:[('工作单位','respondentName')],3:[('文书邮寄送达地址','applicantMailingAddress')],
           4:[('联系电话','phone')],7:[('文书邮寄送达地址','mailingAddress')],8:[('联系电话','companyContactPhone')]})
        replace_span(p[6],0,len(p[6].text),f"用人单位住所地：{value(data,'registeredAddress')}，企业类型：{value(data,'respondentType')}，注册资金：{value(data,'registeredCapital')}")
        p[15]._p.getparent().remove(p[15]._p)
        fill_sections(doc,data)
    elif folder == 'JiangXiSheng':
        p=doc.paragraphs
        # This official form uses sentence placeholders rather than labels for
        # gender/ethnicity; replace only that sentence's template text.
        replace_span(p[1],p[1].text.index('：')+1,len(p[1].text),f"{value(data,'applicantName')}，{value(data,'gender')}，{value(data,'nationality')}，{value(data,'birthDate')}出生，公民身份号码：{value(data,'idNumber')}，住址：{value(data,'idAddress')}，联系电话：{value(data,'phone')}。")
        fill_paragraph_fields(doc,data,{2:[('被申请人','respondentName'),('统一社会信用代码','socialCreditCode'),('单位地址','registeredAddress'),('邮政编码','respondentPostalCode'),('联系电话','companyContactPhone')],3:[('法定代表人','legalRepName'),('职务','legalRepPosition')]})
        fill_sections(doc,data)
    elif folder == 'HeNanSheng':
        p=doc.paragraphs
        replace_span(p[3],p[3].text.index('：')+1,len(p[3].text),f"{value(data,'applicantName')}，{value(data,'gender')}，{value(data,'birthDate')}出生，{value(data,'nationality')}，现住址：{value(data,'currentAddress')}，联系电话：{value(data,'phone')}。")
        replace_span(p[4],p[4].text.index('：')+1,len(p[4].text),f"{value(data,'respondentName')}，法定代表人：{value(data,'legalRepName')}，职务：{value(data,'legalRepPosition')}，办公地址：{value(data,'registeredAddress')}，办公电话：{value(data,'companyContactPhone')}。")
        # A stray '4、' in the official sample belongs to the unused request slot.
        replace_span(p[9],0,len(p[9].text),'事实和理由：')
        fill_sections(doc,data)
    elif folder == 'HaiNanSheng':
        fill_hainan(doc,data)
    else:
        fill_sections(doc,data)
    fill_footer(doc,data,handwritten_date=folder=='HaiNanSheng')
    if folder=='FuJianSheng':
        for p in list(doc.paragraphs):
            if not p.text.strip() and not p._p.xpath('.//w:sectPr'):
                p._p.getparent().remove(p._p)
    # Keep handwritten signature and its associated date on the same page.
    signature=False
    for p in doc.paragraphs:
        if ('签名' in p.text or '签字' in p.text) and len(p.text.strip())<100:
            signature=True
        if signature:
            if re.fullmatch(r'[×xX\d]*年[×xX\d]*月[×xX\d]*日',body_text(p._p)):
                signature=False
            elif not p.text.strip() or '申请人' in p.text:
                p.paragraph_format.keep_with_next=True
            else:signature=False
    remove_empty_edges(doc)
    return doc


def fill_hainan(doc,data):
    p=doc.paragraphs
    if len(data.get('requests',[]))>2:
        raise ValueError('本模板最多容纳2项请求，请合并相关请求或附续页')
    for index,content in {
      2:f"申请人：{value(data,'applicantName')}，{value(data,'gender')}，{value(data,'age')}周岁，公民身份号码：{value(data,'idNumber')}，现住址：{value(data,'currentAddress')}，联系电话：{value(data,'phone')}。",
      3:f"被申请人：{value(data,'respondentName')}，住所/经营场所：{value(data,'registeredAddress')}。",
      4:f"法定代表人：{value(data,'legalRepName')}，职务：{value(data,'legalRepPosition')}，联系电话：{value(data,'legalRepPhone')}。",
      11:f"（一）员工入职时间：{value(data,'employmentDate')}",
      18:f"（三）员工工作岗位：{value(data,'jobTitle')}",20:f"1.实际发放的月工资数额：{value(data,'startingSalary')}元。",
      22:f"3.员工离职前十二个月的月平均工资：{value(data,'averageMonthlySalary')}元。",
      25:f"2.离职时间：{value(data,'departureDate')}；最后正常工作日：{value(data,'lastWorkingDate')}。",
      26:f"3.离职原因：{value(data,'departureReason')}",
      37:value(data,'otherNotes') or '无。',
      44:'申请人：       （日期及签名均由当事人手写）',46:'       年       月       日',
    }.items():
        replace_span(p[index],0,len(p[index].text),content)
    for index in (5,6):
        set_paragraph_text(p[index],'')
    for index,claim in enumerate(data.get('requests',[])[:2]):
        replace_span(p[8+index],0,len(p[8+index].text),f'{index+1}、{text(claim)}')
        start=28+4*index
        for offset,content in enumerate((f'诉求{index+1}：{text(claim)}',
          '1.事实及理由：'+value(data,'factsReason'),'2.证据：详见证据清单。',
          '3.计算方式：'+value(data,'claimCalculation'))):
            replace_span(p[start+offset],0,len(p[start+offset].text),content)
    if len(data.get('requests',[]))<2:
        for index in (9,32,33,34,35):set_paragraph_text(p[index],'')
    for index,selected in ((13,data.get('signedContract')),(21,data.get('salaryPaymentMethod')),(24,'未离职' if data.get('currentlyEmployed')=='是' else '已离职')):
        if selected:
            match=re.search(re.escape(str(selected))+r'\s*□',p[index].text)
            if match:
                replace_span(p[index],match.start(),match.end(),str(selected)+'☑')
    replace_span(p[15],0,len(p[15].text),('☑从'+value(data,'contractStartDate')+'起至'+value(data,'contractEndDate')+'止的固定期限劳动合同；') if data.get('signedContract')=='有' and data.get('contractStartDate') else '')
    # Remove the sample's alternative unused contract durations and claims.
    for index in (16,17):
        set_paragraph_text(p[index],'')
    for paragraph in p:
        if not paragraph.text.strip() and not paragraph._p.xpath('.//w:drawing | .//w:pict | .//w:sectPr'):
            paragraph._p.getparent().remove(paragraph._p)


def remove_empty_edges(doc):
    body=doc._element.body
    for child in reversed(list(body)):
        if child.tag==qn('w:sectPr'):
            continue
        if child.tag==qn('w:p') and not body_text(child):
            body.remove(child)
        else:
            break
    children=[c for c in body if c.tag!=qn('w:sectPr')]
    if children:
        for node in list(children[-1].iter(qn('w:br'))):
            if node.get(qn('w:type'))=='page':
                node.getparent().remove(node)


EVIDENCE_COLUMNS = {
 'BeiJingShi': {0:'number',1:'name',2:'kind',3:'hasOriginal',4:'purpose',5:'pageCount'},
 'TianJinShi': {0:'number',1:'name',2:'purpose',3:'pageCount'},
 'JiangSuSheng': {0:'number',1:'name',2:'source',3:'purpose',4:'originalType',5:'pageCount'},
 'FuJianSheng': {0:'number',1:'name',2:'pageCount',3:'originalType',4:'purpose',5:'pageRange'},
 'JiangXiSheng': {0:'number',1:'name',2:'purpose'},
 'ShanDongSheng': {0:'number',1:'name',2:'purpose',3:'notes'},
 'HuNanSheng': {0:'number',1:'name',2:'pageRange',3:'originalType',4:'source',5:'purpose'},
 'HaiNanSheng': {0:'number',1:'name',3:'pageRange',4:'purpose',6:'hasOriginal'},
 'SiChuanSheng': {0:'number',1:'name',2:'itemCount',3:'pageCount',4:'originalMark',5:'copyMark',6:'purpose',7:'notes'},
 'GuangDongSheng': {0:'number',1:'name',2:'proof',3:'pageCount'},
 'ZheJiangSheng': {0:'number',1:'name',3:'pageCount',4:'source',5:'purpose'},
 'HuBeiSheng': {0:'number',1:'name',2:'proof',3:'pageCount'},
 'YunNanSheng': {0:'number',1:'name',2:'proof',3:'pageCount'},
}


def fill_evidence(data,path, *, profile=None):
    doc=Document(path)
    if not any(text(x.get('name')) for x in data.get('evidenceList',[])) and not any(text(data.get(k)) for k in ('committeeName','submitDate','caseNumber')):
        return doc
    select_form(doc,path,'evidence',profile=profile)
    normalize_smart_tags(doc)
    folder=profile or Path(path).parent.name
    columns=EVIDENCE_COLUMNS[folder]
    table=doc.tables[0]
    start=2 if folder=='ZheJiangSheng' else 1
    items=[dict(x) for x in data.get('evidenceList',[]) if text(x.get('name'))]
    rows=[]
    for row in list(table.rows)[start:]:
        content=''.join(c.text for c in row.cells)
        if any(t in content for t in ('合计','提交人','接收人','接受人','备注：','以上证据')):
            break
        rows.append(row)
    if not rows:
        raise ValueError('未找到官方证据填写行')
    from docx.table import _Row
    while len(rows)<len(items):
        node=deepcopy(rows[-1]._tr)
        rows[-1]._tr.addnext(node)
        rows.append(_Row(node,table))
    for index,item in enumerate(items):
        row=rows[index]
        # Keep the template's minimum height but allow longer evidence to wrap.
        for h in row._tr.iter(qn('w:trHeight')):
            if h.get(qn('w:hRule'))=='exact':
                h.set(qn('w:hRule'),'atLeast')
        item['number']=str(index+1)
        item['proof']='；'.join(filter(None,[text(item.get('source')),text(item.get('purpose'))]))
        item['originalMark']='√' if item.get('originalType')=='原件' else ''
        item['copyMark']='√' if item.get('originalType')=='复印件' else ''
        for col,key in columns.items():
            set_cell_text(row.cells[col],text(item.get(key)))
    # Remove unused evidence slots, preserving totals and human signature rows.
    for row in rows[len(items):]:
        table._tbl.remove(row._tr)
    if folder=='ZheJiangSheng':
        if data.get('caseNumber'):
            set_cell_text(table.rows[0].cells[1],data['caseNumber'])
        set_cell_text(table.rows[0].cells[5],str(data.get('copyCount',2)))
        footer=table.rows[start+len(items)]
        set_cell_text(footer.cells[2],'')
        set_cell_text(footer.cells[5],formatted_date(data.get('submitDate')))
    elif folder in {'HuNanSheng','HaiNanSheng'}:
        for row in table.rows:
            if row.cells[0].text.strip().startswith(('证据提交人','提交人')):
                set_cell_text(row.cells[2],'')
                set_cell_text(row.cells[6 if folder=='HuNanSheng' else 5],formatted_date(data.get('submitDate')))
    elif folder=='FuJianSheng':
        row=table.rows[-1]
        if '合计' in row.cells[0].text:
            counts=[re.fullmatch(r'(\d+)\s*(?:页)?',text(x.get('pageCount'))) for x in items]
            unique={c._tc for c in row.cells}
            total=str(sum(int(m.group(1)) for m in counts)) if all(counts) else ''
            if len(unique)==1:
                set_cell_text(row.cells[0],f'合计：{total}页（{len(items)}项）' if total else f'合计：{len(items)}项')
            else:set_cell_text(row.cells[2],total)
    for p in all_paragraphs(doc):
        compact=body_text(p._p)
        if '仲裁委员会' in compact and len(compact)<55 and data.get('committeeName'):
            replace_span(p,0,len(p.text),data['committeeName'])
        if data.get('caseNumber') and (compact.startswith('案号') or '仲案' in compact):
            colon=re.search('[:：]',p.text)
            replace_span(p,colon.end() if colon else 0,len(p.text),data['caseNumber'])
        if p._p in {q._p for q in doc.paragraphs} and '提交' in compact and any(t in compact for t in ('时间','日期')):
            # Fill only submission date; neither a signature nor receiving date.
            match=re.search(r'提交(?:时间|日期)\s*[:：]?',p.text)
            if match and data.get('submitDate'):
                end=re.search(r'接收日期',p.text[match.end():])
                replace_span(p,match.end(),match.end()+end.start() if end else len(p.text),formatted_date(data['submitDate'])+'  ')
        if folder=='FuJianSheng' and p._p in {q._p for q in doc.paragraphs} and '日期' in compact:
            fill_label_slots(p,[('提交人',''),('日期',formatted_date(data.get('submitDate')))],pad=False)
        elif folder=='BeiJingShi' and '提交日期' in compact and data.get('submitDate'):
            fill_label_slots(p,[('提交日期',formatted_date(data['submitDate'])),('接收日期','')])
        if '副本' in compact or '附证据' in compact or '本表一式' in compact:
            for m in reversed(list(re.finditer(r'(?P<slot>[_＿\s]+)(?=份)',p.text))):
                replace_span(p,m.start('slot'),m.end('slot'),str(data.get('copyCount',2)))
    remove_empty_edges(doc)
    return doc
