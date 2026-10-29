"""Offline own-source Ningbo district guides and strictly bound blank filling.

No network, database, publication or model calls. Source files are public rendered
content, sanitized control observations and cookie-free official originals.
"""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
from urllib.parse import parse_qs, urlsplit

from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from regional_guides import digest
from regional_documents import PreservedPackageDocument
from WenShu.Zhejiang.LAdocx import generate_LAdocx
from WenShu.exact_template_fill import fill_evidence
from WenShu.template_fill import set_cell_text

ROOT=Path(__file__).resolve().parents[1]
REPORT='reports/nationwide-verification-20261029/zhejiang'
MANIFEST_SHA256='dc941a9bedc140e9a18eca58b2d2a390209a9016d337e9b16f194f356ffd6c2e'
SERVICE='劳动人事争议仲裁申请'
STANDARD='fc266994-f669-4861-acdc-799f3ea2decc'
ADAPTER='zhejiang_ningbo_own_blank_oct29_v1'
SPECS={
    'haishu':{'code':'330203','district':'海曙区','local':'a57febe6-cd30-4a53-8b4c-f24edaaf7864','agency':'区人社局','phone':'0574-55889905',
        'window':'海曙区劳动人事争议仲裁院（宁波市海曙区月湖街道解放南路259号房地产大厦805室）'},
    'yinzhou':{'code':'330212','district':'鄞州区','local':'73f99740-5b23-424a-9ae8-be607cdceaeb','agency':'区人力社保局','phone':'0574-89293274',
        'window':'宁波市鄞州区人力资源和社会保障局（宁波市鄞州区首南街道惠风东路225号）信访监察仲裁联合接访服务大厅（宁波市鄞州区惠风东路225号辅楼3楼）'}}
MATERIALS=['用人单位登记注册材料','劳动者身份证','《劳动人事争议仲裁申请书》','《证据清单》','存在劳动关系的证据和可证实申请人主张的证据']
BLANKS={
    'application':('application.docx','21B24C52D2AC3EB007AA7F55CE0AAD79','0eab0f0a2ee98e8ca1e4d2fa5469fcb5738d0c0f24671345bebc8394099ae41f',2),
    'evidence':('evidence-list.docx','872DFF90A8578907CB947197A784AA48','4965bf3d2bf76404ab190d1e47d2fa250c8cb8bd93762949df94be38b2d4ca8c',3)}

class Evidence:
    def __init__(self, root):
        self.root=Path(root).resolve(); raw=(self.root/'SOURCE-MANIFEST.json').read_bytes()
        if hashlib.sha256(raw).hexdigest()!=MANIFEST_SHA256:raise ValueError('固定浙江本轮来源清单不匹配')
        self.files=json.loads(raw)['files']
    def read(self, relative):
        path=(self.root/relative).resolve()
        if not path.is_relative_to(self.root):raise ValueError('来源路径越界')
        data=path.read_bytes(); expected=self.files.get(relative)
        if expected!={'sha256':hashlib.sha256(data).hexdigest(),'bytes':len(data)}:raise ValueError('来源字节已变化：'+relative)
        return data
    def load(self, relative):return json.loads(self.read(relative).decode('utf8'))

def parse_capture(name, capture_root=None):
    if name not in SPECS:raise ValueError('本适配器仅处理海曙和鄞州的本轮来源')
    e=Evidence(capture_root or ROOT/REPORT); s=SPECS[name]
    source='https://www.zjzwfw.gov.cn/zjservice-fe/#/workguide?localInnerCode='+s['local']
    pages={}
    for section in ['conditions','materials','flow','window','faq','basic','law']:
        page=e.load(name+'/'+section+'/page.json');e.read(name+'/'+section+'/page.png')
        if page['url']!=source or SERVICE not in page['text'] or s['phone'] not in page['text']:raise ValueError('栏目不是本区自身来源或仍为切换前地区数据')
        pages[section]=page
    search=e.load(name+'/search/page.json')
    if parse_qs(urlsplit(search['url']).query).get('websiteid')!=[s['code']+'000000'] or ('当前搜索范围为： '+s['district']) not in search['text']:
        raise ValueError('本区目录身份不匹配')
    links=[a for a in search['links'] if a.get('title')==SERVICE and a.get('localInnerCode')==s['local'] and a.get('standard')==STANDARD]
    if len(links)!=1:raise ValueError('缺少本区目录与正式指南对应关系')
    basic=pages['basic']; text=basic['text']; window=pages['window']['text']
    if s['window'] not in window or s['phone'] not in text or ('实施主体：\n'+s['agency']) not in text:
        raise ValueError('本区窗口、咨询电话或实施主体不匹配')
    if not re.search(r'法定办结时限：\n60天',text) or '44\n工作日' not in text:raise ValueError('时限原文已改变')
    own={'source_url':source,'region_code':s['code'],'collected_at':basic['captured_at'],'verified_at':basic['captured_at']}
    def field(value,**extra):return {**own,'raw_value':value,'normalized_value':value,'verification_status':'verified',**extra}
    fields={'实施主体':field(s['agency'],normalization_note='保持指南简称，机关地域由本区目录和自身窗口核对，不自行扩写全称'),
        '法定办结时限':field('60天'),'承诺办结时限':field('44工作日'),'到现场次数':field('0次'),
        '办理地点':field(s['window']),'咨询电话':field(s['phone']),'是否收费':field('不收费')}
    hours=window.split('工作时间\n')[-1].split('联系电话\n')[0].split('好差评\n')[0].strip()
    fields['办理时间']=field(hours)
    flow=pages['flow']['text'].split('申请－受理－审理－出具结案文书－送达－办结\n')[-1].split('更多信息\n')[0].strip().removesuffix('收起').strip()
    if '在5个工作日内作出是否受理的决定' not in flow or '现场或邮寄提交申请材料' not in flow:raise ValueError('受理阶段或办理方式原文缺失')
    fields['受理决定时限']=field('5个工作日',scope='acceptance_decision_only')
    material_text=pages['materials']['text'].split('共 5 项 · 自备材料\n')[-1].split('办理流程\n')[0]
    materials=[]
    for i,label in enumerate(MATERIALS):
        expected=label+'\n纸质材料 2份，A4\n原件'
        if expected not in material_text:raise ValueError('本区材料名称、份数或原件要求改变')
        materials.append({**own,'id':digest([source,label])[:20],'name':label,'required':None,'copies':'2','raw_copies':'纸质材料 2份，A4',
            'format':'纸质材料，A4','original_copy':'原件','necessity_status':'not_found_in_checked_sources','verification_status':'verified'})
    observation=e.load(name+'/download-observation.json')
    if observation['source_url']!=source or [x['name'] for x in observation['downloads']]!=[MATERIALS[2]+'.docx',MATERIALS[3]+'.docx']:
        raise ValueError('缺少本区自己实际空白控件与原名证据')
    templates=[]
    for kind,(filename,annex,checksum,index) in BLANKS.items():
        relative=name+'/originals/blank/'+filename;body=e.read(relative);meta=e.load(relative+'.source.json')
        url='https://portal.zjzwfw.gov.cn/media/oss/FILE/META_DATA/'+annex+'.docx'
        if (hashlib.sha256(body).hexdigest()!=checksum or meta['source_url']!=source or meta['download_url']!=url or meta['final_url']!=url
            or meta['kind']!='blank' or meta['original_name']!=MATERIALS[index]+'.docx' or url not in observation['public_fetches']):
            raise ValueError('原件与本区下载动作、地区或来源不匹配')
        templates.append({**own,'id':digest([source,annex,'blank'])[:24],'official_annex_id':annex,'annex_id':annex,
            'material_id':materials[index]['id'],'original_name':meta['original_name'],'source_filename':annex+'.docx','official_kind':'blank','file_type':'docx',
            'file':REPORT+'/'+relative,'retained_download_file':REPORT+'/'+relative,'sha256':checksum,'bytes':len(body),'download_url':url,
            'final_download_url':url,'downloaded_at':meta['downloaded_at'],'download_status':'downloaded','http_status':200,
            'can_generate':False,'document_types':[],'applicant_scope':'worker' if kind=='application' else 'evidence_list'})
    sample=e.load(name+'/sample/page.json');sample_file=name+'/originals/sample/identity-example.jpg';sample_bytes=e.read(sample_file);sample_meta=e.load(sample_file+'.source.json')
    if ('材料示例预览' not in sample['text'] or sample_meta['source_url']!=source or sample_meta['kind']!='sample'
        or hashlib.sha256(sample_bytes).hexdigest()!='1a992bf994963a6a03f7344e9a766246c0b495756eb4877406b052d039272b09'):
        raise ValueError('本区身份证示例来源不匹配')
    templates.append({**own,'id':digest([source,'identity_sample'])[:24],'material_id':materials[1]['id'],'original_name':sample_meta['original_name'],
        'official_kind':'sample','file_type':'jpg','file':REPORT+'/'+sample_file,'sha256':sample_meta['sha256'],'bytes':len(sample_bytes),
        'download_url':sample_meta['download_url'],'downloaded_at':sample_meta['downloaded_at'],'download_status':'downloaded','can_generate':False,'document_types':[]})
    boundary=e.load(name+'/online-boundary.json');e.read(name+'/online-boundary.png')
    if boundary['origin']!='https://user.zjzwfw.gov.cn' or boundary['path']!='/pc/login' or not boundary['personalLogin'] or boundary['passwordInputs']!=1:
        raise ValueError('在线边界证据不匹配')
    faq=e.load(name+'/faq-expanded/page.json');e.read(name+'/faq-expanded/page.png')
    if faq['url']!=source or s['phone'] not in faq['text']:raise ValueError('展开问答不是本区自身数据')
    sections={'基本信息':text.split('基本信息\n')[-1].split('好差评\n')[0].strip(),
        '受理条件':pages['conditions']['text'].split('该事项共 5 项条件和 1 条特别提示。\n')[-1].split('申报材料\n')[0].strip(),
        '申请材料':material_text.strip(),'办理流程':flow,'窗口办理':s['window']+'\n'+hours,
        '常见问题':faq['text'].split('基本信息\n')[-1].split('好差评\n')[0].strip(),
        '设定依据':pages['law']['text'].split('设定依据\n')[-1]}
    return {'province':'zhejiang','province_name':'浙江省','city':'宁波市','district':s['district'],'region_code':s['code'],
        'service_name':SERVICE,'service_code':s['local'],'service_code_type':'observed_localInnerCode_not_national_implementation_code',
        'service_ids':{'localInnerCode':s['local'],'stdcateguid':STANDARD},'source_id':digest(source)[:24],'source_url':source,
        'source_title':basic['title'],'source_type':'official_service_guide','authority':s['agency'],**{k:own[k] for k in ['collected_at','verified_at']},
        'verification_status':'partial','published':False,'materials_status':'five_rows_captured','fields':fields,'materials':materials,'templates':templates,'sections':sections,
        'channel':{'form':'window_and_online','raw_value':flow,'online_status':'auth_required','entry_url':source,
            'entry_check':{'url':'https://user.zjzwfw.gov.cn/pc/login','status':'auth_required','official_link':True,'checked_at':boundary['captured_at'],'post_login_flow':'unknown','submitted':False}},
        'evidence':{'manifest_sha256':MANIFEST_SHA256,'capture_file':REPORT+'/'+name+'/basic/page.json','capture_sha256':hashlib.sha256(e.read(name+'/basic/page.json')).hexdigest()},
        'warnings':['仅限宁波市本区自身采集事项；不代表全省统一模板或全国完成。','60天、44工作日分别保留原单位；5个工作日仅为受理决定阶段。',
            '申请书填写适配仅支持原件中劳动者申请部分；用人单位申请部分未适配。',
            '材料表写2份原件，申请书正文副本说明为按对方当事人人数提交；两处原文并列保留，须按受理机关要求核对。',
            '网上办理仅到实际登录边界；登录后流程未知，未登录或提交。','当前为未发布候选，候选网页验收未运行。']}

def parse_haishu_capture(capture_root=None):return parse_capture('haishu',capture_root)
def parse_yinzhou_capture(capture_root=None):return parse_capture('yinzhou',capture_root)

def blank_binding(guide, template):
    try:
        name=next(n for n,s in SPECS.items() if s['code']==guide.get('region_code')); retained=parse_capture(name)
        keys=['province','city','district','region_code','source_url','source_id','source_type','authority','service_name','service_code','service_code_type','service_ids','evidence','fields','materials','verification_status','collected_at','verified_at']
        if any(guide.get(k)!=retained[k] for k in keys):return None
        own=next(t for t in retained['templates'] if t['id']==template.get('id'))
        if own['official_kind']!='blank':return None
        if any(template.get(k)!=v for k,v in own.items() if k not in {'can_generate','document_types'}):return None
        return next(kind for kind,spec in BLANKS.items() if spec[1]==own['official_annex_id'])
    except (ValueError,KeyError,TypeError,OSError,StopIteration):return None

def with_capabilities(guide):
    result=deepcopy(guide)
    for template in result['templates']:
        kind=blank_binding(result,template)
        template.update(can_generate=bool(kind),document_types=[kind] if kind else [])
        if kind:template.update(adapter=ADAPTER,content_kind='filled_official_blank',generation_label='依据宁波市'+result['district']+'自身官方空白表格填写')
    return result

def fill_document(guide, template, kind, path, data):
    if kind not in BLANKS or blank_binding(guide,template)!=kind or template.get('adapter')!=ADAPTER or template.get('can_generate') is not True:
        raise ValueError('仅可填写本区严格来源绑定的官方空白文书')
    path=Path(path)
    if path.resolve()!=(ROOT/template['file']).resolve():raise ValueError('不得传入其他地区或任意路径的原件')
    original=path.read_bytes();values=deepcopy(data)
    if values.get('copyCount') is None:values['copyCount']='    '
    if kind=='application':
        document=generate_LAdocx(values,path)
        for cell in document.tables[0].rows[12].cells:
            for paragraph in cell.paragraphs:paragraph.paragraph_format.keep_with_next=True
        tail=document.add_paragraph();tail.paragraph_format.space_before=Pt(0);tail.paragraph_format.space_after=Pt(0);tail.paragraph_format.line_spacing=Pt(1)
        rpr=OxmlElement('w:rPr');size=OxmlElement('w:sz');size.set(qn('w:val'),'2');rpr.append(size);tail._p.get_or_add_pPr().append(rpr)
    else:
        case_number=values.pop('caseNumber',None)
        values['evidenceList']=[{**item,'pageCount':re.sub(r'^\s*(\d+)\s*页\s*$',r'\1',str(item.get('pageCount','')))} for item in values.get('evidenceList',[])]
        document=fill_evidence(values,path,profile='ZheJiangSheng');table=document.tables[0]
        if case_number:set_cell_text(table.rows[0].cells[1],case_number)
        for row in table.rows[:2]:
            properties=row._tr.get_or_add_trPr()
            if properties.find(qn('w:tblHeader')) is None:properties.append(OxmlElement('w:tblHeader'))
        for row in table.rows:
            properties=row._tr.get_or_add_trPr()
            if properties.find(qn('w:cantSplit')) is None:properties.append(OxmlElement('w:cantSplit'))
        for row in table.rows[-2:]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:paragraph.paragraph_format.keep_with_next=True
    return PreservedPackageDocument(document,original)
