"""Exact own-source Zhejiang E worker forms with public-root-only evidence.

The E source bundle is immutable. Only pure slot operations are reused; no D
source, private sample, browser, database, service or model is accessed.
"""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path

from docx import Document
from regional_documents import PreservedPackageDocument
from regional_guides import digest
from regional_zhejiang_oct29d import (
    _application, _evidence, _validate_data, APPLICATION_TEXT_FIELDS,
    APPLICATION_FIELDS, EVIDENCE_FIELDS, EVIDENCE_ITEM_FIELDS, LIMITS as SLOT_LIMITS,
    HOUSEHOLD_OPTIONS,
)

ROOT = Path(__file__).resolve().parents[1]
REPORT = 'reports/nationwide-verification-20261029-e/zhejiang'
ADAPTER = 'zhejiang_own_worker_oct29e_v1'
WHITELIST_SHA256 = '14b0ff6cb9737e2ecb0edb1d6809434449495cb71752aa428e43ccdd95af760c'
SERVICE = '劳动人事争议仲裁申请'
BLANKS = {
    'application': ('21B24C52D2AC3EB007AA7F55CE0AAD79', '0eab0f0a2ee98e8ca1e4d2fa5469fcb5738d0c0f24671345bebc8394099ae41f'),
    'evidence': ('872DFF90A8578907CB947197A784AA48', '4965bf3d2bf76404ab190d1e47d2fa250c8cb8bd93762949df94be38b2d4ca8c'),
}
SPECS = {
    'fuyang': dict(code='330111', district='富阳区', local='31ce6363-7379-40c0-8454-f0715785e193',
                   task='11330183002506152E4002014302002', authority='区人社局', phone='0571-63342646', window_phone='0571-63342646',
                   place='富阳区达夫路121号矛调中心窗口',
                   manifest='b260a7385135e471a05d102c052e3eabff1dd8c67f611b94b4d65b0d30e3e0e7',
                   white='ff1122c98aec38fcdc63d712827ac5c4289368fa712e2b229c3c3ca528c2cb20',
                   materials=['261717A3296E53D46953524DA7562D2C', 'BAFCD2B89394AA02DD2914CA0720DAE0', 'B55E6CC7851FBC4F02F8DB9F4595191B', 'DB7F21FA300654EF781D40C287EC5C75'],
                   kinds={'application': 'BAFCD2B89394AA02DD2914CA0720DAE0', 'evidence': 'B55E6CC7851FBC4F02F8DB9F4595191B'}),
    'jiande': dict(code='330182', district='建德市', local='ed48d6e4-e7ee-433d-97aa-b1352d1fb341',
                   task='1133018205366747424002014302002', authority='市人力资源和社会保障局', phone='0571-64721936', window_phone='0571-64721626',
                   place='建德市新安江街道广兴北路29号劳动保障大楼1层101室',
                   manifest='20f7c2e7b8efa5058c39439f1079e309da2c3922aded10275a4514bf553a542c',
                   white='0bf72f90af6a7ffcc7ee37ac9bec884b1b76de9f41cb2e12f5a6dc560dd76a7d',
                   materials=['6391B4BE14B46C98766FA4275B7BBDD1', '1B8606497E41B76124A672B3BE9F36DB', '46ACCB2DF05B1F8C1F581602D1F346E8', 'F7FF66119F2DDB743C0370239FABF58E', '325D68FBD236A41F682A9A3EAFE440AA'],
                   kinds={'application': '46ACCB2DF05B1F8C1F581602D1F346E8', 'evidence': 'F7FF66119F2DDB743C0370239FABF58E'}),
}
DATA_FIELDS = {'application': APPLICATION_FIELDS, 'evidence': EVIDENCE_FIELDS, 'evidence_item': EVIDENCE_ITEM_FIELDS}
LIMITS = deepcopy(SLOT_LIMITS)
CAPABILITY_FIELDS = frozenset({'can_generate', 'document_types', 'adapter', 'content_kind', 'generation_label'})
GUIDE_METADATA = frozenset({'version', 'content_hash', 'region_id', 'published', 'publication_status', 'candidate_web_qa', 'can_generate', 'warnings'})
OFFLINE_NOTICE = '仅本区自身空白文书完成离线填写；申请默认且仅支持劳动者区段，单位申请未适配。版式及候选网页验收另行记录；不代表已发布。'
PUBLISHED_NOTICE = '本区自身空白文书已通过填写、版式及网页验收；申请仅支持劳动者区段，单位申请未适配。'
FAILURES = (OSError, ValueError, KeyError, TypeError, StopIteration, AttributeError)


def _root(root):
    return Path(ROOT if root is None else root).resolve()


def _path(root, relative):
    relative = Path(relative)
    if relative.is_absolute() or relative.drive or '..' in relative.parts:
        raise ValueError('来源路径必须位于发行根内')
    resolved = (root / relative).resolve()
    if not resolved.is_relative_to(root) or resolved.relative_to(root).as_posix() != relative.as_posix():
        raise ValueError('来源路径越界或符号链接替换')
    return resolved


class Evidence:
    """Verify the entire fixed 56-file public list plus its own fixed bytes."""
    def __init__(self, root=None):
        self.root = _root(root)
        white = _path(self.root, REPORT + '/PUBLIC-WHITELIST.json').read_bytes()
        if sha256(white).hexdigest() != WHITELIST_SHA256:
            raise ValueError('浙江E固定公开来源白名单已变化')
        records = json.loads(white)['files']
        self.files = {r['path']: {'sha256': r['sha256'], 'bytes': r['bytes']} for r in records}
        if len(records) != 56 or len(self.files) != 56:
            raise ValueError('浙江E来源56文件链条不完整')
        for relative in self.files:
            if not relative.startswith(REPORT + '/'):
                raise ValueError('浙江E来源清单越界')
            self.read(relative.removeprefix(REPORT + '/'))
        for name, spec in SPECS.items():
            manifest = self.read(name + '/SOURCE-MANIFEST.json')
            local_white = self.read(name + '/PUBLIC-WHITELIST.json')
            if sha256(manifest).hexdigest() != spec['manifest'] or sha256(local_white).hexdigest() != spec['white']:
                raise ValueError('浙江E本区来源封存链已变化')
            m, w = json.loads(manifest), json.loads(local_white)
            if len(m['files']) != 25 or len(w['files']) != 26 or w['manifest_sha256'] != spec['manifest']:
                raise ValueError('浙江E本区来源文件不完整')
            for r in m['files'] + w['files']:
                if self.files.get(r['path']) != {'sha256': r['sha256'], 'bytes': r['bytes']}:
                    raise ValueError('浙江E本区与整体来源证明不一致')

    def read(self, relative):
        key = REPORT + '/' + relative
        if key not in self.files:
            raise ValueError('不在浙江E固定公开来源内')
        body = _path(self.root, key).read_bytes()
        if self.files[key] != {'sha256': sha256(body).hexdigest(), 'bytes': len(body)}:
            raise ValueError('浙江E封存来源字节已变化：' + relative)
        return body

    def load(self, relative):
        return json.loads(self.read(relative).decode('utf-8-sig'))


def parse_capture(name, root=None):
    """Rebuild Fuyang/Jiande from only their own pinned public project root."""
    if name not in SPECS:
        raise ValueError('只处理富阳和建德自身E来源')
    evidence = Evidence(root)
    spec = SPECS[name]
    source = 'https://www.zjzwfw.gov.cn/zjservice-fe/#/workguide?localInnerCode=' + spec['local']
    guide = evidence.load(name + '/guide.json')
    api = evidence.load(name + '/item-response.json')['data']['item']
    if (guide['province'], guide['city'], guide['district'], guide['region_code'], guide['source_url'], guide['source_id'], guide['service_name'], guide['service_code'], guide['authority']) != (
            'zhejiang', '杭州市', spec['district'], spec['code'], source, digest(source)[:24], SERVICE, spec['task'], spec['authority']):
        raise ValueError('浙江E自身指南身份不匹配')
    if (api['jurisCode'], api['dcode'], api['taskCode'], api['name'], api['qlDept'], api['askphone'], api['addresssList'][0]['addressPhone']) != (
            spec['code'], spec['local'], spec['task'], SERVICE, spec['authority'], spec['phone'], spec['window_phone']):
        raise ValueError('浙江E自身地区、事项、机关或电话不匹配')
    if len(api) != 252 or guide['raw_business_fields'] != api:
        raise ValueError('浙江E完整业务字段缺失')
    search = evidence.load(name + '/search.json')
    if 'websiteid=' + spec['code'] + '000000' not in search['url'] or spec['district'] not in search['text']:
        raise ValueError('浙江E自身区域目录不匹配')
    pages = {section: evidence.load(name + '/' + section + '-dom.json') for section in ['guide', 'basic', 'faq', 'conditions', 'materials', 'flow', 'window', 'law']}
    if any(page['url'] != source or SERVICE not in page['text'] for page in pages.values()):
        raise ValueError('浙江E自身原生栏目不匹配')
    if spec['place'] not in pages['guide']['text'] or spec['place'] != api['addresssList'][0]['address']:
        raise ValueError('浙江E自身窗口地点不匹配')
    materials = api['materialList']
    if [m['impleCode'] for m in materials] != spec['materials'] or [m['raw'] for m in guide['materials']] != materials:
        raise ValueError('浙江E自身材料集合不匹配')
    originals = evidence.load(name + '/originals.json')
    for kind, material_id in spec['kinds'].items():
        annex, checksum = BLANKS[kind]
        material = next(m for m in materials if m['impleCode'] == material_id)
        template = next(t for t in guide['templates'] if t['id'] == material_id + ':blank')
        url = 'https://portal.zjzwfw.gov.cn/media/oss/FILE/META_DATA/' + annex + '.docx'
        relative = name + '/originals/blank/' + annex + '.docx'
        body = evidence.read(relative)
        metadata = evidence.load(relative + '.source.json')
        own = next(r for r in originals if r['material_id'] == material_id and r['official_kind'] == 'blank')
        if (material['blankForm'] != url.removeprefix('https:') or
            (template['material_id'], template['region_code'], template['source_url'], template['download_url'], template['file'], template['sha256'], template['bytes'], template['official_kind'], template['official_annex_id']) !=
            (material_id, spec['code'], source, url, REPORT + '/' + relative, checksum, len(body), 'blank', annex) or
            metadata != own or
            (metadata['source_url'], metadata['region_code'], metadata['local_inner_code'], metadata['service_code'], metadata['download_url'], metadata['sha256'], metadata['http_status'], metadata['native_control_clicked']) !=
            (source, spec['code'], spec['local'], spec['task'], url, checksum, 200, True) or sha256(body).hexdigest() != checksum):
            raise ValueError('浙江E本区材料、附件、下载与原件未完整绑定')
        document = Document(BytesIO(body))
        expected = [(23, 5), (23, 11)] if kind == 'application' else [(13, 6)]
        if [(len(t.rows), len(t.columns)) for t in document.tables] != expected:
            raise ValueError('浙江E自身原件槽位结构不匹配')
        template.update(observed_document_kind=kind, file_type='docx', applicant_scope='worker' if kind == 'application' else 'evidence_list')
    samples = [t for t in guide['templates'] if t['official_kind'] == 'sample']
    if len(samples) != (1 if name == 'jiande' else 0) or any(t['file'] is not None or t['can_generate'] or t['document_types'] for t in samples):
        raise ValueError('浙江E示例不得作为公开原件或填写模板')
    boundary = evidence.load(name + '/online-boundary.json')
    if (boundary['origin'], boundary['path'], boundary['login_visible'], boundary['captcha_visible'], boundary['authenticated'], boundary['submitted'], boundary['query_saved']) != (
            'https://user.zjzwfw.gov.cn', '/pc/login', True, True, False, False, False):
        raise ValueError('浙江E自身原生网办认证边界不匹配')
    guide['sections']['常见问题'] = '\n\n'.join(q['question'] + '\n' + q['answer'] for q in api['qaList'])
    guide['sections']['法定依据'] = api['fdconditions']
    guide['service_ids'] = {'localInnerCode': spec['local'], 'taskCode': spec['task']}
    guide['evidence'].update(manifest_sha256=spec['manifest'], public_whitelist_sha256=WHITELIST_SHA256, own_public_whitelist_sha256=spec['white'])
    guide['channel']['entry_check'] = {'url': boundary['origin'] + boundary['path'], 'status': 'auth_required', 'official_link': True, 'checked_at': boundary['captured_at'], 'submitted': False, 'post_login_flow': 'unknown'}
    guide['warnings'] = [('封存来源采集阶段记录：' + w) if w.startswith('本批只采集来源') else w for w in guide['warnings']]
    if name == 'fuyang':
        guide['warnings'].append('自身材料清单本次列4项未列劳动者身份证，不构成免交身份证结论；本区自身空白申请书仍明确要求提交身份证复印件或其他身份证件。')
    guide['warnings'].append('本区自身申请书户口原文含重复的“外省非农业户口”，不推断改为农业；同名项不自动勾选，留空后本人核对。')
    return guide


def _capability(template, kind):
    return {**template, 'can_generate': True, 'document_types': [kind], 'adapter': ADAPTER,
            'content_kind': 'filled_official_blank', 'generation_label': '依据本区自身官方空白表格填写（申请仅劳动者部分）'}


def _kind(template):
    return template.get('observed_document_kind') if template.get('official_kind') == 'blank' else None


def _template_matches(actual, canonical):
    if not isinstance(actual, dict):
        return False
    kind = _kind(canonical)
    allowed = _capability(canonical, kind) if kind else canonical
    if set(actual) - set(allowed) or not set(canonical).issubset(actual):
        return False
    return all(actual.get(k) == v or (kind and k in CAPABILITY_FIELDS and actual.get(k) == allowed[k]) for k, v in canonical.items()) and all(actual[k] == allowed[k] for k in set(actual) - set(canonical))


def _validate_guide(guide, root=None):
    if not isinstance(guide, dict):
        raise ValueError('浙江E指南必须为自身来源对象')
    name = next((n for n, s in SPECS.items() if s['code'] == guide.get('region_code')), None)
    if name is None:
        raise ValueError('浙江E指南不属于本区')
    canonical = parse_capture(name, root)
    if set(guide) - set(canonical) - GUIDE_METADATA or any(guide.get(k) != v for k, v in canonical.items() if k not in GUIDE_METADATA | {'templates'}):
        raise ValueError('浙江E指南事实、材料或来源不匹配')
    templates = guide.get('templates')
    if not isinstance(templates, list) or len(templates) != len(canonical['templates']) or any(not isinstance(t, dict) for t in templates):
        raise ValueError('浙江E本区附件集合不匹配')
    if len({t.get('id') for t in templates}) != len(templates):
        raise ValueError('浙江E附件重复')
    for own in canonical['templates']:
        if not _template_matches(next((t for t in templates if t.get('id') == own['id']), None), own):
            raise ValueError('浙江E本区附件身份不匹配')
    return canonical


def blank_binding(guide, template, root=None):
    try:
        canonical = _validate_guide(guide, root)
        own = next(t for t in canonical['templates'] if t['id'] == template.get('id'))
        return _kind(own) if _template_matches(template, own) else None
    except FAILURES:
        return None


def with_capabilities(guide, root=None):
    canonical = _validate_guide(guide, root)
    kinds = {t['id']: _kind(t) for t in canonical['templates']}
    result = deepcopy(guide)
    result['templates'] = [_capability(t, kinds[t['id']]) if kinds[t['id']] else deepcopy(t) for t in result['templates']]
    result['can_generate'] = True
    incoming = result.get('warnings', [])
    if not isinstance(incoming, list) or any(not isinstance(w, str) for w in incoming):
        raise ValueError('浙江E警告字段必须为文字列表')
    warnings = [w for w in canonical['warnings'] + incoming if w not in {OFFLINE_NOTICE, PUBLISHED_NOTICE}]
    result['warnings'] = list(dict.fromkeys(warnings))
    result['warnings'].append(PUBLISHED_NOTICE if result.get('published') is True and result.get('candidate_web_qa') == 'PASS' else OFFLINE_NOTICE)
    return result


def checked_adapter(guide, template, kind, root=None):
    if not isinstance(kind, str) or kind not in {'application', 'evidence'} or blank_binding(guide, template, root) != kind or not isinstance(template, dict) or template.get('adapter') != ADAPTER or template.get('can_generate') is not True or template.get('document_types') != [kind]:
        raise ValueError('仅支持浙江E本区严格绑定的真实空白文书')
    return ADAPTER


def public_evidence(guide, root):
    """Explicit public project root required; no fallback to original checkout."""
    try:
        return None if root is None else deepcopy(_validate_guide(guide, root)['evidence'])
    except FAILURES:
        return None


def approved_report_attachment(guide, template, root):
    try:
        if root is None:
            return False
        canonical = _validate_guide(guide, root)
        own = next(t for t in canonical['templates'] if t['id'] == template.get('id'))
        return bool(_kind(own)) and _template_matches(template, own)
    except FAILURES:
        return False


def fill_document(guide, template, kind, path, data, *, root=None, applicant_mode='worker'):
    checked_adapter(guide, template, kind, root)
    if applicant_mode != 'worker':
        raise ValueError('默认且仅支持劳动者申请，单位区段未适配')
    expected = _path(_root(root), template['file'])
    if Path(path).resolve() != expected:
        raise ValueError('不得传入其他地区或任意同SHA原件路径')
    original = expected.read_bytes()
    if sha256(original).hexdigest() != template['sha256']:
        raise ValueError('浙江E自身原件SHA已变化')
    _validate_data(kind, data)
    document = Document(BytesIO(original))
    (_application if kind == 'application' else _evidence)(document, deepcopy(data))
    return PreservedPackageDocument(document, original)
