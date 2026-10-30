"""Chunan own-source forms with separate staging and publication evidence.

The explicit release root must carry the entire pinned public Chunan capture.
Only pure slot operations are shared with earlier Zhejiang adapters; no other
locality's guide, attachment path, identity or requirements are inherited.
"""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path

from docx import Document
from regional_documents import PreservedPackageDocument
from regional_zhejiang_oct29d import _application, _evidence, _validate_data

ROOT = Path(__file__).resolve().parents[1]
REPORT = 'reports/nationwide-verification-20261029-f/zhejiang/chunan'
ADAPTER = 'zhejiang_chunan_own_worker_oct30_v1'
STAGE_RECEIPT = 'reports/production-readiness-20261030/localities/chunan/STAGE-APPROVAL.json'
STAGE_SHA256 = 'f6a6de4522b1af6ed509fcfeb465ba90696c833b0794e677203704fb9fb57876'
PUBLICATION_RECEIPT = 'reports/production-readiness-20261030/localities/chunan/FINAL-BROWSER-QA.json'
PUBLICATION_SHA256 = '442a6faa2aeae6e15a3b6177e098d42d8545c377d3d8344070dbab8af244671d'
WHITELIST_SHA256 = '689740ac4a5827f926beafd708840400fbe72f062a52707471666a50e521e292'
MANIFEST_SHA256 = '78ded402d48d559ece5032e3ca6eef6c775073dc48e0d4c2cd97bc7f82b8e1a5'
LOCAL = 'ee47fff8-558c-4a85-88e9-8a28d2bddbdc'
SOURCE = 'https://www.zjzwfw.gov.cn/zjservice-fe/#/workguide?localInnerCode=' + LOCAL
TASK = '11330127593095219M4002014302002'
BLANKS = {
    'application': dict(material='D508BEE2A05365A7C6C3D2AC490BBD24', annex='21B24C52D2AC3EB007AA7F55CE0AAD79',
                        sha256='0eab0f0a2ee98e8ca1e4d2fa5469fcb5738d0c0f24671345bebc8394099ae41f', structure=[(23, 5), (23, 11)]),
    'evidence': dict(material='FBC9E9D8D9254A4AB07013C05B516870', annex='872DFF90A8578907CB947197A784AA48',
                     sha256='4965bf3d2bf76404ab190d1e47d2fa250c8cb8bd93762949df94be38b2d4ca8c', structure=[(13, 6)]),
}
MATERIALS = ['BEDCFDA4DA66F2785CABA5A77B712D65', 'A76581A1B00195D1F1A6508F63D2E9BC',
             BLANKS['application']['material'], BLANKS['evidence']['material'], '644C2B7C6FF6701766CAAF7272B6785B']


def _path(root, relative):
    if root is None or isinstance(root, str) and not root.strip():
        raise ValueError('淳安离线适配必须提供显式公开根目录')
    root = Path(root).resolve()
    relative = Path(relative)
    if relative.is_absolute() or relative.drive or '..' in relative.parts:
        raise ValueError('淳安来源路径越界')
    resolved = (root / relative).resolve()
    if not resolved.is_relative_to(root) or resolved.relative_to(root).as_posix() != relative.as_posix():
        raise ValueError('淳安来源路径越界或符号链接替换')
    return resolved


class Evidence:
    def __init__(self, root):
        self.root = root
        white = _path(root, REPORT + '/PUBLIC-WHITELIST.json').read_bytes()
        if sha256(white).hexdigest() != WHITELIST_SHA256:
            raise ValueError('淳安固定公开白名单已变化')
        records = json.loads(white)['files']
        self.files = {r['path']: {'sha256': r['sha256'], 'bytes': r['bytes']} for r in records}
        if len(records) != 30 or len(self.files) != 30:
            raise ValueError('淳安30项公开来源链不完整')
        for relative in self.files:
            if not relative.startswith(REPORT + '/'):
                raise ValueError('淳安来源清单含其他地区')
            self.read(relative.removeprefix(REPORT + '/'))
        manifest = self.read('SOURCE-MANIFEST.json')
        if sha256(manifest).hexdigest() != MANIFEST_SHA256:
            raise ValueError('淳安来源清单已变化')
        for r in json.loads(manifest)['files']:
            if self.files.get(r['path']) != {'sha256': r['sha256'], 'bytes': r['bytes']}:
                raise ValueError('淳安封存清单与公开清单不一致')

    def read(self, relative):
        key = REPORT + '/' + relative
        if key not in self.files:
            raise ValueError('不在淳安自身公开清单内')
        raw = _path(self.root, key).read_bytes()
        if self.files[key] != {'sha256': sha256(raw).hexdigest(), 'bytes': len(raw)}:
            raise ValueError('淳安封存来源字节已变化：' + relative)
        return raw

    def load(self, relative):
        return json.loads(self.read(relative))


def parse_capture(*, root):
    """Verify own frozen source; returned guide keeps every capability disabled."""
    e = Evidence(root)
    guide = e.load('guide.json')
    api = e.load('item-response.json')['data']['item']
    if (guide['province'], guide['city'], guide['district'], guide['region_code'], guide['source_id'],
        guide['source_url'], guide['service_code'], guide['authority']) != (
            'zhejiang', '杭州市', '淳安县', '330127', '3fb683d473fe253a3a117d46', SOURCE, TASK, '人社局'):
        raise ValueError('淳安自身指南身份不匹配')
    if (api['jurisCode'], api['dcode'], api['taskCode'], api['qlDept']) != ('330127', LOCAL, TASK, '人社局'):
        raise ValueError('淳安自身业务地区或事项不匹配')
    if len(api) != 252 or guide['raw_business_fields'] != api or [m['impleCode'] for m in api['materialList']] != MATERIALS:
        raise ValueError('淳安完整业务字段或五项材料不匹配')
    if [m['raw'] for m in guide['materials']] != api['materialList']:
        raise ValueError('淳安材料事实不匹配')
    if guide['fields']['窗口联系电话']['raw_value'] != '' or api['addresssList'][0]['addressPhone'] != '':
        raise ValueError('淳安空白窗口电话必须保留')
    originals = e.load('originals.json')
    for kind, spec in BLANKS.items():
        t = next(t for t in guide['templates'] if t['id'] == spec['material'] + ':blank')
        relative = 'originals/blank/' + spec['annex'] + '.docx'
        raw = e.read(relative)
        material = next(m for m in api['materialList'] if m['impleCode'] == spec['material'])
        own = next(o for o in originals if o['material_id'] == spec['material'] and o['official_kind'] == 'blank')
        url = 'https://portal.zjzwfw.gov.cn/media/oss/FILE/META_DATA/' + spec['annex'] + '.docx'
        if (t['region_code'], t['material_id'], t['official_annex_id'], t['source_url'], t['download_url'], t['file'], t['sha256']) != (
                '330127', spec['material'], spec['annex'], SOURCE, url, REPORT + '/' + relative, spec['sha256']):
            raise ValueError('淳安自身材料附件绑定不完整')
        if (material['blankForm'] != url.removeprefix('https:') or sha256(raw).hexdigest() != spec['sha256'] or
            own != e.load(relative + '.source.json') or
            (own['region_code'], own['local_inner_code'], own['service_code'], own['download_url'], own['sha256'], own['http_status']) !=
                ('330127', LOCAL, TASK, url, spec['sha256'], 200)):
            raise ValueError('淳安自身附件下载证明不匹配')
        if [(len(t.rows), len(t.columns)) for t in Document(BytesIO(raw)).tables] != spec['structure']:
            raise ValueError('淳安原件槽位结构不匹配')
    if guide['published'] or guide['can_generate'] or any(t['can_generate'] or t['document_types'] for t in guide['templates']):
        raise ValueError('淳安离线来源不得自动开放生成或发布')
    return guide


def fill_offline_document(guide, template, kind, data, *, root, applicant_mode='worker'):
    """Produce a QA copy, without changing the source-only guide or registry."""
    canonical = parse_capture(root=root)
    if guide != canonical:
        raise ValueError('淳安指南事实或来源不匹配')
    if not isinstance(kind, str) or kind not in BLANKS or applicant_mode != 'worker':
        raise ValueError('仅支持淳安自身空白表劳动者区段及证据清单')
    own = next(t for t in canonical['templates'] if t['id'] == BLANKS[kind]['material'] + ':blank')
    if template != own:
        raise ValueError('淳安本区附件身份、材料或文书类型不匹配')
    validated = deepcopy(data)
    _validate_data(kind, validated)
    original = _path(root, own['file']).read_bytes()
    document = Document(BytesIO(original))
    (_application if kind == 'application' else _evidence)(document, validated)
    return PreservedPackageDocument(document, original)


def _stage_receipt(root):
    raw = _path(root, STAGE_RECEIPT).read_bytes()
    if sha256(raw).hexdigest() != STAGE_SHA256:
        raise ValueError('淳安固定版式验收收据已变化')
    receipt = json.loads(raw)
    if (receipt['region_code'], receipt['source_id'], receipt['public_whitelist_sha256'],
        receipt['source_manifest_sha256'], receipt['can_publish'], receipt['visual_review']['status']) != (
            '330127', '3fb683d473fe253a3a117d46', WHITELIST_SHA256, MANIFEST_SHA256, False, 'PASS_FOR_THIS_SYNTHETIC_FIXTURE'):
        raise ValueError('淳安自身版式验收范围不匹配')
    for record in receipt['files']:
        if not record['path'].startswith('reports/production-readiness-20261030/localities/chunan/'):
            raise ValueError('淳安验收来源路径越界')
        body = _path(root, record['path']).read_bytes()
        if len(body) != record['bytes'] or sha256(body).hexdigest() != record['sha256']:
            raise ValueError('淳安版式验收证据字节已变化')
    return {'receipt': STAGE_RECEIPT, 'sha256': STAGE_SHA256}


def prepare_candidate(*, root):
    """Explicitly derive a candidate after own Word QA; never mutate sources/DB."""
    guide = parse_capture(root=root)
    guide['candidate_adapter_evidence'] = _stage_receipt(root)
    guide['can_generate'] = True
    guide['publication_status'] = 'staged_own_word_qa'
    for template in guide['templates']:
        kind = next((k for k, s in BLANKS.items() if template['id'] == s['material'] + ':blank'), None)
        if kind:
            template.update(can_generate=True, document_types=[kind], adapter=ADAPTER,
                            content_kind='filled_official_blank', applicant_scope='worker' if kind == 'application' else 'evidence_list',
                            generation_label='依据淳安县自身官方空白表格填写（申请仅劳动者部分）')
    guide['warnings'].append('淳安县自身原件已完成离线填写及本次虚构样例逐页Word版式检查；仅为候选，网页生成验收和发布尚未完成。')
    guide['warnings'].append('自身申请书户籍类型原文重复“外省非农业户口”，不擅自更正或自动勾选；申请人手写签名必须留空。')
    return guide


_STORE_METADATA = frozenset({'version', 'content_hash', 'region_id'})
_FAILURES = (OSError, ValueError, KeyError, TypeError, StopIteration, AttributeError)


def _matches(guide, canonical):
    if not isinstance(guide, dict) or set(guide) - set(canonical) - _STORE_METADATA:
        return False
    return all(_same(guide.get(k), v) for k, v in canonical.items())


def _same(actual, expected):
    # Python equates True with 1; source and lifecycle values require JSON types.
    if type(actual) is not type(expected):
        return False
    if isinstance(expected, dict):
        return actual.keys() == expected.keys() and all(_same(actual[k], v) for k, v in expected.items())
    if isinstance(expected, list):
        return len(actual) == len(expected) and all(_same(a, b) for a, b in zip(actual, expected))
    return actual == expected


def _candidate(guide, root):
    canonical = prepare_candidate(root=root)
    if not _matches(guide, canonical):
        raise ValueError('淳安候选事实、附件或固定版式收据不匹配')
    return canonical


def _publication_receipt(root):
    raw = _path(root, PUBLICATION_RECEIPT).read_bytes()
    if sha256(raw).hexdigest() != PUBLICATION_SHA256:
        raise ValueError('淳安最终网页下载及Word验收收据已变化')
    receipt = json.loads(raw)
    if (receipt['region_code'], receipt['source_id'], receipt['service_code'],
        receipt['source_manifest_sha256'], receipt['public_whitelist_sha256'], receipt['stage_receipt_sha256'],
        receipt['native_click_confirmation_download_both_kinds'], receipt['visual_review']['status'],
        receipt['render_pages'], len(receipt['files'])) != (
            '330127', '3fb683d473fe253a3a117d46', TASK, MANIFEST_SHA256, WHITELIST_SHA256, STAGE_SHA256,
            True, 'PASS', {'application': 3, 'evidence': 1}, 11):
        raise ValueError('淳安最终验收收据不属于自身事项及两个文书')
    for record in receipt['files']:
        if not record['path'].startswith('reports/production-readiness-20261030/localities/chunan/browser-generation/'):
            raise ValueError('淳安最终验收路径越界')
        body = _path(root, record['path']).read_bytes()
        if len(body) != record['bytes'] or sha256(body).hexdigest() != record['sha256']:
            raise ValueError('淳安实际网页下载或逐页Word验收字节已变化')
    return {'receipt': PUBLICATION_RECEIPT, 'sha256': PUBLICATION_SHA256}


def _publication_projection(root):
    # Rebuild from both complete source/stage chains before final browser proof.
    guide = prepare_candidate(root=root)
    guide['publication_adapter_evidence'] = _publication_receipt(root)
    guide.update(published=True, publication_status='published', candidate_web_qa='PASS')
    guide['warnings'] = [w for w in guide['warnings'] if not w.startswith('淳安县自身原件已完成离线填写及本次虚构样例')]
    guide['warnings'].append('淳安县自身两份空白文书已完成真实网页填写下载及下载件逐页Word检查；申请仅支持劳动者区段，生成后仍须本人逐项核对并手写签名。')
    return guide


def with_publication_evidence(guide, *, root):
    """Return an attested in-memory projection; caller separately approves/writes.

    This function performs no publication, database, seed or frozen-file write.
    Boolean lifecycle flags cannot substitute for either fixed evidence chain.
    """
    _candidate(guide, root)
    projected = _publication_projection(root)
    for key in _STORE_METADATA & set(guide):
        projected[key] = deepcopy(guide[key])
    return projected


def _bound_guide(guide, root):
    if isinstance(guide, dict) and (guide.get('published') is True or 'publication_adapter_evidence' in guide):
        canonical = _publication_projection(root)
        if not _matches(guide, canonical):
            raise ValueError('淳安发布状态缺少自身固定网页下载及Word验收绑定')
        return canonical
    return _candidate(guide, root)


def with_capabilities(guide, root=None):
    """Keep raw captures disabled; preserve only explicit own attested copies."""
    root = ROOT if root is None else root
    if _matches(guide, parse_capture(root=root)):
        return deepcopy(guide)
    _bound_guide(guide, root)
    return deepcopy(guide)


def checked_adapter(guide, template, kind, root=None):
    canonical = _bound_guide(guide, ROOT if root is None else root)
    if not isinstance(kind, str) or kind not in BLANKS:
        raise ValueError('淳安文书种类不匹配')
    own = next(t for t in canonical['templates'] if t['id'] == BLANKS[kind]['material'] + ':blank')
    if not _same(template, own):
        raise ValueError('淳安自身附件与请求种类不匹配')
    return ADAPTER


def public_evidence(guide, root):
    try:
        source = parse_capture(root=root)
        canonical = source if _matches(guide, source) else _bound_guide(guide, root)
        return deepcopy(canonical['evidence'])
    except _FAILURES:
        return None


def approved_report_attachment(guide, template, root):
    try:
        source = parse_capture(root=root)
        canonical = source if _matches(guide, source) else _bound_guide(guide, root)
        return any(_same(template, t) and t['official_kind'] == 'blank' for t in canonical['templates'])
    except _FAILURES:
        return False


def fill_document(guide, template, kind, path, data, *, root, applicant_mode='worker'):
    checked_adapter(guide, template, kind, root=root)
    original_guide = parse_capture(root=root)
    original_template = next(t for t in original_guide['templates'] if t['id'] == template['id'])
    if Path(path).resolve() != _path(root, original_template['file']):
        raise ValueError('不得传入其他地区或同SHA的任意原件路径')
    return fill_offline_document(original_guide, original_template, kind, data, root=root, applicant_mode=applicant_mode)
