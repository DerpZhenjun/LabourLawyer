"""Own-source Shangcheng binding and field/package behavior, with private-proof skips."""
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import sys,tempfile,unittest
from unittest.mock import patch
from zipfile import ZipFile
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[2];sys.path.insert(0,str(ROOT/'backend'))
from regional_zhejiang_more import parse_shangcheng_capture
try:
    import regional_shangcheng_documents as adapter
except ModuleNotFoundError:
    adapter=None
REPORT=ROOT/'reports/nationwide-localization-20261028-next/zhejiang'

class ShangchengDocumentsTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        manifest=REPORT/'SOURCE-MANIFEST.json'
        if not manifest.is_file():raise unittest.SkipTest('Private sealed own capture unavailable; SSO/raw browser proof is not published')
        inventory=json.loads(manifest.read_text(encoding='utf8'))['files']
        if any(not (REPORT/p).is_file() for p in inventory):raise unittest.SkipTest('Private sealed capture bundle incomplete')
        originals=json.loads((REPORT/'shangcheng/originals.json').read_text(encoding='utf8'))
        if any(not (ROOT/t['canonical_file']).is_file() for t in originals):raise unittest.SkipTest('Own canonical official originals unavailable')
        cls.raw=parse_shangcheng_capture(REPORT,manifest)
    def setUp(self):self.assertIsNotNone(adapter,'Shangcheng own-source document adapter is not implemented')
    def selected(self,kind):
        guide=adapter.with_shangcheng_capabilities(self.raw)
        return guide,next(t for t in guide['templates'] if t['document_types']==[kind])
    def test_only_this_own_source_enables_two_blank_kinds(self):
        # Break: enabling via province or shared SHA instead of the own captured locality.
        guide=adapter.with_shangcheng_capabilities(self.raw)
        self.assertTrue(all(not t['can_generate'] for t in self.raw['templates']))
        self.assertEqual({tuple(t['document_types']) for t in guide['templates']},{('application',),('evidence',)})
        self.assertTrue(all(t['adapter']==adapter.ADAPTER for t in guide['templates']))
    def test_mutated_source_identity_or_annex_disables(self):
        # Break: matching a copied form while ignoring any part of own joint identity.
        guide,template=self.selected('application')
        for target,key,value in [('guide','source_id','wrong'),('guide','source_url','https://example.com/'),
            ('guide','service_code','8f1bc917-6583-4541-be13-231de0b91d0b'),('guide','authority','杭州市人力资源和社会保障局'),
            ('guide','region_code','330110'),('guide','district','余杭区'),('guide','city','宁波市'),
            ('template','official_annex_id','872DFF90A8578907CB947197A784AA48'),('template','sha256','0'*64),
            ('template','file','templates/_docx/_regional/330110/21B24C52D2AC3EB007AA7F55CE0AAD79-劳动人事争议仲裁申请书.docx'),
            ('template','download_url','https://example.com/application.docx'),('template','official_kind','sample')]:
            with self.subTest(target=target,key=key):
                g,t=deepcopy(guide),deepcopy(template);(g if target=='guide' else t)[key]=value
                self.assertIsNone(adapter.shangcheng_blank_binding(g,t))
                g['templates']=[t];self.assertFalse(adapter.with_shangcheng_capabilities(g)['templates'][0]['can_generate'])
    def test_worker_fields_and_human_signature_slots_are_separate(self):
        # Break: wrong field cells, employer form leakage or automatic handwritten signature.
        guide,t=self.selected('application')
        doc=adapter.fill_shangcheng_document(guide,t,'application',ROOT/t['file'],{
            'applicantName':'上城虚构甲','birthDate':'1990-01-02','gender':'女','nationality':'汉族','idNumber':'TEST-ID',
            'phone':'TEST-PHONE','idAddress':'虚构测试住址','respondentName':'上城虚构单位乙','socialCreditCode':'TEST-CREDIT',
            'registeredAddress':'虚构注册地址','mailingAddress':'虚构通讯地址','workAddress':'虚构工作地址',
            'legalRepName':'虚构负责人丙','legalRepPhone':'TEST-REP','legalRepPosition':'测试职务','companyContact':'虚构联系人丁',
            'companyContactPhone':'TEST-CONTACT','committeeName':'杭州市上城区劳动人事争议仲裁委员会',
            'requests':['测试工资100元'],'factsReason':'全部数据均为离线软件测试虚构内容。','copyCount':2,'applicationDate':'2026-10-28'})
        self.assertEqual(len(doc.tables),1);table=doc.tables[0]
        for r,c,expected in [(1,1,'上城虚构甲'),(1,4,'1990年1月2日'),(3,1,'TEST-ID'),(3,4,'TEST-PHONE'),
            (5,2,'虚构测试住址'),(7,2,'上城虚构单位乙'),(8,2,'TEST-CREDIT'),(9,2,'虚构注册地址'),
            (10,2,'虚构通讯地址'),(11,2,'虚构工作地址'),(12,2,'虚构负责人丙'),(14,2,'虚构联系人丁')]:
            self.assertEqual(table.rows[r].cells[c].text,expected)
        footer=table.rows[-1].cells[0].paragraphs
        self.assertIn('杭州市上城区',footer[2].text);self.assertNotIn('上城虚构甲',footer[3].text)
        self.assertIn('2026年',footer[4].text)
    def test_evidence_uses_explicit_profile_and_keeps_signature_blank(self):
        # Break: selecting profile from numeric regional folder or writing submitter signature.
        guide,t=self.selected('evidence')
        doc=adapter.fill_shangcheng_document(guide,t,'evidence',ROOT/t['file'],{
            'committeeName':'杭州市上城区劳动人事争议仲裁委员会','caseNumber':'测试上城号','copyCount':2,
            'submitterName':'不可自动签名','submitDate':'2026-10-28',
            'evidenceList':[{'name':'测试工资记录','pageCount':'2页','source':'虚构保存来源','purpose':'仅用于软件验证'}]})
        table=doc.tables[0]
        self.assertEqual([table.rows[2].cells[i].text for i in [1,3,4,5]],['测试工资记录','2','虚构保存来源','仅用于软件验证'])
        self.assertNotIn('不可自动签名','\n'.join(c.text for r in table.rows for c in r.cells))
    def test_preserved_package_changes_only_main_document_part(self):
        # Break: losing original OOXML parts when saving the filled document.
        for kind in ['application','evidence']:
            guide,t=self.selected(kind);path=ROOT/t['file']
            doc=adapter.fill_shangcheng_document(guide,t,kind,path,{'applicantName':'测试','requests':['测试请求'],'evidenceList':[{'name':'测试证据'}]})
            memory=BytesIO();doc.save(memory)
            with ZipFile(path) as a,ZipFile(memory) as b:
                self.assertEqual(set(a.namelist()),set(b.namelist()))
                self.assertEqual([n for n in a.namelist() if a.read(n)!=b.read(n)],['word/document.xml'])
            self.assertEqual(path.read_bytes(),doc.original)
    def test_wrong_path_kind_or_missing_private_proof_is_rejected(self):
        # Break: accepting same-byte alternate path, a missing kind or absent sealed source.
        guide,t=self.selected('application')
        alternate=REPORT/'shangcheng/originals/application.docx'
        with self.assertRaises(ValueError):adapter.fill_shangcheng_document(guide,t,'application',alternate,{'applicantName':'测试'})
        with self.assertRaises(ValueError):adapter.checked_shangcheng_adapter(guide,t,None)
        with tempfile.TemporaryDirectory() as temp:
            with patch.object(adapter,'ROOT',Path(temp)):
                self.assertIsNone(adapter.shangcheng_blank_binding(guide,t))
    def test_over_capacity_request_list_fails_without_changing_original(self):
        # Break: dropping a fifth request instead of rejecting unsupported form capacity.
        guide,t=self.selected('application');path=ROOT/t['file'];before=path.read_bytes()
        with self.assertRaises(ValueError):adapter.fill_shangcheng_document(guide,t,'application',path,{'applicantName':'测试','requests':['请求']*5})
        self.assertEqual(path.read_bytes(),before)

    def test_evidence_case_number_does_not_replace_static_case_label(self):
        # Break: generic paragraph replacement overwrites both label and value.
        guide,t=self.selected('evidence')
        doc=adapter.fill_shangcheng_document(guide,t,'evidence',ROOT/t['file'],{
            'caseNumber':'上城虚构测试号','evidenceList':[{'name':'测试证据'}]})
        self.assertEqual(doc.tables[0].rows[0].cells[0].text,'案号')
        self.assertEqual(doc.tables[0].rows[0].cells[1].text,'上城虚构测试号')

    def test_expanded_evidence_keeps_rows_whole_and_repeats_own_headers(self):
        # Break: a long record splits its final character onto a continuation page.
        guide,t=self.selected('evidence')
        doc=adapter.fill_shangcheng_document(guide,t,'evidence',ROOT/t['file'],{
            'evidenceList':[{'name':f'虚构证据{i}','purpose':'检查多页证据表的完整记录。'*3} for i in range(10)]})
        table=doc.tables[0]
        for row in table.rows[:2]:self.assertIsNotNone(row._tr.find(qn('w:trPr')+'/'+qn('w:tblHeader')))
        for row in table.rows[2:12]:self.assertIsNotNone(row._tr.find(qn('w:trPr')+'/'+qn('w:cantSplit')))

if __name__=='__main__':unittest.main()
