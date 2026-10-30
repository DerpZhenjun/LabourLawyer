"""Private generated downloads through real province/regional ASGI responses."""
from concurrent.futures import ThreadPoolExecutor
from contextlib import ExitStack
from io import BytesIO
import json
import os
from pathlib import Path
import stat
import sys
import tempfile
import threading
import unittest
from unittest.mock import patch
from urllib.parse import unquote

from docx import Document
from fastapi import BackgroundTasks,FastAPI
from fastapi.testclient import TestClient

ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'backend'))
from routes import province_documents as province
from routes import regional_guides as regional
from regional_guides import RegionalStore


class GeneratedDocumentStorageTests(unittest.TestCase):
    def setUp(self):
        temporary=tempfile.TemporaryDirectory(prefix='generated-storage-');self.addCleanup(temporary.cleanup)
        self.root=Path(temporary.name).resolve();self.data=self.root/'data';self.release=self.root/'release'
        self.release.mkdir();self.stack=ExitStack();self.addCleanup(self.stack.close)
        self.stack.enter_context(patch.dict(os.environ,{'LABOURLAWYER_DATA_DIR':str(self.data),
            'LABOURLAWYER_REGIONAL_DB':str(self.root/'regional.sqlite')}))
        # Old implementation is confined to a fake release while RED runs.
        self.stack.enter_context(patch.object(province,'__file__',str(self.release/'backend/routes/province_documents.py')))
        app=FastAPI();app.include_router(province.router,prefix='/province');app.include_router(regional.router,prefix='/regional')
        self.client=TestClient(app,raise_server_exceptions=False)

    @property
    def generated(self):return self.data/'generated'

    def snapshot(self,folder):
        return {str(p.relative_to(folder)):p.read_bytes() for p in folder.rglob('*') if p.is_file()} if folder.exists() else {}

    def test_province_anonymous_word_stream_preserves_bytes_filename_cache_and_cleans_own_file(self):
        from province_documents import get_province_document
        from WenShu.generic_template_fill import generate_application,generate_evidence
        cases=[('liaoning','application',{'applicantName':'虚构申请人'},generate_application),('zhejiang','evidence',{},generate_evidence)]
        for slug,kind,payload,generate in cases:
            with self.subTest(kind=kind):
                _,source=get_province_document(slug,kind);original=source.read_bytes()
                document=generate(payload,source);expected=BytesIO()
                # ZIP timestamps otherwise make identical docx saves differ at a second boundary.
                with patch('zipfile.time.localtime',return_value=(2026,10,31,0,0,0,5,304,0)):
                    document.save(expected)
                    response=self.client.post(f'/province/{slug}/{kind}/generate',json=payload)
                self.assertEqual(response.status_code,200,response.text[:200] if response.status_code!=200 else '')
                self.assertEqual(response.content,expected.getvalue())
                self.assertEqual(response.headers['cache-control'],'no-store')
                self.assertIn('.docx',unquote(response.headers['content-disposition']))
                self.assertEqual(response.headers['content-type'],'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
                self.assertEqual(source.read_bytes(),original)
        self.assertTrue(self.generated.is_dir())
        self.assertEqual(self.snapshot(self.generated),{})
        self.assertEqual(self.snapshot(self.release),{})
        self.assertFalse((self.data/'labourlawyer.db').exists())

    def test_regional_actual_source_stream_and_headers_use_same_private_storage(self):
        inventory=json.loads((ROOT/'reports/government-guide-localization/SOURCE-INVENTORY.json').read_text(encoding='utf8'))
        store=RegionalStore(self.root/'regional.sqlite')
        store.put_guide(next(g for g in inventory['regions'] if g['district']=='天府新区'))
        guide=store.lookup('sichuan','成都市','天府新区')
        template=next(t for t in guide['templates'] if t.get('can_generate'))
        for kind in ('application','evidence'):
            payload={k:guide[k] for k in ('province','city','district','source_id','version')}
            payload.update(template_id=template['id'],kind=kind,mode='reference_adaptation',data={'applicantName':'虚构申请人'})
            response=self.client.post('/regional/generate',json=payload)
            self.assertEqual(response.status_code,200,response.text[:200] if response.status_code!=200 else '')
            self.assertEqual(response.headers['x-guide-source-id'],guide['source_id'])
            self.assertEqual(response.headers['x-guide-version'],str(guide['version']))
            self.assertEqual(response.headers['x-document-origin'],'official-sample-adaptation-not-blank')
            self.assertEqual(response.headers['cache-control'],'no-store')
            self.assertGreater(len(Document(BytesIO(response.content)).paragraphs),0)
        self.assertTrue(self.generated.is_dir());self.assertEqual(self.snapshot(self.generated),{})
        self.assertEqual(self.snapshot(self.release),{})

    def test_partial_save_failure_cleans_only_own_file_and_never_touches_release(self):
        self.generated.mkdir(parents=True);sentinel=self.generated/'existing-user-document.docx';sentinel.write_bytes(b'USER ORIGINAL')
        old=self.generated/'old-temp.docx';old.write_bytes(b'OLD TEMP')
        class BrokenDocument:
            def save(self,destination):
                if hasattr(destination,'write'):destination.write(b'PARTIAL PRIVATE SYNTHETIC')
                else:Path(destination).write_bytes(b'PARTIAL PRIVATE SYNTHETIC')
                raise RuntimeError('synthetic save failure')
        with patch.object(province,'generate_application',return_value=BrokenDocument()):
            response=self.client.post('/province/liaoning/application/generate',json={})
        self.assertEqual(response.status_code,500)
        self.assertEqual(self.snapshot(self.generated),{'existing-user-document.docx':b'USER ORIGINAL','old-temp.docx':b'OLD TEMP'})
        self.assertEqual(self.snapshot(self.release),{})

    def test_concurrent_requests_own_distinct_files_and_background_cleanup_is_scoped(self):
        barrier=threading.Barrier(2);seen=[];lock=threading.Lock()
        class ConcurrentDocument:
            def save(inner,destination):
                with lock:seen.extend(p.name for p in self.generated.glob('*.docx'))
                barrier.wait(timeout=10)
                doc=Document();doc.add_paragraph('虚构并发文书');doc.save(destination)
        def request():return self.client.post('/province/liaoning/application/generate',json={})
        with patch.object(province,'generate_application',side_effect=lambda *a:ConcurrentDocument()),ThreadPoolExecutor(max_workers=2) as pool:
            responses=list(pool.map(lambda _:request(),range(2)))
        self.assertEqual([r.status_code for r in responses],[200,200])
        self.assertEqual(len(set(seen)),2);self.assertEqual(self.snapshot(self.generated),{})

    def test_each_request_reloads_current_data_dir_and_keeps_existing_files(self):
        from generated_document_storage import save_generated_document
        document=Document();first=save_generated_document(document)
        alternate=self.root/'alternate'
        with patch.dict(os.environ,{'LABOURLAWYER_DATA_DIR':str(alternate)}):second=save_generated_document(document)
        self.assertEqual(first.path.parent,self.generated)
        self.assertEqual(second.path.parent,alternate/'generated')
        first.cleanup();self.assertTrue(second.path.exists());second.cleanup()

    def test_unset_data_dir_matches_platform_fallback_without_initializing_databases(self):
        import generated_document_storage as storage
        fake_backend=self.root/'fallback/backend';fake_backend.mkdir(parents=True)
        with patch.dict(os.environ,{},clear=False),patch.object(storage,'__file__',str(fake_backend/'generated_document_storage.py')):
            os.environ.pop('LABOURLAWYER_DATA_DIR',None)
            saved=storage.save_generated_document(Document())
        self.assertEqual(saved.path.parent,fake_backend/'data/generated');saved.cleanup()
        self.assertFalse((fake_backend/'data/labourlawyer.db').exists())

    def test_new_directories_and_files_are_private_and_cleanup_does_not_delete_replacement(self):
        from generated_document_storage import save_generated_document
        saved=save_generated_document(Document())
        if os.name=='posix':
            self.assertEqual(stat.S_IMODE(saved.path.stat().st_mode),0o600)
            for directory in (self.data,self.generated):self.assertEqual(stat.S_IMODE(directory.stat().st_mode),0o700)
        replacement=saved.path.with_suffix('.replacement');replacement.write_bytes(b'USER REPLACEMENT');os.replace(replacement,saved.path)
        saved.cleanup();self.assertEqual(saved.path.read_bytes(),b'USER REPLACEMENT')

    def test_existing_data_root_permissions_and_user_original_are_unchanged(self):
        from generated_document_storage import save_generated_document
        self.data.mkdir(mode=0o755);before=stat.S_IMODE(self.data.stat().st_mode)
        original=self.data/'user-original.docx';original.write_bytes(b'USER ORIGINAL')
        saved=save_generated_document(Document());saved.cleanup()
        self.assertEqual(stat.S_IMODE(self.data.stat().st_mode),before)
        self.assertEqual(original.read_bytes(),b'USER ORIGINAL')

    def test_save_baseexception_cleans_owned_partial_file(self):
        from generated_document_storage import save_generated_document
        class InterruptedDocument:
            def save(self,destination):destination.write(b'PRIVATE PARTIAL');raise KeyboardInterrupt()
        with self.assertRaises(KeyboardInterrupt):save_generated_document(InterruptedDocument())
        self.assertEqual(self.snapshot(self.generated),{})

    def test_symlink_or_junction_generated_is_rejected_without_outside_writes(self):
        from generated_document_storage import save_generated_document
        outside=self.root/'outside';outside.mkdir();(outside/'user.docx').write_bytes(b'KEEP')
        for segment in ('generated',):
            base=self.root/('links-'+segment);base.mkdir()
            link=base/segment
            if os.name=='nt':
                # A native junction needs no elevated symlink permission; no alternate capture/tool path.
                import subprocess
                result=subprocess.run(['cmd.exe','/d','/c','mklink','/J',str(link),str(outside)],capture_output=True)
                self.assertEqual(result.returncode,0,result.stderr.decode(errors='replace'))
            else:link.symlink_to(outside,target_is_directory=True)
            try:
                with patch.dict(os.environ,{'LABOURLAWYER_DATA_DIR':str(base)}),self.assertRaises((ValueError,OSError)):
                    save_generated_document(Document())
                self.assertEqual(self.snapshot(outside),{'user.docx':b'KEEP'})
            finally:link.rmdir() if os.name=='nt' else link.unlink()

    def test_response_construction_failure_cleans_own_file_only(self):
        self.generated.mkdir(parents=True);(self.generated/'existing.docx').write_bytes(b'KEEP')
        with patch.object(province,'FileResponse',side_effect=RuntimeError('response failed')):
            with self.assertRaises(RuntimeError):province._response(Document(),'test.docx',BackgroundTasks())
        self.assertEqual(self.snapshot(self.generated),{'existing.docx':b'KEEP'})

    def test_release_write_denied_download_still_succeeds_with_external_data(self):
        original_mkdir=Path.mkdir;original_mkstemp=tempfile.mkstemp
        def mkdir(path,*args,**kwargs):
            if path.resolve().is_relative_to(self.release):raise PermissionError('Immutable release')
            return original_mkdir(path,*args,**kwargs)
        def mkstemp(*args,**kwargs):
            if Path(kwargs['dir']).resolve().is_relative_to(self.release):raise PermissionError('Immutable release')
            return original_mkstemp(*args,**kwargs)
        # Windows directory read-only bits are not an ACL. Deny both output creation
        # operations explicitly; POSIX also exercises actual read-only directory mode.
        previous=stat.S_IMODE(self.release.stat().st_mode)
        if os.name=='posix':self.release.chmod(0o500)
        try:
            with patch.object(Path,'mkdir',mkdir),patch.object(tempfile,'mkstemp',mkstemp):
                response=self.client.post('/province/liaoning/application/generate',json={})
            self.assertEqual(response.status_code,200)
            self.assertEqual(self.snapshot(self.release),{});self.assertTrue(self.generated.is_dir())
        finally:
            if os.name=='posix':self.release.chmod(previous)


if __name__=='__main__':unittest.main()
