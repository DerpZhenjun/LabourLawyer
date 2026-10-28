"""Haizhou's own originals may share layout, never sources or annex identity."""
import copy
import hashlib
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from fastapi import FastAPI
from fastapi.testclient import TestClient
from docx import Document
from regional_guides import RegionalStore
from regional_adapters import with_reference_capabilities
from regional_blank_documents import blank_binding
from routes.regional_guides import router

ROOT=Path(__file__).resolve().parents[2]
BASE=ROOT/'reports/nationwide-localization-20261023-night/cross-province/jiangsu-haizhou'
IDS={'application':'50286982-5aaa-4310-b593-5767344cea86','evidence':'2a18398d-be86-4613-a878-e17ed90f0354'}


class HaizhouGenerationTests(unittest.TestCase):
    def guide(self):return json.loads((BASE/'staging-0540/reviewed-guide.json').read_text('utf8'))

    def test_only_own_reviewed_two_blank_forms_enable_filling(self):
        g=with_reference_capabilities(self.guide())
        self.assertEqual({t['id']:t['document_types'] for t in g['templates'] if t.get('can_generate')},
                         {'50286982-5aaa-4310-b593-5767344cea86':['application'],
                          '2a18398d-be86-4613-a878-e17ed90f0354':['evidence']})
        for t in g['templates']:
            if t.get('can_generate'):self.assertIn('海州区',t['generation_label'])

    def test_wrong_source_annex_path_checksum_or_download_identity_rejected(self):
        g=with_reference_capabilities(self.guide())
        own=next(t for t in g['templates'] if t['id']==IDS['application'])
        self.assertEqual(blank_binding(g,own),'application')
        for key,value in [('district','赣榆区'),('region_code','320707'),('source_url','https://lyghz.jszwfw.gov.cn/unverified'),
                          ('source_id','5725c8df9cb686d0517254e8')]:
            changed=copy.deepcopy(g);changed[key]=value
            self.assertIsNone(blank_binding(changed,own))
        for key,value in [('file',own['file'].replace('/320706/','/320904/')),('sha256','0'*64),
                          ('id','8984153f-2c04-4676-a8fd-3a627658fb9d'),('official_kind','sample'),
                          ('download_url',own['download_url'].replace(IDS['application'],'8984153f-2c04-4676-a8fd-3a627658fb9d'))]:
            changed=copy.deepcopy(own);changed[key]=value
            self.assertIsNone(blank_binding(g,changed),key)

    def test_two_real_api_outputs_preserve_own_package_and_reject_neighbour(self):
        with tempfile.TemporaryDirectory() as temp:
            store=RegionalStore(Path(temp)/'public.sqlite');g=store.put_guide(with_reference_capabilities(self.guide()))
            app=FastAPI();app.include_router(router,prefix='/api/regional-guides')
            with patch('routes.regional_guides.store',return_value=store),TestClient(app) as client:
                for kind,data in [('application',{'applicantName':'海州合成甲（示例）','respondentName':'海州合成测试公司（示例）',
                      'requests':['支付合成测试工资100元。'],'factsReason':'本文件为合成测试，不用于真实仲裁申请。',
                      'copyCount':3,'evidenceCopyCount':2}),
                      ('evidence',{'evidenceList':[{'name':'海州合成劳动合同','source':'本人保存','purpose':'说明测试劳动关系',
                                                    'originalType':'复印件','pageCount':'3页'}]})]:
                    p={k:g[k] for k in ['province','city','district','source_id','version']}
                    p.update(template_id=IDS[kind],kind=kind,mode='blank',data=data)
                    response=client.post('/api/regional-guides/generate',json=p)
                    self.assertEqual(response.status_code,200,response.text if response.status_code!=200 else '')
                    self.assertEqual(response.headers['X-Guide-Source-ID'],'f2208816780787a1a046b91f')
                    self.assertEqual(response.headers['X-Document-Origin'],'official-blank-filled')
                    doc=Document(io.BytesIO(response.content));text='\n'.join(p.text for p in doc.paragraphs)
                    if kind=='application':
                        self.assertIn('海州合成甲（示例）',text);self.assertNotIn('用人单位申请',text)
                        self.assertEqual(text.count('被申请人：'),1)
                        self.assertIn('副本3份；相关证据2份',text)
                        self.assertIn('申请人：         （签名、盖章或指模）',text)
                    else:
                        self.assertEqual(len(doc.tables[0].columns),6);self.assertEqual(len(doc.tables[0].rows),2)
                        self.assertEqual(doc.tables[0].rows[1].cells[2].text,'本人保存')
                        self.assertEqual(doc.tables[0].rows[1].cells[5].text,'3')
                    own=next(t for t in g['templates'] if t['id']==IDS[kind])
                    with ZipFile(ROOT/own['file']) as before,ZipFile(io.BytesIO(response.content)) as after:
                        self.assertEqual(set(before.namelist()),set(after.namelist()))
                        for part in before.namelist():
                            if part!='word/document.xml':self.assertEqual(before.read(part),after.read(part))
                    for key,value in [('district','赣榆区'),('source_id','ef1eb897ebf880c6f2a316f5'),
                                      ('template_id','985f4770-1ab2-4264-9482-0f8528b16513'),('mode','reference_adaptation')]:
                        changed=copy.deepcopy(p);changed[key]=value
                        self.assertEqual(client.post('/api/regional-guides/generate',json=changed).status_code,409)


if __name__=='__main__':unittest.main()
