"""D Zhejiang filling uses only sealed own sources and temporary output bytes."""
from copy import deepcopy
from io import BytesIO
import importlib.util
import json
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'backend'))


class ZhejiangDTests(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(importlib.util.find_spec('regional_zhejiang_oct29d'),
                             'D Zhejiang requires its own strict source adapter')
        import regional_zhejiang_oct29d as provider
        self.p = provider

    def own(self, name='xiaoshan', root=None):
        return self.p.with_capabilities(self.p.parse_capture(name, root=root), root=root)

    def template(self, guide, kind):
        return next(t for t in guide['templates'] if t['document_types'] == [kind])

    def fill(self, guide, kind, data, root=None):
        template = self.template(guide, kind)
        return self.p.fill_document(guide, template, kind,
                                   Path(root or ROOT) / template['file'], data, root=root)

    def public_copy(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        root = Path(temporary.name)
        white = json.loads((ROOT / self.p.REPORT / 'PUBLIC-WHITELIST.json').read_text('utf8'))
        for relative in list(white['files']) + [self.p.REPORT + '/PUBLIC-WHITELIST.json']:
            target = root / relative
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(ROOT / relative, target)
        return root

    def test_two_own_sources_keep_different_material_sets_and_three_real_blanks(self):
        guides = [self.own(name) for name in ['xiaoshan', 'linan']]
        self.assertEqual([g['region_code'] for g in guides], ['330109', '330112'])
        self.assertEqual([len(g['materials']) for g in guides], [5, 4])
        self.assertEqual([len([t for t in g['templates'] if t['can_generate']]) for g in guides], [2, 1])
        self.assertEqual(len({g['source_id'] for g in guides}), 2)
        self.assertNotEqual(guides[0]['service_code'], guides[1]['service_code'])
        self.assertTrue(all(not g['published'] and g['candidate_web_qa'] == 'NOT_RUN' for g in guides))
        self.assertEqual(self.template(guides[0], 'application')['sha256'],
                         self.template(guides[1], 'application')['sha256'])
        self.assertFalse(any('evidence' in t['document_types'] for t in guides[1]['templates']))
        self.assertEqual(guides[0]['fields']['咨询电话']['raw_value'], '0571-89653752')
        self.assertEqual(guides[1]['fields']['咨询电话']['raw_value'], '0571-61071571')
        self.assertIn('4.申请书副本份数', self.fill(guides[1], 'application', {}).tables[0].rows[-1].cells[0].text)

    def test_publication_notice_changes_only_after_both_lifecycle_gates_and_preserves_source_warnings(self):
        guide = self.own('linan')
        guide['published'] = True
        unpublished = self.p.with_capabilities(guide)
        self.assertTrue(any('不代表已发布' in warning for warning in unpublished['warnings']))
        guide['candidate_web_qa'] = 'PASS'
        published = self.p.with_capabilities(guide)
        self.assertFalse(any('仅本区自身空白文书完成离线' in warning for warning in published['warnings']))
        self.assertTrue(any('已通过填写、版式及网页验收' in warning for warning in published['warnings']))
        self.assertTrue(any('ISPUBLISH=0' in warning for warning in published['warnings']))
        self.assertTrue(any('法定60天' in warning for warning in published['warnings']))
        self.assertEqual(self.p.with_capabilities(published)['warnings'], published['warnings'])

    def test_each_scope_source_service_material_and_annex_is_required_despite_equal_bytes(self):
        guide = self.own()
        template = self.template(guide, 'application')
        for field in ['district', 'source_url', 'source_id', 'service_code', 'authority']:
            bad = deepcopy(guide)
            bad[field] += '-changed'
            self.assertIsNone(self.p.blank_binding(bad, template))
        for field in ['id', 'material_id', 'file', 'download_url', 'sha256', 'region_code', 'official_kind']:
            bad = deepcopy(template)
            bad[field] += '-changed'
            self.assertIsNone(self.p.blank_binding(guide, bad))
        other = self.own('linan')
        self.assertIsNone(self.p.blank_binding(other, template))
        with self.assertRaises(ValueError):
            self.p.checked_adapter(other, self.template(other, 'application'), 'evidence')
        bad = deepcopy(guide)
        bad['materials'][0]['copies'] = '99'
        self.assertIsNone(self.p.blank_binding(bad, template))

    def test_entire_52_file_chain_is_checked_and_explicit_root_never_falls_back(self):
        root = self.public_copy()
        with patch.object(self.p, 'ROOT', root / 'nonexistent-original'):
            guide = self.own(root=root)
            self.assertTrue(self.fill(guide, 'application', {'phone': '虚构电话'}, root).tables)
        retained = root / self.p.REPORT / 'linan/search.json'
        retained.write_bytes(retained.read_bytes() + b' ')
        with self.assertRaisesRegex(ValueError, '来源|source|封存'):
            self.p.parse_capture('xiaoshan', root=root)
        self.assertIsNone(self.p.blank_binding(guide, self.template(guide, 'application'), root=root))

    def test_missing_manifest_or_whitelist_fails_in_fresh_explicit_root(self):
        root = self.public_copy()
        (root / self.p.REPORT / 'SOURCE-MANIFEST.json').unlink()
        with self.assertRaises((ValueError, OSError)):
            self.p.parse_capture('linan', root=root)
        root2 = self.public_copy()
        white = root2 / self.p.REPORT / 'PUBLIC-WHITELIST.json'
        white.write_bytes(white.read_bytes() + b' ')
        with self.assertRaises(ValueError):
            self.p.parse_capture('xiaoshan', root=root2)

    def test_public_proof_requires_explicit_root_and_only_own_pinned_attachments(self):
        self.assertTrue(callable(getattr(self.p, 'public_evidence', None)))
        self.assertTrue(callable(getattr(self.p, 'approved_report_attachment', None)))
        root = self.public_copy()
        guide = self.own(root=root)
        template = self.template(guide, 'application')
        with patch.object(self.p, 'ROOT', root / 'missing-original'):
            self.assertEqual(self.p.public_evidence(guide, root), guide['evidence'])
            self.assertTrue(self.p.approved_report_attachment(guide, template, root))
        self.assertIsNone(self.p.public_evidence(guide, None))
        wrong = deepcopy(template)
        wrong['file'] = wrong['file'].replace('/xiaoshan/', '/linan/')
        self.assertFalse(self.p.approved_report_attachment(guide, wrong, root))
        guide['fields']['咨询电话']['raw_value'] = 'wrong locality'
        self.assertIsNone(self.p.public_evidence(guide, root))

    def test_worker_application_preserves_every_supported_text_and_only_changes_document_xml(self):
        guide = self.own()
        values = {key: '  虚构' + key + '  ' for key in self.p.APPLICATION_TEXT_FIELDS
                  if key not in {'householdType', 'birthDate', 'applicationDate'}}
        values.update(householdType='外籍人员', birthDate='1990-01-02', applicationDate='2026-10-29',
                      copyCount=3, requests=['  虚构请求甲  ', '', '虚构请求乙', '虚构请求丙'],
                      factsReason='  虚构事实首行\n虚构事实次行  ')
        document = self.fill(guide, 'application', values)
        self.assertEqual(len(document.tables), 1)
        self.assertNotIn('（用于用人单位申请）', '\n'.join(p.text for p in document.paragraphs))
        text = '\n'.join(c.text for row in document.tables[0].rows for c in row.cells)
        for key, value in values.items():
            if key in {'householdType', 'birthDate', 'applicationDate', 'copyCount', 'requests'}:
                continue
            self.assertIn(value, text, key)
        for request in values['requests']:
            self.assertIn(request, text)
        self.assertIn('4.申请书副本份数，应按对方当事人人数提交。', text)
        signature = next(p for p in document.tables[0].rows[-1].cells[0].paragraphs if '申请人：' in p.text)
        self.assertNotIn(values['applicantName'], signature.text)
        self.assertIn('☑', document.tables[0].rows[4].cells[1].text)
        self.assertNotIn('√外籍人员', document.tables[0].rows[4].cells[1].text)
        output = BytesIO()
        document.save(output)
        original = ROOT / self.template(guide, 'application')['file']
        with ZipFile(original) as before, ZipFile(output) as after:
            self.assertEqual(before.namelist(), after.namelist())
            self.assertEqual([n for n in before.namelist() if before.read(n) != after.read(n)], ['word/document.xml'])

    def test_application_without_story_keeps_optional_slots_and_never_invents_copy_or_date(self):
        guide = self.own('linan')
        document = self.fill(guide, 'application', {'phone': '虚构电话', 'copyCount': 4})
        self.assertIn('虚构电话', document.tables[0].rows[3].cells[4].text)
        self.assertIn('4', document.tables[0].rows[-1].cells[0].text)
        blank = self.fill(guide, 'application', {})
        self.assertEqual(len(blank.tables), 1)
        footer = blank.tables[0].rows[-1].cells[0].text
        self.assertIn('申请书副本      份', footer)
        self.assertIn('年   月    日', footer)
        with self.assertRaises(ValueError):
            t = self.template(guide, 'application')
            self.p.fill_document(guide, t, 'application', ROOT / t['file'], {}, applicant_mode='employer')

    def test_evidence_preserves_nameless_rows_labels_signature_blanks_and_submitted_values(self):
        guide = self.own()
        items = [{'name': '  虚构证据甲  ', 'pageCount': '2页', 'source': '  本人虚构保存  ', 'purpose': '虚构证明目的'},
                 {'name': '', 'pageCount': '3', 'source': '虚构无名称来源', 'purpose': '仍须保留此项内容'}]
        document = self.fill(guide, 'evidence', {'committeeName': '虚构委员会', 'caseNumber': '虚构案号',
                                               'copyCount': 3, 'submitDate': '2026-10-29', 'evidenceList': items})
        table = document.tables[0]
        self.assertEqual(table.rows[0].cells[0].text, '案号')
        self.assertEqual(table.rows[0].cells[1].text, '虚构案号')
        for index, item in enumerate(items, 2):
            for key, column in [('name', 1), ('pageCount', 3), ('source', 4), ('purpose', 5)]:
                self.assertEqual(table.rows[index].cells[column].text, item[key])
        self.assertEqual(table.rows[-2].cells[2].text, '')
        self.assertEqual(table.rows[-1].cells[2].text, '')
        self.assertEqual(table.rows[-1].cells[5].text, '')
        self.assertIn('本表一式', document.paragraphs[-1].text)
        blank = self.fill(guide, 'evidence', {})
        self.assertEqual(blank.tables[0].rows[0].cells[5].text, '')

    def test_unsupported_or_overcapacity_values_are_rejected_without_clipping_or_ignoring(self):
        guide = self.own()
        for data in [{'requests': ['一'] * 5}, {'applicantName': '字' * 301},
                     {'householdType': '不在官方选项内'}, {'applicationDate': '\x00'},
                     {'postalCode': '无该槽不得丢弃'}, {'copyCount': 1.5}, {'copyCount': True}]:
            with self.assertRaises(ValueError):
                self.fill(guide, 'application', data)
        for data in [{'submitterName': '签名不得自动填'},
                     {'evidenceList': [{'name': '证据', 'originalType': '原件'}]},
                     {'evidenceList': [{'name': '字' * 301}]}, {'evidenceList': ['bad-shape']}]:
            with self.assertRaises(ValueError):
                self.fill(guide, 'evidence', data)


if __name__ == '__main__':
    unittest.main()
