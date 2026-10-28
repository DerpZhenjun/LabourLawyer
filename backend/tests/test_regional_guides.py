"""Regional isolation, provenance, channels and durable outbox contracts."""
import copy
import json
import sys
import tempfile
import unittest
import io
from zipfile import ZipFile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


def guide(district='天府新区', code='510110000000'):
    return {'province':'sichuan','province_name':'四川省','city':'成都市','district':district,
            'region_code':code,'service_name':'劳动人事争议仲裁申请',
            'source_url':f'https://www.sczwfw.gov.cn/guide?areaCode={code}',
            'collected_at':'2026-10-08T05:00:00Z','verified_at':'2026-10-08T05:00:00Z',
            'verification_status':'verified','fields':{'办理形式':{'raw_value':'窗口办理','normalized_value':'window'}},
            'materials':[{'name':'仲裁申请书','copies':'3'}], 'templates':[], 'sections':{},
            'channel':{'form':'window','online_status':'unavailable','entry_url':None}}


class RegionalContracts(unittest.TestCase):
    def setUp(self):
        from regional_guides import RegionalStore
        self.temp=tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path=Path(self.temp.name)/'test.sqlite'
        self.store=RegionalStore(self.path)

    def test_exact_region_never_inherits_neighbor_or_province(self):
        self.store.put_guide(guide())
        missing=self.store.lookup('sichuan','成都市','武侯区')
        self.assertEqual(missing['verification_status'],'unverified')
        self.assertIsNone(missing['source_id'])
        self.assertEqual(missing['materials'],[])
        self.assertEqual(self.store.lookup('sichuan','','')['verification_status'],'unverified')

    def test_distinct_regions_have_distinct_sources_and_requirements(self):
        self.store.put_guide(guide())
        other=guide('锦江区','510104000000');other['materials']=[{'name':'仲裁申请书','copies':'2'}]
        self.store.put_guide(other)
        a=self.store.lookup('sichuan','成都市','天府新区');b=self.store.lookup('sichuan','成都市','锦江区')
        self.assertNotEqual(a['source_id'],b['source_id'])
        self.assertEqual((a['materials'][0]['copies'],b['materials'][0]['copies']),('3','2'))

    def test_mediation_cannot_be_ingested_as_arbitration(self):
        bad=guide();bad['service_name']='劳动人事争议调解申请'
        with self.assertRaises(ValueError):self.store.put_guide(bad)

    def test_source_region_mismatch_is_rejected(self):
        bad=guide();bad['region_code']='510104000000'
        with self.assertRaises(ValueError):self.store.put_guide(bad)

    def test_content_changes_create_versions_without_overwriting_history(self):
        self.store.put_guide(guide());same=guide();same['collected_at']='2026-10-08T05:30:00Z'
        self.store.put_guide(same)
        self.assertEqual(len(self.store.versions('sichuan','成都市','天府新区')),1)
        changed=guide();changed['materials'][0]['copies']='4';self.store.put_guide(changed)
        versions=self.store.versions('sichuan','成都市','天府新区')
        self.assertEqual(len(versions),2)
        self.assertEqual(versions[0]['materials'][0]['copies'],'4')
        self.assertEqual(versions[1]['materials'][0]['copies'],'3')

    def test_recheck_timestamps_screenshots_and_signed_links_do_not_fake_versions(self):
        d=guide();d['templates']=[{'id':'x','sha256':'abc','collected_at':'old','verified_at':'old','download_url':'https://www.sczwfw.gov.cn/file?signature=old'}]
        d['evidence']={'capture_file':'old.json','capture_sha256':'old','screenshot':'old.png'}
        self.store.put_guide(d)
        d['templates'][0].update(collected_at='new',verified_at='new',download_url='https://www.sczwfw.gov.cn/file?signature=new')
        d['evidence']['capture_sha256']='new';self.store.put_guide(d)
        self.assertEqual(len(self.store.versions('sichuan','成都市','天府新区')),1)

    def test_other_region_name_in_hidden_choices_cannot_validate_wrong_header(self):
        from regional_capture import parse_sichuan_capture
        p=Path(__file__).resolve().parents[2]/'reports/government-guide-localization/tianfu-dom.json'
        d=json.loads(p.read_text(encoding='utf-8-sig'));d=json.loads(d) if isinstance(d,str) else d
        d['html']=d['html'].replace('id="cityName">四川天府新区','id="cityName">武侯区')
        with self.assertRaises(ValueError):parse_sichuan_capture(d,'成都市','天府新区')

    def test_field_provenance_timestamps_survive_content_hashing(self):
        d=guide();d['fields']['办理形式'].update(source_url=d['source_url'],region_code=d['region_code'],collected_at=d['collected_at'],verified_at=d['verified_at'],verification_status='verified')
        result=self.store.put_guide(d)
        self.assertEqual(result['fields']['办理形式']['collected_at'],'2026-10-08T05:00:00Z')

    def test_reverted_official_content_becomes_current_without_losing_versions(self):
        self.store.put_guide(guide());changed=guide();changed['materials'][0]['copies']='4';self.store.put_guide(changed)
        self.store.put_guide(guide())
        self.assertEqual(self.store.lookup('sichuan','成都市','天府新区')['materials'][0]['copies'],'3')
        self.assertEqual(len(self.store.versions('sichuan','成都市','天府新区')),2)


    def test_live_sichuan_capture_keeps_raw_fields_and_material_details(self):
        from regional_capture import parse_sichuan_capture
        p=Path(__file__).resolve().parents[2]/'reports/government-guide-localization/tianfu-dom.json'
        d=json.loads(p.read_text(encoding='utf-8-sig'));d=json.loads(d) if isinstance(d,str) else d
        g=parse_sichuan_capture(d,'成都市','天府新区')
        self.assertEqual(g['fields']['办理形式']['raw_value'],'窗口办理')
        self.assertEqual(g['fields']['是否网办']['raw_value'],'否')
        self.assertEqual(g['channel']['online_status'],'conflict')
        self.assertEqual(len(g['materials']),4)
        self.assertEqual(g['materials'][2]['copies'],'3')
        self.assertEqual(g['materials'][2]['details']['纸质材料规格'],'A4')
        application=next(t for t in g['templates'] if '仲裁申请书' in t['original_name'])
        self.assertEqual(application['official_kind'],'sample')
        self.assertEqual(application['annex_id'],'ec3268e5fda94301ab40a7e086afe07a')

    def test_parser_rejects_wrong_region_text_even_with_valid_official_host(self):
        from regional_capture import parse_sichuan_capture
        with self.assertRaises(ValueError):parse_sichuan_capture({'url':'https://www.sczwfw.gov.cn/?areaCode=510110000000','text':'锦江区 劳动人事争议仲裁申请','html':'','tables':[],'captured_at':'2026-10-08'},'成都市','天府新区')

    def test_events_survive_restart_and_failed_openclaw_does_not_ack(self):
        from regional_guides import RegionalStore
        e=self.store.record_event({'type':'ONLINE_CHANNEL_UNVERIFIED','province':'sichuan','city':'成都市','district':'天府新区','reason':'入口待核验'})
        restarted=RegionalStore(self.path)
        self.assertEqual(restarted.events()[0]['id'],e['id'])
        self.assertEqual(restarted.events()[0]['delivery_status'],'PENDING_OPENCLAW')
        with self.assertRaises(ValueError):restarted.acknowledge([e['id']],{'ok':False})
        self.assertEqual(restarted.events()[0]['delivery_status'],'PENDING_OPENCLAW')

    def test_openclaw_receipt_must_cover_exact_event_ids(self):
        e=self.store.record_event({'type':'SEARCH_NO_RESULT','reason':'短词未找到事项'})
        with self.assertRaises(ValueError):self.store.acknowledge([e['id']],{'ok':True,'accepted_event_ids':['wrong']})
        self.store.acknowledge([e['id']],{'ok':True,'accepted_event_ids':[e['id']],'engine':'OpenClaw','receipt_file':'receipt.json'})
        self.assertEqual(self.store.events()[0]['delivery_status'],'ACKNOWLEDGED_OPENCLAW')

    def test_template_region_mismatch_blocks_generation(self):
        from regional_guides import select_template
        d=guide();d['templates']=[{'id':'x','region_code':'510104000000','official_kind':'blank','can_generate':True,'document_types':['application']}]
        with self.assertRaises(ValueError):select_template(d,'x','application','blank')

    def test_sample_is_not_blank_even_if_name_says_template(self):
        from regional_guides import select_template
        d=guide();d['templates']=[{'id':'x','region_code':d['region_code'],'official_kind':'sample','can_generate':True,'document_types':['application']}]
        with self.assertRaises(ValueError):select_template(d,'x','application','blank')
        self.assertEqual(select_template(d,'x','application','reference_adaptation')['official_kind'],'sample')


class ChannelContracts(unittest.TestCase):
    def test_window_only_has_no_online_action(self):
        from regional_guides import resolve_channel
        c=resolve_channel({'办理形式':'窗口办理','是否网办':'否'},'到窗口提交')
        self.assertEqual(c['online_status'],'unavailable');self.assertIsNone(c['entry_url'])

    def test_conflict_must_not_become_verified_online_or_definitive_window_only(self):
        from regional_guides import resolve_channel
        c=resolve_channel({'办理形式':'窗口办理','是否网办':'否'},'登录http://221.237.108.25:7001/cddtj/login.jsp进行线上申请')
        self.assertEqual(c['form'],'window');self.assertEqual(c['online_status'],'conflict')
        self.assertIsNone(c['entry_url']);self.assertIn('http://221.237.108.25:7001/cddtj/login.jsp',c['mentioned_urls'])

    def test_online_entry_requires_browser_verification(self):
        from regional_guides import resolve_channel
        fields={'办理形式':'窗口办理、网上办理','是否网办':'是'}
        c=resolve_channel(fields,'',{'url':'https://www.sczwfw.gov.cn/apply','status':'auth_required','official_link':True})
        self.assertEqual(c['form'],'window_and_online');self.assertEqual(c['online_status'],'auth_required')
        self.assertEqual(c['entry_url'],'https://www.sczwfw.gov.cn/apply')
        self.assertEqual(resolve_channel(fields,'')['online_status'],'unverified')

    def test_missing_form_never_asserts_official_requirement(self):
        from regional_guides import resolve_channel
        self.assertEqual(resolve_channel({},'')['form'],'unverified')


class RegionalDocumentContracts(unittest.TestCase):
    def test_evidence_has_no_authorization_spacing_tail_and_repeats_header(self):
        from regional_documents import fill_tianfu_reference
        from docx.oxml.ns import qn
        path=Path(__file__).resolve().parents[2]/'reports/government-guide-localization/tianfu-application-reference.docx'
        d=fill_tianfu_reference(path,'evidence',{})
        self.assertTrue(d.paragraphs[-1].text.startswith('备注：'))
        self.assertIsNotNone(d.tables[0].rows[0]._tr.find(qn('w:trPr')+'/'+qn('w:tblHeader')))

    def test_unedited_package_parts_are_preserved_byte_for_byte(self):
        from regional_documents import fill_tianfu_reference
        path=Path(__file__).resolve().parents[2]/'reports/government-guide-localization/tianfu-application-reference.docx'
        buffer=io.BytesIO();fill_tianfu_reference(path,'application',{}).save(buffer)
        with ZipFile(path) as original,ZipFile(buffer) as generated:
            for name in original.namelist():
                if name!='word/document.xml':self.assertEqual(original.read(name),generated.read(name),name)

    def test_empty_cells_keep_official_mark_font_and_facts_keep_value_font(self):
        from docx import Document
        from regional_documents import fill_tianfu_reference
        path=Path(__file__).resolve().parents[2]/'reports/government-guide-localization/tianfu-application-reference.docx'
        original=Document(path)
        d=fill_tianfu_reference(path,'evidence',{'evidenceList':[{'name':'劳动合同','purpose':'说明劳动关系'}]})
        self.assertEqual(d.tables[0].cell(1,1).paragraphs[0].runs[0].font.size.pt,10.5)
        d=fill_tianfu_reference(path,'application',{'factsReason':'这是填写内容'})
        value=next(r for r in d.paragraphs[15].runs if r.text=='这是填写内容')
        self.assertEqual(value._r.rPr.xml,original.paragraphs[15].runs[1]._r.rPr.xml)
    def test_tianfu_reference_adapter_removes_sample_facts_and_other_forms(self):
        from regional_documents import fill_tianfu_reference
        path=Path(__file__).resolve().parents[2]/'reports/government-guide-localization/tianfu-application-reference.docx'
        d=fill_tianfu_reference(path,'application',{'applicantName':'测试劳动者','requests':['请求支付工资100元'],'factsReason':'这是一份自动化测试文书，不用于提交。'})
        text='\n'.join(p.text for p in d.paragraphs)
        self.assertIn('测试劳动者',text);self.assertIn('请求支付工资100元',text)
        self.assertNotIn('申请人于x年',text);self.assertNotIn('xxx公司',text)
        self.assertNotIn('授权委托书',text);self.assertNotIn('证据材料清单',text)
        self.assertIn('副本____份',text)

    def test_tianfu_evidence_generates_only_requested_rows_and_preserves_signing(self):
        from regional_documents import fill_tianfu_reference
        path=Path(__file__).resolve().parents[2]/'reports/government-guide-localization/tianfu-application-reference.docx'
        d=fill_tianfu_reference(path,'evidence',{'evidenceList':[{'name':'测试劳动合同','pageCount':'3','purpose':'证明劳动关系'}]})
        text='\n'.join(p.text for p in d.paragraphs)
        self.assertEqual(len(d.tables),1);self.assertEqual(len(d.tables[0].rows),2)
        self.assertEqual(d.tables[0].cell(1,1).text,'测试劳动合同')
        self.assertEqual(d.tables[0].cell(1,6).text,'证明劳动关系')
        self.assertIn('提交人：',text);self.assertNotIn('授权委托书',text)


if __name__=='__main__':unittest.main()
