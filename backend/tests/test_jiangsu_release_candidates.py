"""Own Dongtai/Sheyang candidate generation: provenance and actual source bytes."""
import copy
import hashlib
from io import BytesIO
from pathlib import Path
import sys
import tempfile
import unittest
import shutil
from unittest.mock import patch
from zipfile import ZipFile
from docx import Document
ROOT=Path(__file__).resolve().parents[2];sys.path.insert(0,str(ROOT/'backend'))
try:import regional_jiangsu_candidates as adapter
except ImportError:adapter=None
OWN={
 '320981':('东台市','盐城市东台市人力资源和社会保障局','11320981014396886W400201430200201',
 'https://ycdt.jszwfw.gov.cn/jszwfw/publicservice/itemlist/showDetail.do?webId=100&iddept_ql_inf=11320981014396886W4002014302002&ql_kind=20&iddept_yw_inf=11320981014396886W400201430200201',
 '5e53e6d3179046433e5fdcb8','d5a8d1f0-ba28-4fa3-8e0d-0c12fb5e8120','5d848c02-a62a-4686-9f25-2e320b1b32ae'),
 '320924':('射阳县','射阳县人力资源和社会保障局','11320924737070337G400201430200201',
 'https://ycsy.jszwfw.gov.cn/jszwfw/publicservice/itemlist/showDetail.do?webId=98&iddept_ql_inf=11320924737070337G4002014302002&ql_kind=20&iddept_yw_inf=11320924737070337G400201430200201',
 '69952614ad2cf51d91ee7460','85e7e12b-18a8-458f-90f5-5a9ea44439a1','c3833b5a-b6ea-40aa-a3ff-905a50e83a4d')}
SHA={'application':'99dfd60d844fa9a4ea3c37df64d447eda980849427a17e1afbba92e61b092dc3','evidence':'c40d42c03b2881fa25144ba895132befbbd23c92b76469dc5a1bc7dd69eefe8c'}
def guide(code):
    district,agency,task,url,source_id,application,evidence=OWN[code]
    g=dict(province='jiangsu',city='盐城市',district=district,region_code=code,authority=agency,service_code=task,
           source_url=url,source_id=source_id,verification_status='verified',templates=[],published=False)
    for kind,annex,name in [('application',application,'劳动人事争议仲裁申请书.docx'),('evidence',evidence,'证据清单.docx')]:
        g['templates'].append(dict(id=annex,annex_id=annex,region_code=code,sha256=SHA[kind],official_kind='blank',
            download_status='downloaded',source_url=url,file=f'templates/_docx/_regional/{code}/{annex}-{name}',
            download_url='https://sxgl.jszwfw.gov.cn:66/audit-qlk-web/rest/frame/pages/basic/attach/attachAction/getContent?isCommondto=true&attachGuid='+annex))
    return g
DATA=dict(applicantName='虚构测试甲',gender='女',birthDate='1990-01-02',respondentName='虚构单位',
    requests=['支付虚构测试工资100元。'],factsReason='仅为软件测试，姓名单位事实全部虚构。',copyCount=3,evidenceCopyCount=2,
    evidenceList=[dict(name='虚构合同',source='虚构来源',purpose='虚构证明目的',originalType='复印件',pageCount='3页')])

class JiangsuCandidateReleaseTests(unittest.TestCase):
    def setUp(self):self.assertIsNotNone(adapter,'Own verified-byte candidate adapter is missing')
    def test_contract_has_only_two_own_blanks_and_grants_no_publication(self):
        for code in OWN:
            with self.subTest(code=code):
                contract=adapter.candidate_contract(guide(code))
                self.assertEqual(set(contract['documents']),{'application','evidence'})
                self.assertEqual(contract['region_code'],code)
                self.assertFalse(contract['published']);self.assertFalse(contract['generation_enabled'])
                self.assertEqual(contract['web_qa_status'],'NOT_RUN')
                for kind,item in contract['documents'].items():self.assertEqual(item['source_sha256'],SHA[kind])

    def test_real_generation_keeps_manual_signature_package_and_one_evidence(self):
        for code in OWN:
            for kind in SHA:
                with self.subTest(code=code,kind=kind):
                    g=guide(code);payload=copy.deepcopy(DATA);payload['committeeName']=g['district']+'劳动人事争议仲裁委员会'
                    source=ROOT/g['templates'][0 if kind=='application' else 1]['file'];before=source.read_bytes()
                    output=BytesIO();adapter.fill_candidate_document(g,kind,payload).save(output)
                    self.assertEqual(source.read_bytes(),before)
                    with ZipFile(BytesIO(before)) as a,ZipFile(BytesIO(output.getvalue())) as b:
                        self.assertEqual(a.namelist(),b.namelist())
                        self.assertEqual([n for n in a.namelist() if a.read(n)!=b.read(n)],['word/document.xml'])
                    doc=Document(BytesIO(output.getvalue()));text='\n'.join(p.text for p in doc.paragraphs)
                    if kind=='application':
                        self.assertIn('出生日期：1990-01-02',text);self.assertIn('副本3份；相关证据2份',text)
                        self.assertRegex(text,r'申请人：\s+（签名、盖章或指模）')
                        self.assertNotIn('用人单位申请',text)
                    else:
                        self.assertEqual(len(doc.tables[0].rows),2)
                        self.assertEqual([c.text for c in doc.tables[0].rows[1].cells],['1','虚构合同','虚构来源','虚构证明目的','复印件','3'])
                        self.assertRegex(text,r'提交人：\s+提交时间：\s+年\s+月\s+日')

    def test_foreign_identity_annex_samples_and_arbitrary_paths_fail(self):
        for code in OWN:
            g=guide(code)
            for key,value in [('province','zhejiang'),('city','南京市'),('district','亭湖区'),('region_code','320902'),
                    ('source_url',g['source_url']+'&foreign=1'),('source_id','foreign'),('authority','foreign'),('service_code','foreign')]:
                with self.subTest(code=code,key=key):
                    changed=copy.deepcopy(g);changed[key]=value
                    with self.assertRaises(ValueError):adapter.fill_candidate_document(changed,'application',DATA)
            for key,value in [('annex_id','foreign'),('id','foreign'),('sha256','0'*64),('official_kind','sample'),
                    ('file','../foreign.docx'),('download_url','https://foreign.example/form.docx')]:
                with self.subTest(code=code,template=key):
                    changed=copy.deepcopy(g);changed['templates'][0][key]=value
                    with self.assertRaises(ValueError):adapter.fill_candidate_document(changed,'application',DATA)
            with self.assertRaises(ValueError):adapter.fill_candidate_document(g,'authorization',DATA)

    def test_changed_bytes_at_correct_canonical_path_are_rejected_before_filling(self):
        g=guide('320981');template=g['templates'][0]
        with tempfile.TemporaryDirectory() as temp:
            target=Path(temp)/template['file'];target.parent.mkdir(parents=True);target.write_bytes((ROOT/template['file']).read_bytes()+b'tampered')
            self.assertNotEqual(hashlib.sha256(target.read_bytes()).hexdigest(),template['sha256'])
            with patch.object(adapter,'ROOT',Path(temp)):
                with self.assertRaisesRegex(ValueError,'SHA'):adapter.fill_candidate_document(g,'application',DATA)

    def test_public_only_copy_needs_no_sealed_capture_and_missing_original_blocks(self):
        for code in OWN:
            with self.subTest(code=code),tempfile.TemporaryDirectory() as temp:
                g=guide(code);public_root=Path(temp)
                # The isolated copy contains only own public blanks: no reports,
                # raw captures, private data directories, model or DB fixtures.
                for t in g['templates']:
                    destination=public_root/t['file'];destination.parent.mkdir(parents=True,exist_ok=True)
                    shutil.copyfile(ROOT/t['file'],destination)
                with patch.object(adapter,'ROOT',public_root):
                    self.assertEqual(set(adapter.candidate_contract(g)['documents']),{'application','evidence'})
                    output=BytesIO();adapter.fill_candidate_document(g,'application',DATA).save(output)
                    self.assertTrue(output.getvalue().startswith(b'PK'))
                    (public_root/g['templates'][0]['file']).unlink()
                    with self.assertRaisesRegex(ValueError,'原件不可用'):adapter.candidate_contract(g)

if __name__=='__main__':unittest.main()
