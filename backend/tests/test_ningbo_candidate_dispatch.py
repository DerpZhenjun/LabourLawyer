"""Real shared generation using own Ningbo sources in an isolated ASGI test DB."""
import io
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient

ROOT=Path(__file__).resolve().parents[2];sys.path.insert(0,str(ROOT/'backend'))
from regional_zhejiang_oct29 import parse_capture
from regional_adapters import with_reference_capabilities
from regional_guides import RegionalStore
from routes import regional_guides as routes


class NingboDispatchTests(unittest.TestCase):
    def test_own_source_generation_download_and_cross_district_rejection(self):
        with tempfile.TemporaryDirectory(dir=ROOT/'.runtime') as temp:
            store=RegionalStore(Path(temp)/'guides.sqlite')
            guides=[store.put_guide(with_reference_capabilities(parse_capture(name))) for name in ['haishu','yinzhou']]
            self.assertNotEqual(guides[0]['source_id'],guides[1]['source_id'])
            app=FastAPI();app.include_router(routes.router,prefix='/regional')
            with patch.object(routes,'store',return_value=store),TestClient(app) as client:
                for g in guides:
                    self.assertEqual(sum(t.get('can_generate') is True for t in g['templates']),2)
                    for kind in ['application','evidence']:
                        t=next(t for t in g['templates'] if t.get('document_types')==[kind])
                        data=({'applicantName':'虚构核验劳动者','requests':['请求支付虚构工资一百元。'],'copyCount':''} if kind=='application'
                              else {'evidenceList':[{'name':'虚构独立证据','source':'虚构来源','purpose':'核对本地填写','pageCount':'3页'}]})
                        payload={k:g[k] for k in ['province','city','district','source_id','version']}
                        payload.update(template_id=t['id'],kind=kind,mode='blank',data=data)
                        r=client.post('/regional/generate',json=payload)
                        self.assertEqual(r.status_code,200,r.text[:200] if r.status_code!=200 else '')
                        self.assertEqual(r.headers['X-Document-Origin'],'official-blank-filled')
                        self.assertEqual(r.headers['X-Guide-Source-ID'],g['source_id'])
                        doc=Document(io.BytesIO(r.content));text='\n'.join(p.text for p in doc.paragraphs)+''.join(c.text for tab in doc.tables for row in tab.rows for c in row.cells)
                        self.assertIn('虚构核验劳动者' if kind=='application' else '虚构独立证据',text)
                        download=client.get('/regional/templates/'+t['id'],params={k:g[k] for k in ['province','city','district']})
                        self.assertEqual(download.status_code,200,download.text[:100] if download.status_code!=200 else '')
                        self.assertEqual(download.content,(ROOT/t['file']).read_bytes())
                        payload['source_id']=next(other['source_id'] for other in guides if other['source_id']!=g['source_id'])
                        self.assertEqual(client.post('/regional/generate',json=payload).status_code,409)


if __name__=='__main__':unittest.main()
