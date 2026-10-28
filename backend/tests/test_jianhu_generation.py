"""Real local files + isolated API; reject locality and blank/sample confusion."""
import copy
import io
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from regional_adapters import with_reference_capabilities
from regional_guides import RegionalStore
from routes.regional_guides import router

ROOT=Path(__file__).resolve().parents[2]
FIXTURE=ROOT/'reports/nationwide-localization-20261023-night/cross-province/jiangsu-jianhu/guide.json'
APP='985f4770-1ab2-4264-9482-0f8528b16513'
EVIDENCE='586cf49f-714b-4338-81ce-08a2956d7aa6'


class JianhuGeneration(unittest.TestCase):
    def guide(self):return json.loads(FIXTURE.read_text('utf8'))

    def test_only_two_own_blank_forms_gain_their_own_document_kind(self):
        guide=with_reference_capabilities(self.guide())
        enabled={t['id']:t['document_types'] for t in guide['templates'] if t.get('can_generate')}
        self.assertEqual(enabled,{APP:['application'],EVIDENCE:['evidence']})
        for field,value in [('city','无锡市'),('district','阜宁县'),('source_url','https://www.jszwfw.gov.cn/other')]:
            changed=self.guide();changed[field]=value
            self.assertFalse(any(t.get('can_generate') for t in with_reference_capabilities(changed)['templates']))

    def test_actual_api_fills_own_files_and_preserves_other_package_parts(self):
        with tempfile.TemporaryDirectory() as temp:
            store=RegionalStore(Path(temp)/'regional.sqlite')
            guide=store.put_guide(with_reference_capabilities(self.guide()))
            app=FastAPI();app.include_router(router,prefix='/api/regional-guides')
            with patch('routes.regional_guides.store',return_value=store),TestClient(app) as client:
                for kind,identity,data in [('application',APP,{'applicantName':'合成测试甲','respondentName':'合成测试公司',
                    'requests':['支付测试工资100元。','支付测试加班费50元。'],'factsReason':'本文件仅用于合成测试，不用于真实申请。',
                    'acceptMediation':'待本人选择','legalAid':'待本人确认','copyCount':2,'evidenceCopyCount':3}),
                    ('evidence',EVIDENCE,{'evidenceList':[{'name':'测试劳动合同','source':'本人保存','purpose':'证明测试劳动关系','originalType':'复印件','pageCount':'3页'}]})]:
                    payload={k:guide[k] for k in ['province','city','district','source_id','version']}
                    payload.update(template_id=identity,kind=kind,mode='blank',data=data)
                    response=client.post('/api/regional-guides/generate',json=payload)
                    self.assertEqual(response.status_code,200,response.text if response.status_code!=200 else '')
                    self.assertEqual(response.headers['X-Document-Origin'],'official-blank-filled')
                    document=Document(io.BytesIO(response.content))
                    text='\n'.join(p.text for p in document.paragraphs)+'\n'+'\n'.join(c.text for t in document.tables for r in t.rows for c in r.cells)
                    self.assertIn('合成测试甲' if kind=='application' else '测试劳动合同',text)
                    if kind=='application':
                        self.assertNotIn('用人单位申请',text)
                        self.assertEqual(text.count('被申请人：'),1)
                        self.assertIn('待本人选择',text)
                        self.assertIn('副本2份；相关证据3份',text)
                        closing=next(i for i,p in enumerate(document.paragraphs) if ''.join(p.text.split())=='此致')
                        signature=next(i for i,p in enumerate(document.paragraphs[closing:],closing) if '申请人' in p.text and '签名' in p.text)
                        for p in document.paragraphs[closing:signature+1]:
                            self.assertTrue(p.paragraph_format.keep_with_next,p.text)
                    else:
                        self.assertEqual(len(document.tables[0].rows),2)
                        self.assertEqual(document.tables[0].rows[1].cells[2].text,'本人保存')
                        self.assertEqual(document.tables[0].rows[1].cells[5].text,'3')
                    source=next(t for t in guide['templates'] if t['id']==identity)
                    with ZipFile(ROOT/source['file']) as original,ZipFile(io.BytesIO(response.content)) as final:
                        self.assertEqual(set(original.namelist()),set(final.namelist()))
                        for part in original.namelist():
                            if part!='word/document.xml':self.assertEqual(original.read(part),final.read(part),part)
                    wrong=copy.deepcopy(payload);wrong['district']='阜宁县'
                    self.assertEqual(client.post('/api/regional-guides/generate',json=wrong).status_code,409)
                    wrong=copy.deepcopy(payload);wrong['template_id']=EVIDENCE if kind=='application' else APP
                    self.assertEqual(client.post('/api/regional-guides/generate',json=wrong).status_code,409)
                    wrong=copy.deepcopy(payload);wrong['mode']='reference_adaptation'
                    self.assertEqual(client.post('/api/regional-guides/generate',json=wrong).status_code,409)
                    wrong=copy.deepcopy(payload);wrong['version']+=1
                    self.assertEqual(client.post('/api/regional-guides/generate',json=wrong).status_code,409)
                    if kind=='application':
                        blank=copy.deepcopy(payload);blank['data']={'applicantName':'仅测试留空份数','copyCount':'','evidenceCopyCount':''}
                        empty=client.post('/api/regional-guides/generate',json=blank)
                        self.assertEqual(empty.status_code,200)
                        footer='\n'.join(p.text for p in Document(io.BytesIO(empty.content)).paragraphs)
                        self.assertRegex(footer,r'副本\s+份；相关证据\s+份')


if __name__=='__main__':unittest.main()
