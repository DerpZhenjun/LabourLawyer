"""Offline Chunan own-source acceptance; never opens a production capability."""
from copy import deepcopy
from io import BytesIO
import importlib.util
import json
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'backend'))


class ChunanTests(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(importlib.util.find_spec('regional_zhejiang_chunan_oct30'),
                             'Chunan requires an independent own-source adapter')
        import regional_zhejiang_chunan_oct30 as provider
        self.p = provider

    def own(self, root=ROOT):
        return self.p.parse_capture(root=root)

    def template(self, guide, kind):
        return next(t for t in guide['templates'] if t['material_id'] == self.p.BLANKS[kind]['material'])

    def fill(self, kind, data, guide=None, root=ROOT):
        guide = guide or self.own(root)
        return self.p.fill_offline_document(guide, self.template(guide, kind), kind, data, root=root)

    def public_copy(self):
        tmp = tempfile.TemporaryDirectory(); self.addCleanup(tmp.cleanup)
        root = Path(tmp.name)
        relative = self.p.REPORT + '/PUBLIC-WHITELIST.json'
        white = json.loads((ROOT / relative).read_text('utf8'))
        for name in [r['path'] for r in white['files']] + [relative]:
            target = root / name; target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(ROOT / name, target)
        return root

    def test_own_scope_materials_and_empty_window_phone_preserved_without_publication(self):
        g = self.own()
        self.assertEqual((g['region_code'], g['district'], g['source_id']),
                         ('330127', '淳安县', '3fb683d473fe253a3a117d46'))
        self.assertEqual(len(g['materials']), 5)
        self.assertEqual(g['fields']['窗口联系电话']['raw_value'], '')
        self.assertEqual(g['fields']['咨询电话']['raw_value'], '0571-89601299')
        self.assertFalse(g['published']); self.assertFalse(g['can_generate'])
        self.assertTrue(all(not t['can_generate'] for t in g['templates']))

    def test_same_annex_bytes_never_authorize_other_scope_or_mutated_facts(self):
        g = self.own(); t = self.template(g, 'application')
        for key in ['district', 'region_code', 'source_id', 'source_url', 'service_code', 'authority']:
            bad = deepcopy(g); bad[key] += '-foreign'
            with self.subTest(key=key), self.assertRaises(ValueError):
                self.p.fill_offline_document(bad, t, 'application', {}, root=ROOT)
        bad = deepcopy(g); bad['fields']['窗口联系电话']['raw_value'] = '0571-89601299'
        with self.assertRaises(ValueError): self.fill('application', {}, bad)
        for key in ['material_id', 'official_annex_id', 'file', 'sha256', 'source_url', 'region_code']:
            bad = deepcopy(t); bad[key] += '-foreign'
            with self.subTest(key=key), self.assertRaises(ValueError):
                self.p.fill_offline_document(g, bad, 'application', {}, root=ROOT)

    def test_explicit_public_root_has_no_checkout_fallback_and_all_30_records_are_checked(self):
        root = self.public_copy()
        self.assertEqual(self.own(root)['source_id'], self.own()['source_id'])
        self.assertTrue(self.fill('application', {'applicantName': '虚构测试劳动者'}, root=root).tables)
        with self.assertRaises(ValueError): self.p.parse_capture(root=None)
        (root / self.p.REPORT / 'faq.png').write_bytes(b'tampered')
        with self.assertRaises(ValueError): self.own(root)

    def test_filled_copies_preserve_other_package_parts_originals_and_hand_signatures(self):
        from docx import Document
        g = self.own()
        payloads = {'application': {'applicantName': '虚构测试劳动者', 'respondentName': '虚构测试公司',
                                    'requests': ['仅用于测试请求'], 'factsReason': '仅用于版式验证，不作真实案件事实。'},
                    'evidence': {'caseNumber': '虚构案号', 'evidenceList': [
                        {'name': '虚构劳动合同', 'pageCount': '2', 'source': '本人提供', 'purpose': '仅用于测试'}]}}
        for kind, data in payloads.items():
            t = self.template(g, kind); original = (ROOT / t['file']).read_bytes()
            out = BytesIO(); self.fill(kind, data, g).save(out)
            with ZipFile(BytesIO(original)) as a, ZipFile(BytesIO(out.getvalue())) as b:
                self.assertEqual(set(a.namelist()), set(b.namelist()))
                for part in a.namelist():
                    if part != 'word/document.xml': self.assertEqual(a.read(part), b.read(part), part)
            doc = Document(BytesIO(out.getvalue()))
            text = '\n'.join(p.text for p in doc.paragraphs) + '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
            self.assertIn('虚构', text)
            self.assertEqual((ROOT / t['file']).read_bytes(), original)
            if kind == 'application':
                self.assertEqual(len(doc.tables), 1, 'Only the worker application is supported')
                self.assertNotIn('申请人（签名）：虚构测试劳动者', text)

    def test_sample_wrong_kind_unknown_field_and_unit_mode_are_rejected(self):
        g = self.own(); sample = next(t for t in g['templates'] if t['official_kind'] == 'sample')
        with self.assertRaises(ValueError): self.p.fill_offline_document(g, sample, 'application', {}, root=ROOT)
        with self.assertRaises(ValueError): self.p.fill_offline_document(g, self.template(g, 'application'), 'evidence', {}, root=ROOT)
        for kind in ['application', 'evidence']:
            with self.assertRaises(ValueError): self.fill(kind, {'unknown': ''})
        with self.assertRaises(ValueError): self.p.fill_offline_document(g, self.template(g, 'application'), 'application', {}, root=ROOT, applicant_mode='employer')


if __name__ == '__main__':
    unittest.main()
