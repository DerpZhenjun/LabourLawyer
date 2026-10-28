"""Exact-locality generation against the retained, independently sourced annexes."""
import io
import json
import os
import sys
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'backend'))
from regional_guides import RegionalStore


class OtherDistrictDocuments(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.store = RegionalStore(Path(temporary.name) / 'regional.sqlite')
        environment = patch.dict(os.environ, {'LABOURLAWYER_REGIONAL_DB': str(self.store.path)})
        environment.start()
        self.addCleanup(environment.stop)
        from routes.regional_guides import router
        app = FastAPI()
        app.include_router(router, prefix='/api/regional-guides')
        self.client = TestClient(app)
        self.sources = json.loads((ROOT / 'reports/government-guide-localization/SOURCE-INVENTORY.json').read_text(encoding='utf8'))['regions']

    def publish(self, district):
        guide = deepcopy(next(g for g in self.sources if g['district'] == district))
        template = next(t for t in guide['templates'] if t['original_name'] == '仲裁申请书(模板).docx')
        # Published adapter capability, not a mock of rendering or HTTP responses.
        template.update(can_generate=True, adapter='sichuan_combined_reference_v1', document_types=['application', 'evidence'])
        return self.store.put_guide(guide)

    def request(self, guide, kind='application', data=None):
        template = next(t for t in guide['templates'] if t.get('can_generate'))
        return {**{k: guide[k] for k in ['province', 'city', 'district', 'source_id', 'version']},
                'template_id': template['id'], 'kind': kind, 'mode': 'reference_adaptation', 'data': data or {}}

    def test_jinjiang_and_wuhou_generate_from_their_own_annex_and_keep_other_parts(self):
        for district, annex in [('锦江区', '8cc65371b3224a63ba25a1af9d49c54d'), ('武侯区', '4452a5a18dc64c0f98b84461765b56d8')]:
            with self.subTest(district=district):
                guide = self.publish(district)
                for kind in ['application', 'evidence']:
                    data = {'applicantName': district + '测试劳动者', 'committeeName': district + '测试委员会', 'requests': ['请求支付测试工资100元'], 'factsReason': '这是合成文书排版测试，不用于实际申请。'} if kind == 'application' else {'evidenceList': [{'name': district + '测试劳动合同', 'itemCount': '1', 'pageCount': '3', 'originalType': '复印件', 'purpose': '证明劳动关系。'}]}
                    payload = self.request(guide, kind, data)
                    self.assertEqual(payload['template_id'], annex)
                    response = self.client.post('/api/regional-guides/generate', json=payload)
                    self.assertEqual(response.status_code, 200, response.text if response.status_code != 200 else '')
                    self.assertEqual(response.headers['x-guide-source-id'], guide['source_id'])
                    self.assertEqual(response.headers['x-document-origin'], 'official-sample-adaptation-not-blank')
                    doc = Document(io.BytesIO(response.content))
                    if kind == 'application':
                        text = '\n'.join(p.text for p in doc.paragraphs)
                        self.assertIn(district + '测试劳动者', text)
                        self.assertIn(district + '测试委员会', text)
                        self.assertNotIn('天府新区', text)
                        self.assertNotIn('授权委托书', text)
                    else:
                        self.assertEqual(doc.tables[0].cell(1, 1).text, district + '测试劳动合同')
                        self.assertEqual(len(doc.tables[0].rows), 2)
                    original = next(t for t in guide['templates'] if t['id'] == annex)
                    with ZipFile(ROOT / original['file']) as before, ZipFile(io.BytesIO(response.content)) as after:
                        self.assertEqual(set(before.namelist()), set(after.namelist()))
                        for part in before.namelist():
                            if part != 'word/document.xml':
                                self.assertEqual(before.read(part), after.read(part), part)

    def test_identical_file_does_not_authorize_a_neighbor_annex_id(self):
        jinjiang, wuhou = self.publish('锦江区'), self.publish('武侯区')
        payload = self.request(jinjiang)
        payload['template_id'] = self.request(wuhou)['template_id']
        self.assertEqual(self.client.post('/api/regional-guides/generate', json=payload).status_code, 409)

    def test_adapter_name_alone_does_not_authorize_unknown_district(self):
        guide = self.publish('锦江区')
        guide['district'] = '未核验测试区'
        guide = self.store.put_guide(guide)
        self.assertEqual(self.client.post('/api/regional-guides/generate', json=self.request(guide)).status_code, 409)

    def test_both_districts_can_generate_empty_reference_but_cannot_claim_official_blank(self):
        for district in ['锦江区', '武侯区']:
            guide = self.publish(district)
            for kind in ['application', 'evidence']:
                payload = self.request(guide, kind)
                self.assertEqual(self.client.post('/api/regional-guides/generate', json=payload).status_code, 200)
                payload['mode'] = 'blank'
                self.assertEqual(self.client.post('/api/regional-guides/generate', json=payload).status_code, 409)

    def test_stale_source_version_still_rejects_generation(self):
        guide = self.publish('武侯区')
        payload = self.request(guide)
        payload['version'] += 1
        self.assertEqual(self.client.post('/api/regional-guides/generate', json=payload).status_code, 409)

    def test_neighbor_file_path_is_rejected_even_when_hash_is_identical(self):
        guide = self.publish('锦江区')
        template = next(t for t in guide['templates'] if t.get('can_generate'))
        template['file'] = 'templates/_docx/_regional/510110000000/ec3268e5fda94301ab40a7e086afe07a-仲裁申请书(模板).docx'
        guide = self.store.put_guide(guide)
        self.assertEqual(self.client.post('/api/regional-guides/generate', json=self.request(guide)).status_code, 409)

    def test_capability_publication_uses_explicit_bindings_not_province_or_hash(self):
        from regional_adapters import with_reference_capabilities
        for district in ['天府新区', '锦江区', '武侯区']:
            source = deepcopy(next(g for g in self.sources if g['district'] == district))
            enabled = with_reference_capabilities(source)
            supported = [t for t in enabled['templates'] if t.get('can_generate')]
            self.assertEqual(len(supported), 1)
            self.assertEqual(supported[0]['document_types'], ['application', 'evidence'])
            enabled['district'] = '未知测试区'
            self.assertFalse(any(t.get('can_generate') for t in with_reference_capabilities(enabled)['templates']))

    def test_application_closing_and_handwritten_signature_are_not_split_apart(self):
        from regional_documents import fill_tianfu_reference
        source = next(g for g in self.sources if g['district'] == '锦江区')
        template = next(t for t in source['templates'] if t['original_name'] == '仲裁申请书(模板).docx')
        doc = fill_tianfu_reference(ROOT / template['file'], 'application', {'factsReason': '检查较长内容与手写签名的分页。' * 20})
        closing = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == '此致')
        date = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip().replace(' ', '') == '年月日')
        self.assertGreater(date, closing)
        self.assertTrue(all(p.paragraph_format.keep_with_next for p in doc.paragraphs[closing:date]))


if __name__ == '__main__':
    unittest.main()
