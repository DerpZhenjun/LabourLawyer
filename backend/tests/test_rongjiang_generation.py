"""Own DOC-derived form filling: no borrowed locality, sample facts or signatures."""
import copy,hashlib,importlib.util,io,json,sys,tempfile,unittest
from pathlib import Path
from zipfile import ZipFile
from unittest.mock import patch
from docx import Document
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
ROOT=Path(__file__).resolve().parents[2]
BASE=ROOT/'reports/nationwide-localization-20261023-night/cross-province/jiangxi-rongjiang'
REF=ROOT/'templates/_docx/_regional/function-jiangxi-ganzhou-rongjiang/853097ba8d9f4036b61e806e44314c21-converted.docx'

class RongjiangGenerationTests(unittest.TestCase):
    def setUp(self):
        self.guide=json.loads((BASE/'guide.json').read_text(encoding='utf8'));self.template=self.guide['templates'][0]

    def fill(self,kind,data):
        self.assertIsNotNone(importlib.util.find_spec('regional_rongjiang_documents'),'Own Rongjiang filler is missing')
        from regional_rongjiang_documents import fill_rongjiang_document
        stream=io.BytesIO();fill_rongjiang_document(self.guide,self.template,kind,ROOT/self.template['file'],data).save(stream)
        return Document(io.BytesIO(stream.getvalue())),stream.getvalue()

    def test_application_maps_own_labels_and_keeps_local_addressee_unsigned_and_notes(self):
        doc,content=self.fill('application',{'applicantName':'蓉江测试甲','gender':'女','birthDate':'1990-02-03',
            'nationality':'汉族','idNumber':'仅用于测试','phone':'测试电话','idAddress':'测试户籍地址',
            'applicantMailingAddress':'测试有效地址','applicantPostalCode':'341000','respondentName':'蓉江测试公司',
            'socialCreditCode':'测试代码','registeredAddress':'测试注册地址','actualOfficeAddress':'测试办公地址',
            'respondentPostalCode':'341001','companyContactPhone':'测试公司电话','legalRepName':'测试负责人',
            'legalRepPosition':'负责人','committeeName':'不允许替换为邻区委员会',
            'requests':['测试第一请求','测试第二请求','测试第三请求'],
            'factsReason':'这是仅用于测试的劳动争议事实，不用于真实申请。','applicationDate':'2026-10-24'})
        text='\n'.join(p.text for p in doc.paragraphs)
        for expected in ['蓉江测试甲','1990年2月3日','测试有效地址','测试注册地址','测试办公地址','341001','测试公司电话','测试第三请求','仅用于测试的劳动争议事实','0797-8168180']:
            self.assertIn(expected,text)
        self.assertIn('赣州蓉江新区劳动人事争议仲裁委员会',text)
        self.assertNotIn('不允许替换为邻区委员会',text);self.assertNotIn('证据目录',text);self.assertFalse(doc.tables)
        closing=next(p.text for p in doc.paragraphs if p.text.lstrip().startswith('申请人：') and '蓉江测试甲' not in p.text)
        self.assertNotIn('蓉江测试甲',closing)
        self.assertIn('人数+1份',text)

    def test_evidence_has_only_user_rows_and_own_five_columns_not_ten_sample_items(self):
        doc,content=self.fill('evidence',{'submitterName':'不能充当手写签名','submitDate':'2026-10-24',
            'evidenceList':[{'name':'合成合同','source':'本人保存','purpose':'合成测试劳动关系','pageCount':'3页'}]})
        table=doc.tables[0]
        self.assertEqual([c.text for c in table.rows[0].cells],['序号','证据名称','证据来源','证明目的','页数'])
        self.assertEqual(len(table.rows),2)
        self.assertEqual([c.text for c in table.rows[1].cells],['1','合成合同','本人保存','合成测试劳动关系','3页'])
        text='\n'.join(p.text for p in doc.paragraphs)
        self.assertNotIn('不能充当手写签名',text);self.assertNotIn('仲裁申请书',text)
        self.assertIn('2026年10月24日',text);self.assertIn('人数+1份',text)
        empty,_=self.fill('evidence',{})
        self.assertEqual(len(empty.tables[0].rows),2)
        self.assertTrue(all(not c.text for c in empty.tables[0].rows[1].cells[1:]))

    def test_more_than_ten_rows_survive_and_preserve_only_package_parts_are_unchanged(self):
        data={'evidenceList':[{'name':f'测试证据{i}','purpose':'长说明'*30,'source':'测试来源','pageCount':str(i)} for i in range(1,13)]}
        doc,content=self.fill('evidence',data)
        self.assertEqual(len(doc.tables[0].rows),13)
        self.assertEqual(doc.tables[0].rows[-1].cells[1].text,'测试证据12')
        with ZipFile(REF) as before,ZipFile(io.BytesIO(content)) as after:
            self.assertEqual(before.namelist(),after.namelist())
            for part in before.namelist():
                if part!='word/document.xml':self.assertEqual(before.read(part),after.read(part),part)
        self.assertEqual(hashlib.sha256(REF.read_bytes()).hexdigest(),'9a7571e55b7260d17931c42936262b58e6e69971efe7254f108f9528d22591f1')

    def test_wrong_region_source_annex_and_original_are_rejected(self):
        for where,key,value in [('guide','district','南康区'),('guide','province','guangdong'),
            ('guide','source_url','https://www.ganzhou.gov.cn/'),('template','sha256','0'*64),
            ('template','download_url','https://www.ganzhou.gov.cn/another.doc'),
            ('template','file','templates/_docx/_regional/440111/other.doc'),('template','id','not-own-annex')]:
            with self.subTest(key=key,value=value):
                saved=copy.deepcopy((self.guide,self.template))
                (self.guide if where=='guide' else self.template)[key]=value
                with self.assertRaises(ValueError):self.fill('application',{})
                self.guide,self.template=saved

    def test_missing_or_replaced_working_reference_is_rejected_without_template_fallback(self):
        self.assertIsNotNone(importlib.util.find_spec('regional_rongjiang_documents'))
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'missing.docx'
            with patch('regional_rongjiang_documents.REFERENCE',str(path)):
                with self.assertRaises(ValueError):self.fill('application',{})
            path.write_bytes(b'foreign document bytes')
            with patch('regional_rongjiang_documents.REFERENCE',str(path)):
                with self.assertRaises(ValueError):self.fill('application',{})

if __name__=='__main__':unittest.main()
