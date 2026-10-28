"""Actual isolated API distinguishes original downloads from DOC-derived filling."""
import copy,hashlib,io,json,sys,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from fastapi import FastAPI
from fastapi.testclient import TestClient
from docx import Document
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from regional_guides import RegionalStore
from regional_adapters import with_reference_capabilities
from routes.regional_guides import router
ROOT=Path(__file__).resolve().parents[2]
BASE=ROOT/'reports/nationwide-localization-20261023-night/cross-province/jiangxi-rongjiang'

class RongjiangGenerationApiTests(unittest.TestCase):
    def test_two_own_components_and_original_download_have_honest_origins(self):
        raw=json.loads((BASE/'guide.json').read_text(encoding='utf8'))
        candidate=with_reference_capabilities(raw)
        self.assertEqual(candidate['templates'][0]['document_types'],['application','evidence'])
        self.assertIn('DOC',candidate['templates'][0]['generation_label'])
        self.assertEqual(candidate['verification_status'],'partial')
        self.assertEqual(candidate['channel']['online_status'],'unverified')
        self.assertIsNone(candidate['channel']['entry_url'])
        with tempfile.TemporaryDirectory() as d:
            store=RegionalStore(Path(d)/'public.sqlite');guide=store.put_guide(candidate)
            app=FastAPI();app.include_router(router,prefix='/api/regional-guides')
            with patch('routes.regional_guides.store',return_value=store),TestClient(app) as client:
                t=guide['templates'][0];params={k:guide[k] for k in ['province','city','district']}
                original=client.get('/api/regional-guides/templates/'+t['id'],params=params)
                self.assertEqual(original.status_code,200)
                self.assertEqual(hashlib.sha256(original.content).hexdigest(),'3729dc3af6b0289db06a29ee3628f6f1bd7bdb3cdab6813103310d48511260c4')
                for kind,data in [('application',{'applicantName':'合成API测试','requests':['测试请求'],'factsReason':'合成事实，不用于申请。'}),
                                  ('evidence',{'evidenceList':[{'name':'合成API证据','source':'测试来源','purpose':'测试目的','pageCount':'2'}]})]:
                    payload={k:guide[k] for k in ['province','city','district','source_id','version']}
                    payload.update(template_id=t['id'],kind=kind,mode='blank',data=data)
                    response=client.post('/api/regional-guides/generate',json=payload)
                    self.assertEqual(response.status_code,200,response.text if response.status_code!=200 else '')
                    self.assertEqual(response.headers['X-Document-Origin'],'official-blank-converted-filled')
                    self.assertEqual(response.headers['X-Guide-Source-ID'],'9324666463f34ff63056367a')
                    doc=Document(io.BytesIO(response.content))
                    if kind=='application':self.assertIn('合成API测试',''.join(p.text for p in doc.paragraphs))
                    else:self.assertEqual(len(doc.tables[0].columns),5);self.assertEqual(len(doc.tables[0].rows),2)
                    for key,value in [('district','南康区'),('source_id','453e9a1c8cb2122fe7fa910d'),('template_id','another-annex'),('mode','reference_adaptation')]:
                        wrong=copy.deepcopy(payload);wrong[key]=value
                        self.assertEqual(client.post('/api/regional-guides/generate',json=wrong).status_code,409)

    def test_forged_or_stale_adaptation_flag_cannot_enable_foreign_source(self):
        raw=json.loads((BASE/'guide.json').read_text(encoding='utf8'))
        raw['templates'][0].update(adapter='rongjiang_own_doc_blank_v1',can_generate=True,document_types=['application','evidence'])
        raw['district']='南康区'
        result=with_reference_capabilities(raw)
        self.assertFalse(result['templates'][0]['can_generate'])
        self.assertEqual(result['templates'][0]['document_types'],[])

    def test_capability_refresh_does_not_create_artificial_guide_versions(self):
        raw=json.loads((BASE/'guide.json').read_text(encoding='utf8'))
        first=with_reference_capabilities(raw)
        refreshed=with_reference_capabilities(first)
        with tempfile.TemporaryDirectory() as d:
            store=RegionalStore(Path(d)/'public.sqlite')
            original=store.put_guide(first);again=store.put_guide(refreshed)
            self.assertEqual(again['version'],original['version'])

if __name__=='__main__':unittest.main()
