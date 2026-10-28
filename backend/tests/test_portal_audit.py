import json
from pathlib import Path
import sys
import tempfile
import unittest
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'backend'))
from portal_audit_capture import KEYWORDS,search_outcome,jurisdiction_choices
from portal_audit_vision import parse_answer,cached_images,completed_review_matches
from portal_audit import MODEL,sha,validate_capture_coverage
import importlib.util
spec=importlib.util.spec_from_file_location('portal_bridge_tests',ROOT/'backend/template_acquisition/openclaw_bridge.py')
bridge=importlib.util.module_from_spec(spec);spec.loader.exec_module(bridge)
_completed_model_output,write_isolated_config=bridge._completed_model_output,bridge.write_isolated_config

class PortalAuditTests(unittest.TestCase):
    def test_capture_completion_requires_each_requested_search_and_verified_screenshot(self):
        with tempfile.TemporaryDirectory(dir=ROOT/'.runtime') as tmp:
            folder=Path(tmp);region=folder/'sample';region.mkdir()
            searches=[]
            for index,keyword in enumerate(KEYWORDS,1):
                picture=region/f'search-{index:02d}.png';picture.write_bytes(f'picture {index}'.encode())
                searches.append(dict(stage=f'search-{index:02d}',keyword=keyword,screenshot=picture.name,screenshot_sha256=sha(picture)))
            item=dict(slug='sample',searches=searches)
            validate_capture_coverage(folder,['sample'],[item])
            with self.assertRaisesRegex(RuntimeError,'地区不完整'):
                validate_capture_coverage(folder,['sample','other'],[item])
            with self.assertRaisesRegex(RuntimeError,'三项检索不完整'):
                validate_capture_coverage(folder,['sample'],[dict(item,searches=searches[:2])])
            (region/'search-03.png').write_bytes(b'changed')
            with self.assertRaisesRegex(RuntimeError,'截图缺失或校验不符'):
                validate_capture_coverage(folder,['sample'],[item])

    def test_complete_cache_requires_raw_provider_model_and_unchanged_picture(self):
        with tempfile.TemporaryDirectory(dir=ROOT/'.runtime') as tmp:
            directory=Path(tmp);picture=directory/'search.png';picture.write_bytes(b'real screenshot bytes')
            digest=sha(picture)
            v=dict(stage='search-01',state='complete',model=MODEL,screenshot_sha256=digest)
            payload=dict(result=dict(ok=True,provider='ollama',model='qwen3.8:27b'),verified_json=[v],input_sha256=[digest],base_url='http://10.66.66.1:11434')
            record=directory/'openclaw-image-search.json';record.write_text(json.dumps(payload),encoding='utf8')
            existing=dict(state='complete',model=MODEL,images=[v]);required=[dict(stage='search-01',screenshot=picture.name)]
            cache=cached_images(directory)
            self.assertTrue(completed_review_matches(directory,existing,required,cache))
            self.assertEqual(cache[digest]['model_server'],'http://10.66.66.1:11434')
            payload['result']['model']='qwen3.5:27b';record.write_text(json.dumps(payload),encoding='utf8')
            self.assertFalse(completed_review_matches(directory,existing,required,cached_images(directory)))
            picture.write_bytes(b'changed');self.assertFalse(completed_review_matches(directory,existing,required,cache))
            picture.unlink();self.assertFalse(completed_review_matches(directory,existing,required,cache))

    def test_unknown_legacy_server_is_not_labeled_local_or_remote(self):
        with tempfile.TemporaryDirectory(dir=ROOT/'.runtime') as tmp:
            directory=Path(tmp);digest='a'*64
            v=dict(stage='search-01',state='complete',model=MODEL,screenshot_sha256=digest)
            (directory/'openclaw-review-01.json').write_text(json.dumps(dict(result=dict(ok=True,provider='ollama',model='qwen3.8:27b'),verified_json=[v],input_sha256=[digest])),encoding='utf8')
            self.assertEqual(cached_images(directory)[digest]['model_server'],'历史记录未注明服务器')

    def test_remote_cache_wins_over_legacy_record_regardless_of_glob_order(self):
        with tempfile.TemporaryDirectory(dir=ROOT/'.runtime') as tmp:
            directory=Path(tmp);digest='c'*64
            entry=dict(stage='search-01',state='complete',model=MODEL,screenshot_sha256=digest)
            payload=dict(result=dict(ok=True,provider='ollama',model='qwen3.8:27b'),verified_json=[entry],input_sha256=[digest])
            remote=directory/'openclaw-images-01.json';legacy=directory/'openclaw-review-01.json'
            remote.write_text(json.dumps(dict(payload,base_url='http://10.66.66.1:11434')),encoding='utf8')
            legacy.write_text(json.dumps(payload),encoding='utf8')
            self.assertEqual(cached_images(directory)[digest]['model_server'],'http://10.66.66.1:11434')
            remote.unlink();remote.write_text(json.dumps(dict(payload,base_url='http://10.66.66.1:11434')),encoding='utf8')
            self.assertEqual(cached_images(directory)[digest]['model_server'],'http://10.66.66.1:11434')

    def test_ai_word_corrects_legacy_guess_from_raw_provider_record(self):
        from portal_audit_report import build_ai_report
        from docx import Document
        with tempfile.TemporaryDirectory(dir=ROOT/'.runtime') as tmp:
            folder=Path(tmp);region=folder/'sample';region.mkdir()
            digest='b'*64
            v=dict(stage='search-01',keyword='劳动仲裁',state='complete',model=MODEL,screenshot_sha256=digest,model_server='http://127.0.0.1:11434',page_kind='搜索结果',entrance_visible=False,template_names=[])
            (region/'openclaw-review-01.json').write_text(json.dumps(dict(result=dict(ok=True,provider='ollama',model='qwen3.8:27b'),verified_json=[v],input_sha256=[digest])),encoding='utf8')
            item=dict(slug='sample',name='样本省',observations=[],ai_review=dict(state='complete',images=[v]))
            doc=Document(build_ai_report(folder,[item]))
            cells='\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
            self.assertIn('服务器未记载',cells)
            self.assertNotIn('前序本机复核',cells)
            self.assertEqual(json.loads((folder/'results.json').read_text(encoding='utf8'))[0]['ai_review']['images'][0]['model_server'],'历史记录未注明服务器')

    def test_public_picker_keeps_each_administrative_level(self):
        self.assertEqual(jurisdiction_choices('江苏省无锡市锡山区'),['江苏省','无锡市','锡山区'])
        self.assertEqual(jurisdiction_choices('内蒙古自治区鄂尔多斯市'),['内蒙古自治区','鄂尔多斯市'])
        self.assertEqual(jurisdiction_choices('北京市朝阳区潘家园街道'),['北京市','朝阳区'])

    def test_failed_and_unsubmitted_searches_cannot_report_no_entrance(self):
        for obs in [dict(page_problem='certificate'),dict(enter_submitted=False),dict(enter_submitted=True,result_page_confirmed=False)]:
            self.assertNotEqual(search_outcome(obs),'本次搜索结果未定位到相关入口')
        self.assertEqual(search_outcome(dict(enter_submitted=True,result_page_confirmed=True,candidate_links=[])),'本次搜索结果未定位到相关入口')

    def test_ai_must_cover_all_supplied_screenshots(self):
        payload={'result':{'outputs':[{'text':json.dumps({'images':[{'stage':'search-01','entrance_visible':False,'template_names':[]}]})}]}}
        with self.assertRaisesRegex(ValueError,'全部输入截图'):parse_answer(payload,['search-01','search-02','search-03'])

    def test_requested_model_must_match_actual_provider_response(self):
        output=json.dumps(dict(ok=True,provider='ollama',model='qwen3.5:27b',outputs=[dict(text='{}')]))
        self.assertIsNone(_completed_model_output(output,'qwen3.8:27b'))

    def test_compact_model_response_maps_image_numbers_without_inventing_fields(self):
        answer={'images':[{'stage':2,'kind':'错误页','entry':False,'names':[],'note':'无法连接'}, {'stage':1,'kind':'搜索结果','entry':True,'names':[],'note':''}]}
        payload={'result':{'outputs':[{'text':json.dumps(answer)}]}}
        reviews=parse_answer(payload,['search-01','search-02'])
        self.assertEqual(reviews[0]['stage'],'search-02')
        self.assertFalse(reviews[0]['entrance_visible'])
        self.assertEqual(reviews[0]['page_kind'],'错误页')
        del answer['images'][0]['entry']
        with self.assertRaises(ValueError):parse_answer({'result':{'outputs':[{'text':json.dumps(answer)}]}},['search-01','search-02'])

    def test_portal_model_isolated_configuration(self):
        with tempfile.TemporaryDirectory(dir=ROOT/'.runtime') as tmp:
            path,env=write_isolated_config(Path(tmp),'qwen3.8:27b');config=json.loads(path.read_text(encoding='utf8'))
            self.assertEqual(config['agents']['defaults']['imageModel']['primary'],'ollama/qwen3.8:27b')
            self.assertFalse(config['browser']['noSandbox']);self.assertTrue(config['browser']['headless'])
            self.assertEqual(config['models']['providers']['ollama']['baseUrl'],'http://10.66.66.1:11434')
            self.assertNotIn('num_gpu',config['models']['providers']['ollama']['models'][0]['params'])
            self.assertTrue(Path(env['TEMP']).is_relative_to(ROOT))

    def test_report_omits_deleted_sections_and_sets_explicit_chinese_fonts(self):
        from portal_audit_report import build_report
        from docx import Document
        with tempfile.TemporaryDirectory(dir=ROOT/'.runtime') as tmp:
            folder=Path(tmp);(folder/'sample').mkdir()
            searches=[dict(stage=f'search-{i:02d}',keyword=k,enter_submitted=True,outcome='本次结果未定位到入口',final_url='https://example.gov.cn/search') for i,k in enumerate(KEYWORDS,1)]
            x=dict(slug='sample',name='样本省',portal_url='https://example.gov.cn/',documents={},scope='样本省',result='待复核',searches=searches,observations=searches)
            doc=Document(build_report(folder,[x]));text='\n'.join(p.text for p in doc.paragraphs)
            for excluded in ('检索方法与复核要求','向主管部门反馈的建议','当时可见文字摘录'):self.assertNotIn(excluded,text)
            for keyword in KEYWORDS:self.assertIn(keyword,text)
            self.assertEqual(doc.styles['Heading 1'].font.name,'黑体')
            self.assertFalse(doc.styles.element.xpath('.//w:pBdr'))
            self.assertTrue(any('https://example.gov.cn/' in r.target_ref for r in doc.part.rels.values()))

    def test_continued_guide_replaces_stale_summary_but_separate_source_does_not(self):
        from portal_audit_report import build_report
        with tempfile.TemporaryDirectory(dir=ROOT/'.runtime') as tmp:
            folder=Path(tmp);(folder/'sample').mkdir()
            searches=[dict(stage=f'search-{i:02d}',keyword=k,enter_submitted=True,result_page_confirmed=True,candidate_links=[]) for i,k in enumerate(KEYWORDS,1)]
            control=dict(stage='district-real-templates',final_url='https://example.gov.cn/guide',template_links=[dict(text='申请书',kind='sample')],templates_detected=True,report_include=True)
            x=dict(slug='sample',name='样本省',portal_url='https://example.gov.cn/',documents={},scope='样本省',result='旧的未定位结论',searches=searches,observations=searches+[control])
            build_report(folder,[x]);self.assertEqual(x['result_code'],'template_candidate_visible')
            control['stage']='official-source-templates'
            build_report(folder,[x]);self.assertEqual(x['result_code'],'not_located_in_limited_search')
            control.update(stage='district-real-templates',template_links=[],template_area_visible=True,material_names=['仲裁申请书'])
            from docx import Document
            doc=Document(build_report(folder,[x]));text='\n'.join(p.text for p in doc.paragraphs)
            self.assertIn('本次没有检测到文书模板。',text);self.assertNotIn('检测到文书下载控件。',text)

    def test_ai_report_distinguishes_servers_and_rejected_visual_claims(self):
        from portal_audit_report import build_ai_report
        from docx import Document
        with tempfile.TemporaryDirectory(dir=ROOT/'.runtime') as tmp:
            folder=Path(tmp);(folder/'sample').mkdir()
            (folder/'manual-reviews.json').write_text(json.dumps({'sample':{'search-01':{'accept_ai_result':False,'conclusion':'其他类型仲裁，未确认劳动仲裁入口'}}}),encoding='utf8')
            reviews=[dict(stage='search-01',keyword='劳动仲裁',state='complete',model='ollama/qwen3.8:27b',entrance_visible=True,template_names=['错误模板名称'],page_kind='搜索结果',model_server='http://127.0.0.1:11434'),dict(stage='search-02',keyword='劳动人事争议仲裁',state='complete',entrance_visible=False,template_names=[],page_kind='搜索结果',model_server='http://10.66.66.1:11434')]
            item=dict(slug='sample',name='样本省',observations=[],ai_review=dict(state='complete',images=reviews))
            doc=Document(build_ai_report(folder,[item]))
            text='\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
            self.assertIn('Ubuntu 复核',text);self.assertIn('前序本机复核',text)
            self.assertIn('原AI判读未采纳',text);self.assertNotIn('错误模板名称',text)

if __name__=='__main__':unittest.main()
