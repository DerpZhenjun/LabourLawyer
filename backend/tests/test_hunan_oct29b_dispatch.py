"""Synthetic own-source Hunan generation, never a province fallback."""
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from regional_guides import RegionalStore
from regional_adapters import with_reference_capabilities
from regional_hunan_oct29b import ROOT, parse_capture, ADAPTER
from routes import regional_guides as routes


class HunanDispatchTests(unittest.TestCase):
    def test_two_own_sources_generate_only_worker_application_and_reject_hidden_fields(self):
        with tempfile.TemporaryDirectory(dir=ROOT/'.runtime') as folder:
            store=RegionalStore(Path(folder)/'guides.sqlite')
            guides=[store.put_guide(with_reference_capabilities(parse_capture(n))) for n in ['tianxin','xiangjiang']]
            self.assertNotEqual(guides[0]['source_id'],guides[1]['source_id'])
            app=FastAPI();app.include_router(routes.router,prefix='/regional')
            with patch.object(routes,'store',return_value=store),TestClient(app) as client:
                for g in guides:
                    ready=[t for t in g['templates'] if t.get('can_generate')]
                    self.assertEqual(len(ready),1)
                    t=ready[0];self.assertEqual(t['adapter'],ADAPTER)
                    payload={k:g[k] for k in ['province','city','district','source_id','version']}
                    payload.update(template_id=t['id'],kind='application',mode='blank',data={
                        'applicantName':'虚构湖南验收劳动者','occupation':'测试职业',
                        'workUnit':'虚构工作单位','requests':['请求支付虚构工资。'],
                        'factsReason':'这是一段虚构软件验收事实。',
                        'evidenceList':[{'name':'虚构劳动合同','source':'虚构来源','purpose':'虚构证明事项'}]})
                    r=client.post('/regional/generate',json=payload)
                    self.assertEqual(r.status_code,200,r.text[:200] if r.status_code!=200 else '')
                    self.assertEqual(r.headers['X-Document-Origin'],'official-blank-converted-filled')
                    self.assertEqual(r.headers['X-Guide-Source-ID'],g['source_id'])
                    doc=Document(io.BytesIO(r.content))
                    text='\n'.join(p.text for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs)
                    for value in ['虚构湖南验收劳动者','测试职业','虚构工作单位','这是一段虚构软件验收事实。','虚构劳动合同']:
                        self.assertIn(value,text)
                    payload['data']['idNumber']='不应被悄悄忽略'
                    self.assertEqual(client.post('/regional/generate',json=payload).status_code,422)
                    del payload['data']['idNumber'];payload['kind']='evidence'
                    self.assertEqual(client.post('/regional/generate',json=payload).status_code,409)
                    payload['kind']='application';payload['source_id']=next(x['source_id'] for x in guides if x['source_id']!=g['source_id'])
                    self.assertEqual(client.post('/regional/generate',json=payload).status_code,409)


if __name__=='__main__':unittest.main()
