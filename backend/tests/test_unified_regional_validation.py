"""Shared regional validation through isolated stores and real reference output."""
from copy import deepcopy
from io import BytesIO
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'backend'))
from regional_guides import RegionalStore
from routes import regional_guides as route


class UnifiedRegionalValidationTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.db = Path(temporary.name) / 'regional.sqlite'
        environment = patch.dict(os.environ, {'LABOURLAWYER_REGIONAL_DB': str(self.db)})
        environment.start()
        self.addCleanup(environment.stop)
        self.store = RegionalStore(self.db)
        inventory = json.loads((ROOT / 'reports/government-guide-localization/SOURCE-INVENTORY.json').read_text(encoding='utf8'))
        self.store.put_guide(next(g for g in inventory['regions'] if g['district'] == '天府新区'))
        self.guide = self.store.lookup('sichuan', '成都市', '天府新区')
        self.template = next(t for t in self.guide['templates'] if t.get('can_generate'))
        self.payload = {k: self.guide[k] for k in ('province', 'city', 'district', 'source_id', 'version')}
        self.payload.update(template_id=self.template['id'], kind='application', mode='reference_adaptation', data={})
        app = FastAPI()
        app.include_router(route.router, prefix='/api/regional-guides')
        self.client = TestClient(app)

    def contract(self, **changes):
        query = {k: v for k, v in self.payload.items() if k not in {'mode', 'data'}}
        return self.client.get('/api/regional-guides/form-contract', params={**query, **changes})

    def generate(self, **changes):
        return self.client.post('/api/regional-guides/generate', json={**self.payload, **changes})

    def test_generate_rejects_neighbor_lookup_with_matching_source_and_version(self):
        # A bad lookup must not bypass own-region checks in either endpoint.
        for changes in ({'province': 'jiangsu'}, {'city': '邻市'}, {'district': '邻区'}):
            with self.subTest(changes=changes), patch.object(RegionalStore, 'lookup', return_value=deepcopy(self.guide)):
                contract = self.contract(**changes)
                self.assertEqual(contract.status_code, 409)
                generated = self.generate(**changes)
                self.assertEqual(generated.status_code, 409)
                self.assertEqual(generated.json()['detail'], '办理地区不匹配，不能使用邻区文书')

    def test_generation_rejects_ambiguous_template_identity(self):
        guide = deepcopy(self.guide)
        guide['templates'].append(deepcopy(self.template))
        with patch.object(RegionalStore, 'lookup', return_value=guide):
            self.assertEqual(self.contract().status_code, 409)
            response = self.generate()
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()['detail'], '模板与办理地区不匹配')

    def test_missing_source_keeps_distinct_read_and_generation_errors(self):
        response = self.contract(district='未核验区')
        self.assertEqual(response.status_code, 404)
        self.assertEqual(response.json()['detail'], '该地区没有已核验文书资料')
        response = self.generate(district='未核验区')
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()['detail'], '办理地区或来源版本已改变，请重新查看指南')

    def test_stale_source_and_version_keep_both_endpoints_closed(self):
        for changes in ({'source_id': 'stale'}, {'version': self.guide['version'] + 1}):
            with self.subTest(changes=changes):
                for response in (self.contract(**changes), self.generate(**changes)):
                    self.assertEqual(response.status_code, 409)
                    self.assertEqual(response.json()['detail'], '办理地区或来源版本已改变，请重新查看指南')

    def test_existing_surrounding_whitespace_in_locality_is_accepted(self):
        changes = {k: '  ' + self.guide[k] + '  ' for k in ('province', 'city', 'district')}
        self.assertEqual(self.contract(**changes).status_code, 200)
        self.assertEqual(self.generate(**changes).status_code, 200)

    def test_missing_template_keeps_endpoint_specific_errors(self):
        response = self.contract(template_id='missing')
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()['detail'], '该地区没有已核验的填写附件')
        response = self.generate(template_id='missing')
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()['detail'], '模板与办理地区不匹配')

    def test_legacy_no_contract_still_ignores_unbound_keys_and_generates_both_kinds(self):
        for kind in ('application', 'evidence'):
            with self.subTest(kind=kind):
                response = self.generate(kind=kind, data={'unbound': '', 'applicantName': '  虚构申请人  '})
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.headers['x-guide-source-id'], self.guide['source_id'])
                self.assertNotIn('x-form-contract-id', response.headers)
                document = Document(BytesIO(response.content))
                self.assertGreater(len(document.paragraphs), 0)
                if kind == 'application':
                    self.assertIn('虚构申请人', '\n'.join(p.text for p in document.paragraphs))

    def test_current_contract_generates_and_stale_contract_rejects_before_filling(self):
        response = self.contract()
        self.assertEqual(response.status_code, 200)
        contract_id = response.json()['contract_id']
        generated = self.generate(contract_id=contract_id, data={'applicantName': '虚构申请人'})
        self.assertEqual(generated.status_code, 200)
        self.assertEqual(generated.headers['x-form-contract-id'], contract_id)
        with patch.object(route, 'fill_tianfu_reference', side_effect=AssertionError('filler must not run')):
            rejected = self.generate(contract_id='0' * 64, data={'unbound': ''})
        self.assertEqual(rejected.status_code, 409)
        self.assertEqual(rejected.json()['detail'], '填写契约已改变，请重新核对当地文书')

    def test_contract_data_and_legacy_model_errors_reject_before_filling(self):
        contract = self.contract()
        self.assertEqual(contract.status_code, 200)
        contract_id = contract.json()['contract_id']
        cases = [
            {'contract_id': contract_id, 'data': {'unbound': ''}},
            {'contract_id': contract_id, 'data': {'applicantName': 'bad\x00xml'}},
            {'contract_id': contract_id, 'data': {'applicantName': 'x' * 301}},
            {'data': {'applicantName': 'x' * 301}},
        ]
        with patch.object(route, 'fill_tianfu_reference', side_effect=AssertionError('filler must not run')):
            for changes in cases:
                with self.subTest(changes=changes):
                    rejected = self.generate(**changes)
                    self.assertEqual(rejected.status_code, 422)
                    self.assertNotIn('x-form-contract-id', rejected.headers)
        self.assertEqual(rejected.json()['detail'], '填写内容过长或格式不正确，请缩短后重试')

    def test_template_adapter_and_requested_mode_keep_own_binding_errors(self):
        guide = deepcopy(self.guide)
        template = next(t for t in guide['templates'] if t['id'] == self.template['id'])
        template['adapter'] = 'unreviewed'
        with patch.object(RegionalStore, 'lookup', return_value=guide), \
             patch.object(route, 'fill_tianfu_reference', side_effect=AssertionError('filler must not run')):
            self.assertEqual(self.contract().status_code, 409)
            self.assertEqual(self.generate().status_code, 409)
        response = self.generate(mode='blank')
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()['detail'], '官方标注为样表，不能当作空白模板生成')


if __name__ == '__main__':
    unittest.main()
