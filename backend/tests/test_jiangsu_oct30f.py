"""Independent Changshu and Zhangjiagang own-source acceptance."""
from copy import deepcopy
from io import BytesIO
import importlib.util
import json
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from zipfile import ZipFile
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'backend'))


class OwnFTests(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(importlib.util.find_spec('regional_jiangsu_oct30f'),
                             'F localities require their independent own-source adapter')
        import regional_jiangsu_oct30f as provider
        self.p = provider

    def test_each_own_source_keeps_identity_materials_and_channel(self):
        for code, district, identity, visits in [
            ('320581', '常熟市', '9a24af792c956feb67bd37c6', '3'),
            ('320582', '张家港市', '95e826085ed0fec9f410cc86', '0'),
        ]:
            g = self.p.parse_capture(code)
            self.assertEqual((g['region_code'], g['district'], g['source_id']), (code, district, identity))
            self.assertEqual(g['fields']['到办事现场次数']['raw_value'], visits)
            self.assertEqual(len(g['materials']), 4)
            self.assertEqual(g['attachment_counts'], {'blank': 4, 'sample': 4, 'instruction': 1})
            self.assertEqual(g['channel']['online_status'], 'auth_required')
            self.assertFalse(g['can_generate'])
            contract = self.p.candidate_contract(g, ROOT)
            self.assertEqual(contract['documents']['application']['annex_id'], self.p.SPECS[code]['application'])

    def test_equal_template_bytes_never_allow_foreign_identity_or_template(self):
        own = self.p.parse_capture('320581'); foreign = self.p.parse_capture('320582')
        for key in ['source_id', 'source_url', 'region_code', 'district', 'authority', 'service_code', 'fields', 'materials', 'templates']:
            bad = deepcopy(own); bad[key] = foreign[key]
            with self.subTest(key=key), self.assertRaises(ValueError):
                self.p.fill_candidate_document(bad, 'application', {}, ROOT)
        for kind in ['application', 'evidence']:
            t = next(t for t in foreign['templates'] if t['id'] == self.p.SPECS['320582'][kind])
            with self.assertRaises(ValueError): self.p.blank_binding(own, t, ROOT)
        with self.assertRaises(ValueError): self.p.parse_capture('320507')

    def test_public_copy_no_checkout_fallback_and_tamper_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); report = ROOT / self.p.REPORT
            white = json.loads((report / 'PUBLIC-WHITELIST.json').read_text('utf8'))['files']
            for name in list(white) + [(self.p.REPORT / 'PUBLIC-WHITELIST.json').as_posix()]:
                dest = root / name; dest.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(ROOT / name, dest)
            g = self.p.parse_capture('320581', root / self.p.REPORT)
            self.p.fill_candidate_document(g, 'evidence', {}, root)
            path = root / g['evidence']['capture_file']; path.write_bytes(b'tampered')
            with self.assertRaises(ValueError): self.p.fill_candidate_document(g, 'evidence', {}, root)
            with self.assertRaises(ValueError): self.p.parse_capture('320581', root / 'missing')

    def test_blank_and_filled_packages_preserve_originals_and_signatures(self):
        for code in ['320581', '320582']:
            g = self.p.parse_capture(code)
            for kind in ['application', 'evidence']:
                t = next(t for t in g['templates'] if t['id'] == self.p.SPECS[code][kind])
                original = (ROOT / t['file']).read_bytes()
                data = {'applicantName': '虚构测试劳动者', 'respondentName': '虚构测试公司',
                        'requests': ['虚构请求'], 'factsReason': '仅供版式测试。', 'copyCount': 2, 'evidenceCopyCount': 3} if kind == 'application' else {
                            'evidenceList': [{'name': '虚构证据', 'purpose': '仅供版式测试。'}]}
                for payload in [{}, data]:
                    out = BytesIO(); self.p.fill_candidate_document(g, kind, payload, ROOT).save(out)
                    with ZipFile(BytesIO(original)) as before, ZipFile(BytesIO(out.getvalue())) as after:
                        self.assertEqual(set(before.namelist()), set(after.namelist()))
                        for part in before.namelist():
                            if part != 'word/document.xml': self.assertEqual(before.read(part), after.read(part), part)
                    doc = Document(BytesIO(out.getvalue()))
                    text = '\n'.join(p.text for p in doc.paragraphs)
                    self.assertNotIn('申请人：虚构测试劳动者（签名', text.replace(' ', ''))
                    if kind == 'application':
                        self.assertNotIn('用人单位申请', text)
                        self.assertNotIn('被申请人：', '\n'.join(p.text for p in doc.paragraphs[13:19]))
                    else:
                        self.assertEqual((len(doc.tables[0].rows), len(doc.tables[0].columns)), (2, 6))
                        self.assertNotIn('虚构', doc.paragraphs[5].text)
                self.assertEqual((ROOT / t['file']).read_bytes(), original)

    def test_capacity_unknown_fields_and_unverified_capabilities_are_rejected(self):
        g = self.p.parse_capture('320581')
        for kind, data in [('application', {'requests': ['x'] * 5}), ('evidence', {'evidenceList': [{}] * 41}),
                           ('application', {'unknown': ''}), ('application', {'copyCount': True}),
                           ('application', {'applicationDate': '2026-02-30'})]:
            with self.subTest(data=data), self.assertRaises(ValueError):
                self.p.fill_candidate_document(g, kind, data, ROOT)
        with self.assertRaises(ValueError): self.p.fill_candidate_document(g, 'employer', {}, ROOT)
        for t in g['templates']:
            if t['id'] not in [self.p.SPECS['320581'][k] for k in ['application', 'evidence']]:
                with self.assertRaises(ValueError): self.p.blank_binding(g, t, ROOT)

    def test_json_types_are_strict_recursively_in_own_facts_and_templates(self):
        g = self.p.parse_capture('320581')
        # False == 0 and True == 1 must never validate as canonical JSON values.
        paths = [('channel', 'submitted'), ('attachment_counts', 'instruction'),
                 ('templates', 0, 'can_generate'), ('materials', 0, 'template_control_count')]
        tested = 0
        for path in paths:
            bad = deepcopy(g); node = bad
            for key in path[:-1]: node = node[key]
            if path[-1] not in node: continue
            value = node[path[-1]]
            if type(value) is bool: node[path[-1]] = int(value)
            elif type(value) is int: node[path[-1]] = float(value)
            else: continue
            tested += 1
            with self.subTest(path=path), self.assertRaises(ValueError):
                self.p.fill_candidate_document(bad, 'application', {}, ROOT)
        self.assertGreaterEqual(tested, 3)

    def test_source_flags_cannot_become_candidate_or_published_evidence(self):
        g = self.p.parse_capture('320581')
        self.assertFalse(self.p.with_capabilities(g, ROOT)['can_generate'])
        for flags in [dict(can_generate=True), dict(published=True, candidate_web_qa='PASS'),
                      dict(generation_adapter=self.p.ADAPTER), dict(candidate_adapter_evidence={'status': 'PASS'})]:
            bad = deepcopy(g); bad.update(flags)
            with self.subTest(flags=flags), self.assertRaises(ValueError): self.p.with_capabilities(bad, ROOT)
        with tempfile.TemporaryDirectory() as tmp, self.assertRaises((ValueError, OSError)):
            self.p.prepare_candidate('320581', root=Path(tmp))

    def test_candidate_removes_only_obsolete_source_lifecycle_warnings(self):
        for code in sorted(self.p.SUPPORTED):
            raw = self.p.parse_capture(code)
            candidate = self.p.prepare_candidate(code, root=ROOT)
            obsolete = [w for w in raw['warnings'] if w.startswith(('本批仅采集自身指南', 'DOCX只读结构与文本检查'))]
            self.assertEqual(len(obsolete), 2)
            self.assertTrue(all(w not in candidate['warnings'] for w in obsolete))
            self.assertTrue(all(w in candidate['warnings'] for w in raw['warnings'] if w not in obsolete))
            self.assertEqual(self.p.parse_capture(code), raw)


if __name__ == '__main__':
    unittest.main()
