"""Field and layout regressions for the captured Shanxi official blank."""
import hashlib
import sys
from pathlib import Path
import unittest
from docx import Document
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[2];sys.path.insert(0,str(ROOT/'backend'))
from WenShu.Shanxi.application import fill_application

class ShanxiTemplateTests(unittest.TestCase):
    source=ROOT/'templates/_docx/ShanXiSheng/仲裁申请书.docx'
    def test_exact_slots_preserve_original_and_leave_signing_blank(self):
        before=hashlib.sha256(self.source.read_bytes()).hexdigest()
        data={'applicantName':'王晨','gender':'男','birthDate':'1992-08-16','nationality':'汉族','applicantWorkUnit':'现在任职的另一单位','currentAddress':'当前住宅','phone':'13800001234','applicantMailingAddress':'有效邮寄地址','respondentName':'发生争议的原单位','registeredAddress':'原单位住所','legalRepName':'张明','legalRepPosition':'负责人','legalRepPhone':'13900005678','requests':['要求一。','要求二。'],'factsReason':'完整事实。','copyCount':2,'committeeName':'受理仲裁委员会','applicationDate':'2026-10-07'}
        doc=fill_application(data,self.source);text='\n'.join(p.text for p in doc.paragraphs)
        for key,value in data.items():
            if key in ('birthDate','applicationDate'):continue
            for v in value if isinstance(value,list) else [str(value)]:self.assertIn(v,text)
        self.assertIn('出生日期：1992年8月16日，民族：汉族，工作单位：现在任职的另一单位，现住址：当前住宅，联系电话：13800001234。',text)
        self.assertIn('被申请人（单位名称）：发生争议的原单位。住所地：原单位住所。',text)
        self.assertIn('法定代表人或负责人：张明。职务：负责人。 联系电话：13900005678。',text)
        self.assertIn('副本2份',text);self.assertIn('2026年10月7日',text)
        signature=next(p for p in doc.paragraphs if '签名并捺印' in p.text);self.assertNotIn('王晨',signature.text)
        self.assertEqual(before,hashlib.sha256(self.source.read_bytes()).hexdigest())
        original=Document(self.source)
        self.assertEqual(len(doc.sections),1)
        self.assertEqual(doc.sections[0]._sectPr.xml,original.sections[0]._sectPr.xml)
        for note in original.paragraphs[18:21]:self.assertIn(note.text,text)
        name_run=next(r for r in doc.paragraphs[3].runs if r.text=='王晨');self.assertTrue(name_run.underline)
        self.assertEqual(name_run._r.rPr.find(qn('w:rFonts')).get(qn('w:eastAsia')),'仿宋_GB2312')

    def test_optional_workplace_not_inferred_from_disputed_employer(self):
        doc=fill_application({'respondentName':'争议单位'},self.source)
        self.assertNotIn('争议单位',doc.paragraphs[3].text)
        self.assertIn('争议单位',doc.paragraphs[4].text)

    def test_long_facts_and_three_requests_retain_content(self):
        facts='具体事实及证据。\n'*400
        doc=fill_application({'requests':['甲','乙','丙'],'factsReason':facts},self.source)
        self.assertEqual(doc.paragraphs[11].text,facts.strip())
        self.assertEqual([p.text for p in doc.paragraphs[7:10]],['1、甲','2、乙','3、丙'])
        with self.assertRaisesRegex(ValueError,'3项'):fill_application({'requests':['甲']*4},self.source)

    def test_wrong_template_version_fails_instead_of_writing_wrong_slots(self):
        from tempfile import TemporaryDirectory
        with TemporaryDirectory(dir=ROOT/'.runtime') as tmp:
            wrong=Document(self.source);wrong.paragraphs[6].text='另一个版式'
            path=Path(tmp)/'wrong.docx';wrong.save(path)
            with self.assertRaisesRegex(ValueError,'版本'):fill_application({},path)

if __name__=='__main__':unittest.main()
