"""The real shared route uses each county's path-free source guard in an isolated DB."""
import io,json,sys,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from fastapi import FastAPI
from fastapi.testclient import TestClient
from docx import Document
ROOT=Path(__file__).resolve().parents[2];sys.path.insert(0,str(ROOT/'backend'))
from regional_adapters import with_reference_capabilities
from regional_guides import RegionalStore
from regional_jiangsu_candidates import fill_candidate_document
import regional_jiangsu_candidates as own_source
from routes import regional_guides as routes

class JiangsuCandidateDispatchTests(unittest.TestCase):
    def test_missing_own_originals_never_advertise_generation(self):
        source=ROOT/'reports/regional-release-20261029/jiangsu/320981/public-guide.json'
        guide=json.loads(source.read_text('utf-8'))
        self.assertEqual(sum(bool(t.get('can_generate')) for t in with_reference_capabilities(guide)['templates']),2)
        with tempfile.TemporaryDirectory() as temp,patch.object(own_source,'ROOT',Path(temp)):
            blocked=with_reference_capabilities(guide)
        self.assertFalse(any(t.get('can_generate') for t in blocked['templates']))

    def test_real_route_uses_own_canonical_guard_and_rejects_cross_county_source(self):
        with tempfile.TemporaryDirectory() as temp:
            store=RegionalStore(Path(temp)/'test-guides.sqlite')
            guides=[]
            for code in ['320981','320924']:
                source=ROOT/f'reports/regional-release-20261029/jiangsu/{code}/public-guide.json'
                guides.append(store.put_guide(with_reference_capabilities(json.loads(source.read_text('utf-8')))))
            self.assertNotEqual(guides[0]['source_id'],guides[1]['source_id'])
            app=FastAPI();app.include_router(routes.router,prefix='/regional')
            with patch.object(routes,'store',return_value=store),patch.object(routes,'fill_candidate_document',wraps=fill_candidate_document,create=True) as own,TestClient(app) as client:
                for g in guides:
                    for kind in ['application','evidence']:
                        t=next(t for t in g['templates'] if t.get('document_types')==[kind])
                        data=({'applicantName':'虚构地区测试','requests':['仅用于软件验证'],'copyCount':'','evidenceCopyCount':''} if kind=='application' else {'evidenceList':[{'name':'虚构材料','purpose':'软件验证','source':'虚构来源','pageCount':'3'}]})
                        payload={k:g[k] for k in ['province','city','district','source_id','version']}
                        payload.update(template_id=t['id'],kind=kind,mode='blank',data=data)
                        r=client.post('/regional/generate',json=payload)
                        self.assertEqual(r.status_code,200,r.text[:200] if r.status_code!=200 else '')
                        self.assertEqual(r.headers['X-Document-Origin'],'official-blank-filled')
                        self.assertEqual(r.headers['X-Guide-Source-ID'],g['source_id'])
                        doc=Document(io.BytesIO(r.content))
                        text='\n'.join(p.text for p in doc.paragraphs)+'\n'+''.join(c.text for tab in doc.tables for row in tab.rows for c in row.cells)
                        self.assertIn('虚构地区测试' if kind=='application' else '虚构材料',text)
                        payload['source_id']=next(other['source_id'] for other in guides if other['source_id']!=g['source_id'])
                        self.assertEqual(client.post('/regional/generate',json=payload).status_code,409)
                self.assertEqual(own.call_count,4,'Shared generation must call the path-free own-source guard')
