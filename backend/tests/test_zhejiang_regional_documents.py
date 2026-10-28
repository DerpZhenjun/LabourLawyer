"""Exercise Zhejiang own-source binding and in-memory field/package preservation."""
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'backend'))
from regional_zhejiang import parse_yuhang_capture, parse_xihu_capture
try:
    import regional_zhejiang_documents as adapter
except ModuleNotFoundError:
    adapter = None
REPORT = ROOT / 'reports/nationwide-localization-20261028/zhejiang'


@unittest.skipUnless((REPORT / 'SHA256-MANIFEST.json').is_file(),
                     'Local sealed capture contains login-flow context; not distributed in public checkout')
class ZhejiangRegionalDocumentsTests(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(adapter, 'Own-source Zhejiang document adapter is not implemented')

    def guide(self, district='yuhang'):
        parser = parse_yuhang_capture if district == 'yuhang' else parse_xihu_capture
        return parser(REPORT, REPORT / 'SHA256-MANIFEST.json')

    def selected(self, kind, district='yuhang'):
        guide = adapter.with_zhejiang_capabilities(self.guide(district))
        template = next(t for t in guide['templates'] if t['document_types'] == [kind])
        return guide, template

    def test_only_two_own_blanks_gain_staging_capabilities(self):
        # Break: treating the identity sample or shared provincial bytes as a fillable form.
        for district in ['yuhang', 'xihu']:
            with self.subTest(district=district):
                raw = self.guide(district)
                guide = adapter.with_zhejiang_capabilities(raw)
                self.assertTrue(all(not t['can_generate'] for t in raw['templates']))
                enabled = [t for t in guide['templates'] if t['can_generate']]
                self.assertEqual({tuple(t['document_types']) for t in enabled}, {('application',), ('evidence',)})
                self.assertEqual(len(enabled), 2)
                self.assertTrue(all(t['file'].startswith('templates/_docx/_regional/' + guide['region_code'] + '/')
                                    and t['adapter'] == adapter.ADAPTER for t in enabled))
                self.assertFalse(next(t for t in guide['templates'] if t['official_kind'] == 'sample')['can_generate'])

    def test_own_source_identity_mutations_disable_binding(self):
        # Break: matching only a SHA while ignoring locality, task, annex, path or source.
        guide, template = self.selected('application')
        cases = [
            ('guide', 'source_id', 'different-source'), ('guide', 'authority', '杭州市人力资源和社会保障局'),
            ('guide', 'service_code', '2510b7af-2402-4502-837c-14eb4a71177d'),
            ('guide', 'source_url', 'https://www.zjzwfw.gov.cn/zjservice-fe/#/workguide?localInnerCode=wrong'),
            ('guide', 'district', '西湖区'), ('guide', 'city', '宁波市'), ('guide', 'region_code', '330106'),
            ('template', 'official_annex_id', '872DFF90A8578907CB947197A784AA48'),
            ('template', 'sha256', '0' * 64), ('template', 'file', 'templates/_docx/ZheJiangSheng/《劳动人事争议仲裁申请书》.docx'),
            ('template', 'download_url', 'https://example.com/application.docx'), ('template', 'official_kind', 'sample'),
        ]
        for target, key, value in cases:
            with self.subTest(target=target, key=key):
                g, t = deepcopy(guide), deepcopy(template)
                (g if target == 'guide' else t)[key] = value
                self.assertIsNone(adapter.zhejiang_blank_binding(g, t))
                g['templates'] = [t]
                disabled = adapter.with_zhejiang_capabilities(g)['templates'][0]
                self.assertFalse(disabled['can_generate'])
                self.assertEqual(disabled['document_types'], [])

    def test_equal_bytes_from_other_district_do_not_bind(self):
        # Break: importing Xihu's equal-byte annex as Yuhang's own copy.
        yuhang, _ = self.selected('application')
        _, xihu_template = self.selected('application', 'xihu')
        self.assertIsNone(adapter.zhejiang_blank_binding(yuhang, xihu_template))

    def test_application_fills_worker_slots_and_leaves_hand_signature_blank(self):
        # Break: using the generic unsupported Zhejiang application profile or filling human signature.
        guide, template = self.selected('application')
        document = adapter.fill_zhejiang_document(guide, template, 'application', ROOT / template['file'], {
            'applicantName': '测试甲', 'birthDate': '1990-01-02', 'gender': '女', 'nationality': '汉族',
            'respondentName': '测试单位乙', 'legalRepName': '测试负责人丙', 'requests': ['支付测试工资123元'],
            'factsReason': '这是独立浙江适配器的内存行为验证。',
            'committeeName': '杭州市余杭区劳动人事争议仲裁委员会', 'applicationDate': '2026-10-28', 'copyCount': 1})
        self.assertEqual(len(document.tables), 1)
        table = document.tables[0]
        self.assertEqual(table.rows[1].cells[1].text, '测试甲')
        self.assertEqual(table.rows[7].cells[2].text, '测试单位乙')
        self.assertIn('支付测试工资123元', '\n'.join(c.text for r in table.rows for c in r.cells))
        footer = table.rows[-1].cells[0].paragraphs
        self.assertIn('杭州市余杭区', footer[2].text)
        self.assertNotIn('测试甲', footer[3].text)
        self.assertIn('2026年', footer[4].text)

    def test_evidence_uses_explicit_profile_with_regional_path(self):
        # Break: deriving a nonexistent profile from the numbered regional directory.
        guide, template = self.selected('evidence', 'xihu')
        doc = adapter.fill_zhejiang_document(guide, template, 'evidence', ROOT / template['file'], {
            'committeeName': '杭州市西湖区劳动人事争议仲裁委员会', 'caseNumber': '测试案号', 'copyCount': 2,
            'submitterName': '不可自动落签名', 'submitDate': '2026-10-28',
            'evidenceList': [{'name': '工资记录', 'pageCount': '2页', 'source': '本人保存', 'purpose': '证明工资'}]})
        table = doc.tables[0]
        self.assertEqual(table.rows[2].cells[1].text, '工资记录')
        self.assertEqual(table.rows[2].cells[3].text, '2')
        self.assertEqual(table.rows[2].cells[4].text, '本人保存')
        self.assertEqual(table.rows[2].cells[5].text, '证明工资')
        self.assertNotIn('不可自动落签名', '\n'.join(c.text for r in table.rows for c in r.cells))

    def test_only_document_xml_changes_when_saved_in_memory(self):
        # Break: dropping/changing original package parts while serializing a typed form.
        for kind in ['application', 'evidence']:
            with self.subTest(kind=kind):
                guide, template = self.selected(kind)
                path = ROOT / template['file']
                data = {'applicantName': '内存验证', 'requests': ['测试请求'], 'evidenceList': [{'name': '测试证据'}]}
                result = adapter.fill_zhejiang_document(guide, template, kind, path, data)
                output = BytesIO(); result.save(output)
                with ZipFile(path) as original, ZipFile(output) as changed:
                    self.assertEqual(set(original.namelist()), set(changed.namelist()))
                    self.assertEqual({name for name in original.namelist() if original.read(name) != changed.read(name)}, {'word/document.xml'})
                self.assertEqual(path.read_bytes(), result.original)

    def test_alternate_equal_byte_path_and_wrong_kind_are_rejected(self):
        # Break: permitting a caller's arbitrary same-byte file in place of the own canonical copy.
        guide, template = self.selected('application')
        alternate = REPORT / 'yuhang/attachments/application.docx'
        with self.assertRaises(ValueError):
            adapter.fill_zhejiang_document(guide, template, 'application', alternate, {'applicantName': '测试'})
        with self.assertRaises(ValueError):
            adapter.fill_zhejiang_document(guide, template, 'evidence', ROOT / template['file'], {})

    def test_changed_canonical_bytes_disable_binding(self):
        # Break: trusting retained SHA metadata without re-reading the canonical file.
        guide, template = self.selected('application')
        tmp_parent = REPORT / 'adapter-test-tmp'; tmp_parent.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=tmp_parent) as temporary:
            root = Path(temporary)
            manifest = json.loads((REPORT / 'SHA256-MANIFEST.json').read_text(encoding='utf8'))
            for relative in manifest['files']:
                destination = root / 'reports/nationwide-localization-20261028/zhejiang' / relative.replace('\\', '/')
                destination.parent.mkdir(parents=True, exist_ok=True)
                shutil.copyfile(REPORT / relative.replace('\\', '/'), destination)
            destination = root / template['file']; destination.parent.mkdir(parents=True, exist_ok=True)
            destination.write_bytes((REPORT / 'yuhang/attachments/evidence-list.docx').read_bytes())
            with patch.object(adapter, 'ROOT', root):
                self.assertIsNone(adapter.zhejiang_blank_binding(guide, template))

    def test_missing_kind_cannot_authorize_a_stale_enabled_sample(self):
        # Break: None binding == None requested kind accidentally authorizes a sample.
        guide = self.guide()
        sample = next(t for t in guide['templates'] if t['official_kind'] == 'sample')
        sample.update(adapter=adapter.ADAPTER, can_generate=True, document_types=[None])
        with self.assertRaises(ValueError):
            adapter.checked_zhejiang_adapter(guide, sample, None)


if __name__ == '__main__':
    unittest.main()
