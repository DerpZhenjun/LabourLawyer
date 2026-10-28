"""Own source API boundary and preserved packages, using only synthetic data."""
import copy
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from fastapi import FastAPI
from fastapi.testclient import TestClient
from docx import Document
from regional_guides import RegionalStore
from regional_adapters import with_reference_capabilities
from routes.regional_guides import router
ROOT=Path(__file__).resolve().parents[2]
BASE=ROOT/'reports/nationwide-localization-20261023-night/cross-province/jiangsu-dafeng'


class DafengGenerationTests(unittest.TestCase):
    def test_two_actual_api_kinds_own_annexes_and_cross_county_rejection(self):
        with tempfile.TemporaryDirectory() as temp:
            store=RegionalStore(Path(temp)/'guide.sqlite')
            guide=store.put_guide(with_reference_capabilities(json.loads((BASE/'guide.json').read_text('utf8'))))
            app=FastAPI();app.include_router(router,prefix='/api/regional-guides')
            with patch('routes.regional_guides.store',return_value=store),TestClient(app) as client:
                for kind,identifier,data in [
                    ('application','8984153f-2c04-4676-a8fd-3a627658fb9d',{'applicantName':'大丰合成测试甲',
                        'respondentName':'大丰合成测试公司','requests':['支付测试工资100元。'],
                        'factsReason':'这是合成排版测试，不用于真实劳动争议申请。','copyCount':2,'evidenceCopyCount':3}),
                    ('evidence','b104e1d9-15a7-489a-ad8b-80e22e0673e7',{'evidenceList':[{
                        'name':'大丰合成测试劳动合同','source':'本人保存','purpose':'说明测试劳动关系','originalType':'复印件','pageCount':'3页'}]})]:
                    payload={k:guide[k] for k in ['province','city','district','source_id','version']}
                    payload.update(template_id=identifier,kind=kind,mode='blank',data=data)
                    response=client.post('/api/regional-guides/generate',json=payload)
                    self.assertEqual(response.status_code,200,response.text if response.status_code!=200 else '')
                    self.assertEqual(response.headers['X-Guide-Source-ID'],'5725c8df9cb686d0517254e8')
                    self.assertEqual(response.headers['X-Document-Origin'],'official-blank-filled')
                    doc=Document(io.BytesIO(response.content));text='\n'.join(p.text for p in doc.paragraphs)
                    if kind=='application':
                        self.assertIn('大丰合成测试甲',text);self.assertNotIn('用人单位申请',text)
                        self.assertEqual(text.count('被申请人：'),1);self.assertIn('副本2份；相关证据3份',text)
                    else:
                        self.assertEqual(len(doc.tables[0].columns),6);self.assertEqual(len(doc.tables[0].rows),2)
                        self.assertEqual(doc.tables[0].rows[1].cells[2].text,'本人保存')
                        self.assertEqual(doc.tables[0].rows[1].cells[5].text,'3')
                    t=next(t for t in guide['templates'] if t['id']==identifier)
                    with ZipFile(ROOT/t['file']) as before,ZipFile(io.BytesIO(response.content)) as after:
                        self.assertEqual(set(before.namelist()),set(after.namelist()))
                        for part in before.namelist():
                            if part!='word/document.xml':self.assertEqual(before.read(part),after.read(part),part)
                    for key,value in [('district','建湖县'),('source_id','ef1eb897ebf880c6f2a316f5'),
                                      ('template_id','985f4770-1ab2-4264-9482-0f8528b16513'),('mode','reference_adaptation')]:
                        wrong=copy.deepcopy(payload);wrong[key]=value
                        self.assertEqual(client.post('/api/regional-guides/generate',json=wrong).status_code,409)


if __name__=='__main__':unittest.main()
