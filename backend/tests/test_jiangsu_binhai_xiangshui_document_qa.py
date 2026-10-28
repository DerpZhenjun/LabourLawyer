"""Real own blanks: one-evidence semantics, manual signatures and package fidelity."""
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import sys
import unittest
from zipfile import ZipFile
from docx import Document
from lxml import etree
ROOT=Path(__file__).resolve().parents[2];sys.path.insert(0,str(ROOT/'backend'))
from regional_blank_documents import blank_binding,fill_blank_document
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

class BinhaiXiangshuiDocumentTests(unittest.TestCase):
    def guides(self):
        for code in ['320922','320921']:
            yield json.loads((ROOT/'reports/nationwide-localization-20261028-next/jiangsu'/code/'normalized-guide.json').read_text('utf8'))
    def filled(self,guide,kind):
        template=next(t for t in guide['templates'] if blank_binding(guide,t)==kind)
        source=ROOT/template['file'];original=source.read_bytes()
        data=dict(committeeName=guide['district']+'劳动人事争议仲裁委员会',applicantName='虚构甲',
            gender='女',birthDate='1990-01-02',idNumber='TEST-ID',idAddress='虚构住址',
            applicantMailingAddress='虚构收件地址',phone='TEST-PHONE',respondentName='虚构单位',
            respondentType='公司',legalRepName='虚构乙',companyContact='虚构丙',
            requests=['支付虚构测试工资100元。'],factsReason='仅为虚构软件测试，不存在真实案件。',
            copyCount=3,evidenceCopyCount=2,evidenceList=[dict(name='虚构合同',source='虚构来源',
                purpose='虚构证明目的',originalType='复印件',pageCount='3页')])
        output=BytesIO();fill_blank_document(guide,template,kind,source,data).save(output)
        self.assertEqual(source.read_bytes(),original)
        self.assertEqual(hashlib.sha256(original).hexdigest(),template['sha256'])
        return original,output.getvalue(),data

    def test_one_evidence_starts_at_one_and_maps_each_official_column(self):
        for guide in self.guides():
            with self.subTest(code=guide['region_code']):
                _,body,_=self.filled(guide,'evidence');doc=Document(BytesIO(body));table=doc.tables[0]
                self.assertEqual(len(table.rows),2)
                self.assertEqual([c.text for c in table.rows[1].cells],['1','虚构合同','虚构来源','虚构证明目的','复印件','3'])
                self.assertEqual(len(table.rows[0].cells),6)
                self.assertEqual(doc.paragraphs[0].text,guide['district']+'劳动人事争议仲裁委员会')
                for p in doc.paragraphs:
                    if p.text.startswith(('提交人','接收人')):
                        self.assertNotIn('虚构',p.text)
                        self.assertNotRegex(p.text,r'\d')

    def test_application_fills_fields_and_keeps_hand_signature_and_date_empty(self):
        for guide in self.guides():
            with self.subTest(code=guide['region_code']):
                _,body,data=self.filled(guide,'application');doc=Document(BytesIO(body));text='\n'.join(p.text for p in doc.paragraphs)
                for value in ['虚构甲','1990-01-02','TEST-ID','虚构住址','虚构收件地址','TEST-PHONE','虚构单位','虚构乙','虚构丙',data['requests'][0],data['factsReason']]:
                    self.assertIn(value,text)
                self.assertIn('本申请书副本3份；相关证据2份',text)
                self.assertIn(guide['district']+'劳动人事争议仲裁委员会',text)
                self.assertNotIn('用人单位申请',text)
                self.assertEqual(sum(p.text.startswith('被申请人') for p in doc.paragraphs),1)
                signatures=[p.text for p in doc.paragraphs if '签名、盖章或指模' in p.text]
                self.assertEqual(len(signatures),1)
                self.assertRegex(signatures[0],r'^申请人：\s+（签名、盖章或指模）$')
                self.assertTrue(any(re.sub(r'\s+','',p.text)=='年月日' for p in doc.paragraphs))
                self.assertTrue(all(str(i)+'、' in text for i in range(1,6)))

    def test_only_document_part_changes_and_section_geometry_is_preserved(self):
        for guide in self.guides():
            for kind in ['application','evidence']:
                with self.subTest(code=guide['region_code'],kind=kind):
                    original,output,_=self.filled(guide,kind)
                    with ZipFile(BytesIO(original)) as source,ZipFile(BytesIO(output)) as final:
                        self.assertEqual(source.namelist(),final.namelist())
                        self.assertEqual([n for n in source.namelist() if source.read(n)!=final.read(n)],['word/document.xml'])
                        def geometry(z):
                            dom=etree.fromstring(z.read('word/document.xml'))
                            return [[(etree.QName(n).localname,sorted(n.attrib.items())) for n in section.iter()] for section in dom.xpath('//w:sectPr',namespaces=NS)]
                        self.assertEqual(geometry(source),geometry(final))

    def test_birth_date_uses_compact_iso_in_the_narrow_official_first_line(self):
        # A real Word render of 1990年1月2日 split 月2日 onto a new line.
        # Keep the source's typography; the same calendar date fits as ISO.
        for guide in self.guides():
            with self.subTest(code=guide['region_code']):
                original,body,_=self.filled(guide,'application');doc=Document(BytesIO(body))
                self.assertIn('出生日期：1990-01-02',doc.paragraphs[1].text)
                self.assertNotIn('1990年1月2日',doc.paragraphs[1].text)
                date_run=next(r for r in doc.paragraphs[1].runs if r.text=='1990-01-02')
                reference=Document(BytesIO(original)).paragraphs[1]
                original_slot=next(r for r in reversed(reference.runs) if r.underline and r.text.strip()=='')
                self.assertTrue(date_run.underline)
                self.assertEqual(date_run.font.size,original_slot.font.size)
                self.assertEqual(etree.tostring(date_run._r.rPr),etree.tostring(original_slot._r.rPr))

if __name__=='__main__':unittest.main()
