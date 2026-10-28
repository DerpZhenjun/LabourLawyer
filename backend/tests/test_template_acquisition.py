from __future__ import annotations

import asyncio
import io
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document


BACKEND_DIR = Path(__file__).resolve().parents[1]
REPOSITORY_ROOT = BACKEND_DIR.parent
if str(REPOSITORY_ROOT) not in sys.path:
    sys.path.insert(0, str(REPOSITORY_ROOT))

from backend.template_acquisition.adapters.guangdong import GuangdongAdapter
from backend.template_acquisition.adapters.jiangsu import JiangsuAdapter
from backend.template_acquisition.adapters.liaoning import LiaoningAdapter
from backend.template_acquisition.adapters.zhejiang import ZhejiangAdapter
from backend.template_acquisition.docx_analysis import analyze_docx, sha256_bytes
from backend.template_acquisition.models import SourceConfig, TemplateCandidate
from backend.template_acquisition.security import validate_docx, validate_download_url, validate_legacy_doc
from backend.template_acquisition.storage import CandidateStore


def make_docx() -> bytes:
    stream = io.BytesIO()
    document = Document()
    document.add_heading("劳动人事争议仲裁申请书")
    document.add_table(rows=1, cols=2)
    document.save(stream)
    return stream.getvalue()


class FakeClient:
    def __init__(self, payload):
        self.payload = payload

    async def request_json(self, method, url, **kwargs):
        return self.payload


class TemplateSecurityTests(unittest.TestCase):
    def test_docx_validation_and_analysis(self):
        data = make_docx()
        validation = validate_docx(data)
        fingerprint, stats = analyze_docx(data)
        self.assertGreater(validation.member_count, 5)
        self.assertEqual(len(fingerprint), 64)
        self.assertEqual(stats["table_count"], 1)

    def test_macro_payload_is_rejected(self):
        source = io.BytesIO(make_docx())
        output = io.BytesIO()
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(output, "w") as modified:
            for item in original.infolist():
                modified.writestr(item, original.read(item.filename))
            modified.writestr("word/vbaProject.bin", b"unsafe")
        with self.assertRaisesRegex(ValueError, "Macro"):
            validate_docx(output.getvalue())

    def test_only_https_official_domains_are_allowed(self):
        domains = ("gdzwfw.gov.cn",)
        validate_download_url("https://static.gdzwfw.gov.cn/guide/file", domains)
        with self.assertRaises(ValueError):
            validate_download_url("https://gdzwfw.gov.cn.evil.example/file", domains)
        with self.assertRaises(ValueError):
            validate_download_url("http://gdzwfw.gov.cn/file", domains)

    def test_legacy_word_signature_is_checked(self):
        data = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64
        self.assertTrue(validate_legacy_doc(data)["valid_legacy_doc"])
        with self.assertRaises(ValueError):
            validate_legacy_doc(b"not-a-word-document")


class AdapterTests(unittest.TestCase):
    def test_zhejiang_material_indices_map_to_blank_forms(self):
        payload = {"data": {"DATAAREA": {"MATERIALS": {"MATERIAL": [
            {}, {},
            {"EMPTYTABLE": {"FILEURL": "//portal.zjzwfw.gov.cn/app.docx"}},
            {"EMPTYTABLE": {"FILEURL": "//portal.zjzwfw.gov.cn/evidence.docx"}},
        ]}}}}
        source = SourceConfig(
            "Zhejiang", "zhejiang", "https://www.zjzwfw.gov.cn/guide", ("zjzwfw.gov.cn",),
            {"material_api": "https://www.zjzwfw.gov.cn/api", "local_inner_code": "id",
             "templates": [
                 {"document_type": "application", "display_name": "申请书", "material_index": 2},
                 {"document_type": "evidence", "display_name": "证据清单", "material_index": 3},
             ]},
        )
        candidates, _ = asyncio.run(ZhejiangAdapter(source, FakeClient(payload)).discover())
        self.assertEqual([item.document_type for item in candidates], ["application", "evidence"])
        self.assertEqual(candidates[0].download_url, "https://portal.zjzwfw.gov.cn/app.docx")

    def test_guangdong_uses_blank_form_not_example(self):
        payload = {"AUDIT_MATERIAL": [
            {"FORM_GUID": [{"FILEPATH": "https://static.gdzwfw.gov.cn/app.docx"}],
             "EXAMPLE_GUID": [{"FILEPATH": "https://static.gdzwfw.gov.cn/example.docx"}]},
            {"FORM_GUID": [{"FILEPATH": "https://static.gdzwfw.gov.cn/evidence.docx"}]},
        ]}
        source = SourceConfig(
            "Guangdong", "guangdong", "https://www.gdzwfw.gov.cn/guide", ("gdzwfw.gov.cn",),
            {"material_api": "https://www.gdzwfw.gov.cn/api", "task_code": "code",
             "templates": [
                 {"document_type": "application", "display_name": "申请书", "material_index": 0},
                 {"document_type": "evidence", "display_name": "证据清单", "material_index": 1},
             ]},
        )
        candidates, _ = asyncio.run(GuangdongAdapter(source, FakeClient(payload)).discover())
        self.assertEqual(candidates[0].download_url, "https://static.gdzwfw.gov.cn/app.docx")

    def test_jiangsu_extracts_official_blank_form_links(self):
        payload = """<!doctype html><html><body><table>
        <tr class="tsize"><td><p>劳动人事争议仲裁申请书</p></td><td>
        <a class="kbbg" onclick="downloadClyb('https://sxgl.jszwfw.gov.cn/app.docx','');">空白表格</a>
        <a class="slyb" onclick="downloadClyb('https://sxgl.jszwfw.gov.cn/example.docx','');">示例</a>
        </td></tr>
        <tr class="tsize"><td><p>劳动人事争议仲裁案件证据清单</p></td><td>
        <a class="kbbg" onclick="downloadClyb('https://sxgl.jszwfw.gov.cn/evidence.docx','');">空白表格</a>
        </td></tr></table></body></html>"""
        source = SourceConfig(
            "Jiangsu", "jiangsu", "https://www.jszwfw.gov.cn/guide", ("jszwfw.gov.cn",),
            {"templates": [
                {"document_type": "application", "material_name": "劳动人事争议仲裁申请书"},
                {"document_type": "evidence", "material_name": "劳动人事争议仲裁案件证据清单"},
            ]},
        )
        client = FakeClient(None)
        client.request_text = lambda *args, **kwargs: None

        async def request_text(*args, **kwargs):
            return payload

        client.request_text = request_text
        candidates, _ = asyncio.run(JiangsuAdapter(source, client).discover())
        self.assertEqual(candidates[0].download_url, "https://sxgl.jszwfw.gov.cn/app.docx")
        self.assertEqual(candidates[1].download_url, "https://sxgl.jszwfw.gov.cn/evidence.docx")

    def test_liaoning_keeps_legacy_official_file_metadata(self):
        payload = """<html><body><a href="https://www.lnzwfw.gov.cn/files/1、仲裁申请书.doc">
        1、仲裁申请书.doc</a></body></html>"""
        source = SourceConfig(
            "Liaoning", "liaoning", "https://center.lnzwfw.gov.cn/guide", ("lnzwfw.gov.cn",),
            {"templates": [{
                "document_type": "application", "display_name": "仲裁申请书",
                "link_text_contains": "仲裁申请书.doc", "file_format": "doc",
                "original_path": "templates/_original/LiaoNingSheng/仲裁申请书.doc",
                "production_path": "templates/_docx/LiaoNingSheng/仲裁申请书.docx",
            }]},
        )
        client = FakeClient(None)

        async def request_text(*args, **kwargs):
            return payload

        client.request_text = request_text
        candidates, _ = asyncio.run(LiaoningAdapter(source, client).discover())
        self.assertEqual(candidates[0].file_format, "doc")
        self.assertTrue(candidates[0].download_url.endswith("仲裁申请书.doc"))


class CandidateStoreTests(unittest.TestCase):
    def test_second_identical_artifact_is_unchanged(self):
        with tempfile.TemporaryDirectory() as directory:
            store = CandidateStore(Path(directory))
            candidate = TemplateCandidate("application", "申请书", "https://official/a", "https://official/guide")
            data = make_docx()
            sha = sha256_bytes(data)
            fingerprint, stats = analyze_docx(data)
            first_dir = store.begin_run("run-one", "Test")
            first = store.save_artifact("run-one", "Test", first_dir, candidate, data, sha, fingerprint, stats)
            second_dir = store.begin_run("run-two", "Test")
            second = store.save_artifact("run-two", "Test", second_dir, candidate, data, sha, fingerprint, stats)
            self.assertEqual(first.status, "new")
            self.assertEqual(second.status, "unchanged")


if __name__ == "__main__":
    unittest.main()
