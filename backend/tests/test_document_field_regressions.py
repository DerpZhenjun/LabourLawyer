"""Regression gates for the concrete errors seen in the user's Word output."""
import hashlib
import json
from pathlib import Path
import sys
import subprocess
import unittest
from docx import Document
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'backend'))
from province_documents import PROVINCE_DOCUMENTS,get_province_document
from WenShu.generic_template_fill import generate_application,generate_evidence
from WenShu.exact_template_fill import all_paragraphs,INLINE,TABLE_FIELDS,derived,value
from routes.province_documents import ApplicationRequest,EvidenceRequest

class DocumentFieldRegressions(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        subprocess.run(['node', str(ROOT/'scripts/export_document_examples.mjs')],
                       cwd=ROOT, check=True, capture_output=True, timeout=30)
        cls.examples=json.loads((ROOT/'.runtime/document-audit/frontend-examples.json').read_text(encoding='utf8'))

    def example(self,slug):return next(x for x in self.examples if x['slug']==slug)

    def test_actual_frontend_payload_survives_api_validation(self):
        data=self.example('jiangsu')['application']
        result=ApplicationRequest(**data).model_dump()
        for key in ('acceptMediation','legalAid','respondentType','departmentHead','departmentHeadPosition'):
            self.assertEqual(result[key],data[key])
        evidence=self.example('hunan')['evidence']
        result=EvidenceRequest(**evidence).model_dump()
        for key in ('originalType','pageRange'):
            self.assertEqual(result['evidenceList'][0][key],evidence['evidenceList'][0][key])

    def test_all_templates_preserve_original_files_and_include_every_claim(self):
        for x in self.examples:
            for kind in PROVINCE_DOCUMENTS[x['slug']]['documents']:
                with self.subTest(region=x['slug'],kind=kind):
                    _,path=get_province_document(x['slug'],kind);before=hashlib.sha256(path.read_bytes()).hexdigest()
                    payload=(ApplicationRequest if kind=='application' else EvidenceRequest)(**x[kind]).model_dump()
                    doc=(generate_application if kind=='application' else generate_evidence)(payload,path)
                    text='\n'.join(p.text for p in all_paragraphs(doc))
                    if kind=='application':
                        for claim in payload['requests']:self.assertIn(claim,text)
                        self.assertIn(payload['applicantName'],text);self.assertIn(payload['respondentName'],text)
                    else:
                        for item in payload['evidenceList']:self.assertIn(item['name'],text)
                        # An evidence submitter name is collected for drafts but
                        # never printed into a handwritten signature field.
                        self.assertNotIn(payload['submitterName'],text)
                    self.assertEqual(before,hashlib.sha256(path.read_bytes()).hexdigest())

    def test_jiangsu_slots_underlines_inline_facts_and_blank_signature(self):
        x=self.example('jiangsu')['application'];_,path=get_province_document('jiangsu','application')
        doc=generate_application(x,path);ps=doc.paragraphs
        for index,fields in INLINE['JiangSuSheng'].items():
            for label,key in fields:
                if x.get(key):
                    self.assertIn(value(x,key),ps[index].text,key)
                    target=next(r for r in ps[index].runs if value(x,key)==r.text.strip())
                    self.assertTrue(target.underline,key)
        facts=next(p for p in ps if p.text.startswith('事实与理由'))
        self.assertEqual(facts.text,'事实与理由：'+x['factsReason'])
        signatures=[p for p in ps if '签名' in p.text]
        self.assertTrue(signatures)
        for p in signatures:self.assertNotIn(x['applicantName'],p.text)
        text='\n'.join(p.text for p in ps)
        self.assertIn('副本2份',text);self.assertNotIn('。22',text)

    def test_birth_dates_are_not_replaced_by_filing_date(self):
        for slug in ('beijing','yunnan','xinjiang'):
            x=self.example(slug)['application'];_,path=get_province_document(slug,'application');doc=generate_application(x,path)
            binding=TABLE_FIELDS[path.parent.name]
            r,c=next(pos for pos,key in binding.items() if key=='birthDate')
            cell=doc.tables[0].rows[r].cells[c]
            self.assertEqual(cell.text,'1992年8月16日',slug)
            raw=''.join(t.text or '' for t in cell._tc.iter(qn('w:t')))
            self.assertEqual(raw,'1992年8月16日',slug)

    def test_evidence_semantic_columns_and_receiver_blank(self):
        for slug in ('beijing','hunan','hainan','zhejiang','fujian','sichuan'):
            x=self.example(slug)['evidence'];_,path=get_province_document(slug,'evidence');doc=generate_evidence(x,path)
            for row in doc.tables[0].rows:
                if '接收人' in row.cells[0].text:
                    self.assertNotIn('2026',row.cells[-1].text,slug)
                for cell in row.cells:
                    if cell.text.strip() in ('证据提交日期','提交日期'):
                        self.assertNotIn('2026',cell.text)
            if slug=='beijing':
                row=doc.tables[0].rows[1];self.assertEqual(row.cells[2].text,'书证');self.assertEqual(row.cells[3].text,'有')
            elif slug=='hunan':self.assertEqual(doc.tables[0].rows[1].cells[2].text,'第1页至第5页')
            elif slug=='fujian':self.assertIn('合计：14页',doc.tables[0].rows[-1].cells[0].text)
            elif slug=='sichuan':
                row=doc.tables[0].rows[1];self.assertEqual(row.cells[4].text,'');self.assertEqual(row.cells[5].text,'√')

    def test_old_xinjiang_example_dates_do_not_leak_into_output(self):
        x=self.example('xinjiang')['application'];_,p=get_province_document('xinjiang','application');doc=generate_application(x,p)
        xml=doc._element.xml
        for old in ('1987年2月25日','2009年1月1日','2013年12月31日'):self.assertNotIn(old,xml)

    def test_filing_date_occurs_once_and_claims_use_same_size(self):
        for slug in ('hunan','hainan','fujian'):
            x=self.example(slug)['evidence'];_,p=get_province_document(slug,'evidence')
            doc=generate_evidence(x,p)
            raw=''.join(t.text or '' for t in doc._element.iter(qn('w:t')))
            self.assertEqual(raw.count('2026年9月23日'),1,slug)
        for slug in ('hunan','shandong'):
            x=self.example(slug)['application'];_,p=get_province_document(slug,'application');doc=generate_application(x,p)
            first=next(r for p in doc.paragraphs for r in p.runs if x['requests'][0] in r.text)
            second=next(r for p in doc.paragraphs for r in p.runs if x['requests'][1] in r.text)
            self.assertEqual(first.font.size,second.font.size)
            self.assertFalse(first.bold);self.assertFalse(second.bold)

if __name__=='__main__':unittest.main()
