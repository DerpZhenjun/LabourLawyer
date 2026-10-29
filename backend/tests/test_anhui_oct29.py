"""Own-source identity, fail-closed originals and lossless filling boundaries."""
import copy
import hashlib
import io
import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from regional_anhui_oct29 import BASE, ROOT, load_guide, fill_feidong_application


class AnhuiOctober29Tests(unittest.TestCase):
    def test_two_independent_guides_keep_materials_and_conflicts(self):
        yaohai, feidong = (load_guide(x) for x in ('yaohai', 'feidong'))
        self.assertEqual((yaohai['region_code'], feidong['region_code']), ('340102', '340122'))
        self.assertEqual((len(yaohai['materials']), len(feidong['materials'])), (3, 4))
        self.assertEqual((len(yaohai['templates']), len(feidong['templates'])), (4, 6))
        self.assertTrue(all(g['publication_status'] == 'unpublished' for g in (yaohai, feidong)))
        self.assertEqual(feidong['online_evidence']['observed_status'], 'auth_required')
        self.assertIn('具体时间视案件具体情况依法办理', feidong['sections']['时限说明'])
        self.assertEqual(feidong['verification_status'], 'conflict')
        self.assertNotEqual(yaohai['source_url'], feidong['source_url'])

    def test_unknown_region_has_no_fallback(self):
        for locality in ('baohe', '340122', 'other'):
            with self.assertRaises(ValueError):
                load_guide(locality)

    def test_cross_region_and_forged_annex_are_rejected(self):
        guide = load_guide('feidong')
        for key, value in [('region_code', '340102'), ('source_url', load_guide('yaohai')['source_url']), ('service_code', 'other')]:
            changed=copy.deepcopy(guide);changed[key]=value
            with self.assertRaises(ValueError):
                fill_feidong_application(changed, {})
        changed=copy.deepcopy(guide)
        changed['templates'][0]['file']='C:/arbitrary.docx'
        with self.assertRaises(ValueError):
            fill_feidong_application(changed, {})

    def test_fill_retains_all_content_and_leaves_signature_date_blank(self):
        guide=load_guide('feidong')
        original=next(t for t in guide['templates'] if t['official_kind']=='blank')
        path=Path(__file__).resolve().parents[2]/original['file']
        before=hashlib.sha256(path.read_bytes()).hexdigest()
        payload={'applicant_lines':['张虚构，女，三十岁','肥东县示例路一号，电话00000000'],
                 'respondent_lines':['肥东县虚构示例有限公司','肥东县示例路二号'],
                 'requests':['请求支付虚构工资一万元。','请求支付虚构经济补偿二千元。'],
                 'facts':'以下全为离线测试虚构事实。申请人曾在被申请人处工作，工资尚未支付。证据为虚构劳动合同及工资记录。'}
        result=Document(io.BytesIO(fill_feidong_application(guide,payload)))
        alltext='\n'.join(p.text for p in result.paragraphs)
        self.assertTrue(all(value in alltext for values in [payload['applicant_lines'],payload['respondent_lines'],payload['requests']] for value in values))
        self.assertIn(payload['facts'],alltext.replace('\n',''))
        self.assertEqual(result.paragraphs[31].text,'                      申请人：         （签名或盖章）')
        self.assertEqual(result.paragraphs[33].text,'年   月   日')
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),before)

    def test_capacity_and_unsupported_fields_fail_instead_of_truncate(self):
        guide=load_guide('feidong')
        for payload in ({'facts':'字'*281}, {'requests':['字'*260]}, {'applicant_lines':['字'*21]},
                        {'applicant_lines':['一','二','三']}, {'signature':'张某'}, {'facts':'bad\x00text'}):
            with self.assertRaises(ValueError):
                fill_feidong_application(guide,payload)

    def test_missing_and_changed_original_cannot_be_reauthorized_by_manifest(self):
        pins=json.loads((ROOT/BASE/'source-lock.json').read_text(encoding='utf8'))
        task_tmp=ROOT/'.runtime/anhui-oct29-tests'
        task_tmp.mkdir(parents=True,exist_ok=True)
        with tempfile.TemporaryDirectory(dir=task_tmp) as folder:
            root=Path(folder)
            paths=[p for p in pins if p.startswith((BASE/'feidong').as_posix()+'/')]
            for relative in paths+[str(BASE/'source-lock.json')]:
                target=root/relative;target.parent.mkdir(parents=True,exist_ok=True)
                shutil.copy2(ROOT/relative,target)
            guide=load_guide('feidong',root=root)
            blank=next(t for t in guide['templates'] if t['official_kind']=='blank')
            target=root/blank['file'];body=target.read_bytes();target.unlink()
            with self.assertRaises(FileNotFoundError):load_guide('feidong',root=root)
            target.write_bytes(body+b'changed')
            with self.assertRaises(ValueError):load_guide('feidong',root=root)
            manifest_path=root/BASE/'feidong/downloads/manifest.json'
            manifest=json.loads(manifest_path.read_text(encoding='utf8'))
            for row in manifest:
                if row['id']==blank['id']:row['sha256']=hashlib.sha256(target.read_bytes()).hexdigest()
            manifest_path.write_text(json.dumps(manifest,ensure_ascii=False),encoding='utf8')
            with self.assertRaises(ValueError):load_guide('feidong',root=root)


if __name__ == '__main__':
    unittest.main()
