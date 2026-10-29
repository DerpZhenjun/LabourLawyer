"""Public export/bootstrap/download integration, using no real application DB."""
from copy import deepcopy
from contextlib import redirect_stdout
import hashlib
import importlib.util
import json
import os
import io
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from types import SimpleNamespace
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'backend'))
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from regional_anhui_oct29 import load_guide, FEIDONG_ADAPTER_KEY
from regional_zhejiang_oct29 import parse_capture as ningbo, ADAPTER as NB_ADAPTER
from regional_guides import digest
from regional_hunan_oct29b import parse_capture as hunan
from regional_oct29_bindings import with_capabilities
from regional_adapters import with_reference_capabilities
from public_seed import bootstrap_public_data
from routes import regional_guides as routes

REPORTS = ['reports/nationwide-verification-20261029/anhui',
           'reports/nationwide-verification-20261029/zhejiang',
           'reports/nationwide-verification-20261029-b/hunan']


class PublicProvenanceIntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        runtime = ROOT / '.runtime'; runtime.mkdir(exist_ok=True)
        cls.temp = tempfile.TemporaryDirectory(dir=runtime)
        cls.fixture = Path(cls.temp.name) / 'public'; cls.fixture.mkdir()
        for report in REPORTS:
            white = json.loads((ROOT / report / 'PUBLIC-WHITELIST.json').read_text('utf8'))
            paths = [x['file'] for x in white] if isinstance(white, list) else list(white['files'])
            for relative in paths + [report + '/PUBLIC-WHITELIST.json']:
                if not relative.startswith(report + '/'): continue
                target = cls.fixture / relative; target.parent.mkdir(parents=True, exist_ok=True)
                shutil.copyfile(ROOT / relative, target)
        cls.guides = [load_guide(n, root=cls.fixture) for n in ['yaohai', 'feidong']]
        cls.guides += [ningbo(n, capture_root=cls.fixture / REPORTS[1]) for n in ['haishu', 'yinzhou']]
        cls.guides += [hunan(n, capture_root=cls.fixture / REPORTS[2]) for n in ['tianxin', 'xiangjiang']]
        for guide in cls.guides:
            guide['source_id'] = digest(guide['source_url'])[:24]
            for template in guide['templates']:
                # Preserve Yaohai's canonical missing region metadata: the
                # release/download boundary must establish it from own proof.
                if guide['district'] != '瑶海区': template['region_code'] = guide['region_code']
                kind = None
                if guide['district'] == '肥东县' and template['official_kind'] == 'blank':
                    kind = 'application'; template['adapter'] = FEIDONG_ADAPTER_KEY
                if guide['province'] == 'zhejiang' and template['official_kind'] == 'blank':
                    kind = 'evidence' if Path(template['file']).name == 'evidence-list.docx' else 'application'
                    template['adapter'] = NB_ADAPTER
                template.update(can_generate=bool(kind), document_types=[kind] if kind else [])
        cls.js_guides = []
        for code in ['320281', '320282']:
            relative = f'reports/nationwide-verification-20261029/jiangsu/{code}/guide.json'
            guide = with_capabilities(json.loads((ROOT / relative).read_text('utf8')))
            cls.js_guides.append(guide)
            # Public release contains only the four approved blanks, not private samples.
            for template in guide['templates']:
                if template.get('can_generate') is True:
                    target = cls.fixture / template['file']; target.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copyfile(ROOT / template['file'], target)
        spec = importlib.util.spec_from_file_location('public_export_integration', ROOT / 'scripts/export-public-data.py')
        cls.exporter = importlib.util.module_from_spec(spec); spec.loader.exec_module(cls.exporter)

    @classmethod
    def tearDownClass(cls): cls.temp.cleanup()

    def release(self):
        path = Path(self.temp.name) / self._testMethodName
        shutil.copytree(self.fixture, path)
        return path

    def bundle(self, root, guides):
        folder = root / 'data/public-regions'; folder.mkdir(parents=True, exist_ok=True)
        files = {}
        for filename, rows in [('guides.json', guides), ('directory.json', [])]:
            path = folder / filename; path.write_text(json.dumps(rows, ensure_ascii=False), encoding='utf8')
            files[filename] = hashlib.sha256(path.read_bytes()).hexdigest()
        (folder / 'manifest.json').write_text(json.dumps({'schema_version': 1, 'files': files}), encoding='utf8')

    def test_export_retains_only_rebuilt_public_evidence_and_canonical_safe_channel(self):
        with patch.object(self.exporter, 'ROOT', self.fixture):
            for guide in self.guides:
                exported = self.exporter.public_guide(guide)
                self.assertEqual(exported.get('evidence'), guide['evidence'])
                self.assertEqual(exported.get('channel'), guide['channel'])
                forged = deepcopy(guide); forged['evidence'] = {'cookie': 'private', 'capture_file': 'raw.har'}
                self.assertNotIn('evidence', self.exporter.public_guide(forged))
        legacy = {'source_url': 'https://example.gov.cn', 'evidence': {'cookies': 'private'},
                  'channel': {'entry_url': 'https://sso.example.gov.cn/login?ticket=secret',
                              'entry_check': {'url': 'https://sso.example.gov.cn/?ticket=secret', 'status': 'auth_required'}}}
        exported = self.exporter.public_guide(legacy)
        self.assertNotIn('evidence', exported)
        self.assertNotIn('secret', json.dumps(exported))

    def test_fresh_release_bootstrap_preserves_report_originals_and_four_js_blanks(self):
        root = self.release()
        exported = [self.exporter.public_guide(g, root=self.fixture) for g in self.guides + self.js_guides]
        self.bundle(root, exported)
        # Relocate actual public Python modules so their ROOT is the clean release,
        # not this repository. No database/session/template outside the allowlist.
        for folder in ['backend', 'WenShu']:
            for source in (ROOT / folder).rglob('*.py'):
                if 'tests' in source.relative_to(ROOT / folder).parts: continue
                target = root / source.relative_to(ROOT); target.parent.mkdir(parents=True, exist_ok=True)
                shutil.copyfile(source, target)
        env = {**os.environ, 'PYTHONPATH': str(root / 'backend') + os.pathsep + str(root), 'PYTHONIOENCODING': 'utf8'}
        script = '''import json
from pathlib import Path
from public_seed import bootstrap_public_data
from regional_guides import RegionalStore
import regional_zhejiang_oct29 as nb
root=Path.cwd();assert nb.ROOT==root
r=bootstrap_public_data(root);rows=RegionalStore(root/'backend/data/regional-guides.sqlite').inventory()
report=[t for g in rows if g['province']!='jiangsu' for t in g['templates']]
js=[t for g in rows if g['province']=='jiangsu' for t in g['templates']]
print(json.dumps({'imports':r['guides_imported'],'report_downloaded':sum(t['download_status']=='downloaded' for t in report),
 'report_generatable':sum(bool(t.get('can_generate')) for t in report),
 'js_blanks_generatable':sum(bool(t.get('can_generate')) for t in js if t['official_kind']=='blank'),
 'js_other_downloaded':sum(t['download_status']=='downloaded' and not t.get('can_generate') for t in js),
 'js_samples_downloaded':sum(t['download_status']=='downloaded' for t in js if t['official_kind']!='blank'),
 'proofs':sum('evidence' in g for g in rows)}))'''
        run = subprocess.run([sys.executable, '-c', script], cwd=root, env=env, capture_output=True, text=True, timeout=60)
        self.assertEqual(run.returncode, 0, run.stderr)
        result = json.loads(run.stdout)
        self.assertEqual(result['imports'], 8)
        self.assertEqual(result['report_downloaded'], 24)
        self.assertEqual(result['report_generatable'], 5)
        self.assertEqual(result['js_blanks_generatable'], 4)
        self.assertEqual(result['js_other_downloaded'], 0)
        self.assertEqual(result['js_samples_downloaded'], 0)
        self.assertEqual(result['proofs'], 6)
        self.assertFalse((root / 'backend/data/labourlawyer.db').exists())

    def test_download_checks_use_given_public_root_and_retain_own_samples(self):
        root = self.release()
        with patch.object(routes, 'ROOT', root):
            for guide in self.guides:
                for template in guide['templates']:
                    self.assertEqual(routes.checked_path(template, guide), (root / template['file']).resolve())
            guide = self.guides[2]
            (root / guide['evidence']['capture_file']).unlink()
            with self.assertRaises(HTTPException): routes.checked_path(guide['templates'][0], guide)

    def test_yaohai_canonical_missing_region_downloads_four_own_originals_without_mutation(self):
        root = self.release(); guide = load_guide('yaohai', root=root)
        self.assertEqual(len(guide['templates']), 4)
        self.assertTrue(all('region_code' not in t for t in guide['templates']))
        before = deepcopy(guide)
        app = FastAPI(); app.include_router(routes.router, prefix='/regional')
        query = {k: guide[k] for k in ['province', 'city', 'district']}
        with patch.object(routes, 'ROOT', root), \
             patch.object(routes, 'store', return_value=SimpleNamespace(lookup=lambda *args: guide)), \
             TestClient(app, raise_server_exceptions=False) as client:
            for template in guide['templates']:
                response = client.get('/regional/templates/' + template['id'], params=query)
                self.assertEqual(response.status_code, 200, response.text[:200] if response.status_code != 200 else '')
                self.assertEqual(response.content, (root / template['file']).read_bytes())
                self.assertEqual(hashlib.sha256(response.content).hexdigest(), template['sha256'])
                self.assertEqual(response.headers.get('cache-control'), 'no-store')
                self.assertIn('attachment;', response.headers.get('content-disposition', ''))
        self.assertEqual(guide, before, 'Download must not mutate retained candidate or source metadata')

    def test_missing_region_metadata_needs_full_own_proof_and_explicit_wrong_code_is_rejected(self):
        root = self.release(); original = load_guide('yaohai', root=root)
        app = FastAPI(); app.include_router(routes.router, prefix='/regional')
        query = {k: original[k] for k in ['province', 'city', 'district']}
        variants = []
        for code in ['340122', None, '']:
            guide = deepcopy(original); guide['templates'][0]['region_code'] = code; variants.append(guide)
        changed = deepcopy(original); changed['authority'] = '其他机关'; variants.append(changed)
        changed = deepcopy(original); changed['materials'] = []; variants.append(changed)
        changed = deepcopy(original); changed['province'] = 'other'; variants.append(changed)
        for guide in variants:
            guide['source_id'] = digest(guide['source_url'])[:24]
            snapshot = deepcopy(guide)
            with patch.object(routes, 'ROOT', root), \
                 patch.object(routes, 'store', return_value=SimpleNamespace(lookup=lambda *args: guide)), \
                 TestClient(app, raise_server_exceptions=False) as client:
                self.assertIn(client.get('/regional/templates/' + guide['templates'][0]['id'], params=query).status_code, [404, 409])
            self.assertEqual(guide, snapshot)
            self.bundle(root, [guide])
            with self.assertRaises(ValueError): bootstrap_public_data(root)
            self.assertFalse((root / 'backend/data').exists())
        proof = root / original['evidence']['基本信息']['capture_file']; proof.unlink()
        with patch.object(routes, 'ROOT', root), \
             patch.object(routes, 'store', return_value=SimpleNamespace(lookup=lambda *args: original)), \
             TestClient(app, raise_server_exceptions=False) as client:
            self.assertIn(client.get('/regional/templates/' + original['templates'][0]['id'], params=query).status_code, [404, 409])
        original['source_id'] = digest(original['source_url'])[:24]
        self.bundle(root, [original])
        with self.assertRaises(ValueError): bootstrap_public_data(root)
        self.assertFalse((root / 'backend/data').exists())

    def test_missing_report_proof_disables_capability_without_original_repo_fallback(self):
        root = self.release(); guide = deepcopy(self.guides[2])
        (root / guide['evidence']['capture_file']).unlink()
        self.bundle(root, [guide]); result = bootstrap_public_data(root)
        from regional_guides import RegionalStore
        actual = RegionalStore(root / 'backend/data/regional-guides.sqlite').inventory()[0]
        self.assertFalse(any(t['can_generate'] for t in actual['templates']))
        self.assertNotIn('evidence', actual)
        self.assertEqual(len(result['missing_attachments']), 3)

    def test_arbitrary_report_bytes_and_undistributed_js_samples_stay_closed(self):
        root = self.release(); guide = deepcopy(self.js_guides[0])
        template = next(t for t in guide['templates'] if t.get('can_generate') is True)
        forged = deepcopy(template); forged['file'] = 'reports/arbitrary/copied.docx'
        path = root / forged['file']; path.parent.mkdir(parents=True); shutil.copyfile(root / template['file'], path)
        with patch.object(routes, 'ROOT', root):
            with self.assertRaises(HTTPException): routes.checked_path(forged, guide)
            for original in guide['templates']:
                if original.get('official_kind') != 'blank':
                    with self.assertRaises(HTTPException): routes.checked_path(original, guide)
            forged_guide = deepcopy(guide)
            next(t for t in forged_guide['templates'] if t['id'] == forged['id']).update(file=forged['file'])
            self.bundle(root, [forged_guide]); result = bootstrap_public_data(root)
            self.assertIn(forged['file'], result['missing_attachments'])
            (root / template['file']).unlink()
            with self.assertRaises(HTTPException): routes.checked_path(template, guide)

    def test_export_main_preserves_canonical_config_and_migrates_only_when_missing(self):
        root = Path(self.temp.name) / self._testMethodName; root.mkdir()
        config = root / 'config/regional_candidates.json'; config.parent.mkdir()
        canonical = [{'region_code': str(index), 'candidate_web_qa': {'status': 'PASS'},
                      'source_contract': 'fixed-proof', 'reviewed_by': 'root'} for index in range(31)]
        config.write_text(json.dumps(canonical, indent=2), encoding='utf8'); before = config.read_bytes()
        report = root / 'reports/nationwide-localization-20261028/CANDIDATES.json'
        report.parent.mkdir(parents=True)
        legacy = {'province': 'test', 'city': '旧市', 'district': '旧区', 'region_code': 'old',
                  'source_id': 'old-source', 'source_url': 'https://example.gov.cn', 'published': False, 'can_generate': False}
        report.write_text(json.dumps([legacy]), encoding='utf8')
        real_connect = self.exporter.sqlite3.connect

        def memory_directory(*args, **kwargs):
            db = real_connect(':memory:')
            db.execute('CREATE TABLE directory_areas(payload TEXT,level INTEGER,province TEXT,city TEXT,district TEXT,id TEXT)')
            return db

        with patch.object(self.exporter, 'ROOT', root), \
             patch.object(self.exporter, 'RegionalStore', return_value=SimpleNamespace(inventory=lambda: [])), \
             patch.object(self.exporter.sqlite3, 'connect', side_effect=memory_directory), redirect_stdout(io.StringIO()):
            self.exporter.main()
            self.assertEqual(config.read_bytes(), before, 'Export must not overwrite canonical QA/source metadata')
            config.unlink()
            self.exporter.main()
            self.assertEqual(json.loads(config.read_text('utf8')), [legacy])

    def test_published_lifecycle_export_clean_bootstrap_and_eleven_actual_generations(self):
        root = self.release(); published = []
        for original in self.guides + self.js_guides:
            guide = deepcopy(original)
            guide.update(published=True, publication_status='published', candidate_web_qa='PASS')
            guide = with_reference_capabilities(guide)
            guide['can_generate'] = any(t.get('can_generate') for t in guide['templates'])
            published.append(guide)
        self.assertEqual(sum(g['can_generate'] for g in published), 7)
        self.assertEqual(sum(len(t.get('document_types', [])) for g in published
                             for t in g['templates'] if t.get('can_generate')), 11)
        exported = [self.exporter.public_guide(g, root=self.fixture) for g in published]
        self.assertTrue(all(g['published'] is True and g['publication_status'] == 'published' for g in exported))
        self.bundle(root, exported)
        # Current product source only; report data was copied from the explicit
        # public allowlists, and JS contains only its four approved blank files.
        for folder in ['backend', 'WenShu']:
            for source in (ROOT / folder).rglob('*.py'):
                if 'tests' in source.relative_to(ROOT / folder).parts: continue
                target = root / source.relative_to(ROOT); target.parent.mkdir(parents=True, exist_ok=True)
                shutil.copyfile(source, target)
        for filename in ['province_official_template_catalog.json', 'document_forms.json']:
            target = root / 'config' / filename; target.parent.mkdir(exist_ok=True)
            shutil.copyfile(ROOT / 'config' / filename, target)
        env = {**os.environ, 'PYTHONPATH': str(root / 'backend') + os.pathsep + str(root), 'PYTHONIOENCODING': 'utf8'}
        script = '''import hashlib,io,json
from pathlib import Path
from fastapi import FastAPI
from fastapi.testclient import TestClient
from docx import Document
from public_seed import bootstrap_public_data
from regional_guides import RegionalStore
from regional_adapters import checked_reference_adapter
from routes import regional_guides as routes
import regional_anhui_oct29 as ah
import regional_zhejiang_oct29 as nb
import regional_hunan_oct29b as hn
import regional_jiangsu_oct29 as js
root=Path.cwd()
assert all(module.ROOT==root for module in [ah,nb,hn,js,routes])
assert not (root/'backend/data').exists()
seed=bootstrap_public_data(root)
store=RegionalStore(root/'backend/data/regional-guides.sqlite')
guides=[store.lookup(g['province'],g['city'],g['district']) for g in store.inventory()]
assert len(guides)==8 and all(g['published'] is True and g['publication_status']=='published' for g in guides)
assert sum(any(t.get('can_generate') for t in g['templates']) for g in guides)==7
assert sum(len(t.get('document_types',[])) for g in guides for t in g['templates'] if t.get('can_generate'))==11
originals={t['file']:hashlib.sha256((root/t['file']).read_bytes()).hexdigest() for g in guides for t in g['templates'] if t['download_status']=='downloaded'}
routes.store=lambda:store
app=FastAPI();app.include_router(routes.router,prefix='/regional')
results=[];downloads=0;unavailable=0;wrong_source_rejected=0
with TestClient(app,raise_server_exceptions=False) as client:
 for g in guides:
  query={k:g[k] for k in ['province','city','district']}
  for t in g['templates']:
   if g['province']=='jiangsu' and not t.get('can_generate'):
    assert t['download_status']!='downloaded'
    assert not (root/t['file']).exists()
    r=client.get('/regional/templates/'+t['id'],params=query)
    assert r.status_code in [404,409,503],(g['district'],t['id'],r.status_code)
    unavailable+=1
   if not t.get('can_generate'):continue
   assert t['official_kind']=='blank'
   for kind in t['document_types']:
    checked_reference_adapter(g,t,kind)
    marker='published-fictional-applicant' if kind=='application' else 'published-fictional-evidence'
    if g['district']=='肥东县':
     data={'applicant_lines':[marker],'requests':['offline fictional wage claim'],'facts':'Fictional public-release verification only.'}
    elif kind=='application':
     data={'applicantName':marker,'respondentName':'fictional-employer','requests':['offline fictional wage claim'],
           'factsReason':'Fictional public-release verification only.','copyCount':''}
    else:
     data={'evidenceList':[{'name':marker,'source':'fictional-source','purpose':'offline verification','pageCount':'2'}]}
    payload={**query,'source_id':g['source_id'],'version':g['version'],'template_id':t['id'],'kind':kind,'mode':'blank','data':data}
    r=client.post('/regional/generate',json=payload)
    assert r.status_code==200,(g['district'],kind,r.status_code,r.text[:300])
    assert r.headers.get('x-guide-source-id')==g['source_id']
    expected_origin='official-blank-converted-filled' if g['province']=='hunan' else 'official-blank-filled'
    assert r.headers.get('x-document-origin')==expected_origin
    document=Document(io.BytesIO(r.content))
    text='\\n'.join(p.text for p in document.paragraphs)+'\\n'+'\\n'.join(c.text for table in document.tables for row in table.rows for c in row.cells)
    assert marker in text,(g['district'],kind,'filled field missing')
    results.append({'code':g['region_code'],'kind':kind,'status':r.status_code,'bytes':len(r.content),'sha256':hashlib.sha256(r.content).hexdigest()})
    r=client.get('/regional/templates/'+t['id'],params=query)
    assert r.status_code==200 and r.content==(root/t['file']).read_bytes()
    downloads+=1
    wrong={**payload,'source_id':next(other['source_id'] for other in guides if other['source_id']!=g['source_id'])}
    assert client.post('/regional/generate',json=wrong).status_code==409
    wrong_source_rejected+=1
assert len(results)==11 and downloads==11 and wrong_source_rejected==11
assert all(hashlib.sha256((root/path).read_bytes()).hexdigest()==checksum for path,checksum in originals.items())
assert not (root/'backend/data/labourlawyer.db').exists()
print(json.dumps({'generatable_localities':7,'outputs':results,'original_downloads':downloads,'wrong_source_rejected':wrong_source_rejected,
                 'js_unavailable':unavailable,'public_originals_unchanged':len(originals),'evidence_profiles':sum('evidence' in g for g in guides)}))'''
        run = subprocess.run([sys.executable, '-c', script], cwd=root, env=env, capture_output=True, text=True, timeout=120)
        self.assertEqual(run.returncode, 0, run.stderr + run.stdout)
        result = json.loads(run.stdout)
        self.assertEqual(result['generatable_localities'], 7)
        self.assertEqual(len(result['outputs']), 11)
        self.assertEqual(result['original_downloads'], 11)
        self.assertEqual(result['wrong_source_rejected'], 11)
        self.assertEqual(result['js_unavailable'], 12)
        self.assertEqual(result['public_originals_unchanged'], 28)
        self.assertEqual(result['evidence_profiles'], 6)


if __name__ == '__main__': unittest.main()
