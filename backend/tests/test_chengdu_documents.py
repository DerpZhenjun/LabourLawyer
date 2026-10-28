"""Every newly captured locality must use its own audited annex, never a neighbor."""
import hashlib
import io
import json
import os
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from regional_adapters import with_reference_capabilities
from regional_guides import RegionalStore

ROOT = Path(__file__).resolve().parents[2]
CAPTURES = ROOT / 'reports/chengdu-district-expansion-20261023/captures'


class ChengduDocumentTests(unittest.TestCase):
    def setUp(self):
        self.guides = [json.loads(p.read_text(encoding='utf8')) for p in sorted(CAPTURES.glob('*/guide.json'))]
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.store = RegionalStore(Path(temporary.name) / 'regional.sqlite')
        env = patch.dict(os.environ, {'LABOURLAWYER_REGIONAL_DB': str(self.store.path)})
        env.start()
        self.addCleanup(env.stop)
        from routes.regional_guides import router
        app = FastAPI()
        app.include_router(router, prefix='/api/regional-guides')
        self.client = TestClient(app)

    def payload(self, guide, template, kind):
        return {**{k: guide[k] for k in ['province', 'city', 'district', 'source_id', 'version']},
                'template_id': template['id'], 'kind': kind, 'mode': 'reference_adaptation',
                'data': {'applicantName': guide['district'] + '合成测试', 'committeeName': guide['district'] + '测试委员会',
                         'requests': ['请求支付测试工资100元', '请求支付测试补偿200元'],
                         'factsReason': '这是地区绑定回归测试，不是真实申请。',
                         'evidenceList': [{'name': guide['district'] + '测试证据', 'purpose': '合成测试，不实际提交。'}]}}

    def test_new_localities_generate_from_independently_downloaded_files(self):
        self.assertGreaterEqual(len(self.guides), 18)
        sources = set()
        for raw in self.guides:
            with self.subTest(district=raw['district']):
                guide = self.store.put_guide(with_reference_capabilities(raw))
                sources.add(guide['source_id'])
                templates = [t for t in guide['templates'] if t.get('can_generate')]
                self.assertEqual(len(templates), 1, raw['district'] + '缺少自己的填写适配')
                template = templates[0]
                self.assertIn('/' + guide['region_code'] + '/', template['file'])
                retained = (ROOT / template['file']).read_bytes()
                self.assertEqual(hashlib.sha256(retained).hexdigest(), template['sha256'])
                for kind in ['application', 'evidence']:
                    response = self.client.post('/api/regional-guides/generate', json=self.payload(guide, template, kind))
                    self.assertEqual(response.status_code, 200)
                    self.assertEqual(response.headers['x-guide-source-id'], guide['source_id'])
                    self.assertEqual(response.headers['x-document-origin'], 'official-sample-adaptation-not-blank')
                    doc = Document(io.BytesIO(response.content))
                    text = '\n'.join(p.text for p in doc.paragraphs)
                    if kind == 'application':
                        self.assertIn(guide['district'] + '合成测试', text)
                        self.assertIn('1. 请求支付测试工资100元', text)
                        self.assertIn('2. 请求支付测试补偿200元', text)
                    else:
                        self.assertEqual(doc.tables[0].cell(1, 1).text, guide['district'] + '测试证据')
                    with ZipFile(io.BytesIO(retained)) as before, ZipFile(io.BytesIO(response.content)) as after:
                        self.assertEqual(set(before.namelist()), set(after.namelist()))
                        for name in before.namelist():
                            if name != 'word/document.xml':
                                self.assertEqual(before.read(name), after.read(name), name)
        self.assertEqual(len(sources), len(self.guides))

    def test_neighbor_source_or_changed_annex_revokes_generation(self):
        for raw in self.guides:
            for mutation in ['district', 'source_url', 'id', 'file', 'sha256']:
                with self.subTest(district=raw['district'], mutation=mutation):
                    guide = with_reference_capabilities(deepcopy(raw))
                    templates = [t for t in guide['templates'] if t.get('can_generate')]
                    self.assertEqual(len(templates), 1)
                    template = templates[0]
                    if mutation in {'district', 'source_url'}:
                        guide[mutation] = '邻区' if mutation == 'district' else self.guides[0]['source_url'] + '&wrong=1'
                    else:
                        template[mutation] = 'changed'
                    self.assertFalse(any(t.get('can_generate') for t in with_reference_capabilities(guide)['templates']))

    def test_collected_public_screenshot_is_accessible_without_arbitrary_files(self):
        url = '/api/regional-guides/evidence/reports/chengdu-district-expansion-20261023/captures/510106000000/materials.png'
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers['content-type'], 'image/png')
        self.assertEqual(self.client.get('/api/regional-guides/evidence/reports/chengdu-district-expansion-20261023/../../README.md').status_code, 404)

    def test_inventory_includes_unverified_target_without_fabricated_guide(self):
        result = self.client.get('/api/regional-guides/inventory').json()
        qingyang = next((a for a in result.get('coverage', []) if a['district'] == '青羊区'), None)
        self.assertIsNotNone(qingyang)
        self.assertIsNone(qingyang['source_id'])
        self.assertEqual(qingyang['verification_status'], 'unverified')


if __name__ == '__main__':
    unittest.main()
