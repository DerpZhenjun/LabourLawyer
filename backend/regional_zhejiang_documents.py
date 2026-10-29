"""Strict own-source blank filling for the two independently captured Zhejiang districts.

``zhejiang_blank_binding(guide, template)`` returns application/evidence or None.
``with_zhejiang_capabilities`` returns a staged copy, never publishes anything.
``checked_zhejiang_adapter`` is the shared-route integration guard; filling is
``fill_zhejiang_document(guide, template, kind, path, data)`` and returns a
PreservedPackageDocument. No database, network or UI operations occur here.
"""
from copy import deepcopy
import hashlib
from pathlib import Path
import re
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from regional_documents import PreservedPackageDocument
from regional_guides import digest
from regional_zhejiang import parse_yuhang_capture, parse_xihu_capture, REPORT
from WenShu.Zhejiang.LAdocx import generate_LAdocx
from WenShu.exact_template_fill import fill_evidence
from WenShu.template_fill import set_cell_text

ROOT = Path(__file__).resolve().parents[1]
ADAPTER = 'zhejiang_own_blank_v1'
SOURCES = {
    '330110': ('余杭区', '8f1bc917-6583-4541-be13-231de0b91d0b', parse_yuhang_capture),
    '330106': ('西湖区', '2b882c41-1fc2-4689-9f56-631917cf63fe', parse_xihu_capture),
}
BLANKS = {
    '21B24C52D2AC3EB007AA7F55CE0AAD79': ('application', '0eab0f0a2ee98e8ca1e4d2fa5469fcb5738d0c0f24671345bebc8394099ae41f'),
    '872DFF90A8578907CB947197A784AA48': ('evidence', '4965bf3d2bf76404ab190d1e47d2fa250c8cb8bd93762949df94be38b2d4ca8c'),
}


def zhejiang_blank_binding(guide, template):
    """Require both own raw evidence and current canonical bytes, never SHA alone."""
    try:
        code = guide.get('region_code')
        if code not in SOURCES:
            return None
        district, local, parse = SOURCES[code]
        source = 'https://www.zjzwfw.gov.cn/zjservice-fe/#/workguide?localInnerCode=' + local
        if (tuple(guide.get(k) for k in ['province', 'city', 'district']) != ('zhejiang', '杭州市', district)
            or guide.get('source_url') != source or guide.get('source_id') != digest(source)[:24]
            or guide.get('authority') != '杭州市' + district + '人力资源和社会保障局'
            or guide.get('service_code') != local or guide.get('service_name') != '劳动人事争议仲裁申请'
            or guide.get('verification_status') not in {'partial', 'verified'}):
            return None
        annex = template.get('official_annex_id')
        if annex not in BLANKS:
            return None
        kind, checksum = BLANKS[annex]
        # Re-parse sealed own DOM, material controls, native request provenance,
        # source originals and identity evidence instead of trusting a JSON label.
        retained = parse(ROOT / REPORT, ROOT / REPORT / 'SHA256-MANIFEST.json')
        if (guide.get('service_ids') != retained['service_ids']
            or guide.get('service_code_type') != retained['service_code_type']
            or guide.get('fields', {}).get('实施主体') != retained['fields']['实施主体']
            or any(guide.get('evidence', {}).get(k) != retained['evidence'][k]
                   for k in ['manifest_digest', 'capture_file', 'capture_sha256'])):
            return None
        own = next(t for t in retained['templates'] if t['official_annex_id'] == annex)
        keys = ['id', 'official_annex_id', 'annex_id', 'material_id', 'region_code', 'source_url',
                'original_name', 'source_filename', 'official_kind', 'file_type', 'file',
                'retained_download_file', 'sha256', 'bytes', 'download_url', 'final_download_url',
                'http_status', 'download_status']
        if any(template.get(k) != own.get(k) for k in keys):
            return None
        path = (ROOT / own['file']).resolve()
        expected_parent = (ROOT / 'templates/_docx/_regional' / code).resolve()
        if (not path.is_relative_to(ROOT.resolve()) or path.parent != expected_parent or not path.is_file()
            or hashlib.sha256(path.read_bytes()).hexdigest() != checksum):
            return None
        return kind
    except (ValueError, KeyError, TypeError, OSError, StopIteration):
        return None


def with_zhejiang_capabilities(guide):
    """Return only these verified own blanks as staged capabilities; samples stay disabled."""
    result = deepcopy(guide)
    for template in result.get('templates', []):
        kind = zhejiang_blank_binding(result, template)
        if kind:
            template.update(can_generate=True, adapter=ADAPTER, document_types=[kind],
                content_kind='filled_official_blank', generation_label=f"依据杭州市{result['district']}官方空白表格填写")
        elif result.get('province') == 'zhejiang' or template.get('adapter') == ADAPTER:
            template.update(can_generate=False, document_types=[])
    if any(t.get('adapter') == ADAPTER and t.get('can_generate') for t in result.get('templates', [])):
        result['warnings'] = [w for w in result.get('warnings', []) if w != '空白文件及身份证示例分别保存；未启用适配填写或生成。']
        notice = '本地区两份官方空白文书已有本地填写适配；当前仅为暂存候选，未发布，浏览器与版式验收尚未完成。'
        if notice not in result['warnings']:
            result['warnings'].append(notice)
    return result


def checked_zhejiang_adapter(guide, template, kind):
    """Convenient integration guard, compatible with the existing adapter dispatch shape."""
    if (kind not in {'application', 'evidence'} or zhejiang_blank_binding(guide, template) != kind or template.get('adapter') != ADAPTER
        or template.get('document_types') != [kind] or template.get('can_generate') is not True):
        raise ValueError('浙江本区官方空白表格与地区、来源或文书种类不匹配')
    return ADAPTER


def fill_zhejiang_document(guide, template, kind, path, data):
    checked_zhejiang_adapter(guide, template, kind)
    path = Path(path)
    expected = (ROOT / template['file']).resolve()
    if path.resolve() != expected:
        raise ValueError('只能填写该地区自己的canonical官方原件')
    original = path.read_bytes()
    if hashlib.sha256(original).hexdigest() != template['sha256']:
        raise ValueError('浙江本区官方原件已变化，需重新核验')
    values = deepcopy(data)
    if values.get('copyCount') is None:
        values['copyCount'] = '    '
    if kind == 'application':
        # The generic application filler has no Zhejiang application profile.
        # Reuse its existing dedicated, inspected 23-row worker form bindings.
        document = generate_LAdocx(values, path)
        # The official two-row merged representative cell otherwise straddles
        # a page boundary. Keep both original rows together; never shrink text.
        for cell in document.tables[0].rows[12].cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next=True
        # Word needs a trailing paragraph after a full-page table. Make only
        # that empty paragraph small enough not to add an empty footer page.
        tail=document.add_paragraph()
        tail.paragraph_format.space_before=Pt(0)
        tail.paragraph_format.space_after=Pt(0)
        tail.paragraph_format.line_spacing=Pt(1)
        rpr=OxmlElement('w:rPr');size=OxmlElement('w:sz');size.set(qn('w:val'),'2');rpr.append(size)
        tail._p.get_or_add_pPr().append(rpr)
    else:
        # This retained original has a separate label/value pair. The generic
        # paragraph heuristic also matches the static 案号 label, so fill the
        # actual value cell only after the inspected Zhejiang profile returns.
        case_number = values.pop('caseNumber', None)
        values['evidenceList'] = [{**item, 'pageCount': re.sub(r'^\s*(\d+)\s*页\s*$', r'\1', str(item.get('pageCount', '')))}
                                  for item in values.get('evidenceList', [])]
        document = fill_evidence(values, path, profile='ZheJiangSheng')
        table = document.tables[0]
        if case_number:
            set_cell_text(table.rows[0].cells[1], case_number)
        for row in table.rows[:2]:
            properties = row._tr.get_or_add_trPr()
            if properties.find(qn('w:tblHeader')) is None:
                properties.append(OxmlElement('w:tblHeader'))
        for row in table.rows:
            properties = row._tr.get_or_add_trPr()
            if properties.find(qn('w:cantSplit')) is None:
                properties.append(OxmlElement('w:cantSplit'))
        for row in table.rows[-2:]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
    return PreservedPackageDocument(document, original)
