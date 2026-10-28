"""Offline filling for Shangcheng's own captured blank originals only.

Bindings require the private sealed own-source capture and canonical originals.
Missing proof disables the binding; equal provincial/other-district bytes confer
no capability. This module never opens a database, calls a model, or publishes.
"""
from copy import deepcopy
import hashlib,re
from pathlib import Path
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from regional_guides import digest
from regional_documents import PreservedPackageDocument
from regional_zhejiang_more import parse_shangcheng_capture,REPORT,MANIFEST_DIGEST
from WenShu.Zhejiang.LAdocx import generate_LAdocx
from WenShu.exact_template_fill import fill_evidence
from WenShu.template_fill import set_cell_text

ROOT=Path(__file__).resolve().parents[1]
ADAPTER='zhejiang_shangcheng_own_blank_v1'
SOURCE='https://www.zjzwfw.gov.cn/zjservice-fe/#/workguide?localInnerCode=01a9e0fa-ab7e-4232-99e1-014d3ad99b45'
SOURCE_ID='b5b5727640f8c28377abd27b'
LOCAL='01a9e0fa-ab7e-4232-99e1-014d3ad99b45'
AUTHORITY='杭州市上城区人力资源和社会保障局'
BLANKS={
    '21B24C52D2AC3EB007AA7F55CE0AAD79':('application','0eab0f0a2ee98e8ca1e4d2fa5469fcb5738d0c0f24671345bebc8394099ae41f'),
    '872DFF90A8578907CB947197A784AA48':('evidence','4965bf3d2bf76404ab190d1e47d2fa250c8cb8bd93762949df94be38b2d4ca8c')}

def shangcheng_blank_binding(guide,template):
    """Return application/evidence only after joint own-source and byte checks."""
    try:
        if (tuple(guide.get(k) for k in ['province','city','district','region_code'])!=('zhejiang','杭州市','上城区','330102')
            or guide.get('source_url')!=SOURCE or guide.get('source_id')!=SOURCE_ID
            or digest(guide.get('source_url',''))[:24]!=SOURCE_ID
            or guide.get('authority')!=AUTHORITY or guide.get('service_code')!=LOCAL
            or guide.get('service_name')!='劳动人事争议仲裁申请'
            or guide.get('verification_status') not in {'partial','verified'}):return None
        annex=template.get('official_annex_id')
        if annex not in BLANKS:return None
        kind,checksum=BLANKS[annex]
        retained=parse_shangcheng_capture(ROOT/REPORT,ROOT/REPORT/'SOURCE-MANIFEST.json')
        if (guide.get('service_ids')!=retained['service_ids']
            or guide.get('service_code_type')!=retained['service_code_type']
            or guide.get('fields',{}).get('实施主体')!=retained['fields']['实施主体']
            or guide.get('evidence',{}).get('manifest_digest')!=MANIFEST_DIGEST
            or any(guide.get('evidence',{}).get(k)!=retained['evidence'][k] for k in ['capture_file','capture_sha256'])):return None
        own=next(t for t in retained['templates'] if t['official_annex_id']==annex)
        for key in ['id','annex_id','official_annex_id','material_id','region_code','source_url','original_name',
            'source_filename','official_kind','file_type','file','retained_download_file','sha256','bytes',
            'download_url','final_download_url','http_status','download_status']:
            if template.get(key)!=own.get(key):return None
        path=(ROOT/own['file']).resolve()
        if (path.parent!=(ROOT/'templates/_docx/_regional/330102').resolve()
            or not path.is_relative_to(ROOT.resolve()) or not path.is_file()
            or hashlib.sha256(path.read_bytes()).hexdigest()!=checksum):return None
        return kind
    except (ValueError,TypeError,KeyError,AttributeError,OSError,StopIteration):return None

def with_shangcheng_capabilities(guide):
    """Return a staged copy; this does not publish or write any shared state."""
    result=deepcopy(guide)
    for template in result.get('templates',[]):
        kind=shangcheng_blank_binding(result,template)
        if kind:
            template.update(adapter=ADAPTER,can_generate=True,document_types=[kind],content_kind='filled_official_blank',
                generation_label='依据上城区自身官方空白表格填写',supported_applicant='worker' if kind=='application' else None)
        elif result.get('region_code')=='330102' or template.get('adapter')==ADAPTER:
            template.update(can_generate=False,document_types=[])
    return result

def checked_shangcheng_adapter(guide,template,kind):
    if (kind not in {'application','evidence'} or shangcheng_blank_binding(guide,template)!=kind
        or template.get('adapter')!=ADAPTER or template.get('can_generate') is not True
        or template.get('document_types')!=[kind]):
        raise ValueError('上城区自身来源、机关、附件或文书种类未通过填写绑定')
    return ADAPTER

def fill_shangcheng_document(guide,template,kind,path,data):
    """Fill only this exact canonical original and preserve every other ZIP part."""
    checked_shangcheng_adapter(guide,template,kind)
    path=Path(path)
    if path.resolve()!=(ROOT/template['file']).resolve():raise ValueError('不是上城区自身canonical原件路径')
    original=path.read_bytes()
    if hashlib.sha256(original).hexdigest()!=template['sha256']:raise ValueError('上城区原件字节改变，需重新核验')
    values=deepcopy(data)
    if values.get('copyCount') is None:values['copyCount']='    '
    if kind=='application':
        # Own OOXML inspection confirms this identical layout; jurisdiction is
        # established above, never by importing another locality's binding.
        document=generate_LAdocx(values,path)
        for cell in document.tables[0].rows[12].cells:
            for paragraph in cell.paragraphs:paragraph.paragraph_format.keep_with_next=True
        # Word requires a paragraph after a full-page table. Only its empty
        # marker is 1pt; no source text or form field font is reduced.
        tail=document.add_paragraph();tail.paragraph_format.space_before=Pt(0);tail.paragraph_format.space_after=Pt(0)
        tail.paragraph_format.line_spacing=Pt(1)
        run=OxmlElement('w:rPr');size=OxmlElement('w:sz');size.set(qn('w:val'),'2');run.append(size)
        tail._p.get_or_add_pPr().append(run)
    else:
        # The generic helper's paragraph heuristic also matches the static
        # two-character 案号 label. Fill only this source's actual value cell.
        case_number=values.pop('caseNumber',None)
        values['evidenceList']=[{**item,'pageCount':re.sub(r'^\s*(\d+)\s*页\s*$',r'\1',str(item.get('pageCount','')))}
            for item in values.get('evidenceList',[])]
        document=fill_evidence(values,path,profile='ZheJiangSheng')
        table=document.tables[0]
        if case_number:set_cell_text(table.rows[0].cells[1],case_number)
        for row in table.rows[:2]:
            properties=row._tr.get_or_add_trPr()
            if properties.find(qn('w:tblHeader')) is None:properties.append(OxmlElement('w:tblHeader'))
        for row in table.rows:
            properties=row._tr.get_or_add_trPr()
            if properties.find(qn('w:cantSplit')) is None:properties.append(OxmlElement('w:cantSplit'))
        for row in table.rows[-2:]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:paragraph.paragraph_format.keep_with_next=True
    return PreservedPackageDocument(document,original)
