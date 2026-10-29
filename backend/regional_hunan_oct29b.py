"""Own-source Tianxin and 430104 Xiangjiang/Yuelu official captures, offline only.

The 430104 double label is proven only for the captured arbitration service:
the official homepage says 湖南湘江新区, the service selector says 岳麓区,
and the own catalogue names the Xiangjiang implementing authority. No general
administrative or template inheritance rule is created.
"""
import hashlib,json,re
from copy import deepcopy
from pathlib import Path
from urllib.parse import parse_qs,urlsplit,unquote
from regional_guides import digest
from regional_documents import PreservedPackageDocument
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from WenShu.template_fill import set_cell_text,replace_text,date_parts

ROOT=Path(__file__).resolve().parents[1]
REPORT='reports/nationwide-verification-20261029-b/hunan'
MANIFEST_SHA256='550a4539c7f94aeb8749721b2de1a693895753b9dde940c645a2eb5b79907f10'
DERIVED_MANIFEST_SHA256='82a18bd0b08bfe3370b2ef31603f515e82ecaba4580ff228b74391a5ee28fc61'
ADAPTER='hunan_own_legacy_worker_oct29b_v1'
SERVICE='劳动人事争议仲裁申请'
STANDARD_ID='aYcZ8ELbR56Me61a8eSG0A'
SPECS={
 'tianxin':{'code':'430103','district':'天心区','portal_label':'天心区','approve':'8d0f29c6783f4fa6b6fe2011ff14bf2e','implementation':'1143010357862420794002014302002',
   'authority':'长沙市天心区人力资源和社会保障局','phone':'0731-85899290','place':'芙蓉南路二段248号天心区综治中心'},
 'xiangjiang':{'code':'430104','district':'岳麓区','portal_label':'湖南湘江新区','approve':'515bd3bf94164cf6b295560fb0ba18e3','implementation':'TE430216HNXJXQMZ014002014302002',
   'authority':'湖南湘江新区管理委员会民政和社会保障局（退役军人事务局）','phone':'0731-88999909','place':'长沙市岳麓区银盆南路273号岳麓区人力资源市场2F'}}
ORIGINAL_SHAS={
 '1624951667887':'c6bbebf6eed3a11f1cbe5ee14cd6de9937f741c9fcda37165e06d0706e18ea97',
 '1624951667935':'756b7ee8f3d4a0d5ad23adaec090b5f1722acc7669c869638b0535dcdf9ae279',
 '1732694895374':'369d38b1bcb1d92e0be6b815d6269967fc7324afbb63cc22a94331bbfcc220cc',
 '1732694895330':'00a47b36fa43d89f8c604d66585ac5e191360dd3d546feecc82ea8d6b4307444'}

class Evidence:
 def __init__(self,root):
  self.root=Path(root).resolve();raw=(self.root/'SOURCE-MANIFEST.json').read_bytes()
  if hashlib.sha256(raw).hexdigest()!=MANIFEST_SHA256:raise ValueError('湖南本轮固定来源清单不匹配')
  self.files=json.loads(raw)['files']
 def read(self,relative):
  p=(self.root/relative).resolve()
  if not p.is_relative_to(self.root):raise ValueError('来源路径越界')
  body=p.read_bytes()
  if self.files.get(relative)!={'sha256':hashlib.sha256(body).hexdigest(),'bytes':len(body)}:raise ValueError('来源原始字节已变化：'+relative)
  return body
 def load(self,relative):return json.loads(self.read(relative).decode('utf8'))

def source_url(name):
 return 'https://zwfw-new.hunan.gov.cn/hnywtb/service/guide.html?id='+STANDARD_ID+'&type=fr&rightsCode=002014302002&ywCode=01&approve_id='+SPECS[name]['approve']
def _pairs(rows):
 result={}
 for row in rows:
  for i in range(0,len(row)-1,2):result[row[i].strip().rstrip('：:')]=row[i+1].strip()
 return result

def parse_capture(name,capture_root=None):
 if name not in SPECS:raise ValueError('仅本轮两地自身事项')
 e=Evidence(capture_root or ROOT/REPORT);s=SPECS[name];source=source_url(name)
 catalogue=e.load(name+'/public-catalogue.json');data=catalogue['public_response'];request=catalogue['observed_request']
 nested=parse_qs(urlsplit(request['url']).query).get('url',[''])[0]
 if (data['approveId']!=s['approve'] or data['approveCode']!=s['implementation'] or data['areaCode']!=s['code']+'999000'
     or data['orgName']!=s['authority'] or data['approveName']!=SERVICE or request['status']!=200 or s['approve'] not in nested):
  raise ValueError('自身公开目录中的编码、地区或实施机关不匹配')
 pages={}
 for section in ['formal-initial','materials-expanded','faq','rights-duties','result-sample','law','template-control','sample-control']:
  page=e.load(name+'/'+section+'/page.json');e.read(name+'/'+section+'/page.png')
  if page['url']!=source or SERVICE not in page['text'] or s['authority'] not in page['text'] or s['phone'] not in page['text']:
   raise ValueError('并非本区自身正式指南，或仍为通用/其他地区正文')
  pages[section]=page
 basic=pages['formal-initial'];raw=_pairs(basic['tables'][0])
 if raw.get('办理部门')!=s['authority'] or s['place'] not in raw.get('办理地点',''):raise ValueError('自身机关和完整窗口没有在界面核实')
 listing=e.load(name+('/own-list/page.json' if name=='tianxin' else '/own-list-settled/page.json'))
 if parse_qs(urlsplit(listing['url']).query).get('areacode')!=[s['code']+'999000'] or SERVICE not in listing['text']:raise ValueError('没有本地区目录检索证据')
 home=e.load(name+'/home/page.json');picker=e.load(name+'/picker/page.json')
 if s['portal_label'] not in home['text'] or s['portal_label'] not in picker['text']:raise ValueError('官方首页地区标签缺失')
 own={'source_url':source,'region_code':s['code'],'collected_at':basic['captured_at'],'verified_at':basic['captured_at']}
 def field(value,status='verified',**extra):return {**own,'raw_value':value,'normalized_value':value,'verification_status':status,**extra}
 fields={k:field(v) for k,v in raw.items()}
 for k,api in [('法定办结时限','approveLimit'),('承诺办结时限','commitmentLimit')]:fields[k]=field(data[api],'conflict',official_context=data[api+'Explain'])
 fields['到办事现场次数']=field('0次','conflict');fields['必须到现场原因']=field(raw['必需现场办理原因'],'conflict')
 fields['咨询方式']=field(raw['咨询服务'])
 expected_names=['劳动仲裁申请书','居民身份证','营业执照','证据材料'];api_materials={m['materialId']:m for m in data['materialList']}
 materials=[]
 if len(basic['materials'])!=4 or [m['text'].splitlines()[1] for m in basic['materials']]!=expected_names:raise ValueError('本地四项材料界面已变化')
 for index,block in enumerate(basic['materials']):
  api=api_materials[block['id']];details=_pairs(basic['tables'][index+1]);title=expected_names[index]
  if api['materialTitle']!=title or not block['text'].startswith('必要\n'):raise ValueError('材料身份或必要性变化')
  if title in {'居民身份证','营业执照'}:
   section='identity-exemption' if title=='居民身份证' else 'license-exemption';exemption=e.load(name+'/'+section+'/page.json')
   if '该材料已关联电子证照，如您已有共享数据，则可免提交' not in exemption['text']:raise ValueError('免交条件缺失')
   details['免提交条件']='该材料已关联电子证照，如您已有共享数据，则可免提交'
  materials.append({**own,'id':digest([source,block['id']])[:20],'official_material_id':block['id'],'name':title,'required':'必要（符合共享条件可免提交）' if title in {'居民身份证','营业执照'} else '必要',
      'copies':details.get('填报须知','') if '份数' in details.get('填报须知','') else ('1份' if '(1份)' in details.get('材料形式','') else None),
      'raw_copies':details.get('材料形式','')+'；'+details.get('填报须知',''),'format':details.get('材料形式',''),'original_copy':details.get('材料类型',''),
      'details':details,'verification_status':'partial' if title in {'劳动仲裁申请书','证据材料'} else 'verified'})
 observation=e.load(name+'/attachment-observation.json');records=e.load(name+'/originals.json')
 if observation['source_url']!=source or len(records)!=4:raise ValueError('本区四份原件或自身下载控件证据缺失')
 templates=[]
 for r in records:
  p=Path(r['file']);identity=p.stem;kind=r['official_control'];body=e.read(r['file']);meta=e.load(r['file']+'.source.json')
  if (r!=meta or r['source_url']!=source or r['approve_id']!=s['approve'] or r['area_code']!=s['code']+'999000'
      or identity not in ORIGINAL_SHAS or hashlib.sha256(body).hexdigest()!=ORIGINAL_SHAS[identity] or r['sha256']!=ORIGINAL_SHAS[identity]
      or r['http_status']!=200 or r['final_url']!=r['download_url'] or kind not in {'blank','sample'}
      or r['attach_code']!=('000002' if kind=='blank' else '000003') or p.parts[:3]!=(name,'originals',kind)
      or not any(unquote(x['url'])==unquote(r['download_url']) and x['status']==200 for x in observation['public_fetches'])):
   raise ValueError('自身原件、地区、来源或控件分类不匹配')
  official=next((a for m in data['materialList'] for a in m.get('attachList',[]) if a.get('attachId')==r['official_annex_id']),None)
  if not official or official['attachName']!=r['original_name']:raise ValueError('本区附件目录元数据不匹配')
  templates.append({**own,'id':digest([source,r['official_annex_id']])[:24],'annex_id':r['official_annex_id'],'official_annex_id':r['official_annex_id'],
      'material_id':next(m['id'] for m in materials if m['name']=='劳动仲裁申请书'),'original_name':r['original_name'],'official_kind':kind,'file_type':'doc',
      'file':REPORT+'/'+r['file'],'retained_download_file':REPORT+'/'+r['file'],'sha256':r['sha256'],'bytes':len(body),'download_url':r['download_url'],
      'final_download_url':r['final_url'],'http_status':200,'downloaded_at':r['downloaded_at'],'download_status':'downloaded','can_generate':False,'document_types':[],
      'content_role':'official_blank_legacy_doc' if kind=='blank' else 'official_sample_never_default_case_facts'})
 boundary=e.load(name+'/online-boundary.json');e.read(name+'/online-boundary.png')
 if boundary['origin']!='https://auth.zwfw.hunan.gov.cn' or boundary['path']!='/oauth2/govLogin.html' or '个人登录' not in boundary['text']:raise ValueError('自身真实在线入口认证边界缺失')
 def section(page,start,end):return page['text'].split(start,1)[-1].split(end,1)[0].strip()
 sections={'基本信息':'\n'.join(k+'：'+v for k,v in raw.items()),'受理条件':data['acceptanceCondition'],'申请材料':'\n'.join(m['text'] for m in pages['materials-expanded']['materials']),
    '流程说明':section(basic,'流程说明\n','特别程序\n'),'办理流程':'\n'.join('｜'.join(row) for row in basic['tables'][-1]),
    '法律依据':section(pages['law'],'法律依据 权利与义务 结果样本 常见问题\n','基本信息\n'),
    '权利与义务':section(pages['rights-duties'],'法律依据 权利与义务 结果样本 常见问题\n','基本信息\n'),
    '结果样本':section(pages['result-sample'],'法律依据 权利与义务 结果样本 常见问题\n','基本信息\n'),
    '常见问题':section(pages['faq'],'法律依据 权利与义务 结果样本 常见问题\n','基本信息\n')}
 return {'province':'hunan','province_name':'湖南省','city':'长沙市','district':s['district'],'region_code':s['code'],'service_name':SERVICE,
    'service_code':s['implementation'],'official_approve_id':s['approve'],'authority':s['authority'],'source_url':source,'source_id':digest(source)[:24],
    'source_title':SERVICE,'source_type':'official_service_guide','collected_at':basic['captured_at'],'verified_at':max(x['captured_at'] for x in pages.values()),
    'verification_status':'conflict','published':False,'fields':fields,'materials':materials,'templates':templates,'sections':sections,
    'scope_identity':{'requested_district':s['district'],'portal_area_label':s['portal_label'],'portal_area_code':s['code']+'999000','selector_area_label':s['district'],
        'selection_for_navigation_only':True,'scope':'仅本轮自己正式仲裁事项；不建立湘江新区与岳麓或其他直管托管区的一般继承'},
    'channel':{'form':'window_and_online','raw_value':raw.get('办理形式'),'online_status':'auth_required','entry_url':source,
       'entry_check':{'url':'https://auth.zwfw.hunan.gov.cn/oauth2/govLogin.html','official_link':True,'status':'auth_required','verified_at':boundary['captured_at'],'submitted':False,'post_login_flow':'unknown'}},
    'evidence_template_status':'not_found_in_checked_sources','evidence':{'manifest_sha256':MANIFEST_SHA256,'capture_file':REPORT+'/'+name+'/formal-initial/page.json','capture_sha256':hashlib.sha256(e.read(name+'/formal-initial/page.json')).hexdigest()},
    'warnings':['仅本区本事项，不按户籍或本轮导航选择替用户判断个案管辖。','抬头60工作日与正文45日及可延长15日单位不同，环节表另列自己的阶段时限；保持原文，不合并承诺结案日期。',
       '基本信息0次、无需到现场，与本区常见问题关于现场办理的回答不一致；应向本区机关确认。','申请书和证据材料表内2份与填报说明当事人数量+1分别保留，原件的副本说明也另行保留。',
       '身份证、营业执照只在有已关联共享数据时可免提交，不是普遍免交。','四原件分别按空白和示例保存，样例不作为默认案件事实。',
       '官方劳动者样例含开福地址及1997.10.32无效日期，只作样例原件留存，不借作本区案件事实。','当前仅未发布候选，候选网页验收NOT_RUN；已核验来源中未找到独立证据清单空白，记not_found_in_checked_sources。']}

def parse_tianxin_capture(capture_root=None):return parse_capture('tianxin',capture_root)
def parse_xiangjiang_capture(capture_root=None):return parse_capture('xiangjiang',capture_root)

def _derived(name):
 raw=(ROOT/REPORT/'DERIVED-MANIFEST.json').read_bytes()
 if hashlib.sha256(raw).hexdigest()!=DERIVED_MANIFEST_SHA256:raise ValueError('转换副本固定清单不匹配')
 rows=json.loads(raw)['files'];item=next(x for x in rows if x['district']==name and x['file']=='derived/'+name+'/1624951667887.docx')
 source=ROOT/REPORT/item['source'];path=ROOT/REPORT/item['file']
 if (item['source_sha256']!=ORIGINAL_SHAS['1624951667887'] or hashlib.sha256(source.read_bytes()).hexdigest()!=item['source_sha256']
     or hashlib.sha256(path.read_bytes()).hexdigest()!=item['sha256'] or item['kind']!='platform_conversion_of_own_official_doc_not_official_docx'):
  raise ValueError('本区转换副本与自身DOC原件的链条不匹配')
 return item,path

def blank_binding(guide,template):
 try:
  name=next(n for n,s in SPECS.items() if s['code']==guide.get('region_code'));retained=parse_capture(name)
  for key in ['province','city','district','region_code','service_code','official_approve_id','authority','source_url','source_id','source_type','fields','materials','evidence','scope_identity','verification_status']:
   if guide.get(key)!=retained[key]:return None
  original=next(t for t in retained['templates'] if t['id']==template.get('id'))
  if original['official_kind']!='blank' or Path(original['file']).stem!='1624951667887':return None
  if any(template.get(k)!=v for k,v in original.items() if k not in {'can_generate','document_types'}):return None
  _derived(name)
  return 'application'
 except (ValueError,OSError,KeyError,TypeError,StopIteration):return None

def with_capabilities(guide):
 result=deepcopy(guide)
 for t in result['templates']:
  kind=blank_binding(result,t);t.update(can_generate=bool(kind),document_types=[kind] if kind else [])
  if kind:t.update(adapter=ADAPTER,content_kind='filled_platform_conversion_of_own_official_doc',source_download_format='doc',
      generation_label='依据自身官方DOC空白表的平台转换副本填写（劳动者申请）')
 result['warnings'].append('填写仅支持本区劳动者申请书的平台转换副本；官方原件为DOC，转换件不是官方DOCX原件，用人单位版与独立证据清单未适配。')
 return result

def fill_document(guide,template,kind,path,data):
 if kind!='application' or blank_binding(guide,template)!=kind or template.get('adapter')!=ADAPTER or template.get('can_generate') is not True:
  raise ValueError('仅支持本区严格来源绑定的劳动者申请书转换副本')
 if Path(path).resolve()!=(ROOT/template['file']).resolve():raise ValueError('不得传入其他地区或任意原件路径')
 name=next(n for n,s in SPECS.items() if s['code']==guide['region_code']);item,derived_path=_derived(name)
 values=deepcopy(data);requests=values.get('requests') or [];items=values.get('evidenceList') or []
 if not isinstance(requests,list) or not isinstance(items,list) or len(requests)>8 or len(items)>10:raise ValueError('本表最多8项请求和10项附随证据')
 if len(str(values.get('factsReason') or ''))>1800:raise ValueError('事实理由超出已验收容量')
 for k,v in values.items():
  if k not in {'requests','evidenceList','factsReason'} and len(str(v or ''))>160:raise ValueError('本表字段过长：'+k)
 original=derived_path.read_bytes();document=Document(derived_path);table=document.tables[0];rows=list(table.rows)
 if (len(document.tables)!=1 or len(rows)!=57 or len(table.columns)!=1 or '申请人姓名' not in rows[0].cells[0].text
     or '请求事项' not in rows[10].cells[0].text or '事实与理由' not in rows[19].cells[0].text or '证据和证据来源' not in rows[34].cells[0].text):
  raise ValueError('该转换副本的已检查字段位置发生变化')
 def v(k):return str(values.get(k) or '')
 def put(i,text):set_cell_text(rows[i].cells[0],text)
 year,month,day=date_parts(values.get('birthDate'));birth=(f'{year}年{month}月{day}日' if year and month and day else v('birthDate'))
 slots={0:'申请人姓名：'+v('applicantName'),1:'性别：'+v('gender')+'    出生年月日：'+birth+'    职业：'+v('occupation')+'    民族：'+v('nationality'),
    2:'住所：'+(v('idAddress') or v('address'))+'    邮编：'+v('postalCode')+'    电话：'+v('phone'),
    3:'工作单位：'+v('workUnit')+'    邮编：'+v('workUnitPostalCode')+'    电话：'+v('workUnitPhone'),
    4:'被申请人名称：'+v('respondentName'),5:'住所：'+(v('registeredAddress') or v('respondentAddress'))+'    邮编：'+v('respondentPostalCode'),
    6:'法定代表人（或主要负责人）：'+v('legalRepName')+'    职务：'+v('legalRepPosition')+'    电话：'+v('legalRepPhone'),
    7:'人力资源管理部门负责人：'+v('hrContactName')+'    职务：'+v('hrContactPosition')+'    电话：'+v('hrContactPhone'),8:'第三人：'+v('thirdParty')}
 for i,text in slots.items():put(i,text)
 for i,request in enumerate(requests):put(11+i,f'{i+1}. {request}')
 put(20,v('factsReason'))
 # Remove only unused fill rows from this exact source, preserving labels,
 # signatures, notes and the original page/table geometry and font sizes.
 for i in list(range(11+max(1,len(requests)),19))+list(range(21,33)):
  table._tbl.remove(rows[i]._tr)
 footer=rows[33].cells[0]
 if v('committeeName'):replace_text(footer.paragraphs[0],'×××劳动争议仲裁委员会',v('committeeName'))
 p=footer.paragraphs[1];text=p.text;year,month,day=date_parts(values.get('applicationDate'))
 if year and month and day:text=re.sub(r'年\s+月\s+日',f'{year}年 {month}月 {day}日',text,count=1)
 if values.get('copyCount') is not None:text=re.sub(r'申请书副本\s+份',lambda m:'申请书副本 '+str(values['copyCount'])+' 份',text,count=1)
 if text!=p.text:replace_text(p,p.text,text)
 for i,evidence in enumerate(items):
  if not isinstance(evidence,dict) or any(len(str(evidence.get(k) or ''))>300 for k in ['name','source','purpose']):raise ValueError('附随证据格式或文字长度不符合已验收容量')
  put(35+i,f"{i+1}. {evidence.get('name') or ''}；来源：{evidence.get('source') or ''}；证明目的：{evidence.get('purpose') or ''}")
 for i in range(35+max(1,len(items)),57):table._tbl.remove(rows[i]._tr)
 for i in [10,19,34]:
  for paragraph in rows[i].cells[0].paragraphs:paragraph.paragraph_format.keep_with_next=True
 for i in range(35,35+len(items)):
  props=rows[i]._tr.get_or_add_trPr()
  if props.find(qn('w:cantSplit')) is None:props.append(OxmlElement('w:cantSplit'))
 return PreservedPackageDocument(document,original)
