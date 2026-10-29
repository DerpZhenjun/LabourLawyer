"""Offline, exact-own-source Zhejiang D worker forms; explicit roots never fall back.

No browser, database, model or service operations. Frozen source files and their
public whitelist are read-only. Only word/document.xml changes in filled copies.
Document layout and candidate browser acceptance are separate deployment gates.
"""
from copy import deepcopy
from datetime import date
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import _Row

from regional_documents import PreservedPackageDocument
from regional_guides import digest
from WenShu.template_fill import replace_span, replace_text, set_cell_text

ROOT = Path(__file__).resolve().parents[1]
REPORT = 'reports/nationwide-verification-20261029-d/zhejiang'
ADAPTER = 'zhejiang_own_worker_oct29d_v1'
MANIFEST_SHA256 = 'c71d92d0d46b59a6708bbb40d75c34eb55f566a3b0b354fde0f6583b4c1c1857'
WHITELIST_SHA256 = '9ff527502ded47094a161caeacbf3e69660dccc88e7062246d20d978902361f3'
SERVICE = '劳动人事争议仲裁申请'
BLANKS = {
    'application': ('21B24C52D2AC3EB007AA7F55CE0AAD79', '0eab0f0a2ee98e8ca1e4d2fa5469fcb5738d0c0f24671345bebc8394099ae41f'),
    'evidence': ('872DFF90A8578907CB947197A784AA48', '4965bf3d2bf76404ab190d1e47d2fa250c8cb8bd93762949df94be38b2d4ca8c'),
}
SPECS = {
    'xiaoshan': dict(code='330109', district='萧山区', local='7c940b86-e89e-47e9-9785-292358209e9a',
                     task='11330109002514048E4002014302002', phone='0571-89653752', place='金城路1558号',
                     materials=['57C6322B03203346C170BAABFB2556BB', 'A9A06FFF502EED9FB73557D9783A5B6B',
                                '34F0BF719BB4A628FF52D29A2CB66C6C', 'DFA6D0E0FAB7998B51EAA2EF0B11737C',
                                'B6325416B6D2EE9E4E4BAD499F60FB2C'],
                     kinds={'application': '34F0BF719BB4A628FF52D29A2CB66C6C', 'evidence': 'DFA6D0E0FAB7998B51EAA2EF0B11737C'}),
    'linan': dict(code='330112', district='临安区', local='41c59062-1bee-45d9-a9ff-645b91059195',
                  task='11330185002508182w4002014302002', phone='0571-61071571', place='衣锦街468号',
                  materials=['E022B99BBAF91A524B2C719D57D4F7D5', '15705B5F59523720AE7166D821DEAD73',
                             '0103BAC26E734E0EAA86044060CCACDA', '51E7CE70366866E21AA9BF7777BBAFD2'],
                  kinds={'application': '0103BAC26E734E0EAA86044060CCACDA'}),
}
APPLICATION_TEXT_FIELDS = frozenset('committeeName applicantName birthDate gender nationality idNumber phone householdType idAddress respondentName socialCreditCode registeredAddress mailingAddress workAddress legalRepName legalRepPhone legalRepPosition companyContact companyContactPhone factsReason applicationDate'.split())
APPLICATION_FIELDS = APPLICATION_TEXT_FIELDS | {'copyCount', 'requests'}
EVIDENCE_FIELDS = frozenset({'committeeName', 'caseNumber', 'copyCount', 'submitDate', 'evidenceList'})
EVIDENCE_ITEM_FIELDS = frozenset({'name', 'pageCount', 'source', 'purpose'})
DATA_FIELDS = {'application': APPLICATION_FIELDS, 'evidence': EVIDENCE_FIELDS, 'evidence_item': EVIDENCE_ITEM_FIELDS}
LIMITS = {'ordinary': 300, 'factsReason': 12000, 'requests': 4, 'request': 1200,
          'evidenceList': 40, 'evidence_name': 300, 'evidence_source': 300,
          'evidence_pageCount': 300, 'evidence_purpose': 1200, 'copyCount_min': 1, 'copyCount_max': 30}
CAPABILITY_FIELDS = frozenset({'can_generate', 'document_types', 'adapter', 'content_kind', 'generation_label'})
GUIDE_METADATA = frozenset({'version', 'content_hash', 'region_id', 'published', 'publication_status', 'candidate_web_qa', 'can_generate', 'warnings'})
HOUSEHOLD_OPTIONS = ('本省非农业户口', '本省农业户口', '外省非农业户口', '外省非农业户口', '港澳台人员', '外籍人员')


def _root(root):
    return Path(ROOT if root is None else root).resolve()


def _path(root, relative):
    relative = Path(relative)
    if relative.is_absolute() or relative.drive or '..' in relative.parts:
        raise ValueError('来源路径必须位于显式发行根内')
    path = root / relative
    resolved = path.resolve()
    if not resolved.is_relative_to(root) or resolved.relative_to(root).as_posix() != relative.as_posix():
        raise ValueError('来源路径越界或被符号链接替换')
    return resolved


class Evidence:
    """Validate the complete fixed 52-file source bundle and exact 53-file delivery."""
    def __init__(self, root=None):
        self.root = _root(root)
        raw = _path(self.root, REPORT + '/SOURCE-MANIFEST.json').read_bytes()
        white_raw = _path(self.root, REPORT + '/PUBLIC-WHITELIST.json').read_bytes()
        if sha256(raw).hexdigest() != MANIFEST_SHA256 or sha256(white_raw).hexdigest() != WHITELIST_SHA256:
            raise ValueError('浙江D固定封存来源或公开清单已变化')
        self.files = json.loads(raw)['files']
        white = json.loads(white_raw)['files']
        expected = {REPORT + '/' + p: meta for p, meta in self.files.items()}
        expected[REPORT + '/SOURCE-MANIFEST.json'] = {'sha256': MANIFEST_SHA256, 'bytes': len(raw)}
        if len(self.files) != 52 or white != expected:
            raise ValueError('浙江D来源52文件与公开53文件链条不完整')
        for relative in self.files:
            self.read(relative)

    def read(self, relative):
        if relative not in self.files:
            raise ValueError('不在浙江D固定来源清单')
        body = _path(self.root, REPORT + '/' + relative).read_bytes()
        if self.files[relative] != {'sha256': sha256(body).hexdigest(), 'bytes': len(body)}:
            raise ValueError('浙江D封存来源字节已变化：' + relative)
        return body

    def load(self, relative):
        return json.loads(self.read(relative).decode('utf-8-sig'))


def _dom(value):
    return value['data']['result']


def parse_capture(name, root=None):
    """Rebuild only Xiaoshan/Linan own facts; root is an explicit project root."""
    if name not in SPECS:
        raise ValueError('只处理萧山和临安自身D来源')
    evidence = Evidence(root)
    spec = SPECS[name]
    source = 'https://www.zjzwfw.gov.cn/zjservice-fe/#/workguide?localInnerCode=' + spec['local']
    guide = evidence.load(name + '/guide.json')
    api = evidence.load(name + '/item-response-raw.json')['data']['item']
    if (guide['province'], guide['city'], guide['district'], guide['region_code'], guide['source_url'],
        guide['source_id'], guide['service_name'], guide['service_code'], guide['authority']) != (
            'zhejiang', '杭州市', spec['district'], spec['code'], source, digest(source)[:24], SERVICE, spec['task'], '区人力社保局'):
        raise ValueError('浙江D自身指南身份不匹配')
    if (api['jurisCode'], api['dcode'], api['taskCode'], api['name'], api['qlDept'], api['askphone']) != (
            spec['code'], spec['local'], spec['task'], SERVICE, '区人力社保局', spec['phone']):
        raise ValueError('浙江D自身正式事项、机关或地区不匹配')
    search = _dom(evidence.load(name + '/search.json'))
    if 'websiteid=' + spec['code'] + '000000' not in search['url'] or '当前搜索范围为： ' + spec['district'] not in search['text']:
        raise ValueError('浙江D自身目录选区证据不匹配')
    pages = {section: _dom(evidence.load(name + '/' + section + '-dom.json'))
             for section in ['guide', 'basic', 'faq']}
    if any(page['url'] != source or SERVICE not in page['text'] or spec['phone'] not in page['text'] for page in pages.values()):
        raise ValueError('浙江D栏目仍为其他地区或来源身份缺失')
    if spec['place'] not in pages['guide']['text'] or not any(spec['place'] in a['address'] for a in api['addresssList']):
        raise ValueError('浙江D自身窗口证据缺失')
    materials = api['materialList']
    conditions = evidence.load(name + '/condition-response-raw.json')['data']['DATAAREA']['MATERIALS']['MATERIAL']
    hidden_materials = {'24CAFAA94DCC085FB5422E15231F88AD'} if name == 'linan' else set()
    if ([m['impleCode'] for m in materials] != spec['materials'] or guide['materials'] != materials or
        {m['MATERIALGUID'] for m in conditions} != set(spec['materials']) | hidden_materials or
        any(m['QLINNERCODE'] != spec['local'] for m in conditions)):
        raise ValueError('浙江D本区材料与目录事项未完整绑定')
    if hidden_materials:
        hidden = next(m for m in conditions if m['MATERIALGUID'] in hidden_materials)
        if hidden['ISPUBLISH'] != 0 or hidden['MATERIALNAME'] != '《证据清单》':
            raise ValueError('临安未发布条件材料的原始标记已变化')
        guide['unpublished_source_materials'] = [deepcopy(hidden)]
        guide['warnings'].append('临安匿名条件响应另有ISPUBLISH=0的证据清单项，原生已发布四材料未列出该控件；不据此开放独立证据下载或生成。')
    observation = evidence.load(name + '/attachment-observation.json')
    originals = evidence.load(name + '/originals.json')
    if observation['source_url'] != source or len(guide['templates']) != len(spec['kinds']) or len(originals) != len(spec['kinds']):
        raise ValueError('浙江D自身下载控件记录不完整')
    for kind, material_id in spec['kinds'].items():
        annex, checksum = BLANKS[kind]
        material = next(m for m in materials if m['impleCode'] == material_id)
        template = next(t for t in guide['templates'] if t['id'] == material_id)
        url = 'https://portal.zjzwfw.gov.cn/media/oss/FILE/META_DATA/' + annex + '.docx'
        relative = name + '/originals/' + annex + '.docx'
        body = evidence.read(relative)
        metadata = evidence.load(relative + '.source.json')
        own = next(r for r in originals if r['material_id'] == material_id)
        if (material['blankForm'] != url.removeprefix('https:') or
            (template['material_id'], template['region_code'], template['source_url'], template['download_url'],
             template['file'], template['sha256'], template['bytes'], template['official_kind'], template['observed_document_kind']) !=
            (material_id, spec['code'], source, url, REPORT + '/' + relative, checksum, len(body), 'blank', kind) or
            metadata != own or metadata['source_url'] != source or metadata['region_code'] != spec['code'] or
            metadata['local_inner_code'] != spec['local'] or metadata['service_code'] != spec['task'] or
            metadata['download_url'] != url or metadata['sha256'] != checksum or sha256(body).hexdigest() != checksum or
            not any(r['material_id'] == material_id and r['observed_get_url'] == url and r['status'] == 200
                    for r in observation['native_controls_clicked'])):
            raise ValueError('浙江D本区材料、空白附件、原件或下载身份不匹配')
        template.update(official_annex_id=annex, file_type='docx', applicant_scope='worker' if kind == 'application' else 'evidence_list')
    boundary = evidence.load(name + '/online-boundary.json')
    if (boundary['origin'], boundary['path'], boundary['login_visible'], boundary['authenticated'], boundary['submitted'], boundary['query_saved']) != (
            'https://user.zjzwfw.gov.cn', '/pc/login', True, False, False, False):
        raise ValueError('浙江D自身原生网办认证边界不匹配')
    guide['materials'] = [dict(id=m['impleCode'], name=m['materialName'], required=m['materialNecessity'],
                             copies=m['paperMaterialNums'], raw_copies='纸质材料 ' + m['paperMaterialNums'] + '份，' + m['paperMaterialSpecific'],
                             original_copy=m['materialType'], format='纸质材料，' + m['paperMaterialSpecific'],
                             source_channel=m['sourceChannel'], instructions=m['instructions'], acceptance_standard=m['acceptCriteria'],
                             source_url=source, region_code=spec['code'], raw=deepcopy(m)) for m in materials]
    for label, value in [('办理地点', '\n'.join(a['address'] for a in api['addresssList'])),
                         ('办理时间', '\n'.join(a['acceptTimeDesc'] for a in api['addresssList'])),
                         ('是否收费', '不收费')]:
        guide['fields'][label] = {'raw_value': value, 'source_url': source, 'source_basis': 'own native basic/window tab and own public item response'}
    guide['sections'] = {'受理条件': api['spconditions'], '办理流程': api['outFlowDesc'],
                         '窗口办理': '\n'.join(a['address'] + '\n' + a['acceptTimeDesc'] for a in api['addresssList']),
                         '常见问题': '\n\n'.join(q['question'] + '\n' + q['answer'] for q in api['qaList']),
                         '法定依据': api['fdconditions']}
    guide['service_ids'] = {'localInnerCode': spec['local'], 'taskCode': spec['task']}
    guide['evidence'].update(manifest_sha256=MANIFEST_SHA256, public_whitelist_sha256=WHITELIST_SHA256)
    guide['channel']['entry_url'] = source
    guide['channel']['entry_check'] = {'url': boundary['origin'] + boundary['path'], 'status': 'auth_required',
                                     'official_link': True, 'checked_at': boundary['captured_at'], 'submitted': False, 'post_login_flow': 'unknown'}
    return guide


def _capability(template, kind):
    return {**template, 'can_generate': True, 'document_types': [kind], 'adapter': ADAPTER,
            'content_kind': 'filled_official_blank', 'generation_label': '依据本区自身官方空白表格填写（申请仅劳动者部分）'}


def _name(guide):
    return next(name for name, spec in SPECS.items() if spec['code'] == guide.get('region_code'))


def _template_matches(actual, canonical, kind):
    if not isinstance(actual, dict):
        return False
    allowed = _capability(canonical, kind)
    if set(actual) - set(allowed):
        return False
    return all(actual.get(key) == value or (key in CAPABILITY_FIELDS and actual.get(key) == allowed[key])
               for key, value in canonical.items()) and all(actual[key] == allowed[key] for key in set(actual) - set(canonical))


def _validate_guide(guide, root=None):
    name = _name(guide)
    canonical = parse_capture(name, root=root)
    if set(guide) - set(canonical) - GUIDE_METADATA:
        raise ValueError('浙江D指南包含非自身来源字段')
    if any(guide.get(key) != value for key, value in canonical.items() if key not in GUIDE_METADATA | {'templates'}):
        raise ValueError('浙江D指南事实、材料或来源不匹配')
    templates = guide.get('templates')
    if not isinstance(templates, list) or len(templates) != len(canonical['templates']):
        raise ValueError('浙江D本区附件集合不匹配')
    if len({t.get('id') for t in templates}) != len(templates):
        raise ValueError('浙江D附件重复')
    for own in canonical['templates']:
        actual = next((t for t in templates if t.get('id') == own['id']), None)
        if not _template_matches(actual, own, own['observed_document_kind']):
            raise ValueError('浙江D本区附件身份不匹配')
    return canonical


def blank_binding(guide, template, root=None):
    try:
        canonical = _validate_guide(guide, root)
        own = next(t for t in canonical['templates'] if t['id'] == template.get('id'))
        kind = own['observed_document_kind']
        return kind if _template_matches(template, own, kind) else None
    except (OSError, ValueError, KeyError, TypeError, StopIteration, AttributeError):
        return None


def with_capabilities(guide, root=None):
    canonical = _validate_guide(guide, root)
    kinds = {t['id']: t['observed_document_kind'] for t in canonical['templates']}
    result = deepcopy(guide)
    result['templates'] = [_capability(t, kinds[t['id']]) for t in result['templates']]
    result['can_generate'] = True
    offline_notice = '仅本区自身空白文书完成离线填写；申请默认且仅支持劳动者区段，单位申请未适配。版式及候选网页验收另行记录；不代表已发布。'
    published_notice = '本区自身空白文书已通过填写、版式及网页验收；申请仅支持劳动者区段，单位申请未适配。'
    notice = published_notice if result.get('published') is True and result.get('candidate_web_qa') == 'PASS' else offline_notice
    result['warnings'] = [w for w in result.get('warnings', []) if '不提供生成' not in w and w not in {offline_notice, published_notice}]
    if notice not in result['warnings']:
        result['warnings'].append(notice)
    return result


def checked_adapter(guide, template, kind, root=None):
    if (kind not in {'application', 'evidence'} or blank_binding(guide, template, root) != kind or
        template.get('adapter') != ADAPTER or template.get('can_generate') is not True or template.get('document_types') != [kind]):
        raise ValueError('仅支持浙江D本区严格来源绑定的实际空白文书')
    return ADAPTER


def public_evidence(guide, root):
    """Return canonical safe proof using only the required explicit public root."""
    try:
        if root is None:
            return None
        return deepcopy(_validate_guide(guide, root)['evidence'])
    except (OSError, ValueError, KeyError, TypeError, StopIteration, AttributeError):
        return None


def approved_report_attachment(guide, template, root):
    """Authorize only this exact own blank, never an identity example or hidden row."""
    try:
        if root is None:
            return False
        canonical = _validate_guide(guide, root)
        own = next(t for t in canonical['templates'] if t['id'] == template.get('id'))
        return own['official_kind'] == 'blank' and _template_matches(template, own, own['observed_document_kind'])
    except (OSError, ValueError, KeyError, TypeError, StopIteration, AttributeError):
        return False


def _text(value, limit=300):
    if not isinstance(value, str) or len(value) > limit or re.search('[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]', value):
        raise ValueError('填写类型、字符或容量不符合本表，不会截断内容')


def _copies(value):
    if value is None or value == '':
        return ''
    if type(value) is int and 1 <= value <= 30:
        return str(value)
    if isinstance(value, str) and value.isascii() and value.isdecimal() and 1 <= int(value) <= 30:
        return value
    raise ValueError('副本份数须为1至30的整数或留空')


def _validate_data(kind, data):
    if not isinstance(data, dict) or set(data) - (APPLICATION_FIELDS if kind == 'application' else EVIDENCE_FIELDS):
        raise ValueError('包含本区表格不支持的字段，不会静默丢弃')
    if 'copyCount' in data:
        _copies(data['copyCount'])
    if kind == 'application':
        for key in APPLICATION_TEXT_FIELDS & data.keys():
            _text(data[key], 12000 if key == 'factsReason' else 300)
        requests = data.get('requests', [])
        if not isinstance(requests, list) or len(requests) > 4:
            raise ValueError('本区原表最多4项请求，不会截断')
        for request in requests:
            _text(request, 1200)
        selected = data.get('householdType', '')
        if selected and HOUSEHOLD_OPTIONS.count(selected) != 1:
            raise ValueError('户籍选项不存在或原件同名选项重复，请留空后本人核对勾选')
    else:
        for key in {'committeeName', 'caseNumber', 'submitDate'} & data.keys():
            _text(data[key])
        items = data.get('evidenceList', [])
        if not isinstance(items, list) or len(items) > 40:
            raise ValueError('最多40项证据，不会截断')
        for item in items:
            if not isinstance(item, dict) or set(item) - EVIDENCE_ITEM_FIELDS:
                raise ValueError('包含本证据表没有的栏目，不会静默丢弃')
            for key, value in item.items():
                _text(value, 1200 if key == 'purpose' else 300)


def _date(value):
    try:
        parsed = date.fromisoformat(value)
        return f'{parsed.year}年{parsed.month}月{parsed.day}日'
    except (TypeError, ValueError):
        return value


def _keep_together(row):
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn('w:cantSplit')) is None:
        properties.append(OxmlElement('w:cantSplit'))


def _application(document, values):
    if len(document.tables) != 2 or len(document.tables[0].rows) != 23 or len(document.tables[0].columns) != 5:
        raise ValueError('本区原件劳动者/单位双区段结构已改变')
    body = document.element.body
    seen = False
    for child in list(body):
        if child.tag == qn('w:tbl') and not seen:
            seen = True
        elif seen and child.tag != qn('w:sectPr'):
            body.remove(child)
    table = document.tables[0]
    rows = list(table.rows)
    slots = {(1, 1): 'applicantName', (1, 4): 'birthDate', (2, 1): 'gender', (2, 4): 'nationality',
             (3, 1): 'idNumber', (3, 4): 'phone', (5, 2): 'idAddress', (7, 2): 'respondentName',
             (8, 2): 'socialCreditCode', (9, 2): 'registeredAddress', (10, 2): 'mailingAddress', (11, 2): 'workAddress',
             (12, 2): 'legalRepName', (12, 4): 'legalRepPhone', (13, 4): 'legalRepPosition',
             (14, 2): 'companyContact', (14, 4): 'companyContactPhone'}
    for (row, column), key in slots.items():
        value = values.get(key, '')
        set_cell_text(rows[row].cells[column], _date(value) if key == 'birthDate' else value)
    selected = values.get('householdType')
    if selected:
        for paragraph in rows[4].cells[1].paragraphs:
            runs = list(paragraph.runs)
            for index, run in enumerate(runs):
                if selected in run.text:
                    if index == 0 or not runs[index - 1]._r.findall(qn('w:sym')):
                        raise ValueError('本区户籍选项前的官方方框位置发生变化')
                    checkbox = runs[index - 1]
                    checkbox.text = '☑'
                    checkbox.font.name = 'Segoe UI Symbol'
                    checkbox._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Segoe UI Symbol')
    requests = values.get('requests', [])
    for index, request in enumerate(requests):
        set_cell_text(rows[16 + index].cells[0], str(index + 1) + '．' + request)
    if requests:
        for row in rows[16 + len(requests):20]:
            table._tbl.remove(row._tr)
    set_cell_text(rows[21].cells[0], values.get('factsReason', ''))
    rows[21].height = rows[16].height
    footer = rows[22].cells[0].paragraphs
    if values.get('committeeName'):
        replace_span(footer[2], 0, len(footer[2].text), values['committeeName'])
    if values.get('applicationDate'):
        replace_text(footer[4], '年   月    日', _date(values['applicationDate']))
    if _copies(values.get('copyCount')):
        replace_text(footer[5], '申请书副本      份', '申请书副本 ' + _copies(values['copyCount']) + ' 份')
    for cell in rows[12].cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True
    footer[3].paragraph_format.keep_with_next = True
    _keep_together(rows[22])
    tail = document.add_paragraph()
    tail.paragraph_format.space_before = Pt(0)
    tail.paragraph_format.space_after = Pt(0)
    tail.paragraph_format.line_spacing = Pt(1)
    properties = OxmlElement('w:rPr')
    size = OxmlElement('w:sz')
    size.set(qn('w:val'), '2')
    properties.append(size)
    tail._p.get_or_add_pPr().append(properties)


def _evidence(document, values):
    if len(document.tables) != 1 or len(document.tables[0].rows) != 13 or len(document.tables[0].columns) != 6:
        raise ValueError('本区原件证据表结构已改变')
    table = document.tables[0]
    rows = list(table.rows)[2:11]
    items = values.get('evidenceList', [])
    while len(rows) < len(items):
        node = deepcopy(rows[-1]._tr)
        rows[-1]._tr.addnext(node)
        rows.append(_Row(node, table))
    for index, item in enumerate(items):
        for column, key in [(0, None), (1, 'name'), (3, 'pageCount'), (4, 'source'), (5, 'purpose')]:
            set_cell_text(rows[index].cells[column], str(index + 1) if key is None else item.get(key, ''))
    if items:
        for row in rows[len(items):]:
            table._tbl.remove(row._tr)
    if values.get('committeeName'):
        replace_span(document.paragraphs[0], 0, len(document.paragraphs[0].text), values['committeeName'])
    if values.get('caseNumber'):
        set_cell_text(table.rows[0].cells[1], values['caseNumber'])
    copies = _copies(values.get('copyCount'))
    if copies:
        set_cell_text(table.rows[0].cells[5], copies)
        replace_text(document.paragraphs[-1], '本表一式     份', '本表一式 ' + copies + ' 份')
    if values.get('submitDate'):
        set_cell_text(table.rows[-2].cells[5], _date(values['submitDate']))
    for row in table.rows:
        _keep_together(row)
    for row in table.rows[:2]:
        properties = row._tr.get_or_add_trPr()
        if properties.find(qn('w:tblHeader')) is None:
            properties.append(OxmlElement('w:tblHeader'))
    for row in table.rows[-2:]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True


def fill_document(guide, template, kind, path, data, *, root=None, applicant_mode='worker'):
    checked_adapter(guide, template, kind, root)
    if applicant_mode != 'worker':
        raise ValueError('默认且仅支持劳动者申请，单位区段未适配')
    base = _root(root)
    expected = _path(base, template['file'])
    if Path(path).resolve() != expected:
        raise ValueError('不得传入其他地区或任意同SHA原件路径')
    original = expected.read_bytes()
    if sha256(original).hexdigest() != template['sha256']:
        raise ValueError('浙江D自身原件SHA已变化')
    _validate_data(kind, data)
    document = Document(BytesIO(original))
    (_application if kind == 'application' else _evidence)(document, deepcopy(data))
    return PreservedPackageDocument(document, original)
