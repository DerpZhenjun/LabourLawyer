"""Preserved-package filling of the inspected Sichuan combined reference sample.

The historical function name remains compatible; the caller must enforce the
separate, exact-locality registry before sharing this structural renderer.
"""
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.table import _Row
from WenShu.template_fill import set_paragraph_text,set_cell_text,replace_span

def fill_sample_slots(paragraph,fields):
    text=paragraph.text;locations=[]
    for label,value in fields:
        start=text.find(label)
        if start<0:raise ValueError('当地样表填写标签变化，请重新核验')
        locations.append((start,start+len(label),str(value or '')))
    locations.sort()
    for index in range(len(locations)-1,-1,-1):
        _,start,value=locations[index]
        end=locations[index+1][0] if index+1<len(locations) else len(text)
        replace_span(paragraph,start,end,value+('   ' if index+1<len(locations) else ''))

def fill_reference_cell(cell,value):
    paragraph=cell.paragraphs[0]
    if not paragraph.runs:
        run=paragraph.add_run('')
        properties=paragraph._p.find(qn('w:pPr')+'/'+qn('w:rPr'))
        if properties is not None:run._r.append(deepcopy(properties))
    set_cell_text(cell,value)

class PreservedPackageDocument:
    """Expose python-docx's object model but patch only the intended OOXML part."""
    def __init__(self,document,original):self.document=document;self.original=original
    def __getattr__(self,name):return getattr(self.document,name)
    def save(self,destination):
        draft=BytesIO();self.document.save(draft)
        with ZipFile(BytesIO(self.original)) as original,ZipFile(draft) as changed,ZipFile(destination,'w') as final:
            for part in original.infolist():
                final.writestr(part,changed.read(part.filename) if part.filename=='word/document.xml' else original.read(part.filename))

def fill_tianfu_reference(path,kind,data):
    original=path.read_bytes();doc=Document(BytesIO(original))
    paragraphs=doc.paragraphs
    if len(paragraphs)<36 or paragraphs[0].text.strip()!='仲裁申请书' or paragraphs[24].text.strip()!='证据材料清单':
        raise ValueError('当地参考样表结构变化，需重新核验填写位置')
    if kind=='application':
        value=lambda k:str(data.get(k) or '').strip()
        slots={2:[('申请人：',value('applicantName')),('性别：',value('gender')),('民族：',value('nationality'))],
               3:[('出生日期：',value('birthDate')),('民身份号码：',value('idNumber'))],
               4:[('户籍所在地：',value('idAddress'))],5:[('现住址：',value('currentAddress'))],
               6:[('联系电话：',value('phone'))],7:[('确认有效的通讯地址：',value('applicantMailingAddress'))],
               8:[('邮政编码：',value('applicantPostalCode'))],9:[('被申请人：',value('respondentName'))],
               10:[('法定代表人：',value('legalRepName')),('职务：',value('legalRepPosition'))],
               11:[('住所：',value('registeredAddress'))],12:[('联系电话：',value('companyContactPhone'))],
               15:[('事实和理由：',value('factsReason'))]}
        for index,fields in slots.items():fill_sample_slots(paragraphs[index],fields)
        replacements={
            14:'\n'.join(f'{i+1}. {str(x).strip()}' for i,x in enumerate(data.get('requests',[])) if str(x).strip()),
            17:value('committeeName') or '劳动人事争议仲裁委员会',
            18:'附件：1.《仲裁申请书》副本____份；',19:'      2.证据清单及有关证据材料____份。',
            21:'申请人：________（请本人手写签名）',22:'                         年    月    日'}
        for i,t in replacements.items():set_paragraph_text(paragraphs[i],t)
        # Let long facts flow naturally, but never strand the signature alone.
        # Keep the closing, addressee, attachments and handwritten signing area
        # together without shrinking the official type or changing the page box.
        for paragraph in paragraphs[16:22]:
            paragraph.paragraph_format.keep_with_next=True
        for paragraph in paragraphs[16:23]:
            paragraph.paragraph_format.keep_together=True
        stop=paragraphs[24]._p
        passed=False
        for node in list(doc._element.body):
            if node is stop:passed=True
            if passed and node.tag!=qn('w:sectPr'):doc._element.body.remove(node)
        # The combined official sample has a tall empty separator before the
        # evidence form. Once that form is removed, it can create a blank page.
        while doc.paragraphs and not doc.paragraphs[-1].text.strip():
            doc._element.body.remove(doc.paragraphs[-1]._p)
    elif kind=='evidence':
        start,stop=paragraphs[24]._p,paragraphs[35]._p;active=False
        for node in list(doc._element.body):
            if node is start:active=True
            if node is stop:active=False
            if not active and node.tag!=qn('w:sectPr'):doc._element.body.remove(node)
        table=doc.tables[0];items=[x for x in data.get('evidenceList',[]) if str(x.get('name','')).strip()]
        # Remove spacing for the following authorization form in the combined sample.
        while doc.paragraphs and not doc.paragraphs[-1].text.strip():
            doc._element.body.remove(doc.paragraphs[-1]._p)
        properties=table.rows[0]._tr.get_or_add_trPr()
        if properties.find(qn('w:tblHeader')) is None:properties.append(OxmlElement('w:tblHeader'))
        template_row=deepcopy(table.rows[1]._tr)
        for row in list(table.rows)[1:]:table._tbl.remove(row._tr)
        for i,item in enumerate(items):
            node=deepcopy(template_row);table._tbl.append(node);row=_Row(node,table)
            values=[str(i+1),item.get('name',''),item.get('itemCount',''),item.get('pageCount',''),
                    '√' if item.get('originalType')=='原件' else '', '√' if item.get('originalType')=='复印件' else '',item.get('purpose',''),item.get('notes','')]
            for cell,value in zip(row.cells,values):fill_reference_cell(cell,str(value or ''))
        if not items:table._tbl.append(deepcopy(template_row))
    else:raise ValueError('不支持的文书种类')
    return PreservedPackageDocument(doc,original)
