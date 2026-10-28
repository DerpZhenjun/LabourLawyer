"""Rongjiang's own inspected DOC-derived reference; never a provincial fallback."""
from copy import deepcopy
import hashlib
from io import BytesIO
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Row
from docx.text.paragraph import Paragraph
from docx.shared import Pt
from regional_documents import PreservedPackageDocument,fill_reference_cell
from regional_guides import digest
from regional_jiangxi import SOURCE,ANNEX,FILE,SHA,HEADERS,DOWNLOAD
from WenShu.template_fill import replace_span,set_paragraph_text,date_parts

ROOT=Path(__file__).resolve().parents[1]
ADAPTER='rongjiang_own_doc_blank_v1'
REFERENCE=FILE.removesuffix('.doc')+'-converted.docx'
REFERENCE_SHA='9a7571e55b7260d17931c42936262b58e6e69971efe7254f108f9528d22591f1'

def rongjiang_source_matches(guide,template):
    if tuple(guide.get(k) for k in ('province','city','district','region_code'))!=('jiangxi','赣州市','蓉江新区','360798'):return False
    if guide.get('source_url')!=SOURCE or guide.get('source_id') not in {None,digest(SOURCE)[:24]}:return False
    if guide.get('verification_status') not in {'partial','verified'}:return False
    expected={'id':ANNEX,'annex_id':ANNEX,'region_code':'360798','source_url':SOURCE,
              'file':FILE,'sha256':SHA,'official_kind':'blank','download_status':'downloaded','download_url':DOWNLOAD}
    return all(template.get(k)==v for k,v in expected.items())

def _slots(paragraph,fields):
    """Exact labels, right to left; retain the source's embedded address newline."""
    text=paragraph.text;locations=[]
    for label,value in fields:
        if text.count(label)!=1:raise ValueError('蓉江原表字段已改变，请重新核验')
        start=text.index(label);locations.append((start,start+len(label),str(value or '').strip()))
    locations.sort()
    for i in range(len(locations)-1,-1,-1):
        _,start,value=locations[i]
        end=locations[i+1][0] if i+1<len(locations) else len(text)
        if value:
            suffix='\n' if '\n' in text[start:end] else '  ' if i+1<len(locations) else ''
            replace_span(paragraph,start,end,value+suffix,pad=False)

def _date(value):
    year,month,day=date_parts(value)
    return f'{year}年{month}月{day}日' if year and month and day else str(value or '')

def _remove(doc,node):
    doc._element.body.remove(node)

def fill_rongjiang_document(guide,template,kind,path,data):
    if not rongjiang_source_matches(guide,template) or kind not in {'application','evidence'}:
        raise ValueError('原DOC与蓉江新区地区、来源或文书种类不匹配')
    if path.resolve()!=(ROOT/FILE).resolve() or hashlib.sha256(path.read_bytes()).hexdigest()!=SHA:
        raise ValueError('蓉江自己的官方DOC原件已变化，请重新核验')
    reference=ROOT/REFERENCE
    if not reference.is_file():raise ValueError('蓉江转换工作副本暂不可用，请下载自己的官方原DOC')
    original=reference.read_bytes()
    if hashlib.sha256(original).hexdigest()!=REFERENCE_SHA:raise ValueError('蓉江转换工作副本已变化，请重新核验')
    doc=Document(BytesIO(original));p=doc.paragraphs
    if (len(p)!=39 or p[0].text.replace(' ','')!='仲裁申请书' or p[29].text!='证据目录'
        or len(doc.tables)!=1 or len(doc.tables[0].rows)!=11
        or [c.text for c in doc.tables[0].rows[0].cells]!=HEADERS):raise ValueError('蓉江自己的原表结构已变化')
    value=lambda k:str(data.get(k) or '').strip()
    if kind=='application':
        slots={1:[('申请人：',value('applicantName')),('性别：',value('gender'))],
            2:[('出生日期：',_date(value('birthDate'))),('民族：',value('nationality'))],
            3:[('身份证号码：',value('idNumber'))],4:[('户籍所在地（或外国人注明国籍）：',value('idAddress'))],
            5:[('确认有效的通讯地址：',value('applicantMailingAddress') or value('currentAddress'))],
            6:[('邮政编码：',value('applicantPostalCode'))],7:[('联系电话：',value('phone'))],
            8:[('被申请人：',value('respondentName')),('统一社会信用代码：',value('socialCreditCode'))],
            9:[('住所地：',value('registeredAddress')),('实际办公所在地：',value('actualOfficeAddress') or value('workAddress')),
               ('邮政编码：',value('respondentPostalCode')),('联系电话：',value('companyContactPhone'))],
            10:[('法定代表人（主要负责人）：',value('legalRepName')),('职务：',value('legalRepPosition'))]}
        for index,fields in slots.items():_slots(p[index],fields)
        requests=[str(x).strip() for x in data.get('requests',[]) if str(x).strip()]
        if requests:
            prototype=deepcopy(p[12]._p)
            _remove(doc,p[12]._p);_remove(doc,p[13]._p)
            for request in requests:
                node=deepcopy(prototype);p[14]._p.addprevious(node)
                set_paragraph_text(Paragraph(node,doc._body),request)
        if value('factsReason'):
            set_paragraph_text(p[15],value('factsReason'));_remove(doc,p[16]._p)
        # Addressee and signature remain the locality's own source text, not input.
        if value('applicationDate'):set_paragraph_text(p[20],_date(value('applicationDate')))
        for paragraph in p[17:20]:paragraph.paragraph_format.keep_with_next=True
        stop=p[29]._p;active=False
        for node in list(doc._element.body):
            if node is stop:active=True
            if active and node.tag!=qn('w:sectPr'):_remove(doc,node)
        for paragraph in p[26:29]:_remove(doc,paragraph._p)
    else:
        start=p[29]._p;active=False
        for node in list(doc._element.body):
            if node is start:active=True
            if not active and node.tag!=qn('w:sectPr'):_remove(doc,node)
        table=doc.tables[0];prototype=deepcopy(table.rows[1]._tr)
        for row in list(table.rows)[1:]:table._tbl.remove(row._tr)
        header=table.rows[0]._tr.get_or_add_trPr()
        if header.find(qn('w:tblHeader')) is None:header.append(OxmlElement('w:tblHeader'))
        items=[x for x in data.get('evidenceList',[]) if str(x.get('name','')).strip()]
        for index,item in enumerate(items or [{}]):
            node=deepcopy(prototype);table._tbl.append(node);row=_Row(node,table)
            properties=node.get_or_add_trPr()
            if properties.find(qn('w:cantSplit')) is None:properties.append(OxmlElement('w:cantSplit'))
            values=[str(index+1),item.get('name',''),item.get('source',''),item.get('purpose',''),item.get('pageCount','')]
            for cell,text in zip(row.cells,values):
                fill_reference_cell(cell,str(text or ''))
                # Blank form cells contain manual first-line placeholder offsets;
                # they are not a useful indent for multiline entered values.
                cell.paragraphs[0].paragraph_format.first_line_indent=Pt(0)
                # The original's 26pt handwriting grid is not appropriate for
                # multiline typed text. Keep its type size and minimum row size.
                cell.paragraphs[0].paragraph_format.line_spacing=1.2
        # Typed names never populate this handwritten submitter signature slot.
        if value('submitDate'):_slots(p[33],[('提交人：',''),('提交时间：',_date(value('submitDate')))])
        for row in table.rows[-1:]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:paragraph.paragraph_format.keep_with_next=True
        for paragraph in p[31:33]:
            paragraph.paragraph_format.keep_with_next=True
            paragraph.paragraph_format.keep_together=True
        p[33].paragraph_format.keep_together=True
        for paragraph in p[34:]:_remove(doc,paragraph._p)
    return PreservedPackageDocument(doc,original)
