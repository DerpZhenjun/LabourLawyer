"""Verify all 38 synthetic browser outputs; render every distinct byte sequence."""
import hashlib
import importlib.util
import json
import shutil
from pathlib import Path
from zipfile import ZipFile
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / 'reports/chengdu-district-expansion-20261023'
QA = REPORT / 'document-qa'


def main():
    checks, representatives, groups = [], {}, {}
    for guide_path in sorted((REPORT / 'captures').glob('*/guide.json')):
        g = json.loads(guide_path.read_text(encoding='utf8'))
        t = next(t for t in g['templates'] if t['original_name'] == '仲裁申请书(模板).docx')
        source = ROOT / t['file']
        assert hashlib.sha256(source.read_bytes()).hexdigest() == t['sha256']
        for kind in ['application', 'evidence']:
            path = REPORT / 'generated' / f"{g['region_code']}-{kind}-replayed.docx"
            doc = Document(path)
            text = '\n'.join(p.text for p in doc.paragraphs)
            assert '授权委托书' not in text
            if kind == 'application':
                assert '（示例）' in text and '1. ' in text and '2. ' in text
                assert '请本人手写签名' in text
            else:
                assert len(doc.tables) == 1 and len(doc.tables[0].rows) == 2
                assert doc.tables[0].cell(1, 1).text == '劳动合同'
            with ZipFile(source) as before, ZipFile(path) as after:
                assert set(before.namelist()) == set(after.namelist())
                for part in before.namelist():
                    if part != 'word/document.xml':
                        assert before.read(part) == after.read(part), part
            section = doc.sections[0]
            assert (section.page_width.twips, section.page_height.twips) == (11906, 16838)
            sha = hashlib.sha256(path.read_bytes()).hexdigest()
            groups.setdefault(sha, []).append(str(path.relative_to(ROOT)))
            if sha not in representatives:
                snapshot = QA / 'verified-inputs' / f'{kind}-{sha[:12]}.docx'
                snapshot.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(path, snapshot)
                representatives[sha] = snapshot
            checks.append({'district': g['district'], 'kind': kind, 'source': t['file'],
                           'source_sha256': t['sha256'], 'generated': str(path.relative_to(ROOT)),
                           'generated_sha256': sha, 'package_preservation': 'PASS', 'geometry': 'PASS'})
    assert len(checks) == 38
    original = ROOT / 'templates/_docx/_regional/510109000000/8c83fd146ca84c98b974b6ddb0814901-仲裁申请书(模板).docx'
    spec = importlib.util.spec_from_file_location('chengdu_word_qa', ROOT / 'scripts/verify-regional-documents.py')
    renderer = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(renderer)
    renderer.OUT = QA
    tasks = renderer.render([original, *representatives.values()])
    for row in checks:
        assert hashlib.sha256((ROOT / row['source']).read_bytes()).hexdigest() == row['source_sha256']
    result = {'engine': 'Microsoft Word own hidden read-only export + bundled Poppler',
              'packaged_renderer': 'failed: soffice.exe absent; see packaged-render.log',
              'checks': checks, 'identical_output_groups': groups, 'tasks': tasks,
              'all_originals_unchanged': True, 'model_calls': 0}
    (QA / 'render-manifest.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf8')
    print(json.dumps({'checks': len(checks), 'unique_outputs': len(groups), 'page_counts': [len(t['pages']) for t in tasks]}))


if __name__ == '__main__':
    main()
