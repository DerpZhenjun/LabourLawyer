"""Read-only official samples QA; never turns a sample into a blank template."""
import hashlib,importlib,json,sqlite3,subprocess,sys
from pathlib import Path
from docx import Document
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'scripts'))
BASE=ROOT/'reports/nationwide-localization-20261023-night/cross-province/beijing-tongzhou'
OUT=BASE/'document-qa'
SKILL=Path('C:/Users/DerpZhenjun/.codex/plugins/cache/openai-primary-runtime/documents/26.1017.11052/skills/documents')
def save(p,v):p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(v,ensure_ascii=False,indent=2),encoding='utf8')
def main():
    if OUT.exists():raise ValueError('Preserve own source review')
    files=json.loads((BASE/'downloads-0700/original-files.json').read_text('utf8'));paths=[];records=[]
    for i,f in enumerate(t for t in files if t['file_type']=='docx'):
        p=ROOT/f['file'];assert hashlib.sha256(p.read_bytes()).hexdigest()==f['sha256']
        q=OUT/'references'/('r'+str(i+1)+'.docx');q.parent.mkdir(parents=True,exist_ok=True);q.write_bytes(p.read_bytes());paths.append(q)
        doc=Document(p)
        records.append({**f,'paragraphs':[p.text for p in doc.paragraphs],
            'tables':[[[c.text for c in r.cells] for r in t.rows] for t in doc.tables],
            'inline_shapes':len(doc.inline_shapes),'sections':[{'width':s.page_width.twips,'height':s.page_height.twips} for s in doc.sections]})
    save(OUT/'source-review.json',records)
    run=subprocess.run([sys.executable,str(SKILL/'render_docx.py'),str(paths[0]),'--output_dir',str(OUT/'packaged-render'),'--emit_pdf'],capture_output=True)
    (OUT/'packaged-render.log').write_bytes(run.stdout+b'\n'+run.stderr)
    if run.returncode==0:raise ValueError('Packaged renderer now available; inspect packaged output first')
    m=importlib.import_module('verify-regional-documents');m.OUT=OUT
    tasks=m.render(paths);save(OUT/'reference-render.json',{'tasks':tasks,'originals_unchanged':all(hashlib.sha256((ROOT/f['file']).read_bytes()).hexdigest()==f['sha256'] for f in files if f.get('file'))})
    directory=ROOT/'backend/data/regional-directory.sqlite'
    with sqlite3.connect(directory.as_uri()+'?mode=ro',uri=True) as db:
        rows=db.execute("SELECT payload FROM directory_areas WHERE province='beijing' AND district='通州区'").fetchall()
    if len(rows)!=1:raise ValueError('Exact own directory identity missing or ambiguous')
    identity=json.loads(rows[0][0]);assert identity['code']=='110112' and identity['city']=='北京市'
    save(BASE/'directory-identity.json',identity)
    print(json.dumps({'pages':[len(t['pages']) for t in tasks],'identity':{k:identity[k] for k in ['provider','code','parent_code','city','district']}}))
if __name__=='__main__':main()
