"""Offline own-source fixtures and Word render queue; no production writes."""
import hashlib
from io import BytesIO
import json
from pathlib import Path
import sys
from zipfile import ZipFile
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'backend'))
import regional_jiangsu_oct30f as own
OUT = ROOT / 'reports/overnight-gap-closure-20261030/localities/jiangsu-f'


def save(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf8')


def run():
    queue, checks, identities = [], [], []
    for code in sorted(own.SUPPORTED):
        g = own.parse_capture(code)
        identities.append({k: g[k] for k in ['region_code', 'district', 'source_id', 'source_url', 'service_code', 'authority', 'materials', 'channel', 'attachment_counts']})
        payloads = {
            'application': {'applicantName': '虚构测试劳动者', 'respondentName': '虚构测试有限公司',
                'gender': '女', 'birthDate': '1990-01-01', 'phone': '虚构联系电话',
                'committeeName': '请核实实际受理的仲裁委员会',
                'requests': ['虚构测试请求一，仅供填写版式检查。', '虚构测试请求二，不用于真实仲裁申请。'],
                'factsReason': '全部姓名、单位、电话和事实均为虚构。' + '仅测试长文字换行、分页及手签留空。' * 24,
                'copyCount': 2, 'evidenceCopyCount': 3},
            'evidence': {'committeeName': '请核实实际受理的仲裁委员会', 'caseNumber': '虚构测试案号',
                'submitDate': '2026-10-30', 'evidenceList': [
                    {'name': '虚构劳动合同', 'source': '虚构测试来源', 'purpose': '仅用于版式检查。' * 10,
                     'originalType': '由本人按实际填写', 'pageCount': '2'},
                    {'name': '虚构工资记录', 'source': '', 'purpose': '', 'originalType': '', 'pageCount': ''}]},
        }
        save(OUT / code / 'synthetic-inputs.json', payloads)
        save(OUT / code / 'source-guide.json', g)
        for kind, payload in payloads.items():
            template = next(t for t in g['templates'] if t['id'] == own.SPECS[code][kind])
            original = (ROOT / template['file']).read_bytes()
            source = OUT / code / ('official-' + kind + '.docx')
            source.parent.mkdir(parents=True, exist_ok=True); source.write_bytes(original)
            queue.append({'source': source.relative_to(ROOT).as_posix(), 'label': code + '-official-' + kind})
            parsed = Document(BytesIO(original))
            save(OUT / code / (kind + '-original-layout.json'), {
                'own_template_id': template['id'], 'source_file': template['file'], 'sha256': template['sha256'],
                'paragraphs': [{'index': i, 'text': p.text} for i, p in enumerate(parsed.paragraphs)],
                'tables': [[[c.text for c in row.cells] for row in table.rows] for table in parsed.tables]})
            target = OUT / code / (kind + '-synthetic.docx')
            own.fill_candidate_document(g, kind, payload, ROOT).save(target)
            with ZipFile(BytesIO(original)) as a, ZipFile(target) as b:
                assert set(a.namelist()) == set(b.namelist())
                assert all(a.read(n) == b.read(n) for n in a.namelist() if n != 'word/document.xml')
            assert original == (ROOT / template['file']).read_bytes()
            checks.append({'region_code': code, 'kind': kind, 'annex_id': template['id'],
                'source_file': template['file'], 'source_sha256': template['sha256'],
                'output_file': target.relative_to(ROOT).as_posix(), 'output_sha256': hashlib.sha256(target.read_bytes()).hexdigest(),
                'original_unchanged': True, 'non_document_xml_parts': 'byte_identical'})
            queue.append({'source': target.relative_to(ROOT).as_posix(), 'label': code + '-filled-' + kind})
    save(OUT / 'RENDER-QUEUE.json', queue)
    save(OUT / 'OFFLINE-VERIFICATION.json', {'identities': identities, 'checks': checks,
        'manifest_sha256': own.MANIFEST_SHA, 'whitelist_sha256': own.WHITELIST_SHA,
        'published': False, 'production_db_or_seed_changed': False, 'model_calls': 0,
        'visual_inspection': 'pending', 'remaining': ['Word all-page inspection', 'root-owned shared integration',
            'isolated native browser generation and download', 'root publication review']})
    print(json.dumps({'own_sources': len(identities), 'filled_outputs': len(checks), 'render_inputs': len(queue)}))


if __name__ == '__main__':
    run()
