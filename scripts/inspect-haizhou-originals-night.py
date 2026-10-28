"""Read-only inspection of every own downloaded blank/sample original."""
import hashlib,importlib,json,sys,subprocess
from pathlib import Path
from docx import Document
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'scripts'))
BASE=ROOT/'reports/nationwide-localization-20261023-night/cross-province/jiangsu-haizhou'
OUT=BASE/'document-qa'
qa=importlib.import_module('qa-xishan-night');qa.BASE=BASE;qa.OUT=OUT


def main():
    if (OUT/'original-render.json').exists():raise ValueError('Preserve own existing original review')
    guide=json.loads((BASE/'guide.json').read_text('utf8'));inventory=[];paths=[];seen=set()
    for t in guide['templates']:
        source=ROOT/t['file'];assert hashlib.sha256(source.read_bytes()).hexdigest()==t['sha256']
        row={k:t[k] for k in ['id','source_url','download_url','file','sha256','original_name','official_kind','material_id']}
        if source.suffix.lower()=='.docx':
            doc=Document(source)
            row.update(paragraphs=[{'index':i,'text':p.text} for i,p in enumerate(doc.paragraphs)],
                sections=[{'width':s.page_width.twips,'height':s.page_height.twips,'margins':[s.top_margin.twips,s.bottom_margin.twips,s.left_margin.twips,s.right_margin.twips]} for s in doc.sections],
                tables=[{'rows':len(tb.rows),'columns':len(tb.columns),'text':[[c.text for c in r.cells] for r in tb.rows]} for tb in doc.tables])
            if t['sha256'] not in seen:
                # Long original names repeated in renderer subdirectories exceed
                # legacy Word's path limit; use byte-identical short QA copies.
                short=OUT/'references'/('r'+str(len(paths)+1)+'.docx')
                short.parent.mkdir(parents=True,exist_ok=True)
                if short.exists():assert short.read_bytes()==source.read_bytes()
                else:short.write_bytes(source.read_bytes())
                row['readonly_render_copy']=short.relative_to(ROOT).as_posix()
                paths.append(short);seen.add(t['sha256'])
        inventory.append(row)
    qa.save(OUT/'original-layout.json',inventory)
    packaged=Path('C:/Users/DerpZhenjun/.codex/plugins/cache/openai-primary-runtime/documents/26.1017.11052/skills/documents/render_docx.py')
    probe=subprocess.run([sys.executable,str(packaged),str(paths[0]),'--output_dir',str(OUT/'packaged-probe'),'--emit_pdf'],capture_output=True,text=True,encoding='utf8',errors='replace',timeout=60)
    (OUT/'packaged-render.log').write_text(probe.stdout+'\n'+probe.stderr,encoding='utf8')
    tasks=qa.render(paths)
    qa.save(OUT/'original-layout.json',inventory)
    qa.save(OUT/'original-render.json',{'at':qa.__name__,'tasks':tasks,'annexes':len(inventory),
        'unique_docx':len(paths),'images':[r['file'] for r in inventory if r['file'].lower().endswith('.jpg')],
        'all_originals_unchanged':all(hashlib.sha256((ROOT/t['file']).read_bytes()).hexdigest()==t['sha256'] for t in guide['templates']),
        'authoring':False,'models':0,'visual_inspection':'pending'})
    print(json.dumps({'originals':len(inventory),'unique_docx':len(paths),'render_pages':[len(t['pages']) for t in tasks],'authoring':False}))


if __name__=='__main__':main()
