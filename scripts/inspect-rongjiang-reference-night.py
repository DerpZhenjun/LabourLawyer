"""Verify conversion against the retained source, before any field adaptation."""
import hashlib,importlib,json,subprocess,sys
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from PIL import Image,ImageChops
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'scripts'))
BASE=ROOT/'reports/nationwide-localization-20261023-night/cross-province/jiangxi-rongjiang'
OUT=BASE/'document-qa'
qa=importlib.import_module('verify-regional-documents');qa.OUT=OUT
SKILL=Path('C:/Users/DerpZhenjun/.codex/plugins/cache/openai-primary-runtime/documents/26.1017.11052/skills/documents')


def main():
    if (OUT/'converted-reference.json').exists():raise ValueError('Preserve completed own conversion verification')
    c=json.loads((OUT/'conversion.json').read_text('utf-8-sig'));source=ROOT/c['reference_file']
    assert hashlib.sha256(source.read_bytes()).hexdigest()==c['reference_sha256']
    doc=Document(source)
    probe=subprocess.run([sys.executable,str(SKILL/'render_docx.py'),str(source),'--output_dir',str(OUT/'packaged-reference'),'--emit_pdf','--verbose'],capture_output=True,text=True,encoding='utf8',errors='replace')
    (OUT/'packaged-reference.log').write_text(probe.stdout+'\n'+probe.stderr,encoding='utf8')
    if probe.returncode and 'soffice.exe' not in probe.stderr:raise ValueError('Unexpected packaged renderer failure')
    tasks=qa.render([source]);old=json.loads((OUT/'original-render.json').read_text('utf8'))['tasks'][0]['pages']
    new=tasks[0]['pages'];assert len(new)==len(old)==2
    pixel_checks=[]
    for before,after in zip(old,new):
        a=Image.open(ROOT/before).convert('RGB');b=Image.open(ROOT/after).convert('RGB')
        assert a.size==b.size
        diff=ImageChops.difference(a,b)
        pixel_checks.append({'original':before,'converted':after,'dimensions':list(a.size),'exact_pixels_equal':diff.getbbox() is None})
    with ZipFile(source) as z:parts={n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()}
    value={'reference_file':c['reference_file'],'sha256':c['reference_sha256'],'original_sha256':c['original_sha256'],
        'sections':[{'width':s.page_width.twips,'height':s.page_height.twips,'margins':[s.top_margin.twips,s.bottom_margin.twips,s.left_margin.twips,s.right_margin.twips]} for s in doc.sections],
        'paragraphs':[{'index':i,'text':p.text,'runs':[{'text':r.text,'font':r.font.name,'size':r.font.size.pt if r.font.size else None} for r in p.runs]} for i,p in enumerate(doc.paragraphs)],
        'tables':[{'rows':len(t.rows),'columns':len(t.columns),'grid':[x.width.twips for x in t.columns],'text':[[c.text for c in r.cells] for r in t.rows]} for t in doc.tables],
        'package_parts':parts,'tasks':tasks,'pixel_comparisons':pixel_checks,'visual_review':'pending','models':0}
    (OUT/'converted-reference.json').write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf8')
    assert hashlib.sha256((ROOT/c['source']).read_bytes()).hexdigest()==c['original_sha256']
    print(json.dumps({'converted_pages':2,'pixel_equal':[p['exact_pixels_equal'] for p in pixel_checks],'reference_sha':c['reference_sha256']}))


if __name__=='__main__':main()
