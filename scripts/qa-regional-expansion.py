"""Package fidelity and page rendering of captured synthetic browser requests."""
import hashlib
import importlib.util
import json
import shutil
import sys
import uuid
from pathlib import Path
from zipfile import ZipFile
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / 'reports/regional-expansion-20261023'
sys.path.insert(0, str(ROOT / 'backend'))
from regional_adapters import BINDINGS, REFERENCE_SHA256


def run():
    references = []
    for code, (district, _, annex, _) in BINDINGS.items():
        source = ROOT / f'templates/_docx/_regional/{code}/{annex}-仲裁申请书(模板).docx'
        assert hashlib.sha256(source.read_bytes()).hexdigest() == REFERENCE_SHA256
        references.append(source)
    original = Document(references[1])
    metadata = {
        'references': [str(p) for p in references], 'sha256': REFERENCE_SHA256,
        'sections': [{'width_twips': s.page_width.twips, 'height_twips': s.page_height.twips,
                      'margins_twips': [s.top_margin.twips, s.bottom_margin.twips, s.left_margin.twips, s.right_margin.twips]} for s in original.sections],
        'paragraphs': [{'index': i, 'text': p.text, 'font_pt': [r.font.size.pt if r.font.size else None for r in p.runs]} for i, p in enumerate(original.paragraphs)],
        'tables': [{'rows': len(t.rows), 'columns': len(t.columns), 'column_width_twips': [c.width.twips for c in t.columns]} for t in original.tables],
    }
    qa = REPORT / 'document-qa'
    qa.mkdir(exist_ok=True)
    (qa / 'reference-layout.json').write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding='utf8')
    (qa / 'artifact.md').write_text(
        '# 地区官方参考样表填写契约\n\n'
        '三份独立官方附件 SHA256 均为 ' + REFERENCE_SHA256 + '。\n'
        '沿用既有 2026-10-08 全页检查与结构渲染器；本轮对两区独立附件重新公开下载核验。\n'
        '精确页尺寸、边距、字号和表格见 reference-layout.json。官方源文件不改写。\n'
        '申请书允许替换第 2–12 段当事人槽、第 14 段请求、第 15 段经过、第 17 段受理机构；'
        '第 18–19 段副本份数留空，21–22 段签名日期留手写位置。删除后续证据/委托书部分。\n'
        '证据清单只保留第 24–34 段所属内容及其中表格，按需要克隆原行；删除申请/委托书，保留提交人手签。\n'
        '第 16–22 段允许增加 keep-with-next / keep-together，避免签名孤立分页；不缩小字号。\n'
        '除 word/document.xml 外，全部包部件必须逐字节保留；不改样式、关系、页尺寸或固定行高。\n'
        '所有最终页必须渲染检查；相同 SHA 的原附件只需渲染一次。\n', encoding='utf8')
    generated = []
    checks = []
    for code, (district, _, annex, _) in BINDINGS.items():
        if district not in {'锦江区', '武侯区'}:
            continue
        source = ROOT / f'templates/_docx/_regional/{code}/{annex}-仲裁申请书(模板).docx'
        for kind in ['application', 'evidence']:
            path = REPORT / 'generated' / f'{district}-{kind}-replayed.docx'
            doc = Document(path)
            text = '\n'.join(p.text for p in doc.paragraphs)
            assert '授权委托书' not in text
            if kind == 'application':
                assert district + '测试劳动者' in text and district + '测试委员会' in text
                assert '24000' in text and '天府新区' not in text
            else:
                assert len(doc.tables) == 1 and len(doc.tables[0].rows) == 2
                assert doc.tables[0].cell(1, 1).text == district + '测试劳动合同'
            with ZipFile(source) as before, ZipFile(path) as after:
                assert set(before.namelist()) == set(after.namelist())
                for part in before.namelist():
                    if part != 'word/document.xml':
                        assert before.read(part) == after.read(part), part
            for section in doc.sections:
                assert section.page_width == original.sections[0].page_width
                assert section.page_height == original.sections[0].page_height
            checks.append({'district': district, 'kind': kind, 'source_file': str(source), 'generated': str(path),
                           'generated_sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'package_preservation': 'PASS'})
            generated.append(path)
    spec = importlib.util.spec_from_file_location('regional_word_qa', ROOT / 'scripts/verify-regional-documents.py')
    renderer = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(renderer)
    renderer.OUT = qa
    tasks = renderer.render([references[1], *generated])
    for source in references:
        assert hashlib.sha256(source.read_bytes()).hexdigest() == REFERENCE_SHA256
    result = {'engine': 'Microsoft Word hidden read-only export + bundled Poppler',
              'packaged_renderer': 'unavailable: bundled soffice.exe missing', 'checks': checks,
              'tasks': tasks, 'model_calls': 0, 'all_originals_unchanged': True}
    manifest = qa / 'render-manifest.json'
    if manifest.exists():
        shutil.copy2(manifest, qa / ('render-manifest-before-' + uuid.uuid4().hex[:8] + '.json'))
    manifest.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf8')
    print(json.dumps({'manifest': str(qa / 'render-manifest.json'), 'page_counts': [len(t['pages']) for t in tasks]}))


if __name__ == '__main__':
    run()
