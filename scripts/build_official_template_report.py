"""Export the local acquisition inventory to Word without visiting official sites.

Historical source records and files are reported separately from live portal observations.
"""
from __future__ import annotations

from datetime import datetime
import hashlib
import json
from pathlib import Path
import sys

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'backend'))
from province_documents import PROVINCE_DOCUMENTS, SLUGS


def read(path: Path, default):
    return json.loads(path.read_text(encoding='utf-8')) if path.is_file() else default


def hyperlink(doc, label: str, url: str):
    if not url:
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.add_run(label + '：').bold = True
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    properties = OxmlElement('w:rPr')
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '18'); properties.append(size)
    run.append(properties)
    text = OxmlElement('w:t')
    # Visible wraps do not alter the URL stored in the actual hyperlink target.
    text.text = '\u200b'.join(url[i:i+36] for i in range(0, len(url), 36))
    run.append(text); link.append(run); p._p.append(link)


def main():
    catalog = read(ROOT / 'config/province_official_template_catalog.json', {})['provinces']
    acquisition = read(ROOT / 'templates/province-template-acquisition-report.json', {})
    downloads = {p['province']: p for p in acquisition.get('provinces', [])}
    adapters = read(ROOT / 'config/template_sources.json', {}).get('sources', [])
    adapter_sources = {p['province']: p for p in adapters}
    adapter_runs = {}
    for file in sorted((ROOT / 'templates/_acquisition/runs').glob('*/*/report.json'), reverse=True):
        result = read(file, {})
        if result.get('province') and result['province'] not in adapter_runs:
            adapter_runs[result['province']] = {**result, 'record_path': str(file.relative_to(ROOT))}
    now = datetime.now().astimezone()
    out = ROOT / 'templates/_reports/template-ledgers' / now.strftime('%Y%m%d-%H%M%S-%f')
    out.mkdir(parents=True, exist_ok=True)
    doc = Document()
    # The bundled default template carries decorative blue paragraph borders.
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.8)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]
        style.font.name = '宋体'; style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
    normal = doc.styles['Normal']
    normal.font.size = Pt(10)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.space_after = Pt(6)
    doc.styles['Title'].font.size = Pt(23)
    doc.styles['Heading 1'].font.size = Pt(16)
    footer = section.footer.paragraphs[0]
    footer.add_run('官方模板来源记录  |  第 ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    footer.add_run(' 页')
    doc.add_heading('官方文书模板获取报告', 0)
    doc.add_paragraph('中国大陆31个省级行政区 · 本地库存与来源记录')
    doc.add_paragraph(f'导出时间：{now.isoformat(timespec="seconds")}')
    available = [p for p in PROVINCE_DOCUMENTS.values() if p['documents']]
    doc.add_heading('记录范围', 1)
    doc.add_paragraph(f'目前可生成文书的地区为 {len(available)} 个，已配置文书类型 {sum(len(p["documents"]) for p in available)} 种。其他地区没有可用的填充配置。香港、澳门、台湾不在本次大陆目录范围内。')
    doc.add_paragraph('本报告读取来源目录、最近保存的下载记录和当前本地文件，不重新访问网站。目录中的状态是历史记录，不能据此断言今天没有劳动仲裁入口。下载时间不等于本次导出时间。')
    doc.add_paragraph('各模板适用范围按发布机构记录；区县模板不能自动视作全省统一模板。DOC 转换的 DOCX 是项目用于填写的副本，官方原文件保存在 _original；PDF 不自动重制为 Word。')
    doc.add_paragraph('本报告包含来源链接、登记的下载链接、文件路径和当前 SHA-256。校验值用于识别文件版本，不代表主管机关认定或法律效力。重新获取模板后，还需检查字段映射及生成文书。')
    doc.add_heading('如何查证', 1)
    doc.add_paragraph('实时门户访问与截图证据请查 templates/_reports/portal-audits 中的“劳动仲裁政务服务门户观察报告”。文书对比检查请查 templates/_validation/generated-documents。OpenClaw 输出是辅助意见，需要人工逐页复核。')
    doc.add_paragraph(f'全国附件目录最近保存的下载批次时间：{acquisition.get("finished_at", "未记录")}。四省接口适配器另存运行记录于 templates/_acquisition。')
    payload = {'exported_at': now.isoformat(), 'live_collection': False, 'regions': []}
    for province in catalog:
        doc.add_page_break()
        name = province['province_zh']; folder = province['folder']
        metadata = PROVINCE_DOCUMENTS.get(SLUGS[province['province']], {})
        doc.add_heading(name, 1)
        doc.add_paragraph('发布机构：' + (province.get('authority') or '目录未记载'))
        doc.add_paragraph('适用范围：' + (province.get('scope') or '目录未记载'))
        doc.add_paragraph('历史目录状态：' + province.get('status', '未记录'))
        kinds = metadata.get('documents', {})
        doc.add_paragraph('当前生成入口：' + ('、'.join('仲裁申请书' if k == 'application' else '证据清单' for k in kinds) if kinds else '尚未启用'))
        hyperlink(doc, '官方来源页面', province.get('source_url', ''))
        registered = province.get('documents', []) or adapter_sources.get(province['province'], {}).get('options', {}).get('templates', [])
        urls = set()
        for item in registered:
            url = item.get('url') or item.get('archive_url')
            if url and url not in urls:
                hyperlink(doc, '登记下载链接', url); urls.add(url)
        if province['province'] in adapter_sources:
            hyperlink(doc, '官方材料接口', adapter_sources[province['province']].get('options', {}).get('material_api', ''))
        adapter_run = adapter_runs.get(province['province'])
        if adapter_run:
            doc.add_paragraph(f'最近接口采集记录：{adapter_run.get("run_id")}；状态：{adapter_run.get("status")}；写入方式：{adapter_run.get("promotion")}')
            for artifact in adapter_run.get('artifacts', []):
                url = artifact.get('download_url', '')
                if url and url not in urls:
                    hyperlink(doc, '接口取得的下载链接', url); urls.add(url)
        download = downloads.get(province['province'], {})
        for error in download.get('errors', []):
            doc.add_paragraph(f'最近下载错误：{error.get("document", "")}；{error.get("error", "")}')
        files = []
        for base in ('_docx', '_original'):
            files.extend(p for p in sorted((ROOT / 'templates' / base / folder).rglob('*')) if p.is_file())
        doc.add_heading(f'当前本地文件（{len(files)} 份）', 2)
        entries = []
        for file in files:
            digest = hashlib.sha256(file.read_bytes()).hexdigest()
            relative = file.relative_to(ROOT).as_posix()
            p = doc.add_paragraph(relative)
            p.paragraph_format.space_after = Pt(2)
            p = doc.add_paragraph('SHA-256：' + digest)
            for run in p.runs: run.font.size = Pt(8)
            entries.append({'path': relative, 'bytes': file.stat().st_size, 'sha256': digest})
        if not files:
            doc.add_paragraph('未取得本地文件；本项目未自行制作替代官方空白模板。需要继续核对官方指南、公开附件或主管机关网站。')
        record = {**province, 'current_documents': kinds, 'local_files': entries,
                  'last_catalog_download': download, 'primary_adapter': adapter_sources.get(province['province']),
                  'last_primary_adapter_record': adapter_run}
        payload['regions'].append(record)
    path = out / '官方文书模板获取报告.docx'
    doc.save(path)
    (out / 'sources.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    (out.parent / 'latest.json').write_text(json.dumps({'report': str(path), 'sources': str(out / 'sources.json'),
                                                      'exported_at': now.isoformat()}, ensure_ascii=False, indent=2), encoding='utf-8')
    print(path, flush=True)


if __name__ == '__main__':
    main()
