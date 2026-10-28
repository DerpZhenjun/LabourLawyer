"""Own blank reference render and package inventory, no authoring or model."""
import hashlib
import importlib
import json
from pathlib import Path
import sys
from zipfile import ZipFile
from docx import Document
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'scripts'))
BASE=ROOT/'reports/nationwide-localization-20261023-night/cross-province/jiangsu-dafeng'
OUT=BASE/'document-qa'
IDS={'application':'8984153f-2c04-4676-a8fd-3a627658fb9d','evidence':'b104e1d9-15a7-489a-ad8b-80e22e0673e7'}
qa=importlib.import_module('qa-xishan-night');qa.BASE=BASE;qa.OUT=OUT


def main():
    if (OUT/'reference-render.json').exists():raise ValueError('Preserve approved own source rendering')
    guide=json.loads((BASE/'guide.json').read_text('utf8'));refs={};paths=[]
    for kind,identifier in IDS.items():
        t=next(t for t in guide['templates'] if t['id']==identifier);source=ROOT/t['file']
        assert hashlib.sha256(source.read_bytes()).hexdigest()==t['sha256'] and t['official_kind']=='blank'
        doc=Document(source)
        with ZipFile(source) as z:parts={n:{'bytes':len(z.read(n)),'sha256':hashlib.sha256(z.read(n)).hexdigest()} for n in z.namelist()}
        refs[kind]={'source':t['file'],'annex_id':identifier,'sha256':t['sha256'],'source_url':guide['source_url'],
            'sections':[{'width':s.page_width.twips,'height':s.page_height.twips,'margins':[s.top_margin.twips,s.bottom_margin.twips,s.left_margin.twips,s.right_margin.twips]} for s in doc.sections],
            'paragraphs':[{'index':i,'text':p.text,'style':p.style.name,'runs':[{'font':r.font.name,'size':r.font.size.pt if r.font.size else None} for r in p.runs]} for i,p in enumerate(doc.paragraphs)],
            'tables':[{'rows':len(t.rows),'grid':[c.width.twips for c in t.columns],'text':[[c.text for c in r.cells] for r in t.rows]} for t in doc.tables],
            'package_parts':parts}
        paths.append(source)
    tasks=qa.render(paths);qa.save(OUT/'reference-layout.json',refs);qa.save(OUT/'reference-render.json',tasks)
    assert all(hashlib.sha256((ROOT/t['file']).read_bytes()).hexdigest()==t['sha256'] for t in guide['templates'])
    print(json.dumps({'own_blank_reference_pages':[len(t['pages']) for t in tasks],'models':0,'authoring':False}))


if __name__=='__main__':main()
