"""Native Word read-only rendering fallback and actual regional API document QA."""
import argparse,hashlib,importlib.util,json,subprocess,sys,uuid
from pathlib import Path
from urllib.request import Request,urlopen
from docx import Document
from zipfile import ZipFile
ROOT=Path(__file__).resolve().parents[1]
REPORT=ROOT/'reports/government-guide-localization';OUT=REPORT/'document-qa'

def render(paths):
    # Never attach to an existing user's Word session. Own hidden COM exporter only.
    probe=subprocess.run(['powershell.exe','-NoProfile','-Command','@(Get-Process WINWORD -ErrorAction SilentlyContinue).Count'],capture_output=True,text=True)
    if probe.stdout.strip()!='0':raise RuntimeError('Word is already running; leave user session untouched and retry later')
    spec=importlib.util.spec_from_file_location('regional_word_renderer',ROOT/'scripts/validate_acquired_templates.py')
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module);module.OUTPUT_ROOT=OUT/'render'/uuid.uuid4().hex[:12]
    tasks=[{'source':str(p),'format':'docx','province_folder':p.stem} for p in paths]
    module.convert_office_files(tasks)
    for item in tasks:
        if item.get('conversion')=='failed':raise RuntimeError('Read-only Word export failed')
        pdf=Path(item['pdf'])
        subprocess.run([str(module.PDFTOPPM),'-scale-to','1600','-png',str(pdf),str(pdf.parent/'page')],check=True,creationflags=subprocess.CREATE_NO_WINDOW)
        item['pages']=[str(p.relative_to(ROOT).as_posix()) for p in sorted(pdf.parent.glob('page-*.png'))]
    return tasks

def run(reference):
    OUT.mkdir(parents=True,exist_ok=True)
    source=REPORT/'tianfu-application-reference.docx';before=hashlib.sha256(source.read_bytes()).hexdigest()
    if reference:
        doc=Document(source)
        metadata={'reference':str(source),'sha256':before,'sections':[{'width_cm':s.page_width.cm,'height_cm':s.page_height.cm,'margins_cm':[s.top_margin.cm,s.bottom_margin.cm,s.left_margin.cm,s.right_margin.cm]} for s in doc.sections],
                  'paragraphs':[{'index':i,'text':p.text,'style':p.style.name,'font_pt':[r.font.size.pt if r.font.size else None for r in p.runs]} for i,p in enumerate(doc.paragraphs)],
                  'tables':[{'columns':len(t.columns),'rows':len(t.rows),'grid_width_cm':[c.width.cm for c in t.columns]} for t in doc.tables]}
        with ZipFile(source) as z:metadata['package_sha256']={n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()}
        (OUT/'reference-layout.json').write_text(json.dumps(metadata,ensure_ascii=False,indent=2),encoding='utf8')
        tasks=render([source]);dest=OUT/'reference-render.json'
    else:
        with urlopen('http://127.0.0.1:8000/api/regional-guides/lookup?province=sichuan&city=%E6%88%90%E9%83%BD%E5%B8%82&district=%E5%A4%A9%E5%BA%9C%E6%96%B0%E5%8C%BA') as r:g=json.load(r)
        t=next(t for t in g['templates'] if t.get('can_generate'));files=[]
        samples={'application':{'applicantName':'视觉测试劳动者','phone':'','respondentName':'排版测试公司','requests':['请求支付测试工资100元。'],'factsReason':'本文件仅用于自动化排版验证，不用于实际申请。'+('测试较长内容换行与分页。'*35)},
                 'evidence':{'evidenceList':[{'name':'排版测试劳动合同','pageCount':'3','itemCount':'1','originalType':'复印件','purpose':'本文件仅用于测试。'+('说明劳动关系及工资约定，检查较长文字换行。'*8)}]}}
        for kind,data in samples.items():
            payload={k:g[k] for k in ['province','city','district','source_id','version']}
            payload.update(template_id=t['id'],kind=kind,mode='reference_adaptation',data=data)
            with urlopen(Request('http://127.0.0.1:8000/api/regional-guides/generate',data=json.dumps(payload,ensure_ascii=False).encode(),headers={'Content-Type':'application/json'})) as r:
                if r.headers.get('X-Document-Origin')!='official-sample-adaptation-not-blank':raise RuntimeError('Missing explicit sample-adaptation origin')
                path=OUT/(kind+'-api-test.docx');path.write_bytes(r.read());files.append(path)
        tasks=render(files);dest=OUT/'generated-render.json'
        with ZipFile(source) as z:
            for path in files:
                with ZipFile(path) as final:
                    for name in z.namelist():
                        if name!='word/document.xml':
                            if z.read(name)!=final.read(name):raise AssertionError('Untouched package changed: '+name)
    assert before==hashlib.sha256(source.read_bytes()).hexdigest(),'Official sample changed'
    dest.write_text(json.dumps({'engine':'Microsoft Word hidden read-only PDF export + bundled Poppler','source_sha256':before,'tasks':tasks},ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps({'manifest':str(dest),'pages':[len(t['pages']) for t in tasks]}))

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--reference',action='store_true');run(p.parse_args().reference)
