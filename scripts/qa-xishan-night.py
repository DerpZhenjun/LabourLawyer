"""Read-only source inventory/render, then own-source synthetic output QA."""
import argparse
import hashlib
import importlib.util
import json
import sys
from pathlib import Path
from zipfile import ZipFile
from urllib.request import Request,urlopen
from docx import Document

ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'backend'))
BASE=ROOT/'reports/nationwide-localization-20261023-night/cross-province/jiangsu-xishan'
OUT=BASE/'document-qa'
IDS={'application':'e1ae7409-3195-4eeb-aace-c9de16282335','evidence':'2d7b1fe9-8bdb-4e9b-b0ef-7616a6b87c97'}


def save(path,data):
    path.parent.mkdir(parents=True,exist_ok=True)
    path.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf8')


def references():
    guide=json.loads((BASE/'guide.json').read_text('utf8'))
    result={}
    for kind,identity in IDS.items():
        template=next(t for t in guide['templates'] if t['id']==identity)
        source=ROOT/template['file'];checksum=hashlib.sha256(source.read_bytes()).hexdigest()
        assert checksum==template['sha256'] and template['official_kind']=='blank'
        file=OUT/'references'/(kind+'.docx');file.parent.mkdir(parents=True,exist_ok=True)
        if file.exists():assert file.read_bytes()==source.read_bytes()
        else:file.write_bytes(source.read_bytes())
        doc=Document(source)
        with ZipFile(source) as package:parts={n:{'bytes':len(package.read(n)),'sha256':hashlib.sha256(package.read(n)).hexdigest()} for n in package.namelist()}
        result[kind]={'own_source':str(source.relative_to(ROOT)),'render_copy':str(file.relative_to(ROOT)),
            'source_url':guide['source_url'],'annex_id':identity,'sha256':checksum,
            'sections':[{'width':s.page_width.twips,'height':s.page_height.twips,'margins':[s.top_margin.twips,s.bottom_margin.twips,s.left_margin.twips,s.right_margin.twips],
                         'header_distance':s.header_distance.twips,'footer_distance':s.footer_distance.twips} for s in doc.sections],
            'paragraphs':[{'index':i,'text':p.text,'style':p.style.name,'alignment':str(p.alignment),
                           'runs':[{'text':r.text,'font':r.font.name,'size':r.font.size.pt if r.font.size else None,'underline':str(r.underline)} for r in p.runs]} for i,p in enumerate(doc.paragraphs)],
            'tables':[{'rows':len(t.rows),'grid':[col.width.twips for col in t.columns],
                       'text':[[c.text for c in row.cells] for row in t.rows]} for t in doc.tables], 'package':parts}
    save(OUT/'reference-layout.json',result)
    return result


def render(paths):
    spec=importlib.util.spec_from_file_location('xishan_word_renderer',ROOT/'scripts/verify-regional-documents.py')
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module);module.OUT=OUT
    return module.render(paths)


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--reference',action='store_true');parser.add_argument('--browser-payloads',type=Path);args=parser.parse_args()
    refs=references()
    if args.reference:
        tasks=render([ROOT/v['render_copy'] for v in refs.values()]);save(OUT/'reference-render.json',tasks)
        print(json.dumps({'references':[{k:v[k] for k in ['annex_id','sha256','sections']} for v in refs.values()],'pages':[len(t['pages']) for t in tasks]},ensure_ascii=False))
    else:
        from regional_adapters import with_reference_capabilities
        from regional_blank_documents import fill_blank_document
        guide=with_reference_capabilities(json.loads((BASE/'guide.json').read_text('utf8')))
        samples={'application':{'applicantName':'合成测试甲（示例）','gender':'女','birthDate':'1992-08-16',
            'idNumber':'证件号码请本人填写','idAddress':'无锡市示例路88号（示例）','applicantMailingAddress':'无锡市示例路88号（示例）',
            'respondentName':'无锡示例科技有限公司（示例）','registeredAddress':'无锡市示例路120号（示例）',
            'acceptMediation':'待本人选择','legalAid':'待本人确认','committeeName':'请填写实际受理的仲裁委员会',
            'requests':['请求支付测试工资100元。','请求支付测试加班工资50元。'],
            'factsReason':'本文件仅用于排版测试，不用于真实仲裁申请。劳动者与测试公司之间的工资支付存在争议。','copyCount':1},
            'evidence':{'committeeName':'请填写实际受理的仲裁委员会','evidenceList':[{'name':'劳动合同（示例）',
                'source':'本人保存','pageCount':'3','purpose':'说明测试劳动关系与工资约定。','originalType':'复印件'}]}}
        payloads={}
        if args.browser_payloads:
            for kind in IDS:
                payload=json.loads((args.browser_payloads/(guide.get('source_id','b1f8ab307b0a4f634081c052')+'-'+kind+'.network-request')).read_text('utf8'))
                assert payload['source_id']=='b1f8ab307b0a4f634081c052' and payload['template_id']==IDS[kind] and payload['mode']=='blank'
                payloads[kind]=payload;samples[kind]=payload['data']
        checks=[];files=[]
        for kind,data in samples.items():
            template=next(t for t in guide['templates'] if t['id']==IDS[kind]);source=ROOT/template['file']
            file=OUT/'generated'/(kind+'.docx');file.parent.mkdir(parents=True,exist_ok=True)
            if payloads:
                with urlopen(Request('http://127.0.0.1:8001/api/regional-guides/generate',data=json.dumps(payloads[kind]).encode(),headers={'Content-Type':'application/json'})) as response:
                    assert response.headers['X-Document-Origin']=='official-blank-filled'
                    file.write_bytes(response.read())
            else:fill_blank_document(guide,template,kind,source,data).save(file)
            with ZipFile(source) as original,ZipFile(file) as final:
                assert set(original.namelist())==set(final.namelist())
                for part in original.namelist():
                    if part!='word/document.xml':assert original.read(part)==final.read(part),part
            parsed=Document(file);section=parsed.sections[0]
            assert (section.page_width.twips,section.page_height.twips)==(11906,16838)
            assert hashlib.sha256(source.read_bytes()).hexdigest()==template['sha256']
            checks.append({'kind':kind,'annex_id':template['id'],'source':template['file'],'source_sha256':template['sha256'],
                'output':str(file.relative_to(ROOT)),'sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'preserved_package':'PASS','A4_geometry':'PASS'})
            files.append(file)
        tasks=render(files)
        save(OUT/'render-manifest.json',{'source_id':'b1f8ab307b0a4f634081c052','checks':checks,'tasks':tasks,
            'engine':'own hidden Word read-only export + bundled Poppler','browser_payloads':str(args.browser_payloads) if args.browser_payloads else None,
            'all_originals_unchanged':True,'visual_inspection':'pending'})
        print(json.dumps({'checked_outputs':len(checks),'pages':[len(t['pages']) for t in tasks]},ensure_ascii=False))


if __name__=='__main__':main()
