"""Bundled-runtime QA of six local sample adaptations, every unique output page."""
import argparse
import hashlib
import importlib.util
import json
import sys
from pathlib import Path
from zipfile import ZipFile

ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'backend'))
from regional_adapters import EXTRA_BINDINGS,with_reference_capabilities
from regional_documents import fill_tianfu_reference
from docx import Document
OUT=ROOT/'reports/nationwide-localization-20261023-night/document-qa'
NIGHT=OUT.parent


def main():
    global OUT
    parser=argparse.ArgumentParser();parser.add_argument('--codes',required=True);parser.add_argument('--output-dir',type=Path,required=True);args=parser.parse_args()
    codes=args.codes.split(',');OUT=args.output_dir.resolve()
    if not OUT.is_relative_to((NIGHT/'batches').resolve()):raise ValueError('New QA batches must not overwrite completed QA')
    if not all(c in EXTRA_BINDINGS for c in codes):raise ValueError('Unbound locality')
    output=OUT/'generated';output.mkdir(parents=True,exist_ok=True)
    checks=[];unique={};groups={};reference=None
    samples={'application':{'applicantName':'测试劳动者（示例）','phone':'','respondentName':'测试用人单位（示例）',
                            'requests':['请求支付测试工资100元。','请求支付测试加班工资50元。'],
                            'factsReason':'本文件仅用于排版测试，不用于真实仲裁申请。劳动者与测试公司之间的工资支付存在争议。',
                            'committee':'排版测试委员会'},
             'evidence':{'evidenceList':[{'name':'劳动合同','pageCount':'3','itemCount':'1','originalType':'复印件','purpose':'说明测试劳动关系和工资约定。'}]}}
    for code in codes:
        guide=with_reference_capabilities(json.loads((NIGHT/'sichuan/captures'/code/'guide.json').read_text('utf8')))
        template=next(t for t in guide['templates'] if t.get('can_generate'))
        source=ROOT/template['file'];reference=reference or source
        assert hashlib.sha256(source.read_bytes()).hexdigest()==template['sha256']
        for kind,data in samples.items():
            document=fill_tianfu_reference(source,kind,data)
            file=output/f'{code}-{kind}.docx';document.save(file)
            with ZipFile(source) as before,ZipFile(file) as after:
                assert set(before.namelist())==set(after.namelist())
                for name in before.namelist():
                    if name!='word/document.xml':assert before.read(name)==after.read(name),name
            parsed=Document(file);section=parsed.sections[0]
            assert (section.page_width.twips,section.page_height.twips)==(11906,16838)
            text='\n'.join(p.text for p in parsed.paragraphs)
            assert '授权委托书' not in text
            if kind=='application':assert '1. ' in text and '2. ' in text and '请本人手写签名' in text
            else:assert len(parsed.tables)==1 and len(parsed.tables[0].rows)==2
            sha=hashlib.sha256(file.read_bytes()).hexdigest();unique.setdefault(sha,file);groups.setdefault(sha,[]).append(str(file.relative_to(ROOT)))
            checks.append({'code':code,'kind':kind,'source':template['file'],'source_sha256':template['sha256'],
                           'output':str(file.relative_to(ROOT)),'sha256':sha,'preserved_package':'PASS','A4_geometry':'PASS'})
    spec=importlib.util.spec_from_file_location('night_word_renderer',ROOT/'scripts/verify-regional-documents.py')
    renderer=importlib.util.module_from_spec(spec);spec.loader.exec_module(renderer);renderer.OUT=OUT
    tasks=renderer.render([reference,*unique.values()])
    for item in checks:assert hashlib.sha256((ROOT/item['source']).read_bytes()).hexdigest()==item['source_sha256']
    result={'checks':checks,'identical_output_groups':groups,'tasks':tasks,'engine':'own hidden Word read-only export + bundled Poppler',
            'all_originals_unchanged':True,'model_calls':0,'visual_inspection':'pending'}
    (OUT/'render-manifest.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),'utf8')
    print(json.dumps({'checked_outputs':len(checks),'unique_outputs':len(unique),'pages':[len(t['pages']) for t in tasks]}))


if __name__=='__main__':main()
