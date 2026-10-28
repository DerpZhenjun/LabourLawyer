"""Concise, province-by-province Word reports from saved program/AI evidence."""
import argparse,json,sys
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'backend'))
from portal_audit_report import document,font,url_line
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

def picture(doc,folder,obs):
    filename=obs.get('screenshot','');path=folder/filename
    if not filename or not path.is_file():
        doc.add_paragraph('本次未取得可用截图。');return
    # Portrait PDF captures otherwise move alone onto the following page,
    # leaving a largely empty province page. Keep the actual pixels intact.
    with Image.open(path) as image:width=min(15,12.5*image.width/image.height)
    doc.add_picture(str(path),width=Cm(width))
    doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph('本次实际页面截图  '+filename,style='Caption')
    p.paragraph_format.keep_with_next=False

def create(run):
    data=json.loads((run/'experiences.json').read_text(encoding='utf8'));status=json.loads((run/'status.json').read_text(encoding='utf8'))
    paths=[]
    for ai_report in (False,True):
        title='各省劳动争议办事指南更新报告' if not ai_report else '各省劳动争议办事指南AI复核报告'
        d=document(title,run)
        for name,size in [('Title',20),('Heading 1',17),('Heading 2',12),('Normal',10),('Caption',9)]:font(d.styles[name],'微软雅黑',size)
        d.add_paragraph('本报告逐省列出本次自动访问记录、官方入口和实际截图。办事指南的地区不等于全省统一模板的适用范围。')
        counts=status['counts'];d.add_paragraph(f"31 个地区：本次取得有效指南 {counts['fresh']} 个，沿用此前记录 {counts['historical']} 个，尚未取得可确认指南 {counts['not_detected']} 个。")
        if ai_report:d.add_paragraph(f"OpenClaw 使用 Ubuntu http://10.66.66.1:11434 的 qwen3.8:27b；完成图像复核 {status['ai_reviewed']}/31。AI 意见与程序记录分别保留。")
        else:d.add_paragraph('入口与表格优先采用省级政务服务门户；程序优先选择省会，需要区县时选择第一项。未确认省会的记录会明确标出。')
        for slug,p in data['provinces'].items():
            d.add_page_break();d.add_heading(p['name'],1)
            folder=run/slug;observations=json.loads((folder/'observations.json').read_text(encoding='utf8'))
            state={'fresh':'本次已打开有效办事指南','historical':'本次未取得更新指南  沿用此前记录','not_detected':'本次没有取得可确认的办事指南'}[p['state']]
            d.add_paragraph(state)
            attempt=p.get('openclaw_attempt') or {}
            if attempt and p['state']!='fresh':
                d.add_paragraph('本次补采结果：'+str(attempt.get('error') or attempt.get('state',''))[:160])
                if attempt.get('url'):url_line(d,'本次访问网址',attempt['url'])
            url_line(d,'省级门户',p['portal_url'])
            if p.get('source_url'):url_line(d,'指南网址',p['source_url'])
            recorded_at=attempt.get('updated_at') if p['state']=='not_detected' and attempt else p.get('captured_at') or p.get('checked_at','')
            d.add_paragraph('指南地区：'+p.get('scope','未确认')+'  记录日期：'+(recorded_at or '')[:10])
            if not p.get('capital_confirmed'):d.add_paragraph('本条未确认省会指南。'+('其来源日期保留原记录。' if p['state']=='historical' else ''))
            guide=p.get('selected_guide') or {}
            if guide.get('source_kind')=='registered_official_pdf_fallback':
                d.add_paragraph('门户路线未取得指南；本条采用重新获取的官方部门 PDF，第'+ '、'.join(str(x) for x in guide.get('pdf_pages',[]))+'页。适用范围按该部门或区县来源标注。')
            elif guide.get('source_kind')=='capital_official_department_fallback':
                d.add_paragraph('本条采用省会人社部门公开指南，适用地区为'+guide['region_name']+'。')
            service={'mediation':'劳动争议调解  不是仲裁申请','arbitration':'劳动人事争议仲裁申请','unknown':'尚未确认事项'}[p.get('service_kind','unknown')]
            d.add_paragraph('办理事项：'+service)
            if ai_report:
                ai=p.get('ai_review',{});d.add_heading('截图复核结果',2)
                if ai.get('state')!='reviewed':d.add_paragraph('图像复核未完成：'+ai.get('error',ai.get('state','未知状态'))[:180])
                elif ai.get('parse_state')!='valid':d.add_paragraph('模型已返回意见，但结构校验未通过。详见该省 ai.json。')
                else:
                    for item in ai.get('findings',[]):
                        d.add_paragraph(item['stage']+'：'+str(item.get('notes',''))[:210])
                for message in ai.get('discrepancies',[]):d.add_paragraph(message)
                for message in ai.get('resolutions',[]):d.add_paragraph(message)
            else:
                d.add_heading('办理步骤',2)
                if guide.get('workflow_representation')=='official_flowchart_image':d.add_paragraph('办理步骤见官方流程图，原图已保存在本省证据目录。')
                elif not guide.get('steps') and not guide.get('workflow_text'):d.add_paragraph('本次尚未提取到完整办理步骤，请查看官方指南。')
                else:d.add_paragraph(' → '.join(s['title'] for s in p.get('steps',[])[:7]) or '请在官方页面查看办理流程。')
                templates=p.get('templates',[])
                d.add_paragraph('已检测表格：'+'、'.join(dict.fromkeys(t.get('material_name') or t.get('text') or '文书下载控件' for t in templates)) if templates else '本次没有检测到文书模板。')
                for t in templates:
                    if t.get('content_note'):d.add_paragraph(t['content_note'])
                if p.get('service_kind')=='mediation':d.add_paragraph('需要仲裁时，请另找明确标明仲裁申请的入口。')
            candidates=[o for o in observations if o['stage'].endswith('-templates')]
            obs=candidates[-1] if candidates else next((o for o in reversed(observations) if o.get('screenshot')), {})
            if obs:picture(d,folder,obs)
            else:d.add_paragraph('本省没有取得新截图，旧证据仍保存在原批次。')
        path=run/(title+'.docx');d.save(path);paths.append(path)
    print('\n'.join(str(p) for p in paths))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('run',type=Path);create(parser.parse_args().run.resolve())
