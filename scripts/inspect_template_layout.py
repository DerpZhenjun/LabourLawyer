"""Read-only OOXML inventory used to author exact template field mappings."""
from pathlib import Path
import json
import sys
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))
from province_documents import PROVINCE_DOCUMENTS, get_province_document

def paragraph(p):
    return {"text": p.text, "style": p.style.name,
            "runs": [{"text": r.text, "underline": str(r.underline), "bold": r.bold,
                      "font": r.font.name, "size": r.font.size.pt if r.font.size else None,
                      "xml": r._r.xml} for r in p.runs], "xml": p._p.xml}

def inspect():
    result = {}
    for slug, item in PROVINCE_DOCUMENTS.items():
        for kind in item["documents"]:
            _, path = get_province_document(slug, kind)
            doc = Document(path)
            tables = []
            for table in doc.tables:
                rows = []
                for row in table.rows:
                    seen, cells = set(), []
                    for index, cell in enumerate(row.cells):
                        if cell._tc in seen:
                            continue
                        seen.add(cell._tc)
                        cells.append({"index": index, "text": cell.text,
                                      "paragraphs": [paragraph(p) for p in cell.paragraphs]})
                    rows.append(cells)
                tables.append(rows)
            result[f"{slug}/{kind}"] = {"path": str(path),
                "paragraphs": [paragraph(p) for p in doc.paragraphs], "tables": tables}
    output = ROOT / ".runtime" / "document-audit" / "template_inventory.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return output

if __name__ == "__main__":
    print(inspect())
