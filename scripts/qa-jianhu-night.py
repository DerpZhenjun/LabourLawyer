"""Own retained blanks, own synthetic outputs and fresh rendering for Jianhu."""
import argparse
import hashlib
import importlib
import json
import sys
from pathlib import Path
from zipfile import ZipFile
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'backend'));sys.path.insert(0,str(ROOT/'scripts'))
from docx import Document
from regional_adapters import with_reference_capabilities
from regional_blank_documents import fill_blank_document
from regional_guides import digest
qa=importlib.import_module('qa-xishan-night')
BASE=ROOT/'reports/nationwide-localization-20261023-night/cross-province/jiangsu-jianhu'
OUT=BASE/'document-qa'
IDS={'application':'985f4770-1ab2-4264-9482-0f8528b16513','evidence':'586cf49f-714b-4338-81ce-08a2956d7aa6'}
qa.BASE=BASE;qa.OUT=OUT;qa.IDS=IDS


def main():
    p=argparse.ArgumentParser();p.add_argument('--reference',action='store_true');a=p.parse_args()
    refs=qa.references()
    if a.reference:
        if (OUT/'reference-render.json').exists():raise ValueError('Preserve prior reference render')
        tasks=qa.render([ROOT/r['render_copy'] for r in refs.values()])
        qa.save(OUT/'reference-render.json',tasks)
        print(json.dumps({'pages':[len(t['pages']) for t in tasks]}));return
    if not (OUT/'artifact.md').is_file():raise ValueError('Own template contract required before authoring')
    if (OUT/'render-manifest.json').exists():raise ValueError('Preserve completed QA; new run required for changes')
    g=with_reference_capabilities(json.loads((BASE/'guide.json').read_text('utf8')))
    samples={'application':{'applicantName':'建湖排版测试劳动者（示例）','gender':'女','birthDate':'1992-08-16',
            'idNumber':'证件号码请本人填写','idAddress':'建湖县示例路88号（示例）','applicantMailingAddress':'建湖县示例路88号（示例）',
            'respondentName':'建湖示例科技有限公司（示例）','registeredAddress':'建湖县示例路120号（示例）',
            'acceptMediation':'待本人选择','legalAid':'待本人确认','committeeName':'请填写实际受理的仲裁委员会',
            'requests':['请求支付测试工资100元。','请求支付测试加班工资50元。'],
            'factsReason':'本文件仅用于排版测试，不用于真实仲裁申请。劳动者与测试公司之间的工资支付存在争议。','copyCount':1},
            'evidence':{'committeeName':'请填写实际受理的仲裁委员会','evidenceList':[{'name':'劳动合同（示例）',
                'source':'本人保存','pageCount':'3','purpose':'说明测试劳动关系与工资约定。','originalType':'复印件'}]}}
    checks=[];files=[]
    for kind,data in samples.items():
        t=next(t for t in g['templates'] if t['id']==IDS[kind]);source=ROOT/t['file']
        if hashlib.sha256(source.read_bytes()).hexdigest()!=t['sha256']:raise ValueError('Own original SHA changed')
        dest=OUT/'generated'/(kind+'.docx');dest.parent.mkdir(parents=True,exist_ok=True)
        if dest.exists():raise ValueError('Never overwrite generated QA file')
        fill_blank_document(g,t,kind,source,data).save(dest)
        with ZipFile(source) as before,ZipFile(dest) as after:
            assert set(before.namelist())==set(after.namelist())
            for name in before.namelist():
                if name!='word/document.xml':assert before.read(name)==after.read(name),name
        doc=Document(dest);section=doc.sections[0]
        assert (section.page_width.twips,section.page_height.twips)==(11906,16838)
        if kind=='application':
            text='\n'.join(p.text for p in doc.paragraphs)
            assert '建湖排版测试劳动者' in text and '用人单位申请' not in text
            assert text.count('被申请人：')==1
        else:assert len(doc.tables)==1 and len(doc.tables[0].columns)==6 and len(doc.tables[0].rows)==2
        checks.append({'kind':kind,'annex_id':t['id'],'source':t['file'],'source_sha256':t['sha256'],
                       'output':dest.relative_to(ROOT).as_posix(),'sha256':hashlib.sha256(dest.read_bytes()).hexdigest(),
                       'preserved_package':'PASS','A4_geometry':'PASS'})
        files.append(dest)
    tasks=qa.render(files)
    for t in g['templates']:
        assert hashlib.sha256((ROOT/t['file']).read_bytes()).hexdigest()==t['sha256']
    qa.save(OUT/'render-manifest.json',{'source_id':digest(g['source_url'])[:24],'checks':checks,'tasks':tasks,
            'engine':'own hidden Word read-only export + bundled Poppler','all_originals_unchanged':True,'visual_inspection':'pending'})
    print(json.dumps({'outputs':2,'pages':[len(t['pages']) for t in tasks]}))


if __name__=='__main__':main()
