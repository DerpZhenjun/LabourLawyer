"""Generate two synthetic own-template outputs, preserve originals, render all pages."""
import argparse,hashlib,importlib,json,re,subprocess,sys
from pathlib import Path
from zipfile import ZipFile
from docx import Document
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'backend'));sys.path.insert(0,str(ROOT/'scripts'))
from regional_adapters import with_reference_capabilities
from regional_blank_documents import fill_blank_document
BASE=ROOT/'reports/nationwide-localization-20261023-night/cross-province/jiangsu-haizhou'
IDS={'application':'50286982-5aaa-4310-b593-5767344cea86','evidence':'2a18398d-be86-4613-a878-e17ed90f0354'}
qa=importlib.import_module('qa-xishan-night')


def main():
    p=argparse.ArgumentParser();p.add_argument('--iteration',default='generated-0555');args=p.parse_args()
    if not re.fullmatch(r'generated-[a-z0-9-]+',args.iteration):raise ValueError('Invalid own iteration')
    out=BASE/'document-qa'/args.iteration
    if out.exists():raise ValueError('Preserve old output QA evidence')
    out.mkdir();qa.OUT=out
    g=with_reference_capabilities(json.loads((BASE/'staging-0540/reviewed-guide.json').read_text('utf8')))
    samples={'application':{'applicantName':'海州合成甲（示例）','gender':'女','birthDate':'1993-06-15',
        'idNumber':'请本人填写真实证件号','idAddress':'连云港市海州区示例路88号（示例）',
        'applicantMailingAddress':'连云港市海州区示例路88号（示例）','phone':'请本人填写',
        'respondentName':'海州合成测试公司（示例）','respondentType':'请按实际填写',
        'registeredAddress':'连云港市海州区示例路120号（示例）','mailingAddress':'连云港市海州区示例路120号（示例）',
        'socialCreditCode':'请按企业登记填写','legalRepName':'请按企业登记填写','legalRepPosition':'请按实际填写',
        'departmentHead':'合成主管（示例）','departmentHeadPosition':'部门负责人（示例）',
        'companyContact':'合成联系人（示例）','companyContactPhone':'请按实际填写',
        'acceptMediation':'由本人选择','legalAid':'由本人确认','committeeName':'请核实实际受理的仲裁委员会',
        'requests':['请求支付合成测试工资100元。','请求支付合成测试加班工资50元。'],
        'factsReason':'本文件仅用于本地版式测试，不用于真实劳动仲裁申请。'+('劳动者与测试公司对工资支付存在争议，相关事实由本人按真实情况填写。'*7),
        'copyCount':3,'evidenceCopyCount':2},
        'evidence':{'committeeName':'请核实实际受理的仲裁委员会','submitterName':'海州合成甲（示例）','submitDate':'2026-10-24',
        'evidenceList':[{'name':'劳动合同（合成示例）','source':'本人保存','originalType':'复印件','pageCount':'3页',
                         'purpose':'说明测试劳动关系及工资约定。'+('本条仅为合成版式验证，实际证明内容由本人确认。'*3)}]}}
    checks=[];files=[]
    for short,(kind,data) in zip(['a','e'],samples.items()):
        t=next(t for t in g['templates'] if t['id']==IDS[kind]);source=ROOT/t['file']
        assert hashlib.sha256(source.read_bytes()).hexdigest()==t['sha256'] and t['can_generate']
        dest=out/(short+'.docx');fill_blank_document(g,t,kind,source,data).save(dest)
        with ZipFile(source) as before,ZipFile(dest) as after:
            assert set(before.namelist())==set(after.namelist())
            for part in before.namelist():
                if part!='word/document.xml':assert before.read(part)==after.read(part),part
        doc=Document(dest)
        assert (doc.sections[0].page_width.twips,doc.sections[0].page_height.twips)==(11906,16838)
        if kind=='application':assert not any('用人单位申请' in p.text for p in doc.paragraphs)
        else:assert len(doc.tables[0].columns)==6 and len(doc.tables[0].rows)==2
        checks.append({'kind':kind,'annex_id':t['id'],'original':t['file'],'original_sha256':t['sha256'],
            'output':dest.relative_to(ROOT).as_posix(),'sha256':hashlib.sha256(dest.read_bytes()).hexdigest(),
            'non_document_xml_parts':'byte_identical','page_geometry':'A4_source_preserved'})
        files.append(dest)
    skill=Path('C:/Users/DerpZhenjun/.codex/plugins/cache/openai-primary-runtime/documents/26.1017.11052/skills/documents')
    probe=subprocess.run([sys.executable,str(skill/'render_docx.py'),str(files[0]),'--output_dir',str(out/'packaged'),'--emit_pdf'],
        capture_output=True,text=True,encoding='utf8',errors='replace',timeout=60)
    (out/'packaged-render.log').write_text(probe.stdout+'\n'+probe.stderr,encoding='utf8')
    tasks=qa.render(files)
    assert all(hashlib.sha256((ROOT/t['file']).read_bytes()).hexdigest()==t['sha256'] for t in g['templates'])
    qa.save(out/'render-manifest.json',{'source_id':'f2208816780787a1a046b91f','checks':checks,'tasks':tasks,
        'all_originals_unchanged':True,'visual_inspection':'pending','published':False,'models':0})
    print(json.dumps({'outputs':2,'pages':[len(t['pages']) for t in tasks],'visual_inspection':'pending','published':False}))


if __name__=='__main__':main()
