"""Own-source synthetic outputs; preserved package and complete rendering QA."""
import hashlib
import importlib
import json
from pathlib import Path
import sys
from zipfile import ZipFile
from docx import Document
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'backend'));sys.path.insert(0,str(ROOT/'scripts'))
from regional_adapters import with_reference_capabilities
from regional_blank_documents import fill_blank_document
from regional_guides import digest
BASE=ROOT/'reports/nationwide-localization-20261023-night/cross-province/jiangsu-dafeng'
OUT=BASE/'document-qa';qa=importlib.import_module('qa-xishan-night');qa.OUT=OUT
IDS={'application':'8984153f-2c04-4676-a8fd-3a627658fb9d','evidence':'b104e1d9-15a7-489a-ad8b-80e22e0673e7'}


def main():
    if not (OUT/'artifact.md').is_file() or not (OUT/'reference-render.json').is_file():raise ValueError('Own reference contract required')
    if (OUT/'render-manifest.json').exists():raise ValueError('Preserve prior QA; changes need a new run')
    g=with_reference_capabilities(json.loads((BASE/'guide.json').read_text('utf8')))
    samples={'application':{'applicantName':'大丰排版测试劳动者（示例）','gender':'女','birthDate':'1992-08-16',
        'idNumber':'证件号码请本人填写','idAddress':'大丰区示例路88号（示例）','applicantMailingAddress':'大丰区示例路88号（示例）',
        'respondentName':'大丰示例科技有限公司（示例）','registeredAddress':'大丰区示例路120号（示例）',
        'acceptMediation':'待本人选择','legalAid':'待本人确认','committeeName':'请填写实际受理的仲裁委员会',
        'requests':['请求支付测试工资100元。','请求支付测试加班工资50元。'],
        'factsReason':'本文件仅用于排版测试，不用于真实仲裁申请。劳动者与测试公司之间的工资支付存在争议。','copyCount':1},
        'evidence':{'committeeName':'请填写实际受理的仲裁委员会','evidenceList':[{'name':'劳动合同（示例）',
            'source':'本人保存','pageCount':'3','purpose':'说明测试劳动关系与工资约定。','originalType':'复印件'}]}}
    files=[];checks=[]
    for kind,data in samples.items():
        t=next(t for t in g['templates'] if t['id']==IDS[kind]);source=ROOT/t['file']
        assert hashlib.sha256(source.read_bytes()).hexdigest()==t['sha256']
        dest=OUT/'generated'/(kind+'.docx');dest.parent.mkdir(parents=True,exist_ok=True)
        if dest.exists():raise ValueError('Never overwrite generated evidence')
        fill_blank_document(g,t,kind,source,data).save(dest)
        with ZipFile(source) as before,ZipFile(dest) as after:
            assert set(before.namelist())==set(after.namelist())
            for part in before.namelist():
                if part!='word/document.xml':assert before.read(part)==after.read(part),part
        doc=Document(dest);s=doc.sections[0];assert (s.page_width.twips,s.page_height.twips)==(11906,16838)
        if kind=='application':
            body='\n'.join(p.text for p in doc.paragraphs)
            assert '大丰排版测试劳动者' in body and '用人单位申请' not in body and body.count('被申请人：')==1
        else:assert len(doc.tables[0].columns)==6 and len(doc.tables[0].rows)==2
        checks.append({'kind':kind,'annex_id':t['id'],'source':t['file'],'source_sha256':t['sha256'],
            'output':dest.relative_to(ROOT).as_posix(),'sha256':hashlib.sha256(dest.read_bytes()).hexdigest(),
            'preserved_package':'PASS','A4_geometry':'PASS'})
        files.append(dest)
    tasks=qa.render(files)
    assert all(hashlib.sha256((ROOT/t['file']).read_bytes()).hexdigest()==t['sha256'] for t in g['templates'])
    qa.save(OUT/'render-manifest.json',{'source_id':digest(g['source_url'])[:24],'checks':checks,'tasks':tasks,
        'engine':'own hidden Word read-only export + bundled Poppler','all_originals_unchanged':True,'visual_inspection':'pending',
        'published':False,'browser_acceptance':'NOT_RUN'})
    print(json.dumps({'own_synthetic_outputs':2,'pages':[len(t['pages']) for t in tasks],'published':False}))


if __name__=='__main__':main()
