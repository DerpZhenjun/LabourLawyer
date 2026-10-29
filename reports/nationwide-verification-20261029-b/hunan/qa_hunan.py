"""Own legacy DOC read-only rendering and explicitly derived DOCX conversions."""
from pathlib import Path
import hashlib,json,subprocess,sys
from zipfile import ZipFile
from pypdf import PdfReader
from docx import Document
OUT=Path(__file__).resolve().parent;ROOT=OUT.parents[2]
POPPLER=Path('C:/Users/DerpZhenjun/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/poppler/Library/bin/pdftoppm.exe')
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def save(name,data):
 p=OUT/name;p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf8')
def prepare():
 tasks=[]
 for name in ['tianxin','xiangjiang']:
  records=json.loads((OUT/name/'originals.json').read_text(encoding='utf8'))
  for r in records:
   source=OUT/r['file'];assert sha(source)==r['sha256'];identity=source.stem
   folder=OUT/'render-originals'/name/(r['official_control']+'-'+identity);folder.mkdir(parents=True,exist_ok=True)
   task={'source':str(source),'source_sha256':r['sha256'],'pdf':str(folder/'document.pdf'),'district':name,'role':r['official_control'],'identity':identity}
   if r['official_control']=='blank':
    derived=OUT/'derived'/name/(identity+'.docx');derived.parent.mkdir(parents=True,exist_ok=True);task['derived_docx']=str(derived)
   tasks.append(task)
 save('render-originals/inputs.json',tasks)
def render():
 manifest=OUT/'render-originals/inputs.json'
 p=subprocess.run(['powershell.exe','-NoProfile','-NonInteractive','-File',str(OUT/'render-isolated.ps1'),'-Manifest',str(manifest)],capture_output=True,text=True,timeout=180,creationflags=subprocess.CREATE_NO_WINDOW)
 (manifest.parent/'word-render.log').write_text(p.stdout+'\n'+p.stderr,encoding='utf8')
 if p.returncode:raise RuntimeError(p.stdout+p.stderr)
 print('Word complete; original DOCs preserved and blank-only derived DOCX created.')
def rasterize():
 tasks=json.loads((OUT/'render-originals/inputs.json').read_text(encoding='utf8'));derived=[]
 for t in tasks:
  assert sha(t['source'])==t['source_sha256'];pdf=Path(t['pdf'])
  subprocess.run([str(POPPLER),'-r','105','-png',str(pdf),str(pdf.parent/'page')],check=True,timeout=60,creationflags=subprocess.CREATE_NO_WINDOW)
  reader=PdfReader(pdf);text='\n\f\n'.join(p.extract_text(extraction_mode='layout') for p in reader.pages);(pdf.parent/'text.txt').write_text(text,encoding='utf8')
  t.update(pdf_sha256=sha(pdf),page_count=len(reader.pages),pages=[str(p.relative_to(OUT)).replace('\\','/') for p in sorted(pdf.parent.glob('page-*.png'))])
  if t.get('derived_docx'):
   d=Document(t['derived_docx']);derived.append({'district':t['district'],'source':str(Path(t['source']).relative_to(OUT)).replace('\\','/'),'source_sha256':t['source_sha256'],
     'file':str(Path(t['derived_docx']).relative_to(OUT)).replace('\\','/'),'sha256':sha(t['derived_docx']),'bytes':Path(t['derived_docx']).stat().st_size,
     'kind':'platform_conversion_of_own_official_doc_not_official_docx','method':'isolated Word SaveAs2 format16 after original PDF export','table_shapes':[[len(tb.rows),len(tb.columns)] for tb in d.tables],
     'paragraphs':[p.text for p in d.paragraphs],'tables':[[[c.text for c in row.cells] for row in tb.rows] for tb in d.tables]})
 isolation=json.loads((OUT/'render-originals/word-instance-isolation.json').read_text(encoding='utf-8-sig'))
 assert isolation['completed']==8 and isolation['new_pid_verified'] and isolation['original_pids_still_alive']
 save('DERIVED-MANIFEST.json',{'scope':'Platform conversions; not official DOCX originals','files':derived})
 save('original-render-result.json',{'originals_unchanged':True,'tasks':tasks,'isolation':isolation,'visual_review':'PENDING_INDIVIDUAL_VIEW'})
 print(json.dumps([{'district':t['district'],'identity':t['identity'],'role':t['role'],'pages':t['page_count']} for t in tasks]))

def build_filled():
 sys.path.insert(0,str(ROOT/'backend'))
 from regional_hunan_oct29b import parse_capture,with_capabilities,fill_document
 tasks=[];packages=[];payloads={}
 for name,label in [('tianxin','天心'),('xiangjiang','湘江新区')]:
  g=with_capabilities(parse_capture(name));save(name+'/guide.json',g);t=next(t for t in g['templates'] if t['can_generate'])
  values={'applicantName':'虚构'+label+'测试甲','respondentName':'虚构'+label+'测试公司乙','gender':'女','birthDate':'1990-01-02','nationality':'汉族','occupation':'虚构测试职员',
    'idAddress':'虚构测试住所一号','phone':'','workUnit':'虚构工作单位','registeredAddress':'虚构单位住所二号','legalRepName':'虚构负责人丙','legalRepPosition':'测试负责人',
    'hrContactName':'虚构联系人丁','hrContactPosition':'虚构测试职员','committeeName':'虚构'+label+'测试仲裁委员会','applicationDate':'2026-10-29',
    'requests':['请求支付虚构测试工资100元。','请求支付虚构测试补偿200元。'],
    'factsReason':'本文件全部内容均为软件排版测试虚构示例，不用于真实仲裁。申请人与被申请人的名称、地址、日期、金额和证据均为虚构，仅用于检查本区自己官方表格转换副本的换行及分页。'*2,
    'submitterName':'不得自动填写手签','evidenceList':[{'name':'虚构测试证据甲','source':'虚构测试保存材料','purpose':'仅用于测试本表附随证据说明区域。'},{'name':'虚构测试证据乙','source':'虚构测试保存材料','purpose':'用于核对第二条证据的完整换行。'}]}
  payloads[name]=values;source=ROOT/t['file'];before=sha(source);dest=OUT/'qa'/name/'worker-application-synthetic.docx';dest.parent.mkdir(parents=True,exist_ok=True)
  if dest.exists():raise ValueError('Preserve retained fictional QA output')
  fill_document(g,t,'application',source,values).save(dest);derived=OUT/'derived'/name/'1624951667887.docx'
  with ZipFile(derived) as a,ZipFile(dest) as b:
   assert set(a.namelist())==set(b.namelist());changed=[n for n in a.namelist() if a.read(n)!=b.read(n)]
  assert changed==['word/document.xml'] and sha(source)==before
  packages.append({'district':name,'original_doc':t['file'],'original_sha256':before,'derived_docx':str(derived.relative_to(OUT)).replace('\\','/'),'derived_sha256':sha(derived),
      'output':str(dest.relative_to(OUT)).replace('\\','/'),'output_sha256':sha(dest),'changed_derived_package_parts':changed})
  folder=OUT/'render-filled'/name;folder.mkdir(parents=True,exist_ok=True);tasks.append({'source':str(dest),'source_sha256':sha(dest),'pdf':str(folder/'document.pdf'),'district':name})
 save('qa/package-checks.json',packages);save('qa/synthetic-payloads.json',payloads);save('render-filled/inputs.json',tasks)
def render_filled():
 manifest=OUT/'render-filled/inputs.json'
 p=subprocess.run(['powershell.exe','-NoProfile','-NonInteractive','-File',str(OUT/'render-isolated.ps1'),'-Manifest',str(manifest)],capture_output=True,text=True,timeout=120,creationflags=subprocess.CREATE_NO_WINDOW)
 (manifest.parent/'word-render.log').write_text(p.stdout+'\n'+p.stderr,encoding='utf8')
 if p.returncode:raise RuntimeError(p.stdout+p.stderr)
 tasks=json.loads(manifest.read_text(encoding='utf8'))
 for t in tasks:
  assert sha(t['source'])==t['source_sha256'];pdf=Path(t['pdf'])
  subprocess.run([str(POPPLER),'-r','110','-png',str(pdf),str(pdf.parent/'page')],check=True,timeout=60,creationflags=subprocess.CREATE_NO_WINDOW)
  reader=PdfReader(pdf);text='\n\f\n'.join(p.extract_text(extraction_mode='layout') for p in reader.pages);(pdf.parent/'text.txt').write_text(text,encoding='utf8')
  compact=''.join(text.split());assert '虚构测试证据甲' in compact and '虚构测试证据乙' in compact and '100元' in compact and '200元' in compact and '2026年10月29日' in compact and '不得自动填写手签' not in compact
  t.update(pdf_sha256=sha(pdf),page_count=len(reader.pages),pages=[str(p.relative_to(OUT)).replace('\\','/') for p in sorted(pdf.parent.glob('page-*.png'))])
 isolation=json.loads((OUT/'render-filled/word-instance-isolation.json').read_text(encoding='utf-8-sig'));assert isolation['completed']==2 and isolation['new_pid_verified'] and isolation['original_pids_still_alive']
 save('filled-render-result.json',{'tasks':tasks,'isolation':isolation,'visual_review':'PENDING_INDIVIDUAL_VIEW','source_originals_unchanged':True})
 print(json.dumps([{'district':t['district'],'pages':t['page_count']} for t in tasks]))
if __name__=='__main__':{'prepare':prepare,'render':render,'rasterize':rasterize,'build-filled':build_filled,'render-filled':render_filled}[sys.argv[1]]()
