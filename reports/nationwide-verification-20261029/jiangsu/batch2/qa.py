"""Own-source originals/generation QA only. No DB, model or service action."""
import hashlib
import json
from pathlib import Path
import subprocess
import sys
from zipfile import ZipFile

ROOT=Path(__file__).resolve().parents[4]
OUT=Path(__file__).resolve().parent
BASE=OUT.parent
sys.path.insert(0,str(ROOT/'backend'))
POPPLER=Path('C:/Users/DerpZhenjun/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/poppler/Library/bin/pdftoppm.exe')


def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def save(name,value):(OUT/name).write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n',encoding='utf8')


def originals():
    tasks=[]
    for code in ['320281','320282']:
        guide=json.loads((BASE/code/'guide.json').read_text(encoding='utf8'))
        for template in guide['templates']:
            source=ROOT/template['file']
            assert sha(source)==template['sha256']
            if source.suffix!='.docx':continue
            folder=OUT/'render-originals'/(code+'-'+template['id']);folder.mkdir(parents=True,exist_ok=True)
            tasks.append(dict(source=str(source),pdf=str(folder/'document.pdf'),sha256=sha(source),
                              label=folder.name,region_code=code,annex_id=template['id']))
    save('originals-render-input.json',tasks)
    render('originals')


def audit():
    from docx import Document
    from regional_jiangsu_oct29 import candidate_contract
    records=[]
    skill=Path('C:/Users/DerpZhenjun/.codex/plugins/cache/openai-primary-runtime/documents/26.1017.22136/skills/documents')
    for code in ['320281','320282']:
        guide=json.loads((BASE/code/'guide.json').read_text(encoding='utf8'))
        contract=candidate_contract(guide)
        save(code+'-adaptation-contract.json',contract)
        for kind,item in contract['documents'].items():
            source=ROOT/item['source_file'];doc=Document(source)
            sections=[dict(width=s.page_width,height=s.page_height,top=s.top_margin,bottom=s.bottom_margin,
                           left=s.left_margin,right=s.right_margin,header=s.header_distance,footer=s.footer_distance)
                      for s in doc.sections]
            with ZipFile(source) as zipped:
                parts=[{'part':p.filename,'bytes':len(zipped.read(p.filename)),
                        'sha256':hashlib.sha256(zipped.read(p.filename)).hexdigest(),
                        'editable':p.filename=='word/document.xml'} for p in zipped.infolist()]
            records.append(dict(region_code=code,kind=kind,source=item['source_file'],sha256=sha(source),
                                sections_emu=sections,package_parts=parts,
                                paragraphs=[p.text for p in doc.paragraphs],
                                tables=[[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]))
            prefix=OUT/(code+'-'+kind)
            result=subprocess.run([sys.executable,str(skill/'scripts/section_audit.py'),str(source)],
                                   capture_output=True,text=True,timeout=30)
            prefix.with_suffix('.sections.txt').write_text(result.stdout+result.stderr,encoding='utf8')
            subprocess.run([sys.executable,str(skill/'scripts/style_lint.py'),str(source),'--json',str(prefix.with_suffix('.styles.json'))],
                           capture_output=True,text=True,timeout=30,check=True)
    save('source-package-audit.json',records)
    print(json.dumps([{'region_code':r['region_code'],'kind':r['kind'],'sections':r['sections_emu']} for r in records],ensure_ascii=False))


def generate(group='filled'):
    from regional_jiangsu_oct29 import candidate_contract,fill_candidate_document
    records=[];tasks=[]
    for code in ['320281','320282']:
        g=json.loads((BASE/code/'guide.json').read_text(encoding='utf8'))
        contract=candidate_contract(g)
        payload=dict(committeeName=g['district']+'劳动人事争议仲裁委员会',applicantName='虚构测试甲',gender='女',birthDate='1990-01-02',
            idNumber='000000199001020000',idAddress='虚构测试路1号',applicantMailingAddress='虚构收件路2号',phone='00000000000',
            acceptMediation='是',legalAid='否',respondentName='虚构测试单位有限公司',respondentType='有限责任公司',
            registeredAddress='虚构单位路3号',socialCreditCode='TESTONLY0000000001',legalRepName='虚构测试乙',legalRepPosition='经理',
            departmentHead='虚构测试丙',departmentHeadPosition='主管',mailingAddress='虚构单位收件路4号',companyContact='虚构测试丁',companyContactPhone='00000000001',
            requests=['请求支付虚构测试工资100元。'],factsReason='仅用于软件测试，姓名、单位、号码、地址及事实全部虚构。双方对虚构测试工资存在争议。',
            copyCount=3,evidenceCopyCount=2,evidenceList=[dict(name='虚构劳动合同',source='虚构本人保存',purpose='仅为软件测试，说明虚构劳动关系。',originalType='复印件',pageCount='3页')])
        for kind,item in contract['documents'].items():
            target=OUT/(code+'-'+kind+('-final' if group=='final' else '')+'.docx');assert not target.exists()
            fill_candidate_document(g,kind,payload).save(target)
            with ZipFile(ROOT/item['source_file']) as a,ZipFile(target) as b:
                assert a.namelist()==b.namelist()
                assert [n for n in a.namelist() if a.read(n)!=b.read(n)]==['word/document.xml']
                preserved=len(a.namelist())-1
            assert sha(ROOT/item['source_file'])==item['source_sha256']
            records.append(dict(region_code=code,kind=kind,source=item['source_file'],source_sha256=item['source_sha256'],
                output=target.relative_to(ROOT).as_posix(),output_sha256=sha(target),changed_parts=['word/document.xml'],preserved_parts=preserved,synthetic_only=True))
            folder=OUT/('render-'+group)/(code+'-'+kind);folder.mkdir(parents=True,exist_ok=True)
            tasks.append(dict(source=str(target),pdf=str(folder/'document.pdf'),sha256=sha(target),label=folder.name,region_code=code,kind=kind))
    save(group+'-package-checks.json',records);save(group+'-render-input.json',tasks)
    render(group)


def render(group):
    manifest=OUT/(group+'-render-input.json')
    tasks=json.loads(manifest.read_text(encoding='utf8'))
    command=['powershell.exe','-NoLogo','-NoProfile','-NonInteractive','-STA','-File',
             str(OUT/'render-isolated.ps1'),'-Manifest',str(manifest)]
    result=subprocess.run(command,capture_output=True,text=True,timeout=240,creationflags=subprocess.CREATE_NO_WINDOW)
    (OUT/(group+'-word-render.log')).write_text(result.stdout+'\n'+result.stderr,encoding='utf8')
    assert result.returncode==0,result.stderr
    isolation=json.loads((OUT/'word-instance-isolation.json').read_text(encoding='utf-8-sig'))
    assert isolation['new_pid_verified'] and isolation['original_pids_still_alive']
    assert isolation['completed']==len(tasks) and isolation['link_updates_disabled']
    shutil_target=OUT/(group+'-word-isolation.json')
    shutil_target.write_bytes((OUT/'word-instance-isolation.json').read_bytes())
    from pypdf import PdfReader
    rendered=[]
    for task in tasks:
        pdf=Path(task['pdf']);assert sha(Path(task['source']))==task['sha256']
        result=subprocess.run([str(POPPLER),'-png','-scale-to','1600',str(pdf),str(pdf.parent/'page')],
                              capture_output=True,text=True,timeout=120,creationflags=subprocess.CREATE_NO_WINDOW)
        assert result.returncode==0,result.stderr
        pages=sorted(pdf.parent.glob('page-*.png'));assert len(pages)==len(PdfReader(pdf).pages)
        rendered.append(dict(**task,pages=len(pages),page_images=[p.relative_to(ROOT).as_posix() for p in pages],
                             engine='distinct new PID/Hwnd read-only Word PDF + bundled Poppler'))
    save(group+'-render-manifest.json',rendered)
    print(json.dumps({'group':group,'files':len(tasks),'pages':sum(r['pages'] for r in rendered)}))


if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf8')
    if sys.argv[1]=='originals':originals()
    elif sys.argv[1]=='audit':audit()
    elif sys.argv[1]=='generate':generate()
    elif sys.argv[1]=='final':generate('final')
    else:raise ValueError('Unknown action')
