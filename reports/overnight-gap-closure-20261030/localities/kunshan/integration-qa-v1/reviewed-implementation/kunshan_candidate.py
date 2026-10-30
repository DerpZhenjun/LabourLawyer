"""Private Kunshan own-source candidate. Never registers or enables public generation.

All locality facts and writable slots were mapped against Kunshan's own frozen
guide and four own annexes. Only locality-free Word editing utilities are reused.
"""
from copy import deepcopy
from datetime import date
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re
import sys

ROOT=Path(__file__).resolve().parents[3]
sys.dont_write_bytecode=True
sys.path.insert(0,str(ROOT/'backend'))
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Row
from regional_documents import PreservedPackageDocument
from WenShu.template_fill import fill_label_slots,replace_span,set_cell_text
from WenShu.exact_template_fill import keep_body,fill_sections,remove_empty_edges
sys.path.pop(0)

ADAPTER='kunshan_own_blank_oct31g_private_v1'
FOLDER='reports/overnight-gap-closure-20261030/localities/kunshan/source-discovery/continuation-20261031-native'
FREEZE_SHA='daa6ecda9f78f398c196ba579a4d870591e02a328aca3d5665ae9177151f2c6d'
SERVICE='1132058301418949XF400201430200201'
SOURCE_ID='d39444ab7f4bcaea2ea91473'
ANNEX={'application':'3c17b2ea-434e-4b29-8fe4-40e88ad11c57',
       'evidence':'de3a0ee9-0827-4f44-b9ad-ba05010a03d9'}
APPLICATION_FIELDS=frozenset('committeeName applicantName gender birthDate idNumber idAddress otherDocumentType otherDocumentNumber applicantMailingAddress phone acceptMediation legalAid respondentName respondentType registeredAddress socialCreditCode legalRepName legalRepPosition departmentHead departmentHeadPosition mailingAddress companyContact companyContactPhone requests factsReason applicationDate copyCount evidenceCopyCount'.split())
EVIDENCE_FIELDS=frozenset({'committeeName','caseNumber','submitDate','evidenceList'})
ITEM_FIELDS=frozenset({'name','source','purpose','originalType','pageCount'})
# Measured from this own original's worker block, not another locality profile.
OWN_SLOTS={1:[('申请人','applicantName'),('性别','gender'),('出生日期','birthDate')],
 2:[('公民身份号码','idNumber'),('居住地址','idAddress')],
 3:[('其他证件类型','otherDocumentType'),('证件号码','otherDocumentNumber')],
 4:[('通讯地址','applicantMailingAddress'),('移动电话','phone')],
 5:[('是否接受调解','acceptMediation'),('是否法律援助','legalAid')],
 6:[('被申请人','respondentName'),('性质','respondentType')],
 7:[('注册地址（住所）','registeredAddress')],8:[('统一社会信用代码','socialCreditCode')],
 9:[('法定代表人或者主要负责人姓名','legalRepName'),('职务','legalRepPosition')],
 10:[('劳动者所在部门负责人','departmentHead'),('职务','departmentHeadPosition')],
 11:[('通讯地址','mailingAddress')],12:[('联系人','companyContact'),('移动电话','companyContactPhone')]}


def require(condition,message):
    if not condition:raise ValueError(message)


def exact(actual,expected):
    if type(actual) is not type(expected):return False
    if isinstance(expected,dict):return actual.keys()==expected.keys() and all(exact(actual[k],v) for k,v in expected.items())
    if isinstance(expected,list):return len(actual)==len(expected) and all(exact(a,b) for a,b in zip(actual,expected))
    return actual==expected


def path_at(root,relative):
    require(root is not None and str(root).strip(),'Explicit own-source root required')
    root=Path(root).resolve(strict=True);relative=Path(relative)
    require(not relative.is_absolute() and '..' not in relative.parts,'Relative own-source path required')
    p=(root/relative).resolve(strict=True)
    require(p.is_relative_to(root) and p.relative_to(root).as_posix()==relative.as_posix(),'Own-source path escaped root')
    return p


def read_json(root,relative):return json.loads(path_at(root,relative).read_text('utf-8-sig'))


def prepare_candidate(root):
    freeze_raw=path_at(root,FOLDER+'/SOURCE-FREEZE.json').read_bytes()
    require(sha256(freeze_raw).hexdigest()==FREEZE_SHA,'Own source freeze differs')
    freeze=json.loads(freeze_raw)
    require(freeze['region_code']=='320583' and freeze['source_id']==SOURCE_ID
            and freeze['service_code']==SERVICE and freeze['online_status']=='auth_required'
            and freeze['published'] is False,'Own source identity or boundary differs')
    for file,pin in freeze['files'].items():
        body=path_at(root,file).read_bytes()
        require(len(body)==pin['bytes'] and sha256(body).hexdigest()==pin['sha256'],'Frozen own source file differs')
    bindings=read_json(root,FOLDER+'/own-native-bindings.json')
    downloads=read_json(root,FOLDER+'/own-source-downloads.json')
    source_url=bindings['source_url']
    require(sha256(json.dumps(source_url,ensure_ascii=False,sort_keys=True).encode()).hexdigest()[:24]==SOURCE_ID,
            'Own guide source identity differs')
    captured=read_json(root,'reports/overnight-gap-closure-20261030/localities/kunshan/source-discovery/own-guide.json')
    fields={}
    for table in captured['tables']:
        for row in table:
            if len(row) in {2,4}:
                for i in range(0,len(row),2):fields[row[i]]=row[i+1]
    materials=[];templates=[]
    for index,material in enumerate(bindings['materials']):
        controls=material['controls'];blank=next(c for c in controls if c['kind']=='空白表格')
        url=re.search(r"downloadClyb\('([^']+)'",blank['handler']).group(1)
        annex=url.split('attachGuid=')[1]
        own=next(d for d in downloads if d.get('annex_id')==annex)
        require(own['source_url']==source_url and own['url']==url and own['status']==200,'Own blank binding differs')
        details={}
        for row in material['detail_fields']:
            for i in range(0,len(row)-1,2):details[row[i]]=row[i+1]
        materials.append({'id':annex,'name':material['name'],'paper_copies':material['paper_copies'],'details':details})
        templates.append({'id':annex,'region_code':'320583','material_id':annex,'official_kind':'blank',
            'source_url':source_url,'download_url':url,'file':FOLDER+'/'+own['file'],'sha256':own['sha256'],
            'bytes':own['bytes'],'can_generate':False,'candidate_document_types':[k for k,v in ANNEX.items() if v==annex],
            'content_role':'fillable_own_blank' if annex in ANNEX.values() else 'heading_only_copy_placeholder'})
    return {'province':'jiangsu','city':'苏州市','district':'昆山市','region_code':'320583',
        'source_id':SOURCE_ID,'source_url':source_url,'service_code':SERVICE,
        'service_name':'劳动人事争议仲裁申请','authority':'昆山市人力资源和社会保障局',
        'source_version':FREEZE_SHA,'verification_status':'source_verified_candidate_only',
        'fields':fields,'materials':materials,'templates':templates,
        'channel':{'raw_value':fields['办理形式'],'online_status':'auth_required',
            'observed_login_origin':'https://www.jszwfw.gov.cn','submission_verified':False},
        'candidate_adapter':ADAPTER,'can_generate':False,'published':False,
        'publication_status':'private_offline_candidate','Word_stage':'NOT_REVIEWED_BY_ROOT',
        'warnings':['只支持本区自身劳动者申请书单一被申请人及六列证据清单。',
                    '身份与证据复印件原件仅有标题，不提供填写生成。',
                    '网上办理已到官方登录及验证码边界；未登录、未提交申请。',
                    '源指南申请书4份与证据清单2份要求保留；原件副本栏须本人核对。',
                    '签名、提交人与接收人由本人或经办人手写；调解、法律援助不自动选择。']}


def template_for(guide,kind):
    require(kind in ANNEX,'Unsupported own document kind')
    return next(t for t in guide['templates'] if t['id']==ANNEX[kind])


def default_data(kind):
    require(kind in ANNEX,'Unsupported own document kind')
    return {'evidenceList':[{k:'' for k in sorted(ITEM_FIELDS)}]} if kind=='evidence' else {'requests':['']}


def checked_text(value,maxlen=300):
    require(type(value) is str and len(value)<=maxlen and not re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',value),'Text value/type/length refused')
    return value.strip()


def validate_data(kind,data):
    require(type(data) is dict and kind in ANNEX,'Own form object required')
    allowed=APPLICATION_FIELDS if kind=='application' else EVIDENCE_FIELDS
    require(not set(data)-allowed,'Unknown own field refused')
    output={}
    for key,value in data.items():
        if key in {'requests','evidenceList'}:continue
        if key in {'copyCount','evidenceCopyCount'}:
            require(type(value) is int and 1<=value<=30,'Explicit copy count must be1..30');output[key]=value
        else:output[key]=checked_text(value,12000 if key=='factsReason' else 300)
        if key.endswith('Date') and output[key]:
            require(re.fullmatch(r'\d{4}-\d{2}-\d{2}',output[key]),'ISO date required')
            date.fromisoformat(output[key])
    if kind=='application':
        require(all(output.get(k) for k in ('applicantName','respondentName','factsReason')),'Required own fields blank')
        requests=data.get('requests');require(type(requests) is list and 1<=len(requests)<=4,'One to four explicit requests required')
        output['requests']=[checked_text(x,1200) for x in requests];require(all(output['requests']),'Blank request refused')
        for key in ('acceptMediation','legalAid'):
            require(output.get(key,'') in {'','是','否'},'Unrecognized explicit choice refused')
    else:
        rows=data.get('evidenceList');require(type(rows) is list and 1<=len(rows)<=40,'One to forty evidence items required')
        output['evidenceList']=[]
        for row in rows:
            require(type(row) is dict and not set(row)-ITEM_FIELDS,'Unknown evidence item field refused')
            parsed={key:checked_text(row.get(key,''),1200 if key=='purpose' else 300) for key in ITEM_FIELDS}
            require(parsed['name'],'Evidence name required')
            require(parsed['originalType'] in {'','原件','复印件'},'Explicit original/copy choice refused')
            output['evidenceList'].append(parsed)
    return output


def written_date(value):
    if not value:return ''
    d=date.fromisoformat(value);return f'{d.year}年{d.month}月{d.day}日'


def application_fill(doc,data):
    require(len(doc.paragraphs)==118 and doc.paragraphs[0].text=='劳动人事争议仲裁申请书（劳动者申请）'
            and doc.paragraphs[59].text=='劳动人事争议仲裁申请书（用人单位申请）','Own application layout changed')
    # Drop employer form and source-metadata footer based on this original map.
    keep_body(doc,0,57);ps=list(doc.paragraphs)
    for index,slots in OWN_SLOTS.items():
        fill_label_slots(ps[index],[(label,data.get(key,'')) for label,key in slots],pad=False)
    if data.get('committeeName'):replace_span(ps[44],0,len(ps[44].text),data['committeeName'])
    if data.get('applicationDate'):replace_span(ps[48],0,len(ps[48].text),written_date(data['applicationDate']))
    for label,key in [('本申请书副本','copyCount'),('相关证据','evidenceCopyCount')]:
        match=re.search(re.escape(label)+r'(?P<blank>\s+)(?=份)',ps[51].text)
        require(match,'Own footer slot changed')
        if key in data:replace_span(ps[51],match.start('blank'),match.end('blank'),str(data[key]))
    for p in ps[13:20]:p._p.getparent().remove(p._p)
    fill_sections(doc,data)
    # A signature belongs with its date; its name is never filled automatically.
    for p in ps[43:48]:p.paragraph_format.keep_with_next=True
    ps[48].paragraph_format.keep_with_next=False
    for p in ps[43:49]:p.paragraph_format.keep_together=True
    remove_empty_edges(doc)


def evidence_fill(doc,data):
    require(len(doc.paragraphs)==9 and len(doc.tables)==1 and len(doc.tables[0].rows)==13
            and len(doc.tables[0].columns)==6,'Own evidence layout changed')
    table=doc.tables[0];items=data['evidenceList'];rows=list(table.rows)[1:]
    require([c.text for c in table.rows[0].cells]==['序号','证据名称','证据来源','证  明  内  容','原件或复印件','页数'],
            'Own evidence columns changed')
    header_properties=table.rows[0]._tr.get_or_add_trPr()
    if header_properties.find(qn('w:tblHeader')) is None:header_properties.append(OxmlElement('w:tblHeader'))
    while len(rows)<len(items):
        node=deepcopy(rows[-1]._tr);rows[-1]._tr.addnext(node);rows.append(_Row(node,table))
    for i,(row,item) in enumerate(zip(rows,items),1):
        properties=row._tr.get_or_add_trPr()
        if properties.find(qn('w:cantSplit')) is None:properties.append(OxmlElement('w:cantSplit'))
        for height in row._tr.iter(qn('w:trHeight')):
            if height.get(qn('w:hRule'))=='exact':height.set(qn('w:hRule'),'atLeast')
        for cell,value in zip(row.cells,[str(i),item['name'],item['source'],item['purpose'],item['originalType'],item['pageCount']]):set_cell_text(cell,value)
    for row in rows[len(items):]:table._tbl.remove(row._tr)
    if data.get('committeeName'):replace_span(doc.paragraphs[0],0,len(doc.paragraphs[0].text),data['committeeName'])
    if data.get('caseNumber'):replace_span(doc.paragraphs[3],0,len(doc.paragraphs[3].text),data['caseNumber'])
    if data.get('submitDate'):
        p=doc.paragraphs[5];start=p.text.index('提交时间：')+len('提交时间：')
        replace_span(p,start,len(p.text),written_date(data['submitDate']))
    remove_empty_edges(doc)


def fill_candidate(guide,template,kind,data,*,root,source_version):
    require(kind in ANNEX and source_version==FREEZE_SHA,'Unsupported document or stale own source version')
    canonical=prepare_candidate(root)
    require(exact(guide,canonical),'Own locality/source/capability context differs')
    own=template_for(canonical,kind)
    require(exact(template,own) and own['official_kind']=='blank' and own['content_role']=='fillable_own_blank',
            'Own annex binding refused; samples/headings cannot generate')
    validated=validate_data(kind,data)
    raw=path_at(root,own['file']).read_bytes()
    require(sha256(raw).hexdigest()==own['sha256'],'Own original changed')
    doc=Document(BytesIO(raw))
    (application_fill if kind=='application' else evidence_fill)(doc,validated)
    return PreservedPackageDocument(doc,raw)
