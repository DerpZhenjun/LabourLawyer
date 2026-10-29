"""Offline source packaging only; no DB, network, Word, models or publication."""
from datetime import datetime,timezone
import hashlib,json,re,sys
from pathlib import Path
from urllib.parse import parse_qs,urlsplit
from docx import Document

HERE=Path(__file__).resolve().parent
BASE=Path('reports/nationwide-verification-20261029-e/jiangsu')
NOW=datetime.now(timezone.utc).isoformat()
SPECS={
 'xiangcheng':('相城区','320507','28','11320507555810103L400201430200201','苏州市相城区人力资源和社会保障局'),
 'wujiang':('吴江区','320509','29','1132058401420395XF400201430200201','苏州市吴江区人力资源和社会保障局'),
}
def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def read(relative):return json.loads((HERE/relative).read_text('utf8'))
def save(relative,value):
 p=HERE/relative;p.parent.mkdir(parents=True,exist_ok=True)
 p.write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
def pin(relative):
 p=HERE/relative
 return {'sha256':sha(p),'bytes':p.stat().st_size}
if sys.argv[1:]==['--verify']:
 # Frozen public verification does not require undistributed sample bytes and
 # must not regenerate timestamps, captures, guides or source-manifest bytes.
 white=read('PUBLIC-WHITELIST.json')
 for relative,expected in white['files'].items():
  relative=Path(relative).relative_to(BASE)
  path=(HERE/relative).resolve()
  assert path.is_relative_to(HERE.resolve()) and pin(relative)==expected
 manifest=read('SOURCE-MANIFEST.json')
 for relative,expected in manifest['files'].items():
  path=(HERE/relative).resolve()
  assert path.is_relative_to(HERE.resolve()) and pin(relative)==expected
 print(json.dumps({'public_files_sha_passed':len(white['files']),'source_files_sha_passed':len(manifest['files']),
                   'source_manifest_sha256':sha(HERE/'SOURCE-MANIFEST.json'),'read_only':True}))
 raise SystemExit(0)
if sys.argv[1:]:raise SystemExit('Use no arguments to build before freezing, or --verify for frozen read-only checks')
def pairs(table):
 result={}
 for row in table:
  for index in range(0,len(row)-1,2):result[row[index].strip()]=row[index+1].strip()
 return result
def section(text,label,following):return text.rsplit(label+'\n',1)[1].split(following+'\n',1)[0].strip()
public=[]
for locality,(district,code,webid,service,authority) in SPECS.items():
 initial=read(locality+'/formal-controls/page.json');source=initial['url'];query=parse_qs(urlsplit(source).query)
 assert query['webId']==[webid] and query['iddept_yw_inf']==[service]
 assert initial['title']=='劳动人事争议仲裁申请' and authority in initial['text']
 search_folder='search'
 search=read(locality+'/'+search_folder+'/page.json')
 assert parse_qs(urlsplit(search['url']).query)['websiteid']==[code+'000000000']
 assert '当前搜索范围：'+district in search['text'] and '可办地区:\n'+district in search['text']
 fields={}
 for table in initial['tables']:
  if table and table[0] and table[0][0] in {'事项类型','实施主体','是否进驻政务大厅','办件类型','办理结果类型','咨询方式'}:
   for key,value in pairs(table).items():
    fields[key]={'raw_value':value,'normalized_value':value or None,'verification_status':'verified' if value else 'missing',
                 'source_url':source,'capture_file':(BASE/locality/'formal-controls/page.json').as_posix(),'capture_sha256':sha(HERE/locality/'formal-controls/page.json')}
 assert fields['业务办理项编码']['raw_value']==service and fields['实施主体']['raw_value']==authority
 material_table=next(t for t in initial['tables'] if t and t[0][:2]==['序号','材料名称'])
 assert len(material_table)==5
 materials=[]
 for index,row in enumerate(material_table[1:],1):
  detail_folder=f'material-{index}-detail'
  detail=read(f'{locality}/{detail_folder}/page.json');name=row[1].strip()
  assert detail['url']==source and '办理材料详情\n材料名称\t'+name in detail['text']
  tables=[t for t in detail['tables'] if t and t[0][:2]==['材料名称',name]]
  assert len(tables)==1
  values=pairs(tables[0])
  assert values['纸质材料份数']+'份'==row[4] and values['材料必要性']==row[5]
  materials.append({'id':code+'-material-'+str(index),'name':name,'required':row[5],'copies':values['纸质材料份数'],
    'raw_copies':row[4],'original_copy':values['材料类型'],'format':values['材料形式'],'paper_size':values['纸质材料规格'],
    'source_channel':values['来源渠道'],'details':values,'source_url':source,'region_code':code,'verification_status':'verified',
    'detail_capture':(BASE/locality/detail_folder/'page.json').as_posix(),
    'detail_capture_sha256':sha(HERE/locality/detail_folder/'page.json')})
 records=read(locality+'/DOWNLOADS.json');assert len(records)==9
 controls=initial['download_controls'];assert len(controls)==8
 observed={parse_qs(urlsplit(re.search(r"downloadClyb\('([^']+)'",c['onclick'])[1]).query)['attachGuid'][0]:c for c in controls}
 templates=[];review=[]
 for record in records:
  path=HERE/record['file'];assert record['status']=='downloaded' and record['http_status']==200 and sha(path)==record['sha256']
  assert record['source_url']==source and record['requested_url']==record['final_url']
  kind=record['official_kind'];published=kind!='sample'
  if kind!='instruction':
   control=observed[record['annex_id']]
   assert record['material_name']==control['material'] and record['control_label']==control['label']
  else:assert record['requested_url']==read(locality+'/flow-image/page.json')['url']
  own={'file':record['file'],'sha256':record['sha256'],'official_kind':kind,'render_status':'NOT_RUN' if record['file_type']=='docx' else 'not_a_docx',
       'public_release_status':'included_blank_or_instruction' if published else 'not_distributed_reference_sample',
       'generation_allowed':False}
  if record['file_type']=='docx':
   document=Document(path);text='\n'.join(p.text for p in document.paragraphs)+'\n'+'\n'.join(c.text for t in document.tables for row in t.rows for c in row.cells)
   own.update(read_status='docx_text_and_structure_only',table_shapes=[{'rows':len(t.rows),'columns':len(t.columns)} for t in document.tables],
              numeric_personal_id_like_count=len(re.findall(r'(?<!\d)\d{17}[\dXx](?!\d)',text)),
              phone_like_count=len(re.findall(r'(?<!\d)1[3-9]\d{9}(?!\d)',text)))
   if kind=='blank':
    own['content_role']=('blank_worker_application' if record['material_name']=='劳动人事争议仲裁申请书' else
                         'blank_evidence_list' if record['material_name']=='劳动人事争议仲裁案件证据清单' else 'placeholder_heading_only')
    assert own['numeric_personal_id_like_count']==0 and own['phone_like_count']==0
   else:own['content_role']='official_sample_never_default_case_facts'
  else:own.update(read_status='own_flow_image_viewed' if kind=='instruction' else 'official_sample_not_distributed',content_role='workflow_instruction' if kind=='instruction' else 'official_sample')
  review.append(own)
  templates.append({'id':record['annex_id'],'annex_id':record['annex_id'],'region_code':code,
     'material_id':next((m['id'] for m in materials if m['name']==record['material_name']),None),
     'official_kind':kind,'original_name':record['original_name'],'source_url':source,'download_url':record['requested_url'],
     'file':(BASE/record['file']).as_posix(),'file_type':record['file_type'],'sha256':record['sha256'],'bytes':record['bytes'],
     'downloaded_at':record['downloaded_at'],'retained_capture_status':'downloaded',
     'download_status':'downloaded' if published else 'reference_sample_not_distributed',
     'public_release_status':own['public_release_status'],'content_role':own['content_role'],
     'content_verification_status':own['read_status'],'can_generate':False,'document_types':[]})
  if published:public.append(record['file'])
 boundary=read(locality+'/online-boundary.json')
 assert boundary['source_url']==source and boundary['status']=='auth_required' and boundary['path']=='/jsjis/front/login.do'
 assert boundary['captcha_visible'] and boundary['login_visible'] and boundary['submitted'] is False
 text=initial['text']
 sections={label:section(text,label,next_label) for label,next_label in [('受理标准','办理材料'),('办理材料','办理流程'),
   ('办理流程','办理时限'),('办理时限','收费情况'),('收费情况','办理结果'),('办理结果','法律依据'),('法律依据','常见问题'),('常见问题','咨询投诉')]}
 sections.update({'基本信息':'\n'.join(k+'：'+v['raw_value'] for k,v in fields.items()),'咨询投诉':'\n'.join(k+'：'+fields[k]['raw_value'] for k in ['咨询方式','监督投诉方式'])})
 issues=[{'type':'COPY_COUNT_CONTEXT','reason':'材料表申请书为4份；自身空白原件注明申请书副本份数应按被申请人人数提交。两处分别保留，不替受理机关统一份数。'},
         {'type':'SOURCE_LAW_TEXT_RETAINED','reason':'法律依据原文写办案规则第17号第三十条；本批没有另行核定所引规范的现行效力。'},
         {'type':'EMPTY_RENDERED_SECTION','reason':'常见问题栏目存在，已检查页面未渲染具体问答；不声称官方没有常见问题。'}]
 guide={'province':'jiangsu','province_name':'江苏省','city':'苏州市','district':district,'region_code':code,
    'service_name':initial['title'],'service_code':service,'authority':authority,'source_url':source,'source_title':initial['title'],
    'source_type':'official_service_guide','source_id':hashlib.sha256(json.dumps(source,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest()[:24],
    'capture_version':'sha256:'+sha(HERE/locality/'formal-controls/page.json'),
    'source_update_date':None,'source_update_date_status':'not_displayed_in_checked_guide',
    'collected_at':initial['captured_at'],'verified_at':initial['captured_at'],
    'verification_status':'partial','fields':fields,'materials':materials,'templates':templates,'sections':sections,
    'scope_identity':{'own_portal_host':urlsplit(source).hostname,'search_websiteid':code+'000000000','native_search_area':district,
                      'native_guide_area_selection':district,'webId':webid,'no_cross_locality_template_inheritance':True},
    'evidence':{'capture_file':(BASE/locality/'formal-controls/page.json').as_posix(),'capture_sha256':sha(HERE/locality/'formal-controls/page.json'),
                'search_capture_file':(BASE/locality/search_folder/'page.json').as_posix(),'search_capture_sha256':sha(HERE/locality/search_folder/'page.json')},
    'channel':{'form':'window_online_and_post','raw_value':fields['办理形式']['raw_value'],'online_status':'auth_required','submitted':False,'post_login_flow':'unknown'},
    'online_boundary':boundary,'source_issues':issues,'template_search_status':'found_in_own_checked_source',
    'faq_status':'no_items_rendered_in_checked_source','attachment_counts':{'blank':4,'sample':4,'instruction':1},
    'generation_adapter':None,'can_generate':False,'published':False,'publication_status':'unpublished','candidate_web_qa':'NOT_RUN',
    'warnings':['本批仅采集自身指南与公开来源，未做新适配或候选网页验收，未发布。','官方样例按原控件分类记录与单独留存，不带入默认案件事实，也不列入本批公开包。','DOCX只读结构与文本检查不等于逐页版式验收，render_status仍为NOT_RUN。']}
 save(locality+'/guide.json',guide)
 save(locality+'/CONTENT-REVIEW.json',{'checked_at':NOW,'files':review,'originals_unchanged':True,'docx_rendered':False,
       'note':'仅文本/结构初查与流程图阅读；不声明DOCX逐页版式已验收，不以同SHA推导地区授权。'})
 folders=['home',search_folder,'guide-area-dialog','formal-controls','flow-image']+[f'material-{i}-detail' for i in range(1,5)]
 if locality=='xiangcheng':folders+=['region-picker']
 for folder in folders:
  for filename in ['page.json','page.txt','page.png','provenance.json']:
   relative=locality+'/'+folder+'/'+filename;assert (HERE/relative).is_file();public.append(relative)
  provenance=read(locality+'/'+folder+'/provenance.json')
  for filename,expected in provenance['files'].items():assert pin(locality+'/'+folder+'/'+filename)==expected
 public += [locality+'/guide.json',locality+'/CONTENT-REVIEW.json',locality+'/DOWNLOADS.json',locality+'/online-boundary.json']
for folder in ['province-home','province-picker','suzhou-home','suzhou-picker']:
 for filename in ['page.json','page.txt','page.png','provenance.json']:public.append(folder+'/'+filename)
save('DIAGNOSTICS.json',{'recorded_at':NOW,'browser':'Agent Browser','session':'js-oct29e','closed':True,
 'private_browser_journal_initialized_before_navigation':True,
 'observations':[{'action':'PowerShell reference quoting','outcome':'an unquoted element reference was interpreted as shell splatting','resolution':'quoted actual observed references; no bypass'},
                 {'action':'material detail selection','outcome':'one local assertion selected a hidden earlier detail table; a nonexistent selector and a shell-encoded text selector also failed','resolution':'retained actual own visible detail captures, selected tables by exact material name and used observed native references/selectors; no DOM modification'},
                 {'action':'Wujiang guide wait','outcome':'wait used the wrong label 申请材料 and timed out','resolution':'own guide rendered 办理材料; complete capture and tables verified'},
                 {'action':'Wujiang online stale reference','outcome':'first reference after tab switch was unknown, so no entry click occurred and the diagnostic boundary was unverified','resolution':'kept initial diagnostic privately; refreshed snapshot, actual entry click then reached login/password/captcha; accepted boundary records the successful native click'},
                 {'action':'public attachment acquisition','outcome':'18 own observed public control/flow HTTPS URLs returned HTTP200 using normal unauthenticated GET; no SSL/port/auth workaround'},
                 {'action':'online native entry','outcome':'both localities stopped at actual user login with captcha; no input, authentication or submission'},
                 {'action':'browser doctor wrapper','outcome':'doctor response did not have the normal data envelope; wrapper raised KeyError','resolution':'private journal retains original response; no doctor success/failure total inferred; own session closed normally'}],
 'public_raw_HAR_cookies_SSO_queries_saved':False})
public += ['PREFLIGHT.json','DIAGNOSTICS.json']
public=sorted(set(public))
save('SOURCE-MANIFEST.json',{'scope':'Own official public captures, safe summaries and blank/instruction originals only; sample original bytes not distributed.',
 'files':{relative:pin(relative) for relative in public}})
public += ['SOURCE-MANIFEST.json','FINAL-REPORT.md','collect.py','finalize.py','verify_source.py','SOURCE-STAGE-CHECK.json']
save('PUBLIC-WHITELIST.json',{'checked_at':NOW,'scope':'Only exact listed files; no recursive directory publication. This is source collection only, not capability/publication approval.',
 'public_originals':10,'collected_originals':18,'samples_not_distributed':8,
 'files':{(BASE/relative).as_posix():pin(relative) for relative in sorted(set(public))},
 'excluded':['all */originals/sample/* original bytes (reference samples not needed for blank-source collection)',
             'all private browser journals, diagnostic files and client temporary output',
             'no raw HAR, cookies, SSO/session queries or login DOM were captured']})
print(json.dumps({'public_files':len(set(public)),'own_guides':2,'collected_originals':18,'public_originals':10,
                  'sample_originals_not_distributed':8,'source_manifest_sha256':sha(HERE/'SOURCE-MANIFEST.json'),
                  'generation_adapters_added':0,'published':False},ensure_ascii=False))
