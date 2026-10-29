"""Read-only package/slot evidence for the retained public originals."""
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile
from docx import Document

ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parent
SOURCE = ROOT / 'reports/nationwide-verification-20261029-d/jiangsu'
records = []
for name in ['gusu', 'wuzhong']:
    g = json.loads((SOURCE / name / 'guide.json').read_text(encoding='utf8'))
    for t in g['templates']:
        if t['official_kind'] != 'blank' or t['content_role'] == 'placeholder_heading_only':
            continue
        p = ROOT / t['file']; raw = p.read_bytes()
        assert hashlib.sha256(raw).hexdigest() == t['sha256']
        d = Document(p)
        with ZipFile(p) as z:
            parts = {n: dict(bytes=len(z.read(n)), sha256=hashlib.sha256(z.read(n)).hexdigest(),
                            edit_policy='editable' if n == 'word/document.xml' else 'preserve_only')
                     for n in z.namelist()}
        records.append(dict(region=g['region_code'], source_file=t['file'], sha256=t['sha256'],
                            paragraphs=[dict(index=i, text=p.text, xml=p._p.xml) for i, p in enumerate(d.paragraphs)],
                            tables=[dict(rows=len(tb.rows), columns=len(tb.columns), xml=tb._tbl.xml) for tb in d.tables],
                            sections=[s._sectPr.xml for s in d.sections], parts=parts))
(OUT / 'source-package-audit.json').write_text(json.dumps(records, ensure_ascii=False, indent=2) + '\n', encoding='utf8')
print(f'{len(records)} public source packages audited; no samples read')
